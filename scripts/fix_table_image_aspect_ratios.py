#!/usr/bin/env python3
"""Repair stretched think-cell table-image pictures in generated decks."""

from __future__ import annotations

import argparse
import json
import shutil
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation

from _directors import canonical_directors
from period_context import DEFAULT_PERIOD, context_for_period


ROOT = Path(__file__).resolve().parent.parent
EMU_PER_INCH = 914400
EMU_PER_POINT = 12700
CONTENT_BOTTOM_MARGIN = int(0.62 * EMU_PER_INCH)
CONTENT_SIDE_MARGIN = int(0.57 * EMU_PER_INCH)
TABLE_PIC_NAME = "Pic"
TABLE_DATA_NAME = "think-cell data - do not delete"
MIN_TABLE_WIDTH_RATIO = 0.25
MIN_TABLE_HEIGHT_RATIO = 0.07
TABLE_TARGET_BOUNDS_PT = {
    4: (41.00, 149.00, 877.76, 137.80),
    5: (41.00, 242.43, 409.45, 192.91),
    6: (41.00, 149.00, 877.76, 334.93),
    7: (41.00, 149.00, 877.76, 335.60),
    8: (41.00, 149.00, 877.76, 335.60),
    9: (41.00, 285.00, 877.76, 78.00),
    11: (41.00, 244.79, 877.76, 236.94),
    12: (41.00, 243.21, 877.76, 127.56),
    13: (586.41, 149.00, 301.84, 173.23),
    16: (41.00, 232.19, 877.76, 251.90),
    18: (471.00, 258.17, 417.25, 225.91),
    19: (41.00, 149.00, 877.76, 185.04),
    21: (468.50, 161.42, 437.01, 181.10),
    22: (41.00, 149.00, 877.76, 185.04),
    23: (41.00, 149.00, 877.76, 208.66),
    24: (41.00, 238.49, 877.76, 245.60),
    25: (41.00, 149.00, 877.76, 335.60),
    26: (41.00, 149.00, 877.76, 190.00),
    27: (41.00, 149.00, 877.76, 335.60),
}


@dataclass
class AspectFix:
    slide: int
    image_px: str
    old_bounds: tuple[int, int, int, int]
    new_bounds: tuple[int, int, int, int]
    old_ratio: float
    image_ratio: float
    distortion_before: float
    distortion_after: float


def _slug(value: str) -> str:
    return value.replace(" ", "-")


def _horizontal_overlap(a_left: int, a_width: int, b_left: int, b_width: int) -> int:
    return max(0, min(a_left + a_width, b_left + b_width) - max(a_left, b_left))


def _available_bottom(prs: Presentation, slide: Any, shape: Any) -> int:
    bottom = int(prs.slide_height) - CONTENT_BOTTOM_MARGIN
    left, top, width = int(shape.left), int(shape.top), int(shape.width)
    for candidate in slide.shapes:
        if candidate is shape:
            continue
        if candidate.name in {TABLE_PIC_NAME, TABLE_DATA_NAME}:
            continue
        candidate_top = int(candidate.top)
        if candidate_top <= top + int(5 * EMU_PER_POINT):
            continue
        overlap = _horizontal_overlap(left, width, int(candidate.left), int(candidate.width))
        if overlap > min(int(width * 0.12), int(0.75 * EMU_PER_INCH)):
            bottom = min(bottom, candidate_top - int(8 * EMU_PER_POINT))
    return max(top + int(0.35 * EMU_PER_INCH), bottom)


def _target_bounds(slide_idx: int) -> tuple[int, int, int, int] | None:
    target = TABLE_TARGET_BOUNDS_PT.get(slide_idx)
    if not target:
        return None
    left, top, width, height = target
    return (
        int(left * EMU_PER_POINT),
        int(top * EMU_PER_POINT),
        int(width * EMU_PER_POINT),
        int(height * EMU_PER_POINT),
    )


def _fit_bounds(
    prs: Presentation, slide: Any, shape: Any, *, slide_idx: int
) -> tuple[int, int, int, int, float, float, float]:
    image_width, image_height = shape.image.size
    if not image_width or not image_height:
        raise ValueError("picture has no readable native image size")
    image_ratio = image_width / image_height
    old_left, old_top = int(shape.left), int(shape.top)
    old_width, old_height = int(shape.width), int(shape.height)
    old_ratio = old_width / old_height if old_height else image_ratio

    is_tiny_placeholder = (
        old_width / int(prs.slide_width) < MIN_TABLE_WIDTH_RATIO
        or old_height / int(prs.slide_height) < MIN_TABLE_HEIGHT_RATIO
    )
    target = _target_bounds(slide_idx) if is_tiny_placeholder else None
    if target:
        old_left, old_top, max_width, max_height = target
    elif is_tiny_placeholder:
        max_width = int(prs.slide_width) - old_left - CONTENT_SIDE_MARGIN
        if max_width < int(prs.slide_width * MIN_TABLE_WIDTH_RATIO):
            old_left = CONTENT_SIDE_MARGIN
            max_width = int(prs.slide_width) - (CONTENT_SIDE_MARGIN * 2)
        max_height = max(1, _available_bottom(prs, slide, shape) - old_top)
    else:
        max_width = old_width
        max_height = max(1, _available_bottom(prs, slide, shape) - old_top)
    new_width = min(max_width, int(max_height * image_ratio))
    new_height = int(new_width / image_ratio)
    if new_height > max_height:
        new_height = max_height
        new_width = int(new_height * image_ratio)

    # Center only when the table is height-constrained inside a wide original
    # placeholder. This removes the visibly stretched strip without shifting
    # full-width tables that already use the content margin.
    new_left = old_left + max(0, int((max_width - new_width) / 2))
    new_top = old_top
    new_ratio = new_width / new_height if new_height else image_ratio
    return new_left, new_top, new_width, new_height, old_ratio, image_ratio, new_ratio


def _set_bounds(shape: Any, bounds: tuple[int, int, int, int]) -> None:
    left, top, width, height = bounds
    shape.left = left
    shape.top = top
    shape.width = width
    shape.height = height


def fix_deck(path: Path, *, backup: bool = False, min_delta: float = 0.03) -> dict[str, Any]:
    path = path.expanduser().resolve()
    prs = Presentation(path)
    fixes: list[AspectFix] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        pics = [shape for shape in slide.shapes if shape.name == TABLE_PIC_NAME]
        if not pics:
            continue
        for pic in pics:
            old_bounds = (int(pic.left), int(pic.top), int(pic.width), int(pic.height))
            new_left, new_top, new_width, new_height, old_ratio, image_ratio, new_ratio = _fit_bounds(
                prs, slide, pic, slide_idx=slide_idx
            )
            before = old_ratio / image_ratio if image_ratio else 1.0
            after = new_ratio / image_ratio if image_ratio else 1.0
            if abs(1.0 - before) < min_delta and old_bounds == (new_left, new_top, new_width, new_height):
                continue
            new_bounds = (new_left, new_top, new_width, new_height)
            _set_bounds(pic, new_bounds)
            for candidate in slide.shapes:
                if candidate.name == TABLE_DATA_NAME:
                    _set_bounds(candidate, new_bounds)
            image_width, image_height = pic.image.size
            fixes.append(
                AspectFix(
                    slide=slide_idx,
                    image_px=f"{image_width}x{image_height}",
                    old_bounds=old_bounds,
                    new_bounds=new_bounds,
                    old_ratio=old_ratio,
                    image_ratio=image_ratio,
                    distortion_before=before,
                    distortion_after=after,
                )
            )
    if fixes:
        if backup:
            backup_path = path.with_suffix(path.suffix + ".aspect-bak")
            if not backup_path.exists():
                shutil.copy2(path, backup_path)
        prs.save(path)
    return {
        "deck": str(path),
        "status": "pass",
        "fix_count": len(fixes),
        "fixes": [asdict(fix) for fix in fixes],
    }


def _linked_deck(period: str, slug: str) -> Path:
    return ROOT / "state" / period / slug / f"{slug}-LAND-{period}-table-image-linked.pptx"


def _meeting_spine(period: str, slug: str) -> Path:
    return ROOT / "state" / period / slug / "factory" / "meeting-spine" / f"{slug}-LAND-{period}-meeting-spine.pptx"


def _selected_directors(director_slug: str | None) -> list[dict[str, Any]]:
    directors = canonical_directors()
    if not director_slug:
        return directors
    return [director for director in directors if _slug(str(director["name"])) == director_slug]


def _deck_paths(args: argparse.Namespace) -> list[Path]:
    paths = [Path(path) for path in args.decks]
    if args.linked_decks or args.meeting_spine_decks:
        context_for_period(args.period)
        for director in _selected_directors(args.director_slug):
            slug = _slug(str(director["name"]))
            if args.linked_decks:
                paths.append(_linked_deck(args.period, slug))
            if args.meeting_spine_decks:
                paths.append(_meeting_spine(args.period, slug))
    if args.package_dir:
        package_dir = args.package_dir.expanduser()
        context = context_for_period(args.period)
        paths.extend(sorted(package_dir.glob(context.meeting_spine_pattern)))
    seen: set[Path] = set()
    unique: list[Path] = []
    for path in paths:
        resolved = path.expanduser().resolve()
        if resolved in seen or not resolved.exists():
            continue
        seen.add(resolved)
        unique.append(resolved)
    return unique


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("decks", nargs="*", type=Path)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--director-slug")
    parser.add_argument("--linked-decks", action="store_true")
    parser.add_argument("--meeting-spine-decks", action="store_true")
    parser.add_argument("--package-dir", type=Path)
    parser.add_argument("--backup", action="store_true")
    parser.add_argument("--jobs", type=int, default=1)
    parser.add_argument("--json-output", type=Path)
    args = parser.parse_args()

    paths = _deck_paths(args)
    if not paths:
        raise SystemExit("no decks selected")
    if args.jobs > 1 and len(paths) > 1:
        with ThreadPoolExecutor(max_workers=min(args.jobs, len(paths))) as executor:
            results = list(executor.map(lambda path: fix_deck(path, backup=args.backup), paths))
    else:
        results = [fix_deck(path, backup=args.backup) for path in paths]
    payload = {
        "schema": "table-image-aspect-fix/v1",
        "status": "pass",
        "period": args.period,
        "deck_count": len(results),
        "total_fixes": sum(int(result["fix_count"]) for result in results),
        "results": results,
    }
    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
