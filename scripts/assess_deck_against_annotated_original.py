#!/usr/bin/env python3
"""Assess a generated monthly deck against an annotated original deck.

The annotated PowerPoint comments are treated as review requirements only. They
are not copied into the generated deck.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation

from _directors import canonical_directors
from period_context import DEFAULT_PERIOD


ROOT = Path(__file__).resolve().parent.parent
HOME = Path.home()
CRM_ANALYTICS = HOME / "crm-analytics"
HEADER_FILL = PatternFill("solid", fgColor="083EA7")
PASS_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
FAIL_FILL = PatternFill("solid", fgColor="FCE4D6")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass(frozen=True)
class CommentRule:
    theme: str
    priority: str
    keywords: tuple[str, ...]
    target_slide_hints: tuple[int, ...]
    evidence_terms: tuple[str, ...]
    requirement: str
    omission_ok: bool = False
    force_manual_review: bool = False


RULES = (
    CommentRule(
        theme="Period and header labeling",
        priority="high",
        keywords=("period", "header", "time stamp", "timestamp", "march 2026", "q1 2026"),
        target_slide_hints=(1, 2, 13),
        evidence_terms=("may 2026", "friday, may 1", "2026-q2", "q2 forward"),
        requirement="Generated deck must state the current review period/date clearly.",
    ),
    CommentRule(
        theme="Pipeline overview data quality",
        priority="high",
        keywords=("top deals", "0 arr", "pipeline overview", "qlik", "trial license", "data we want", "pipeline volume"),
        target_slide_hints=(5, 8, 11, 12),
        evidence_terms=("nonzero", "top 7", "top four", "apac", "amova", "concentration"),
        requirement="Top-deal and overview slides must be APAC-scoped and avoid zero-ARR artifacts.",
    ),
    CommentRule(
        theme="Stage and candidate definition",
        priority="high",
        keywords=("candidates are stage", "stage 2", "stage 3", "progressed"),
        target_slide_hints=(5, 6),
        evidence_terms=("current-quarter", "approval", "readiness", "stage", "q2 forward"),
        requirement="Candidate logic must define the stage/progression universe behind the slide.",
    ),
    CommentRule(
        theme="Renewal watchlist usability",
        priority="medium",
        keywords=("axioma", "annually", "probability", "comment field", "sort on close date", "close date"),
        target_slide_hints=(7,),
        evidence_terms=("fy26 renewal watchlist", "prob", "close", "next step", "renewal"),
        requirement="Renewal view must expose timing/probability/action fields and handle Axioma scope deliberately.",
    ),
    CommentRule(
        theme="Reference visual adaptation",
        priority="medium",
        keywords=("picture", "emea", "apac of course"),
        target_slide_hints=(10,),
        evidence_terms=("apac", "renewal", "acv", "finance"),
        requirement="Visual reference from original comments must be checked manually against the APAC replacement slide.",
        force_manual_review=True,
    ),
    CommentRule(
        theme="Churn source governance",
        priority="medium",
        keywords=("churn", "finance", "alex p"),
        target_slide_hints=(10,),
        evidence_terms=("finance", "churn"),
        requirement="Do not invent churn reporting; use Finance as source of truth before publishing churn claims.",
        omission_ok=True,
    ),
    CommentRule(
        theme="Commercial approval source coverage",
        priority="high",
        keywords=("conditionally approved", "approved cases", "commercial approval"),
        target_slide_hints=(6,),
        evidence_terms=("approval", "approved", "gate", "commercial"),
        requirement="Commercial approval slide must preserve approved/conditional approval source coverage.",
    ),
    CommentRule(
        theme="Push versus slip clarity",
        priority="high",
        keywords=("pushed", "slip", "slipped", "quarter or month", "from when"),
        target_slide_hints=(3, 9, 12),
        evidence_terms=("push", "pushed", "slip", "q1 accountability", "original apac push"),
        requirement="Push/slip slides must distinguish timing movement from quarter slippage.",
    ),
    CommentRule(
        theme="Executive summary placement",
        priority="medium",
        keywords=("exec summary", "summary on top"),
        target_slide_hints=(2,),
        evidence_terms=("operating summary", "summary"),
        requirement="Items called out for executive summary should be visible in the opening operating summary.",
    ),
    CommentRule(
        theme="Pipeline slide alignment",
        priority="medium",
        keywords=("pipeline slide",),
        target_slide_hints=(4, 5, 11),
        evidence_terms=("pipeline", "forecast", "concentration", "q2 forward"),
        requirement="Pipeline story should map to the expected pipeline/forecast surfaces, not a disconnected appendix.",
    ),
    CommentRule(
        theme="Manual visual sanity",
        priority="medium",
        keywords=("looks off",),
        target_slide_hints=(4, 5, 11),
        evidence_terms=("forecast", "pipeline", "concentration"),
        requirement="Reviewer flagged the original visual as suspicious; generated replacement needs human visual review.",
        force_manual_review=True,
    ),
)


def slugify(value: str) -> str:
    return value.replace(" ", "-")


def lower_slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def selected_director(director_slug: str) -> dict[str, Any]:
    for director in canonical_directors():
        if slugify(str(director["name"])) == director_slug:
            return director
    raise SystemExit(f"unknown director slug: {director_slug}")


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def text_content(element: ET.Element) -> str:
    values = [
        child.text or ""
        for child in element.iter()
        if local_name(child.tag) == "t" and child.text
    ]
    return " ".join(" ".join(values).split())


def deck_comment_part_count(path: Path) -> int:
    if not path.exists():
        return 0
    try:
        with ZipFile(path) as zf:
            return len(
                [
                    name
                    for name in zf.namelist()
                    if name.startswith("ppt/comments/") and name.endswith(".xml")
                ]
            )
    except BadZipFile:
        return 0


def default_original_candidates(director: dict[str, Any]) -> list[Path]:
    name = str(director["name"])
    scope = str(director.get("scope_label", ""))
    name_slug = lower_slug(name)
    candidates: list[Path] = []
    downloads = HOME / "Downloads"
    if downloads.exists():
        patterns = [
            f"Sales Director Monthly - {name}*.pptx",
            f"*{name}*.pptx",
            f"*{name_slug}*.pptx",
        ]
        for pattern in patterns:
            candidates.extend(downloads.glob(pattern))
    if scope:
        candidates.append(downloads / f"Sales Director Monthly - {name} ({scope}).pptx")

    candidates.extend(
        [
            CRM_ANALYTICS
            / "output"
            / "sales_director_monthly_runs"
            / "2026-04-10"
            / f"Sales Director Monthly - {name} ({scope}).pptx",
            CRM_ANALYTICS
            / "output"
            / "simcorp_director_decks"
            / "2026-04-20"
            / "land-only"
            / f"{name_slug}-LAND.pptx",
        ]
    )

    seen: set[Path] = set()
    existing: list[Path] = []
    for candidate in candidates:
        resolved = candidate.expanduser()
        if resolved in seen or not resolved.exists():
            continue
        seen.add(resolved)
        existing.append(resolved)
    return sorted(existing, key=lambda path: (deck_comment_part_count(path) == 0, str(path)))


def default_original_deck(director: dict[str, Any]) -> Path:
    candidates = default_original_candidates(director)
    if not candidates:
        name = str(director["name"])
        raise SystemExit(f"no original deck candidate found for {name}")
    return candidates[0]


def default_target_deck(period: str, director_slug: str) -> Path:
    return (
        ROOT
        / "state"
        / period
        / director_slug
        / "factory"
        / "meeting-spine"
        / f"{director_slug}-LAND-{period}-meeting-spine.pptx"
    )


def slide_id_map(path: Path) -> dict[str, int]:
    namespace = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    with ZipFile(path) as zf:
        root = ET.fromstring(zf.read("ppt/presentation.xml"))
    mapping: dict[str, int] = {}
    for index, node in enumerate(root.findall(".//p:sldId", namespace), start=1):
        slide_id = node.attrib.get("id")
        if slide_id:
            mapping[slide_id] = index
    return mapping


def deck_slides(path: Path) -> list[dict[str, Any]]:
    prs = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                texts.append(" ".join(shape.text.split()))
        slides.append(
            {
                "slide_number": slide_number,
                "title": texts[0] if texts else "",
                "text": "\n".join(texts),
            }
        )
    return slides


def extract_comments(path: Path) -> list[dict[str, Any]]:
    id_map = slide_id_map(path)
    comments: list[dict[str, Any]] = []
    with ZipFile(path) as zf:
        part_names = sorted(
            name
            for name in zf.namelist()
            if name.startswith("ppt/comments/") and name.endswith(".xml")
        )
        for part_name in part_names:
            root = ET.fromstring(zf.read(part_name))
            for comment in root.iter():
                if local_name(comment.tag) != "cm":
                    continue
                slide_id = ""
                for child in comment.iter():
                    if local_name(child.tag) == "sldMk" and child.attrib.get("sldId"):
                        slide_id = str(child.attrib["sldId"])
                        break
                tx_body = next(
                    (child for child in list(comment) if local_name(child.tag) == "txBody"),
                    None,
                )
                main_text = text_content(tx_body) if tx_body is not None else ""
                replies: list[str] = []
                for reply in comment.iter():
                    if local_name(reply.tag) == "reply":
                        reply_text = text_content(reply)
                        if reply_text:
                            replies.append(reply_text)
                title = comment.attrib.get("title", "")
                if not main_text and title:
                    main_text = title
                comments.append(
                    {
                        "part": part_name,
                        "comment_id": comment.attrib.get("id", ""),
                        "author_id": comment.attrib.get("authorId", ""),
                        "created": comment.attrib.get("created", ""),
                        "slide_id": slide_id,
                        "slide_number": id_map.get(slide_id),
                        "title": title,
                        "text": main_text,
                        "replies": replies,
                    }
                )
    return comments


def normalized(value: str) -> str:
    value = value.lower()
    value = re.sub(r"\[@[^]]+\]|\B@\w+(?:\s+\w+)?", " ", value)
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return " ".join(value.split())


def classify_comment(comment_text: str, original_slide_title: str = "") -> CommentRule:
    norm = normalized(f"{comment_text} {original_slide_title}")
    for rule in RULES:
        if any(keyword in norm for keyword in rule.keywords):
            return rule
    return CommentRule(
        theme="General reviewer note",
        priority="medium",
        keywords=(),
        target_slide_hints=(),
        evidence_terms=(),
        requirement="Manual review required because no deterministic comment rule matched.",
    )


def target_slide_lookup(slides: list[dict[str, Any]], slide_numbers: tuple[int, ...]) -> list[dict[str, Any]]:
    if not slide_numbers:
        return []
    by_number = {slide["slide_number"]: slide for slide in slides}
    return [by_number[number] for number in slide_numbers if number in by_number]


def assess_comment(comment: dict[str, Any], target_slides: list[dict[str, Any]]) -> dict[str, Any]:
    rule = classify_comment(comment.get("text", ""), comment.get("original_slide_title", ""))
    mapped = target_slide_lookup(target_slides, rule.target_slide_hints)
    mapped_text = "\n".join(str(slide["text"]) for slide in mapped).lower()
    whole_deck_text = "\n".join(str(slide["text"]) for slide in target_slides).lower()
    evidence_hits = [
        term
        for term in rule.evidence_terms
        if term.lower() in mapped_text or term.lower() in whole_deck_text
    ]

    if rule.omission_ok and "churn" not in whole_deck_text:
        status = "covered_by_omission"
        evidence = "Generated deck does not publish unsupported churn claims."
    elif evidence_hits and rule.force_manual_review:
        status = "needs_manual_review"
        evidence = "Text evidence exists, but this visual/screenshot requirement needs human review: " + ", ".join(evidence_hits)
    elif evidence_hits:
        status = "covered"
        evidence = "Matched evidence terms: " + ", ".join(evidence_hits)
    elif mapped:
        status = "needs_manual_review"
        evidence = "Relevant target slide exists, but text extraction did not prove the requirement."
    else:
        status = "gap"
        evidence = "No mapped target slide or evidence term found."

    return {
        **comment,
        "theme": rule.theme,
        "priority": rule.priority,
        "requirement": rule.requirement,
        "target_slide_hints": list(rule.target_slide_hints),
        "target_slide_titles": [slide["title"] for slide in mapped],
        "status": status,
        "evidence": evidence,
    }


def write_rows(ws: Any, rows: list[list[Any]], table_name: str | None = None) -> None:
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    if table_name and ws.max_row > 1 and ws.max_column > 1:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def style_sheet(ws: Any) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            value = str(cell.value or "").lower()
            if value in {"covered", "covered_by_omission", "pass"}:
                cell.fill = PASS_FILL
            elif value in {"needs_manual_review", "review", "partial"}:
                cell.fill = WARN_FILL
            elif value in {"gap", "missing", "fail"}:
                cell.fill = FAIL_FILL
    for index in range(1, ws.max_column + 1):
        letter = get_column_letter(index)
        width = 12
        for cell in ws[letter]:
            if cell.value is None:
                continue
            width = max(width, min(68, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def build_workbook(payload: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True

    summary = payload["summary"]
    summary_rows = [["Field", "Value"]]
    for key, value in summary.items():
        summary_rows.append([key, json.dumps(value) if isinstance(value, dict) else value])
    ws = wb.create_sheet("Assessment_Summary")
    write_rows(ws, summary_rows, table_name="tbl_Assessment_Summary")
    ws.sheet_properties.tabColor = "083EA7"

    comment_rows = [
        [
            "Status",
            "Priority",
            "Theme",
            "Original slide",
            "Original slide title",
            "Comment",
            "Replies",
            "Requirement",
            "Target slides",
            "Target slide titles",
            "Evidence",
            "Created",
            "Comment part",
        ]
    ]
    original_titles = {
        slide["slide_number"]: slide["title"] for slide in payload["original_slides"]
    }
    for item in payload["assessments"]:
        comment_rows.append(
            [
                item["status"],
                item["priority"],
                item["theme"],
                item.get("slide_number") or "",
                original_titles.get(item.get("slide_number"), ""),
                item.get("text", ""),
                "\n".join(item.get("replies", [])),
                item["requirement"],
                ", ".join(str(slide) for slide in item["target_slide_hints"]),
                "\n".join(item["target_slide_titles"]),
                item["evidence"],
                item.get("created", ""),
                item.get("part", ""),
            ]
        )
    ws = wb.create_sheet("Comment_Assessment")
    write_rows(ws, comment_rows, table_name="tbl_Comment_Assessment")

    original_rows = [["Slide", "Title", "Comment count"]]
    comments_by_slide = Counter(
        item.get("slide_number") for item in payload["assessments"] if item.get("slide_number")
    )
    for slide in payload["original_slides"]:
        original_rows.append(
            [slide["slide_number"], slide["title"], comments_by_slide.get(slide["slide_number"], 0)]
        )
    ws = wb.create_sheet("Original_Slides")
    write_rows(ws, original_rows, table_name="tbl_Original_Slides")

    target_rows = [["Slide", "Title", "Text excerpt"]]
    for slide in payload["target_slides"]:
        target_rows.append([slide["slide_number"], slide["title"], str(slide["text"])[:700]])
    ws = wb.create_sheet("Target_Slides")
    write_rows(ws, target_rows, table_name="tbl_Target_Slides")

    wb.save(path)


def markdown_report(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    lines = [
        "# Annotated Original Deck Assessment",
        "",
        f"- Director: {summary['director']}",
        f"- Period: {summary['period']}",
        f"- Original: `{summary['original_deck']}`",
        f"- Target: `{summary['target_deck']}`",
        f"- Comments assessed: {summary['comment_count']}",
        f"- Status counts: {summary['status_counts']}",
        "",
        "| Status | Original slide | Theme | Comment | Evidence |",
        "|---|---:|---|---|---|",
    ]
    for item in payload["assessments"]:
        comment = str(item.get("text", "")).replace("|", "\\|")
        evidence = str(item.get("evidence", "")).replace("|", "\\|")
        if len(comment) > 120:
            comment = comment[:117] + "..."
        lines.append(
            f"| {item['status']} | {item.get('slide_number') or ''} | "
            f"{item['theme']} | {comment} | {evidence} |"
        )
    return "\n".join(lines) + "\n"


def build_assessment(
    *,
    director: dict[str, Any],
    period: str,
    original_deck: Path,
    target_deck: Path,
    out_dir: Path,
) -> dict[str, Path]:
    if not original_deck.exists():
        raise SystemExit(f"missing original deck: {original_deck}")
    if not target_deck.exists():
        raise SystemExit(f"missing target deck: {target_deck}")

    original_slides = deck_slides(original_deck)
    target_slides = deck_slides(target_deck)
    comments = extract_comments(original_deck)
    original_titles = {slide["slide_number"]: slide["title"] for slide in original_slides}
    for comment in comments:
        comment["original_slide_title"] = original_titles.get(comment.get("slide_number"), "")
    assessments = [assess_comment(comment, target_slides) for comment in comments]
    status_counts = dict(Counter(item["status"] for item in assessments))
    summary = {
        "schema": "sales-ops-annotated-original-assessment/v1",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "director": str(director["name"]),
        "period": period,
        "original_deck": str(original_deck),
        "target_deck": str(target_deck),
        "original_slide_count": len(original_slides),
        "target_slide_count": len(target_slides),
        "comment_count": len(comments),
        "comment_part_count": deck_comment_part_count(original_deck),
        "status_counts": status_counts,
        "assessment_mode": "comments-as-requirements-not-comment-migration",
    }
    payload = {
        "summary": summary,
        "original_slides": original_slides,
        "target_slides": target_slides,
        "assessments": assessments,
    }

    slug = slugify(str(director["name"]))
    out_dir.mkdir(parents=True, exist_ok=True)
    base = out_dir / f"{slug}-{period}-annotated-original-assessment"
    json_path = base.with_suffix(".json")
    md_path = base.with_suffix(".md")
    xlsx_path = base.with_suffix(".xlsx")
    json_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    md_path.write_text(markdown_report(payload), encoding="utf-8")
    build_workbook(payload, xlsx_path)
    return {"json": json_path, "md": md_path, "xlsx": xlsx_path}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--director-slug", default="Jesper-Tyrer")
    parser.add_argument("--original-deck", type=Path)
    parser.add_argument("--target-deck", type=Path)
    parser.add_argument("--out-dir", type=Path)
    parser.add_argument("--copy-to-downloads", action="store_true")
    args = parser.parse_args()

    director = selected_director(args.director_slug)
    original_deck = args.original_deck or default_original_deck(director)
    target_deck = args.target_deck or default_target_deck(args.period, args.director_slug)
    out_dir = args.out_dir or (
        ROOT
        / "state"
        / args.period
        / args.director_slug
        / "factory"
        / "original-comment-assessment"
    )
    outputs = build_assessment(
        director=director,
        period=args.period,
        original_deck=original_deck,
        target_deck=target_deck,
        out_dir=out_dir,
    )

    if args.copy_to_downloads:
        for path in list(outputs.values()):
            target = HOME / "Downloads" / path.name
            shutil.copy2(path, target)
            outputs[f"downloads_{path.suffix.lstrip('.')}"] = target

    print("annotated_original_assessment:")
    for kind, path in outputs.items():
        print(f"{kind}={path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
