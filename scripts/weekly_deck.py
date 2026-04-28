#!/usr/bin/env python3
"""
sales-ops-copilot — Weekly PowerPoint deck builder.

Renders a 7-slide weekly Sales Ops review deck using the SimCorp template.
Mirrors deck.py shape but uses 7-day rollup data (snapshots/*.json + per-day
alert sidecars) for the headline + alert-resolution slides, and falls back
gracefully when snapshot history is insufficient.

Slide structure:
    1. Title — "Weekly Sales Ops Rollup — week ending YYYY-MM-DD"
    2. What changed this week (4 bullets)
    3. Multi-quarter weighted forecast (table + bar chart) — current state
    4. Alert resolution — count Δ over 7 days
    5. Account concentration (top 10) — current state, Δ when computable
    6. Owner concentration (top 10) — current state, Δ when computable
    7. Recommended actions for the week

Usage:
    python3 scripts/weekly_deck.py [--out PATH] [--no-llm] [--onedrive-publish]
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import pathlib
import shutil
import sys
from typing import Any

from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

# Reuse helpers + slide builders from deck.py — we are NOT reimplementing
# tables, charts, fonts, colors. Same look + feel guaranteed.
from deck import (  # type: ignore[import-not-found]
    COLOR_BODY_TEXT,
    FOOTER_COLOR,
    LAYOUT_BLANK,
    LAYOUT_TITLE,
    SIMCORP_TEMPLATE,
    _add_bullets,
    _add_footer,
    _add_table,
    _add_textbox,
    _add_title,
    _build_account_concentration_slide,
    _build_actions_slide,
    _build_forecast_slide,
    _build_owner_concentration_slide,
    _fallback_actions,
    _fmt_money,
    _strip_template_demo_slides,
    llm_recommended_actions,
)

ROOT = pathlib.Path(__file__).resolve().parent.parent
REPORTS_DIR = ROOT / "reports"
SNAP_DIR = ROOT / "state" / "snapshots"


# --- Snapshot loaders (mirror weekly_brief.py) ------------------------------


def _snap_path(date: dt.date) -> pathlib.Path:
    return SNAP_DIR / f"{date.isoformat()}.json"


def _alert_path(date: dt.date) -> pathlib.Path:
    return SNAP_DIR / f"{date.isoformat()}_alerts.json"


def _load_json(path: pathlib.Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    try:
        import json as _json

        return _json.loads(path.read_text())
    except Exception:
        return None


def collect_window(today: dt.date, days: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for offset in range(days, -1, -1):
        d = today - dt.timedelta(days=offset)
        snap = _load_json(_snap_path(d))
        if snap is None:
            continue
        rows.append(
            {
                "date": d.isoformat(),
                "snapshot": snap,
                "alerts": _load_json(_alert_path(d)),
            }
        )
    return rows


def _totals(snap: dict[str, Any]) -> dict[str, float]:
    t = (snap or {}).get("totals", {}) or {}
    return {
        "new_business_arr_open": float(t.get("new_business_arr_open_this_quarter") or 0),
        "weighted_new_business_arr": float(t.get("weighted_new_business_arr") or 0),
        "renewal_acv_open": float(t.get("renewal_acv_open_this_quarter") or 0),
        "weighted_renewal_acv": float(t.get("weighted_renewal_acv") or 0),
    }


def _alert_diff(window: list[dict[str, Any]]) -> list[dict[str, Any]]:
    have = [w for w in window if w.get("alerts")]
    if len(have) < 2:
        return []
    first_a = have[0]["alerts"]
    last_a = have[-1]["alerts"]
    first_counts = first_a.get("counts") or {}
    last_counts = last_a.get("counts") or {}
    first_arr = first_a.get("total_arr") or {}
    last_arr = last_a.get("total_arr") or {}
    names = sorted(set(first_counts) | set(last_counts))
    rows: list[dict[str, Any]] = []
    for n in names:
        fc = int(first_counts.get(n) or 0)
        lc = int(last_counts.get(n) or 0)
        fa = float(first_arr.get(n) or 0)
        la = float(last_arr.get(n) or 0)
        rows.append(
            {
                "name": n,
                "first_count": fc,
                "last_count": lc,
                "count_delta": lc - fc,
                "first_arr": fa,
                "last_arr": la,
                "arr_delta": la - fa,
            }
        )
    rows.sort(key=lambda r: abs(r["arr_delta"]), reverse=True)
    return rows


# --- Slide builders specific to the weekly deck -----------------------------


def _build_weekly_title_slide(
    prs: Presentation, today: dt.date, window: list[dict[str, Any]]
) -> Any:
    layout = prs.slide_layouts[LAYOUT_TITLE]
    slide = prs.slides.add_slide(layout)
    if slide.shapes.title is not None:
        slide.shapes.title.text = f"Weekly Sales Ops Rollup — week ending {today.isoformat()}"
    else:
        _add_title(slide, f"Weekly Sales Ops Rollup — week ending {today.isoformat()}")
    if window:
        sub = (
            f"Window: {window[0]['date']} → {window[-1]['date']} "
            f"({len(window)} snapshot{'s' if len(window) != 1 else ''})"
        )
    else:
        sub = "No snapshots in window"
    _add_textbox(
        slide,
        Inches(0.5),
        Inches(2.4),
        Inches(12.33),
        Inches(0.6),
        sub,
        Pt(20),
        color=COLOR_BODY_TEXT,
    )
    _add_textbox(
        slide,
        Inches(0.5),
        Inches(3.2),
        Inches(12.33),
        Inches(0.5),
        "Prepared for: Sales Operations leadership",
        Pt(14),
        color=COLOR_BODY_TEXT,
    )
    _add_textbox(
        slide,
        Inches(0.5),
        Inches(3.7),
        Inches(12.33),
        Inches(0.5),
        "Source: Salesforce + apro-openai (Microsoft Agent Framework)",
        Pt(12),
        color=FOOTER_COLOR,
    )
    return slide


def _what_changed_bullets(window: list[dict[str, Any]]) -> list[str]:
    if len(window) < 2:
        return [
            "Insufficient snapshot history (need ≥2 days). Trend bullets will populate as snapshots accumulate."
        ]
    first = _totals(window[0]["snapshot"])
    last = _totals(window[-1]["snapshot"])
    delta_open = last["new_business_arr_open"] - first["new_business_arr_open"]
    delta_w = last["weighted_new_business_arr"] - first["weighted_new_business_arr"]
    delta_acv = last["renewal_acv_open"] - first["renewal_acv_open"]
    delta_wacv = last["weighted_renewal_acv"] - first["weighted_renewal_acv"]

    def _signed(v: float) -> str:
        sign = "+" if v >= 0 else "-"
        return f"{sign}${abs(v):,.0f}"

    bullets: list[str] = []
    bullets.append(
        f"Open new-business ARR: {_fmt_money(first['new_business_arr_open'])} → "
        f"{_fmt_money(last['new_business_arr_open'])} ({_signed(delta_open)})."
    )
    bullets.append(
        f"Weighted new-business ARR: {_fmt_money(first['weighted_new_business_arr'])} → "
        f"{_fmt_money(last['weighted_new_business_arr'])} ({_signed(delta_w)})."
    )
    bullets.append(
        f"Open renewal ACV: {_fmt_money(first['renewal_acv_open'])} → "
        f"{_fmt_money(last['renewal_acv_open'])} ({_signed(delta_acv)})."
    )
    bullets.append(
        f"Weighted renewal ACV: {_fmt_money(first['weighted_renewal_acv'])} → "
        f"{_fmt_money(last['weighted_renewal_acv'])} ({_signed(delta_wacv)})."
    )
    return bullets


def _build_what_changed_slide(prs: Presentation, window: list[dict[str, Any]]) -> Any:
    layout = prs.slide_layouts[LAYOUT_BLANK]
    slide = prs.slides.add_slide(layout)
    _add_title(slide, "What changed this week")
    bullets = _what_changed_bullets(window)
    _add_bullets(
        slide,
        Inches(0.6),
        Inches(1.5),
        Inches(12.13),
        Inches(5.0),
        bullets,
        font_size=Pt(18),
    )
    _add_textbox(
        slide,
        Inches(0.5),
        Inches(6.5),
        Inches(12.33),
        Inches(0.4),
        "Land+Expand reported in ARR. Renewals reported in ACV. Never blended.",
        Pt(10),
        color=FOOTER_COLOR,
    )
    return slide


def _build_alert_resolution_slide(prs: Presentation, window: list[dict[str, Any]]) -> Any:
    layout = prs.slide_layouts[LAYOUT_BLANK]
    slide = prs.slides.add_slide(layout)
    _add_title(slide, "Alert resolution — 7-day Δ")

    rows_data = _alert_diff(window)
    if not rows_data:
        _add_textbox(
            slide,
            Inches(0.6),
            Inches(2.0),
            Inches(12.13),
            Inches(2.0),
            "Insufficient alert sidecar history (need ≥2 days of "
            "<date>_alerts.json). Trend will populate as the daily brief "
            "runs over the coming week.",
            Pt(16),
            color=COLOR_BODY_TEXT,
        )
        return slide

    headers = ["Alert", "Count Start → End (Δ)", "ARR Start → End (Δ)"]
    rows: list[list[str]] = []
    for r in rows_data[:10]:
        name = str(r.get("name") or "—")
        if len(name) > 60:
            name = name[:57] + "..."
        cd = int(r.get("count_delta") or 0)
        ad = float(r.get("arr_delta") or 0)
        cd_s = f"{cd:+d}"
        ad_s = f"{'+' if ad >= 0 else '-'}${abs(ad):,.0f}"
        rows.append(
            [
                name,
                f"{r['first_count']} → {r['last_count']} ({cd_s})",
                f"${r['first_arr']:,.0f} → ${r['last_arr']:,.0f} ({ad_s})",
            ]
        )
    _add_table(
        slide,
        Inches(0.4),
        Inches(1.4),
        Inches(12.5),
        Inches(5.0),
        headers,
        rows,
        right_align_cols={1, 2},
    )
    _add_textbox(
        slide,
        Inches(0.5),
        Inches(6.5),
        Inches(12.33),
        Inches(0.4),
        "Negative Δ = alert population shrinking (improving). Sorted by absolute ARR movement.",
        Pt(10),
        color=FOOTER_COLOR,
    )
    return slide


# --- Orchestration ----------------------------------------------------------


def build_weekly_deck(
    out_path: pathlib.Path, *, use_llm: bool = True, days: int = 7
) -> pathlib.Path:
    if not SIMCORP_TEMPLATE.exists():
        raise FileNotFoundError(f"SimCorp template not found at {SIMCORP_TEMPLATE}")

    today = dt.date.today()

    print(f"→ Loading snapshots over last {days} days...")
    window = collect_window(today, days)
    print(f"  Found {len(window)} snapshot(s)")

    print("→ Pulling current Salesforce snapshot for forecast slide...")
    from brief import pull_salesforce_snapshot  # type: ignore[import-not-found]

    sf = pull_salesforce_snapshot()

    print("→ Pulling alerts + concentration (current state)...")
    from alerts import (  # type: ignore[import-not-found]
        pull_account_concentration,
        pull_all_alerts,
        pull_owner_concentration,
    )

    alerts = pull_all_alerts()
    owners = pull_owner_concentration(top_n=10)
    accounts = pull_account_concentration(top_n=15)

    print(f"→ Generating recommended actions ({'LLM' if use_llm else 'fallback'})...")
    if use_llm:
        actions = llm_recommended_actions(sf, alerts, owners, accounts, weekly=True)
    else:
        actions = _fallback_actions(sf, alerts, owners, accounts)

    print("→ Building deck from SimCorp template...")
    prs = Presentation(str(SIMCORP_TEMPLATE))
    _strip_template_demo_slides(prs)

    _build_weekly_title_slide(prs, today, window)
    _build_what_changed_slide(prs, window)
    _build_forecast_slide(prs, sf)
    _build_alert_resolution_slide(prs, window)
    _build_account_concentration_slide(prs, accounts)
    _build_owner_concentration_slide(prs, owners)
    _build_actions_slide(prs, actions)

    new_slides = list(prs.slides)[-7:]
    total_new = len(new_slides)
    run_date = today.isoformat()
    for i, s in enumerate(new_slides, 1):
        _add_footer(s, i, total_new, run_date)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"✓ Wrote {out_path}")
    return out_path


def publish_onedrive(deck_path: pathlib.Path) -> pathlib.Path:
    target_dir = (
        pathlib.Path.home() / "Library" / "CloudStorage" / "OneDrive-SimCorp" / "Sales Ops Briefs"
    )
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / deck_path.name
    tmp = target.with_suffix(target.suffix + ".tmp")
    shutil.copyfile(deck_path, tmp)
    os.replace(tmp, target)
    return target


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--out",
        help="Output .pptx path; defaults to reports/weekly-deck-<today>.pptx",
    )
    ap.add_argument("--days", type=int, default=7)
    ap.add_argument(
        "--no-llm",
        action="store_true",
        help="Skip apro-openai recommended-actions generation; use fallback bullets.",
    )
    ap.add_argument(
        "--onedrive-publish",
        action="store_true",
        help="Atomic-write to OneDrive-SimCorp/Sales Ops Briefs/ for Power Automate pickup.",
    )
    args = ap.parse_args()

    today = dt.date.today()
    REPORTS_DIR.mkdir(exist_ok=True)
    out_path = (
        pathlib.Path(args.out)
        if args.out
        else REPORTS_DIR / f"weekly-deck-{today.isoformat()}.pptx"
    )

    deck_path = build_weekly_deck(out_path, use_llm=not args.no_llm, days=args.days)

    if args.onedrive_publish:
        try:
            target = publish_onedrive(deck_path)
            print(f"✓ Published to OneDrive: {target}")
        except Exception as e:
            print(f"  ⚠ OneDrive publish failed: {e}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
