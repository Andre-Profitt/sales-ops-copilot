#!/usr/bin/env python3
"""Build reduced regional meeting-spine deck candidates from linked decks."""

from __future__ import annotations

import argparse
import json
import shutil
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from pptx import Presentation

from _directors import canonical_directors
from fix_table_image_aspect_ratios import fix_deck
from meeting_spine_action_layer import enhance_meeting_spine_deck
from period_context import DEFAULT_PERIOD, context_for_period
from scrub_stale_thinkcell_metadata import scrub_deck


ROOT = Path(__file__).resolve().parent.parent
KEEP_SLIDES = [
    1,   # Cover
    2,   # May operating summary
    4,   # Original targets / Q1 accountability
    5,   # Forecast quality
    7,   # May deal readiness
    9,   # Commercial approval gaps
    11,  # FY26 renewal watchlist
    15,  # Owner pipeline coverage
    16,  # Owner coaching focus
    18,  # QTD wins and losses
    21,  # Concentration risk
    22,  # Q2-Q3 risk triage
    24,  # May territory plan
    26,  # May action register
    27,  # May decision checklist
    28,  # Thank you
]
FORBIDDEN_TEXT = [
    "SC Test",
    "Test Account",
    "Test MASB",
    "5000%",
    "9000%",
    "#NAME",
    "#NULL",
    "Click to add subtitle",
    "Title of the section",
]


def _slug(value: str) -> str:
    return value.replace(" ", "-")


def _selected(director_slug: str | None) -> list[dict[str, Any]]:
    directors = canonical_directors()
    if not director_slug:
        return directors
    return [director for director in directors if _slug(str(director["name"])) == director_slug]


def _merge_manifest_results(manifest_path: Path, results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not manifest_path.exists():
        return results
    try:
        existing = json.loads(manifest_path.read_text(encoding="utf-8")).get("results") or []
    except (json.JSONDecodeError, OSError):
        existing = []
    merged = {str(result.get("slug")): result for result in existing if result.get("slug")}
    for result in results:
        merged[str(result.get("slug"))] = result
    order = {_slug(str(director["name"])): index for index, director in enumerate(canonical_directors())}
    return sorted(merged.values(), key=lambda result: order.get(str(result.get("slug")), 999))


def _source_deck(period: str, slug: str) -> Path:
    return ROOT / "state" / period / slug / f"{slug}-LAND-{period}-table-image-linked.pptx"


def _out_dir(period: str, slug: str) -> Path:
    return ROOT / "state" / period / slug / "factory" / "meeting-spine"


def _out_deck(period: str, slug: str) -> Path:
    return _out_dir(period, slug) / f"{slug}-LAND-{period}-meeting-spine.pptx"


def _slide_title(slide: object) -> str:
    texts: list[str] = []
    for shape in slide.shapes:  # type: ignore[attr-defined]
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            texts.append(" ".join(shape.text.split()))
    return texts[0] if texts else ""


def _deck_text(prs: Presentation) -> str:
    return "\n".join(
        " ".join(shape.text.split())
        for slide in prs.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text
    )


def _delete_slides_except(path: Path, keep_slides: list[int], *, required_text: list[str]) -> dict[str, Any]:
    prs = Presentation(path)
    original_count = len(prs.slides)
    keep_indexes = {slide_no - 1 for slide_no in keep_slides}
    kept_titles = {
        str(slide_no): _slide_title(prs.slides[slide_no - 1])
        for slide_no in keep_slides
        if 1 <= slide_no <= original_count
    }
    slide_id_list = prs.slides._sldIdLst  # noqa: SLF001
    for index in range(original_count - 1, -1, -1):
        if index in keep_indexes:
            continue
        rel_id = slide_id_list[index].rId
        prs.part.drop_rel(rel_id)
        del slide_id_list[index]
    prs.save(path)

    checked = Presentation(path)
    text = _deck_text(checked)
    forbidden_found = [needle for needle in FORBIDDEN_TEXT if needle in text]
    required_missing = [needle for needle in required_text if needle not in text]
    return {
        "original_slide_count": original_count,
        "output_slide_count": len(checked.slides),
        "kept_source_slides": keep_slides,
        "kept_titles": kept_titles,
        "forbidden_found": forbidden_found,
        "required_missing": required_missing,
        "status": "pass" if len(checked.slides) == len(keep_slides) and not forbidden_found and not required_missing else "fail",
    }


def build_one(period: str, director: dict[str, Any], *, required_text: list[str]) -> dict[str, Any]:
    slug = _slug(str(director["name"]))
    source = _source_deck(period, slug)
    output_dir = _out_dir(period, slug)
    output = _out_deck(period, slug)
    if not source.exists():
        raise SystemExit(f"missing linked source deck: {source}")
    output_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)
    result = _delete_slides_except(output, KEEP_SLIDES, required_text=required_text)
    linked_action_result = enhance_meeting_spine_deck(period, slug, output)
    aspect_result = fix_deck(output)
    stale_thinkcell_scrub = scrub_deck(output)
    if stale_thinkcell_scrub.status != "pass":
        result["status"] = "fail"
    result.update(
        {
            "director": str(director["name"]),
            "slug": slug,
            "territory": str(director["scope_label"]),
            "source_deck": str(source),
            "output_deck": str(output),
            "linked_action_layer": linked_action_result,
            "table_image_aspect_fixes": aspect_result["fix_count"],
            "stale_thinkcell_metadata_scrub": asdict(stale_thinkcell_scrub),
        }
    )
    (output_dir / "meeting_spine_manifest.json").write_text(json.dumps(result, indent=2, default=str) + "\n", encoding="utf-8")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--director-slug")
    parser.add_argument("--jobs", type=int, default=1, help="Parallel director builds for local file-only work.")
    args = parser.parse_args()
    try:
        period_context = context_for_period(args.period)
    except ValueError as exc:
        print(f"error: {exc}")
        return 2

    directors = _selected(args.director_slug)
    if not directors:
        raise SystemExit(f"unknown director slug: {args.director_slug}")

    if args.jobs > 1 and len(directors) > 1:
        with ThreadPoolExecutor(max_workers=min(args.jobs, len(directors))) as executor:
            results = list(
                executor.map(
                    lambda director: build_one(
                        args.period,
                        director,
                        required_text=list(period_context.required_deck_text),
                    ),
                    directors,
                )
            )
    else:
        results = [
            build_one(
                args.period,
                director,
                required_text=list(period_context.required_deck_text),
            )
            for director in directors
        ]
    output_dir = ROOT / "state" / args.period / "__regional__" / "meeting_spine"
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output_dir / "meeting_spine_manifest.json"
    manifest_results = _merge_manifest_results(manifest_path, results) if args.director_slug else results
    manifest = {
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "period": args.period,
        "status": "pass" if all(result["status"] == "pass" for result in manifest_results) else "fail",
        "results": manifest_results,
    }
    manifest_path.write_text(json.dumps(manifest, indent=2, default=str) + "\n", encoding="utf-8")
    for result in results:
        print(f"{str(result['status']).upper():>7} {result['slug']:<22} {result['output_deck']}")
    print(f"manifest={manifest_path}")
    return 0 if manifest["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
