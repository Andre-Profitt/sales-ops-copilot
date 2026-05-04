"""Build the Template Atlas corpus from sources.yml.

For each source:
  - pptx       : extract per-slide text + shape names + layout name
  - markdown   : split on H2 headings; one chunk per section
  - markdown_glob : same, expanded
  - yaml_registry : one chunk per slide entry
  - yaml_slotmap  : one chunk per slide entry

Each chunk gets:
  id, source, role, type, slide (or section), text, weight, embedding

Embeddings via Azure OpenAI text-embedding-3-small. Set
APRO_OPENAI_ENDPOINT and run `az login` first.

Idempotent: rebuilds corpus.jsonl from scratch each run.
"""

from __future__ import annotations

import argparse
import glob
import hashlib
import json
import os
import re
import sys
from pathlib import Path
from typing import Iterable

import yaml
from pptx import Presentation

REPO = Path(__file__).resolve().parent.parent.parent


def _chunk_id(*parts: str) -> str:
    return hashlib.sha1("|".join(parts).encode("utf-8")).hexdigest()[:16]


def _pptx_slide_chunks(src: dict) -> Iterable[dict]:
    path = REPO / src["path"] if not Path(src["path"]).is_absolute() else Path(src["path"])
    if not path.exists():
        if src.get("optional"):
            return
        raise FileNotFoundError(path)
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, 1):
        chunks: list[str] = []
        names: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "name", None):
                names.append(shape.name)
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()  # type: ignore[attr-defined]
                if t:
                    chunks.append(t)
        text = " ".join(chunks)
        if not text and not names:
            continue
        layout_name = getattr(slide.slide_layout, "name", "")
        body = f"layout={layout_name} | shapes=[{','.join(names)}] | text={text}"
        yield {
            "id": _chunk_id(src["id"], "slide", str(i)),
            "source": src["id"],
            "role": src["role"],
            "type": "pptx_slide",
            "slide": i,
            "text": body[:4000],
            "weight": src.get("weight", 1.0),
        }


def _md_chunks(src: dict, path: Path) -> Iterable[dict]:
    if not path.exists():
        if src.get("optional"):
            return
        raise FileNotFoundError(path)
    raw = path.read_text(errors="ignore")
    sections = re.split(r"\n## ", raw)
    for j, sec in enumerate(sections):
        body = sec.strip()
        if not body:
            continue
        first_line = body.splitlines()[0][:200]
        yield {
            "id": _chunk_id(src["id"], str(path), str(j)),
            "source": src["id"],
            "role": src["role"],
            "type": "md_section",
            "section_idx": j,
            "section_heading": first_line,
            "path": str(path),
            "text": body[:4000],
            "weight": src.get("weight", 1.0),
        }


def _yaml_chunks(src: dict) -> Iterable[dict]:
    path = REPO / src["path"]
    if not path.exists():
        if src.get("optional"):
            return
        raise FileNotFoundError(path)
    doc = yaml.safe_load(path.read_text())
    for slide in doc.get("slides", []):
        sid = slide.get("slide_id", "")
        text = json.dumps(slide, ensure_ascii=False)
        yield {
            "id": _chunk_id(src["id"], "slide", sid),
            "source": src["id"],
            "role": src["role"],
            "type": src["type"],
            "slide_id": sid,
            "text": text[:4000],
            "weight": src.get("weight", 1.0),
        }


def _embed_batch(texts: list[str]) -> list[list[float]]:
    from azure.identity import AzureCliCredential
    from openai import AzureOpenAI

    cred = AzureCliCredential()
    token = cred.get_token("https://cognitiveservices.azure.com/.default")
    client = AzureOpenAI(
        azure_endpoint=os.environ.get(
            "APRO_OPENAI_ENDPOINT", "https://apro-openai.openai.azure.com/"
        ),
        azure_ad_token=token.token,
        api_version="2024-08-01-preview",
    )
    resp = client.embeddings.create(
        input=texts,
        model=os.environ.get("APRO_EMBED_DEPLOYMENT", "text-embedding-3-small"),
    )
    return [list(d.embedding) for d in resp.data]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sources", default="config/atlas/sources.yml", type=Path)
    parser.add_argument("--out", default="state/atlas/corpus.jsonl", type=Path)
    parser.add_argument("--manifest", default="state/atlas/manifest.json", type=Path)
    parser.add_argument(
        "--skip-embeddings",
        action="store_true",
        help="Skip Azure OpenAI calls; emit chunks without 'embedding'.",
    )
    parser.add_argument("--batch", type=int, default=64)
    args = parser.parse_args(argv)

    sources = yaml.safe_load(args.sources.read_text())["sources"]
    chunks: list[dict] = []
    for src in sources:
        try:
            t = src["type"]
            if t == "pptx":
                chunks.extend(_pptx_slide_chunks(src))
            elif t == "markdown":
                chunks.extend(_md_chunks(src, REPO / src["path"]))
            elif t == "markdown_glob":
                for p in glob.glob(str(REPO / src["path"]), recursive=True):
                    chunks.extend(_md_chunks(src, Path(p)))
            elif t in ("yaml_registry", "yaml_slotmap"):
                chunks.extend(_yaml_chunks(src))
            else:
                sys.stderr.write(f"skipping unknown source type {t}\n")
        except FileNotFoundError as exc:
            sys.stderr.write(f"missing required source: {exc}\n")
            if not src.get("optional"):
                return 1

    if not args.skip_embeddings:
        for i in range(0, len(chunks), args.batch):
            batch = chunks[i : i + args.batch]
            vecs = _embed_batch([c["text"] for c in batch])
            for c, v in zip(batch, vecs):
                c["embedding"] = v
            sys.stdout.write(f"  embedded {i + len(batch)}/{len(chunks)}\n")
            sys.stdout.flush()

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w") as f:
        for c in chunks:
            f.write(json.dumps(c) + "\n")

    manifest = {
        "sources": sources,
        "chunk_count": len(chunks),
        "by_source": {s["id"]: sum(1 for c in chunks if c["source"] == s["id"]) for s in sources},
        "by_role": {},
        "embedded": not args.skip_embeddings,
    }
    for c in chunks:
        manifest["by_role"][c["role"]] = manifest["by_role"].get(c["role"], 0) + 1
    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    args.manifest.write_text(json.dumps(manifest, indent=2))
    print(f"OK: {len(chunks)} chunks -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
