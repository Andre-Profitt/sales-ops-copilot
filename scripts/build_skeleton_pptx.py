"""Build the clean 28-slide LAND review skeleton from canonical layouts.

Output is a brand-correct, debris-free PPTX with no Think-Cell objects.
The Windows VM steward inserts Think-Cell elements manually and saves
as LAND_review_full_28.tcseed.pptx.
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

from pptx import Presentation

REPO = Path(__file__).resolve().parent.parent
CANONICAL = REPO / "assets" / "golden" / "LAND_canonical.pptx"

# (slide_id, layout_name_substring) — 28 entries matching the registry.
# Substrings are matched case-insensitively against canonical layout names.
# Live SimCorp LAND layouts (per HANDOFF): 0/2/6/10/31.
#   0  "Title 1"                          → cover / closing-of-cover
#   2  "Divider 1"                        → section dividers
#   6  "Title and Content"                → analytic content (chart, table)
#   10 "2 x content w/ gradient line"     → two-column / KPI strip
#   31 "End slide with disclaimer 1"      → closing
SLIDE_PLAN = [
    ("S01", "Title 1"),  # cover
    ("S02", "2 x content w/ gradient line"),  # exec summary
    ("S03", "Divider 1"),  # pipeline divider
    ("S04", "Title and Content"),  # one chart
    ("S05", "Title and Content"),  # one chart
    ("S06", "Title and Content"),  # one chart
    ("S07", "Title and Content"),  # table
    ("S08", "Title and Content"),  # table
    ("S09", "Title and Content"),  # table
    ("S10", "Divider 1"),  # retention divider
    ("S11", "Title and Content"),  # table
    ("S12", "Title and Content"),  # table
    ("S13", "Title and Content"),  # one chart
    ("S14", "Divider 1"),  # territory divider
    ("S15", "Title and Content"),  # one chart
    ("S16", "Title and Content"),  # one chart
    ("S17", "Title and Content"),  # one chart
    ("S18", "Title and Content"),  # one chart
    ("S19", "Title and Content"),  # one chart
    ("S20", "Divider 1"),  # risks divider
    ("S21", "Title and Content"),  # one chart
    ("S22", "Title and Content"),  # one chart
    ("S23", "2 x content w/ gradient line"),  # KPI cards
    ("S24", "Title and Content"),  # table
    ("S25", "Title and Content"),  # one chart
    ("S26", "Title and Content"),  # table
    ("S27", "2 x content w/ gradient line"),  # narrative
    ("S28", "End slide with disclaimer 1"),  # closing
]


def _layout_by_substring(prs, sub: str):
    for layout in prs.slide_layouts:
        if sub.lower() in layout.name.lower():
            return layout
    # fallback to layout 0 if no match
    return prs.slide_layouts[0]


def _strip_placeholder_text(slide) -> None:
    """Clear any default placeholder text that bleeds through from layouts.

    Empty layout placeholders sometimes carry prompt text like
    "Click to add title" — fine for live editing but noise in a skeleton.
    """
    for shape in slide.placeholders:
        if shape.has_text_frame:
            tf = shape.text_frame
            # leave the structure, just blank the text
            tf.text = ""


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--canonical", default=CANONICAL, type=Path)
    args = parser.parse_args()

    if not args.canonical.exists():
        print(f"canonical not found: {args.canonical}", file=sys.stderr)
        return 2

    # Start from canonical to preserve master/theme/brand
    args.out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(args.canonical, args.out)

    prs = Presentation(str(args.out))

    # Drop existing slides — keep master/layouts only.
    # We must drop both the <p:sldIdLst> entries AND the rels from the
    # presentation part to the slide parts. python-pptx serializes only
    # parts reachable from rels, so dropping the rels also drops the slide
    # parts from the output zip — preventing duplicate slide1.xml entries
    # when add_slide re-allocates names.
    R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    xml_slides = prs.slides._sldIdLst  # internal but stable
    sld_ids = list(xml_slides)
    pres_part = prs.part
    for sld_id in sld_ids:
        rId = sld_id.attrib[R_NS + "id"]
        pres_part.drop_rel(rId)
        xml_slides.remove(sld_id)

    # Add 28 slides using planned layouts
    for _slide_id, sub in SLIDE_PLAN:
        layout = _layout_by_substring(prs, sub)
        slide = prs.slides.add_slide(layout)
        _strip_placeholder_text(slide)

    prs.save(str(args.out))
    print(f"OK: wrote 28-slide skeleton to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
