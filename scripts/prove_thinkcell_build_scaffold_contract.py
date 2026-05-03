#!/usr/bin/env python3
"""Write an L5 proof artifact for a think-cell build scaffold contract."""

from __future__ import annotations

import argparse
import datetime as dt
import json
import subprocess
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image, ImageStat
from pptx import Presentation

from ppttc_template import template_named_elements


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_PERIOD = "2026-Q2"
DEFAULT_CONTRACT = "QTR10_ActionDecisionRegister_TableImage"
QTR11_CONTRACT = "QTR11_CommercialApprovalGap_TableImage"
EMU_PER_INCH = 914400


@dataclass(frozen=True)
class TargetProofSpec:
    name: str
    slide: int
    workbook_name: str
    required_terms: tuple[str, ...]
    min_width_in: float
    min_height_in: float


@dataclass
class Check:
    name: str
    status: str
    details: dict[str, Any] = field(default_factory=dict)


CONTRACT_TARGETS: dict[str, tuple[TargetProofSpec, ...]] = {
    DEFAULT_CONTRACT: (
        TargetProofSpec(
            name="S26_ActionItems",
            slide=26,
            workbook_name="S26_ActionItems",
            required_terms=("Zombie ARR", "Approval Gap", "Jesper Tyrer"),
            min_width_in=5.0,
            min_height_in=1.0,
        ),
        TargetProofSpec(
            name="S27_DecisionChecklist",
            slide=27,
            workbook_name="S27_DecisionChecklist",
            required_terms=("Forecast call", "Renewal ACV", "Land+Expand ARR"),
            min_width_in=5.0,
            min_height_in=2.0,
        ),
    ),
    QTR11_CONTRACT: (
        TargetProofSpec(
            name="S09_PendingCommercialApproval",
            slide=9,
            workbook_name="S09_PendingCommercialApproval",
            required_terms=("PT Bank Mandiri", "3 - Engagement", "Land", "0.5 mEUR"),
            min_width_in=5.0,
            min_height_in=0.8,
        ),
    ),
}


def _slug(value: str) -> str:
    return value.replace(" ", "-")


def _director_name(slug: str) -> str:
    return slug.replace("-", " ")


def _linked_deck(period: str, director_slug: str) -> Path:
    return ROOT / "state" / period / director_slug / f"{director_slug}-LAND-{period}-table-image-linked.pptx"


def _workbook(period: str, director_slug: str) -> Path:
    return ROOT / "state" / period / director_slug / "factory" / "connected" / "connected_factory_table_images.xlsx"


def _scaffold_work_dir(period: str, contract: str) -> Path:
    return ROOT / "state" / "thinkcell_bridge" / "build_scaffold" / period / "work" / contract


def _run(command: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(command, cwd=ROOT, text=True, capture_output=True, check=False)


def _cell_text(values: list[list[Any]]) -> str:
    chunks: list[str] = []
    for row in values:
        for value in row:
            if value is not None:
                chunks.append(str(value))
    return "\n".join(chunks)


def _workbook_range_values(path: Path, name: str) -> tuple[str, str, list[list[Any]]]:
    wb = load_workbook(path, data_only=True, read_only=False)
    defined = wb.defined_names.get(name)
    if defined is None:
        raise KeyError(name)
    destinations = list(defined.destinations)
    if not destinations:
        raise ValueError(f"defined name has no destinations: {name}")
    sheet, ref = destinations[0]
    ws = wb[sheet]
    values = [[cell.value for cell in row] for row in ws[ref]]
    return sheet, ref, values


def _proof_workbook(path: Path, targets: tuple[TargetProofSpec, ...]) -> Check:
    details: dict[str, Any] = {"workbook": str(path), "targets": []}
    if not path.exists():
        return Check("workbook_named_ranges", "fail", details | {"error": "missing workbook"})
    failures: list[str] = []
    for target in targets:
        try:
            sheet, ref, values = _workbook_range_values(path, target.workbook_name)
        except Exception as exc:  # noqa: BLE001 - proof output should preserve exact failure.
            failures.append(target.workbook_name)
            details["targets"].append({"name": target.workbook_name, "status": "fail", "error": str(exc)})
            continue
        text = _cell_text(values)
        missing_terms = [term for term in target.required_terms if term not in text]
        if missing_terms:
            failures.append(target.workbook_name)
        details["targets"].append(
            {
                "name": target.workbook_name,
                "status": "pass" if not missing_terms else "fail",
                "sheet": sheet,
                "range": ref,
                "rows": len(values),
                "cols": len(values[0]) if values else 0,
                "required_terms": target.required_terms,
                "missing_terms": missing_terms,
                "preview_rows": values[:3],
            }
        )
    return Check("workbook_named_ranges", "pass" if not failures else "fail", details)


def _find_table_shapes(slide: Any) -> tuple[Any | None, Any | None]:
    pic = None
    data = None
    for shape in slide.shapes:
        if shape.name == "Pic":
            pic = shape
        if shape.name == "think-cell data - do not delete":
            data = shape
    return pic, data


def _shape_bounds(shape: Any) -> dict[str, float]:
    return {
        "left_in": round(int(shape.left) / EMU_PER_INCH, 3),
        "top_in": round(int(shape.top) / EMU_PER_INCH, 3),
        "width_in": round(int(shape.width) / EMU_PER_INCH, 3),
        "height_in": round(int(shape.height) / EMU_PER_INCH, 3),
    }


def _proof_deck(path: Path, targets: tuple[TargetProofSpec, ...]) -> Check:
    details: dict[str, Any] = {"deck": str(path), "targets": []}
    if not path.exists() or path.stat().st_size == 0:
        return Check("deck_linked_table_images", "fail", details | {"error": "missing or empty deck"})
    names = set(template_named_elements(path))
    prs = Presentation(path)
    failures: list[str] = []
    for target in targets:
        slide = prs.slides[target.slide - 1]
        pic, data = _find_table_shapes(slide)
        target_details = {
            "target": target.name,
            "slide": target.slide,
            "has_named_thinkcell_element": target.name in names,
            "pic_shape": _shape_bounds(pic) if pic is not None else None,
            "data_shape": _shape_bounds(data) if data is not None else None,
            "native_image_px": getattr(getattr(pic, "image", None), "size", None) if pic is not None else None,
        }
        ok = bool(target_details["has_named_thinkcell_element"] and pic is not None and data is not None)
        if pic is not None:
            width_in = int(pic.width) / EMU_PER_INCH
            height_in = int(pic.height) / EMU_PER_INCH
            ok = ok and width_in >= target.min_width_in and height_in >= target.min_height_in
            target_details["min_width_in"] = target.min_width_in
            target_details["min_height_in"] = target.min_height_in
        if not ok:
            failures.append(target.name)
        target_details["status"] = "pass" if ok else "fail"
        details["targets"].append(target_details)
    return Check("deck_linked_table_images", "pass" if not failures else "fail", details)


def _render_deck(deck: Path, render_dir: Path) -> Check:
    command = [
        sys.executable,
        "scripts/render_deck_for_review.py",
        str(deck),
        "--output-dir",
        str(render_dir),
    ]
    result = _run(command)
    return Check(
        "render_deck",
        "pass" if result.returncode == 0 else "fail",
        {
            "command": command,
            "returncode": result.returncode,
            "stdout_tail": result.stdout[-4000:],
            "stderr_tail": result.stderr[-4000:],
            "png_dir": str(render_dir / "png"),
            "pdf_dir": str(render_dir / "pdf"),
        },
    )


def _proof_render(render_dir: Path, targets: tuple[TargetProofSpec, ...]) -> Check:
    png_dir = render_dir / "png"
    details: dict[str, Any] = {"png_dir": str(png_dir), "targets": []}
    failures: list[str] = []
    for target in targets:
        path = png_dir / f"slide-{target.slide:02d}.png"
        if not path.exists():
            failures.append(target.name)
            details["targets"].append({"target": target.name, "slide": target.slide, "status": "fail", "error": "missing render"})
            continue
        with Image.open(path) as image:
            rgb = image.convert("RGB")
            width, height = rgb.size
            crop = rgb.crop((int(width * 0.04), int(height * 0.20), int(width * 0.96), int(height * 0.88)))
            gray = crop.convert("L")
            stat = ImageStat.Stat(gray)
            pixel_data = getattr(gray, "get_flattened_data", gray.getdata)()
            dark_pixels = sum(1 for pixel in pixel_data if pixel < 190)
        ok = stat.stddev[0] >= 8.0 and dark_pixels >= 5000
        if not ok:
            failures.append(target.name)
        details["targets"].append(
            {
                "target": target.name,
                "slide": target.slide,
                "status": "pass" if ok else "fail",
                "image_size": [width, height],
                "content_crop_stddev": round(stat.stddev[0], 3),
                "content_crop_dark_pixels": dark_pixels,
                "min_stddev": 8.0,
                "min_dark_pixels": 5000,
            }
        )
    return Check("rendered_table_visibility", "pass" if not failures else "fail", details)


def _write_markdown(path: Path, payload: dict[str, Any]) -> None:
    lines = [
        f"# {payload['contract']} L5 Proof",
        "",
        f"- Status: `{payload['status']}`",
        f"- Period: `{payload['period']}`",
        f"- Director: `{payload['director_slug']}`",
        f"- Deck: `{payload['deck']}`",
        f"- Workbook: `{payload['workbook']}`",
        f"- Render dir: `{payload['render_dir']}`",
        "",
        "## Checks",
        "",
        "| Check | Status |",
        "|---|---|",
    ]
    for check in payload["checks"]:
        lines.append(f"| `{check['name']}` | `{check['status']}` |")
    lines.extend(["", "## Targets", "", "| Target | Slide | Status |", "|---|---:|---|"])
    target_status: dict[tuple[str, int], str] = {}
    for check in payload["checks"]:
        for target in check.get("details", {}).get("targets", []):
            key = (str(target.get("target") or target.get("name")), int(target.get("slide") or 0))
            if target.get("status") == "fail":
                target_status[key] = "fail"
            else:
                target_status.setdefault(key, "pass")
    for target in payload["targets"]:
        key = (target["name"], target["slide"])
        lines.append(f"| `{target['name']}` | {target['slide']} | `{target_status.get(key, 'unknown')}` |")
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--director-slug", default="Jesper-Tyrer")
    parser.add_argument("--contract", default=DEFAULT_CONTRACT)
    parser.add_argument("--deck", type=Path)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--render", action="store_true", help="Render the deck before checking visible table output.")
    args = parser.parse_args()

    targets = CONTRACT_TARGETS.get(args.contract)
    if not targets:
        raise SystemExit(f"unsupported contract: {args.contract}")

    deck = (args.deck or _linked_deck(args.period, args.director_slug)).expanduser().resolve()
    workbook = (args.workbook or _workbook(args.period, args.director_slug)).expanduser().resolve()
    work_dir = _scaffold_work_dir(args.period, args.contract)
    render_dir = work_dir / "rendered"
    work_dir.mkdir(parents=True, exist_ok=True)

    checks = [
        _proof_workbook(workbook, targets),
        _proof_deck(deck, targets),
    ]
    if args.render:
        checks.append(_render_deck(deck, render_dir))
    checks.append(_proof_render(render_dir, targets))

    status = "pass" if all(check.status == "pass" for check in checks) else "fail"
    payload = {
        "schema": "thinkcell-build-scaffold-proof/v1",
        "created_at_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "status": status,
        "contract": args.contract,
        "period": args.period,
        "director_slug": args.director_slug,
        "director_name": _director_name(args.director_slug),
        "deck": str(deck),
        "workbook": str(workbook),
        "render_dir": str(render_dir),
        "targets": [asdict(target) for target in targets],
        "checks": [asdict(check) for check in checks],
    }
    proof_json = work_dir / f"{args.contract}-proof.json"
    proof_md = work_dir / f"{args.contract}-proof.md"
    proof_json.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    _write_markdown(proof_md, payload)
    print(f"status={status}")
    print(f"proof_json={proof_json}")
    print(f"proof_md={proof_md}")
    return 0 if status == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
