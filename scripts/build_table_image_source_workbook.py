#!/usr/bin/env python3
"""Create Excel source ranges for think-cell table-image refresh.

This keeps the table-image lane tied to the connected factory workbook while
avoiding destructive edits to the main workbook. The output is a workbook copy
with `TC_*` sheets and defined names matching the PowerPoint donor names.
"""

from __future__ import annotations

import argparse
import re
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName

from sales_director_row_filters import is_internal_sales_record


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_PERIOD = "2026-Q2"
DEFAULT_DIRECTOR = "Jesper-Tyrer"

HEADER_FILL = PatternFill("solid", fgColor="083EA7")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BODY_FILL = PatternFill("solid", fgColor="FFFFFF")
THIN = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass(frozen=True)
class TableImageSpec:
    name: str
    sheet: str
    rows: list[list[Any]]
    widths: tuple[float, ...] = ()
    render_scale: float = 1.0


def _ref_rows(source_sheet: str, rows: int, cols: int) -> list[list[str]]:
    return [
        [f"='{source_sheet}'!{get_column_letter(col)}{row}" for col in range(1, cols + 1)]
        for row in range(1, rows + 1)
    ]


def _ref_project_rows(source_sheet: str, rows: int, cols: tuple[str, ...]) -> list[list[str]]:
    return [[f"='{source_sheet}'!{col}{row}" for col in cols] for row in range(1, rows + 1)]


def _lookup_text_formula(key: str) -> str:
    return f'=IFERROR(VLOOKUP("{key}",tbl_Raw_Original_Intel,2,FALSE),"")'


def _lookup_meur_formula(key: str, suffix: str = "") -> str:
    return f'=IFERROR("EUR "&TEXT(VLOOKUP("{key}",tbl_Raw_Original_Intel,2,FALSE)/1000000,"0.0")&"M{suffix}","")'


def _signal_count_formula(match_text: str, *, min_pushes: int | None = None) -> str:
    if min_pushes is not None:
        return f'=COUNTIF(Raw_Current_Q2_Readiness!K2:K13,">={min_pushes}")'
    return f'=COUNTIF(Raw_Current_Q2_Readiness!M2:M13,"*{match_text}*")'


def _signal_amount_formula(match_text: str, *, min_pushes: int | None = None) -> str:
    amount = (
        'IF(ISNUMBER(Raw_Current_Q2_Readiness!H2:H13),Raw_Current_Q2_Readiness!H2:H13,'
        'IFERROR(NUMBERVALUE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(Raw_Current_Q2_Readiness!H2:H13,"EUR ","")," mEUR",""),"M",""))*1000000,0))'
    )
    if min_pushes is not None:
        condition = f"--(Raw_Current_Q2_Readiness!K2:K13>={min_pushes})"
    else:
        condition = f'--ISNUMBER(SEARCH("{match_text}",Raw_Current_Q2_Readiness!M2:M13))'
    return f'=SUMPRODUCT({condition},{amount})'


def _signal_deals_formula(match_text: str, *, min_pushes: int | None = None) -> str:
    if min_pushes is not None:
        condition = f"Raw_Current_Q2_Readiness!K2:K13>={min_pushes}"
    else:
        condition = f'ISNUMBER(SEARCH("{match_text}",Raw_Current_Q2_Readiness!M2:M13))'
    return (
        '=IFERROR(TEXTJOIN(", ",TRUE,TAKE(FILTER('
        f"Raw_Current_Q2_Readiness!C2:C13,{condition}"
        "),4)),\"\")"
    )


def _number_from_display(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "")
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", text)
    if not match:
        return 0.0
    number = float(match.group(0).replace(",", ""))
    if "meur" in text.lower() or " m" in text.lower():
        return number * 1_000_000
    return number


def _signal_rows(wb: Any) -> list[list[Any]]:
    raw_rows = list(wb["Raw_Current_Q2_Readiness"].iter_rows(min_row=2, values_only=True))

    def summarize(label: str, predicate: Any) -> list[Any]:
        matched = [row for row in raw_rows if row and predicate(row)]
        amount = sum(_number_from_display(row[7]) for row in matched)
        deals = ", ".join(str(row[2] or "") for row in matched[:4])
        return [label, len(matched), _fmt_meur(amount), deals]

    return [
        ["Signal", "# deals", "Unweighted ARR", "Named deals"],
        summarize("Silent 60d+", lambda row: "Silent 60d+" in str(row[12] or "")),
        summarize("No next step", lambda row: "No next step" in str(row[12] or "")),
        summarize("4+ pushes", lambda row: _number_from_display(row[10]) >= 4),
        summarize("Approval gap", lambda row: "Commercial approval gap" in str(row[12] or "")),
    ]


def _watchlist_rows(limit: int = 5) -> list[list[str]]:
    rows = [["Deal", "Owner", "Unweighted ARR", "Fcst", "Last activity", "Signal", "Next step"]]
    for row in range(2, 2 + limit):
        rows.append(
            [
                f"=Raw_Current_Q2_Readiness!C{row}",
                f"=Raw_Current_Q2_Readiness!D{row}",
                f"=Raw_Current_Q2_Readiness!H{row}",
                f"=Raw_Current_Q2_Readiness!I{row}",
                f"=Raw_Current_Q2_Readiness!L{row}",
                f"=Raw_Current_Q2_Readiness!M{row}",
                f"=Raw_Current_Q2_Readiness!N{row}",
            ]
        )
    return rows


def _hygiene_rows(wb: Any, limit: int = 4) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for row in _signal_rows(wb):
        rows.append([*row, "", "", ""])
    rows.extend(_readiness_watchlist_rows(wb, limit))
    return rows


def _operating_rhythm_rows() -> list[list[str]]:
    return [
        ["Cadence", "Gate", "Population", "Required evidence"],
        ["1", "Deal evidence", "Every Commit/Pipeline deal", "Customer-owned next step, dated decision, and sponsor proof."],
        ["2", "Forecast hygiene", "Commit + Pipeline", "Downgrade stale/no-next-step deals unless proof is supplied."],
        ["3", "Governance", "Stage 3+ approval gaps", "Submit Commercial Approval before the next forecast call."],
        ["4", "Loss discipline", "Closed-lost QTD", "Attach reason codes and preserve territory lessons."],
        ["5", "Push discipline", "Current workbook push flags", "Review pushed deals from readiness rows."],
        ["6", "Renewal ACV", "FY26 renewal watchlist", "Validate ACV basis before quoting outside this review."],
    ]


def _next_14_days_rows() -> list[list[str]]:
    return [
        ["Date", "Forum", "Focus", "Output"],
        ["May 1", "Director kickoff", "Confirm Q2 Commit/Pipeline roster and owner asks.", "Top deals have named owners and required proof."],
        ["May 4-6", "Deal inspection", "Highest-value current readiness rows.", "Customer next step, sponsor, close proof, approval status, and close-date realism."],
        ["May 7-8", "Governance gate", "Supported current-quarter approval gaps.", "Submissions logged; stale/no-next-step rows refreshed."],
        ["May 11-13", "Forecast reset", "Deals missing customer proof or realistic timing.", "Downgrade, re-date, or confirm with evidence."],
        ["May 14-15", "May checkpoint", "Q2 close view, creation plan, Renewal ACV basis.", "One evidence-backed May forecast and next pipeline-creation list."],
        ["Every week", "Owner hygiene", "All Q2 readiness rows.", "LastActivityDate and NextStep updated before forecast call."],
    ]


def _decision_checklist_rows() -> list[list[str]]:
    return [
        ["#", "Decision", "Question", "Required closure"],
        ["1", "Forecast call", "Is there named customer evidence behind current Commit and Pipeline?", "Deal owners refresh next steps and close proof."],
        ["2", "Governance", "Are supported Q2 approval gaps submitted?", "Track current gaps; do not import another territory's approval context."],
        ["3", "Q1 learning", "Are territory-specific losses, reason codes, and slipped deals included?", "Use the original territory pack where available."],
        ["4", "Push discipline", "Are current workbook push patterns owned as coaching actions?", "Review current readiness rows with push flags."],
        ["5", "Renewals", "Is the current FY26 Renewal ACV basis validated?", "Keep Renewal ACV separate from Land+Expand ARR."],
    ]


def _fmt_meur(value: float) -> str:
    return f"EUR {value / 1_000_000:.1f}M"


def _original_value(wb: Any, key: str, default: Any = "") -> Any:
    ws = wb["Raw_Original_Intel"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == key:
            return row[1]
    return default


def _original_label(wb: Any, noun: str = "ETL") -> str:
    territory = str(_original_value(wb, "territory", "director")).strip()
    if territory:
        return f"Original {territory} {noun}"
    return f"Original director {noun}"


def _date_value(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return None
    return None


def _current_q2_closeable_arr(wb: Any) -> float:
    total = 0.0
    start = datetime(2026, 4, 1)
    end = datetime(2026, 7, 1)
    for row in wb["Raw_Current_Model_Data"].iter_rows(min_row=2, values_only=True):
        close_date = _date_value(row[4])
        if row[1] in {"Land", "Expand"} and close_date and start <= close_date < end:
            total += float(row[10] or 0)
    return total


def _new_pipeline_since_q2_start(wb: Any) -> tuple[int, float]:
    count = 0
    total = 0.0
    start = datetime(2026, 4, 1)
    for row in wb["Raw_Current_Model_Data"].iter_rows(min_row=2, values_only=True):
        created = _date_value(row[3])
        if row[1] in {"Land", "Expand"} and created and created >= start:
            count += 1
            total += float(row[10] or 0)
    return count, total


def _raw_renewal_totals(wb: Any) -> tuple[int, float]:
    count = 0
    total = 0.0
    for row in wb["Raw_Renewals"].iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        count += 1
        total += float(row[5] or 0)
    return count, total


def _closed_renewal_totals(wb: Any) -> tuple[int, float]:
    count = 0
    total = 0.0
    for row in wb["Raw_Current_Closed_CFQ"].iter_rows(min_row=2, values_only=True):
        if row[2] != "Renewal":
            continue
        count += 1
        total += float(row[10] or 0)
    return count, total


def _forecast_row_count(wb: Any) -> int:
    return max(0, wb["Raw_Forecast_Items"].max_row - 1)


def _clean_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith("#") and value not in {
        "#",
        "# Opps",
        "# Open opps",
        "# deals",
    }:
        return ""
    return value


def _compact_currency_text(value: Any) -> Any:
    if not isinstance(value, str):
        return value

    def repl(match: re.Match[str]) -> str:
        amount = float(match.group(1).replace(",", ""))
        return f"EUR {amount / 1_000_000:.1f}M"

    return re.sub(r"EUR\s*([0-9][0-9,]{5,})", repl, value)


def _compact_action_text(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = _compact_currency_text(value)
    text = text.replace("Commercial Approval", "approval")
    text = text.replace("Land/Expand", "L+E")
    text = text.replace("Land+Expand", "L+E")
    text = text.replace("Task/Event", "task/event")
    text = text.replace("ALL Land deals", "all Land deals")
    text = text.replace("approval is mandatory", "Approval is mandatory")
    return text


def _leading_count(value: Any) -> str:
    match = re.search(r"\b(\d+)\b", str(value or ""))
    return match.group(1) if match else ""


def _first_meur(value: Any) -> str:
    text = _compact_currency_text(str(value or ""))
    match = re.search(r"EUR\s+([0-9]+(?:\.[0-9])?M)", text)
    return f"EUR {match.group(1)}" if match else ""


def _action_register_display(rule: Any, claim: Any, action: Any) -> tuple[str, str]:
    rule_key = str(rule or "").strip().lower()
    count = _leading_count(claim)
    amount = _first_meur(claim)
    amount_suffix = f"{amount.replace('EUR ', '')} EUR" if amount else ""
    if rule_key == "zombie_arr":
        return (
            f"{count} stale opps | {amount_suffix}".strip(" |"),
            "Inspect with reps; close or disqualify by EOM.",
        )
    if rule_key == "approval_gap":
        return (
            f"{count} approval gaps | {amount_suffix}".strip(" |"),
            "Submit gaps before EOM; all Land deals need approval.",
        )
    if rule_key == "coverage_gap":
        return (
            f"{count} Tier-1 accounts without open opps",
            "Assign opener coverage for starved Tier-1 accounts.",
        )
    if rule_key == "activity_drought":
        return (
            f"{count} inactive Q2 opps | {amount_suffix}".strip(" |"),
            "Reset 14-day task/event SLA on active Q2 opps.",
        )
    return str(_compact_action_text(claim)), str(_compact_action_text(action))


def _meaningful(value: Any) -> bool:
    value = _clean_cell(value)
    return value is not None and value != ""


def _as_float(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return _number_from_display(value)


def _raw_rows(wb: Any, sheet: str) -> list[dict[str, Any]]:
    ws = wb[sheet]
    headers = [_clean_cell(cell.value) for cell in ws[1]]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        record = {str(header): _clean_cell(value) for header, value in zip(headers, values)}
        if any(_meaningful(value) for value in record.values()) and not is_internal_sales_record(record):
            rows.append(record)
    return rows


def _date_in_q2(value: Any) -> bool:
    close_date = _date_value(value)
    return bool(close_date and datetime(2026, 4, 1) <= close_date < datetime(2026, 7, 1))


def _q1_accountability_rows(wb: Any) -> list[list[Any]]:
    return [
        ["Q1 accountability lens", "Original director baseline", "Current proof", "Director action"],
        [
            "Opened / promised",
            _fmt_meur(_as_float(_original_value(wb, "q1_promised_opened_arr_eur"))),
            "Original deck commentary",
            "Use as baseline, not current pipeline.",
        ],
        [
            "Won",
            f'{_original_value(wb, "q1_land_wins", 0)} / {_fmt_meur(_as_float(_original_value(wb, "q1_land_wins_arr_eur")))}',
            "Original sidecar",
            "Show delivered Land ARR separately from May QTD.",
        ],
        [
            "Lost",
            f'{_original_value(wb, "q1_land_lost", 0)} / {_fmt_meur(_as_float(_original_value(wb, "q1_land_lost_arr_eur")))}',
            "Original sidecar",
            "Keep loss accountability visible.",
        ],
        [
            "Slipped",
            f'{_original_value(wb, "q1_slipped_deals", 0)} / {_fmt_meur(_as_float(_original_value(wb, "q1_slipped_arr_eur")))}',
            "Original sidecar",
            "Use only where original pack has evidence.",
        ],
        [
            "Q2 original open",
            f'{_original_value(wb, "q2_original_deals", 0)} / {_fmt_meur(_as_float(_original_value(wb, "q2_original_arr_eur")))}',
            "Original sidecar",
            "Compare to refreshed May Q2 view.",
        ],
    ]


def _forecast_category_rows(wb: Any) -> list[list[Any]]:
    totals: dict[str, dict[str, float]] = {}
    for row in _raw_rows(wb, "Raw_Current_Model_Data"):
        if row.get("Type") not in {"Land", "Expand"} or not _date_in_q2(row.get("CloseDate")):
            continue
        category = str(row.get("ForecastCategoryName") or "Uncategorized")
        entry = totals.setdefault(category, {"count": 0, "arr": 0.0})
        entry["count"] += 1
        entry["arr"] += _as_float(row.get("ARR_EUR"))

    rows: list[list[Any]] = [["Category", "# Opps", "ARR (EUR)", "Readout"]]
    for category in ("Commit", "Pipeline", "Best Case", "Omitted"):
        entry = totals.pop(category, {"count": 0, "arr": 0.0})
        if category == "Omitted" and not entry["count"] and not entry["arr"]:
            continue
        readout = "Track separately; not headline closeable." if category == "Omitted" else "Land+Expand unweighted ARR only."
        rows.append([category, int(entry["count"]), _fmt_meur(entry["arr"]), readout])
    for category, entry in sorted(totals.items()):
        if entry["count"] or entry["arr"]:
            rows.append([category, int(entry["count"]), _fmt_meur(entry["arr"]), "Land+Expand unweighted ARR only."])
    return rows


def _readiness_rows(
    wb: Any,
    *,
    columns: tuple[str, ...],
    limit: int,
    flagged_only: bool = False,
) -> list[list[Any]]:
    rows = _raw_rows(wb, "Raw_Current_Q2_Readiness")
    if flagged_only:
        filtered = [row for row in rows if "No obvious hygiene flag" not in str(row.get("Readiness") or "")]
        rows = filtered or rows
    output = [[*columns]]
    for row in rows[:limit]:
        values = [_clean_cell(row.get(column)) for column in columns]
        if any(_meaningful(value) for value in values[1:]):
            output.append(values)
    if len(output) == 1:
        output.append(["No current rows found", *["" for _ in columns[1:]]])
    return output


def _readiness_watchlist_rows(wb: Any, limit: int = 4) -> list[list[Any]]:
    rows = _readiness_rows(
        wb,
        columns=("Opportunity", "Owner", "ARR", "Forecast", "Last Activity", "Readiness", "Next Step"),
        limit=limit,
        flagged_only=True,
    )
    rows[0] = ["Deal", "Owner", "Unweighted ARR", "Fcst", "Last activity", "Signal", "Next step"]
    return rows


def _pending_approval_rows(wb: Any) -> list[list[Any]]:
    columns = ("#", "Account", "Opportunity", "Owner", "Stage", "Close Date", "Type", "ARR (EUR)")
    rows = _raw_rows(wb, "Raw_Current_Pending_Approval")
    output = [[*columns]]
    for row in rows:
        values = [_clean_cell(row.get(column)) for column in columns]
        if any(_meaningful(value) for value in values[1:]):
            output.append(values)
    if len(output) == 1:
        output.append(["", "No current Q2 commercial approval gaps found", "", "", "", "", "", ""])
    return output


def _renewal_pipeline_rows(wb: Any, limit: int = 5) -> list[list[Any]]:
    columns = ("Close Date", "Account", "Opportunity", "Owner", "Stage", "ACV Unweighted (EUR)", "Probability %", "Comments")
    output: list[list[Any]] = [["#", *columns]]
    for index, row in enumerate(_raw_rows(wb, "Raw_Renewals")[:limit], 1):
        values = [index, *[_clean_cell(row.get(column)) for column in columns]]
        if any(_meaningful(value) for value in values[1:]):
            output.append(values)
    if len(output) == 1:
        output.append(["", "No FY26 renewal rows found", "", "", "", "", "", "", ""])
    return output


def _owner_coaching_rows(wb: Any, limit: int = 7) -> list[list[Any]]:
    owners: dict[str, dict[str, float]] = {}
    for row in _raw_rows(wb, "Raw_Current_Model_Data"):
        if row.get("Type") not in {"Land", "Expand"}:
            continue
        stage = str(row.get("StageName") or "")
        if stage.startswith("0") or "Won" in stage:
            continue
        owner = str(row.get("OwnerName") or "Unassigned")
        entry = owners.setdefault(owner, {"count": 0, "arr": 0.0})
        entry["count"] += 1
        entry["arr"] += _as_float(row.get("ARR_EUR"))

    rows: list[list[Any]] = [["Owner", "# Open opps", "Open unweighted ARR", "Action cue", "Original director risk"]]
    original_risk = (
        f'{_original_value(wb, "owner_push_count", 0)} pushes / '
        f'{_fmt_meur(_as_float(_original_value(wb, "owner_push_arr_eur")))}'
    )
    if str(_original_value(wb, "owner_push_count", 0)) in {"", "0", "0.0"}:
        original_risk = "No original owner-push summary"
    for index, (owner, entry) in enumerate(sorted(owners.items(), key=lambda item: item[1]["arr"], reverse=True)[:limit]):
        action = "Capacity/risk review" if entry["arr"] >= 3_000_000 else "Coach active Q2 actions"
        rows.append([owner, int(entry["count"]), _fmt_meur(entry["arr"]), action, original_risk if index == 0 else ""])
    if len(rows) == 1:
        rows.append(["No owner rows found", "", "", "", original_risk])
    return rows


def _closed_le_rows(wb: Any, is_won: bool) -> tuple[int, float]:
    count = 0
    arr = 0.0
    for row in _raw_rows(wb, "Raw_Current_Closed_CFQ"):
        if row.get("Type") not in {"Land", "Expand"}:
            continue
        if bool(row.get("IsWon")) is not is_won:
            continue
        count += 1
        arr += _as_float(row.get("ARR_EUR"))
    return count, arr


def _qtd_loss_spine_rows(wb: Any) -> list[list[Any]]:
    won_count, won_arr = _closed_le_rows(wb, True)
    lost_count, lost_arr = _closed_le_rows(wb, False)
    return [
        ["Loss / close spine", "#", "Unweighted ARR / ACV", "Action"],
        [
            "Original Q1 Land wins",
            _original_value(wb, "q1_land_wins", 0),
            _fmt_meur(_as_float(_original_value(wb, "q1_land_wins_arr_eur"))),
            "Show delivered Land ARR separately from current QTD.",
        ],
        [
            "Original Q1 Land losses",
            _original_value(wb, "q1_land_lost", 0),
            _fmt_meur(_as_float(_original_value(wb, "q1_land_lost_arr_eur"))),
            "Keep loss accountability visible in the May readout.",
        ],
        [
            "QTD closed-won L+E",
            won_count,
            _fmt_meur(won_arr),
            "Land+Expand closed-won unweighted ARR only.",
        ],
        [
            "QTD closed-lost L+E",
            lost_count,
            _fmt_meur(lost_arr),
            "Land+Expand closed-lost unweighted ARR only.",
        ],
    ]


def _named_risk_rows(wb: Any, limit: int = 4) -> list[list[Any]]:
    return _readiness_rows(
        wb,
        columns=("Account", "Opportunity", "Stage", "Close Date", "ARR", "Readiness"),
        limit=limit,
        flagged_only=True,
    )


def _action_item_rows(wb: Any, limit: int = 7) -> list[list[Any]]:
    columns = ("#", "Priority", "Rule", "Claim", "Suggested action", "Owner")
    output = [["#", "Priority", "Rule", "Signal", "Suggested action", "Owner"]]
    for row in _raw_rows(wb, "Raw_Current_Action_Items")[:limit]:
        values = [_clean_cell(row.get(column)) for column in columns]
        if len(values) >= 5:
            values[1] = str(values[1]).title() if values[1] else ""
            values[2] = str(values[2]).replace("_", " ").title() if values[2] else ""
            values[2] = str(values[2]).replace("Arr", "ARR")
            values[3], values[4] = _action_register_display(row.get("Rule"), values[3], values[4])
        if any(_meaningful(value) for value in values[1:]):
            output.append(values)
    if len(output) == 1:
        output.append(["", "No action rows found", "", "", "", ""])
    return output


def _specs(wb: Any) -> list[TableImageSpec]:
    q2_arr = _current_q2_closeable_arr(wb)
    new_count, new_arr = _new_pipeline_since_q2_start(wb)
    raw_renewal_count, raw_renewal_acv = _raw_renewal_totals(wb)
    closed_renewal_count, closed_renewal_acv = _closed_renewal_totals(wb)
    original_etl = _original_label(wb, "ETL")
    original_baseline = _original_label(wb, "baseline")
    return [
        TableImageSpec(
            "S04_ReviewDeltaTargets",
            "TC_S04_ReviewDelta",
            _q1_accountability_rows(wb),
            (20, 24, 24, 32),
        ),
        TableImageSpec(
            "S05_ForecastQualityTable",
            "TC_S05_ForecastQuality",
            [row[:3] for row in _forecast_category_rows(wb)],
            (42, 14, 28),
        ),
        TableImageSpec(
            "S06_HygieneSignals",
            "TC_S06_HygieneSignals",
            _hygiene_rows(wb),
            (16, 9, 13, 22, 12, 18, 34),
        ),
        TableImageSpec(
            "S07_TopDealsLand",
            "TC_S07_TopDealsLand",
            _readiness_rows(
                wb,
                columns=("#", "Account", "Opportunity", "Owner", "Type", "Stage", "Close Date", "ARR", "Forecast", "Prob"),
                limit=9,
            ),
            (4, 18, 18, 13, 9, 13, 11, 13, 10, 26),
            # think-cell Table-as-Image uses the Excel range's native rendered
            # size. Scaling this source range keeps the linked table full-size
            # without a PowerPoint-side resize that triggers carryover warnings.
            render_scale=2.50,
        ),
        TableImageSpec(
            "S08_TopDealsExpand",
            "TC_S08_ClosePlan",
            _readiness_rows(
                wb,
                columns=("Account", "Opportunity", "Owner", "Stage", "Forecast", "ARR", "Readiness"),
                limit=7,
                flagged_only=True,
            ),
            (4, 22, 22, 13, 12, 20, 30),
        ),
        TableImageSpec(
            "S09_PendingCommercialApproval",
            "TC_S09_CommercialApproval",
            _pending_approval_rows(wb),
            (4, 22, 26, 13, 12, 11, 9, 10),
            render_scale=2.50,
        ),
        TableImageSpec(
            "S11_RenewalPipeline",
            "TC_S11_Renewals",
            _renewal_pipeline_rows(wb),
            (4, 12, 20, 20, 13, 13, 11, 10, 18),
        ),
        TableImageSpec(
            "S12_GRRProxyTable",
            "TC_S12_RenewalRecon",
            [
                ["Source", "Scope", "#", "Amount", "Use in review"],
                [
                    original_etl,
                    "FY26 open renewals",
                    _original_value(wb, "fy26_renewals_count_original"),
                    f'{_fmt_meur(_as_float(_original_value(wb, "fy26_renewals_acv_eur_original", 0)))} ACV',
                    "Prior-deck baseline; do not blend with ARR.",
                ],
                [
                    original_etl,
                    "Q2 renewal subset",
                    _original_value(wb, "q2_renewals"),
                    f'{_fmt_meur(_as_float(_original_value(wb, "q2_renewals_acv_eur", 0)))} ACV',
                    "Current-quarter renewal context only.",
                ],
                [
                    "Refreshed workbook",
                    "FY26 renewal rows",
                    raw_renewal_count,
                    f"{_fmt_meur(raw_renewal_acv)} ACV",
                    "Operational watchlist basis.",
                ],
                [
                    "Refreshed workbook",
                    "QTD closed renewal",
                    closed_renewal_count,
                    f"{_fmt_meur(closed_renewal_acv)} ACV",
                    "Closed renewal motion only.",
                ],
            ],
            (18, 22, 8, 14, 30),
        ),
        TableImageSpec(
            "S13_ForecastCategoryDetail",
            "TC_S13_ForecastDetail",
            [row[:3] for row in _forecast_category_rows(wb)],
            (42, 14, 28),
        ),
        TableImageSpec(
            "S16_OwnerCoaching",
            "TC_S16_OwnerCoaching",
            _owner_coaching_rows(wb),
            (18, 10, 16, 36, 32),
        ),
        TableImageSpec(
            "S18_QTDLossSpine",
            "TC_S18_QTDLossSpine",
            _qtd_loss_spine_rows(wb),
            (24, 8, 18, 42),
        ),
        TableImageSpec(
            "S19_DealHygieneSignals",
            "TC_S19_DealHygiene",
            _signal_rows(wb),
            (18, 10, 16, 56),
        ),
        TableImageSpec(
            "S21_ConcentrationTable",
            "TC_S21_Concentration",
            [
                [original_baseline, "Value", "How to use it"],
                [
                    "Top 7 open deals",
                    _fmt_meur(_as_float(_original_value(wb, "concentration_top7_arr_eur", 0))),
                    "Preserve as prior concentration spine.",
                ],
                [
                    "Top 5 share",
                    f'{_as_float(_original_value(wb, "concentration_top5_share", 0)):.0%}',
                    "Use as a prior concentration flag.",
                ],
                [
                    "Largest original deal",
                    f'{_original_value(wb, "largest_original_deal_account", "")} {_fmt_meur(_as_float(_original_value(wb, "largest_original_deal_arr_eur", 0)))}'.strip(),
                    "Compare current Q2 readiness.",
                ],
                [
                    "Current Q2 closeable",
                    _fmt_meur(q2_arr),
                    "Current Land+Expand unweighted ARR only.",
                ],
            ],
            (22, 16, 30),
        ),
        TableImageSpec(
            "S22_NamedRiskTriage",
            "TC_S22_NamedRisk",
            _named_risk_rows(wb),
            (18, 26, 16, 12, 14, 32),
        ),
        TableImageSpec(
            "S23_OperatingRhythm",
            "TC_S23_OperatingRhythm",
            _operating_rhythm_rows(),
            (9, 18, 24, 49),
        ),
        TableImageSpec(
            "S24_AccountExpansion",
            "TC_S24_MayPlan",
            [
                ["May lane", "Fact base", "May output"],
                [
                    "Confirm or reset Q2 close",
                    f"{_fmt_meur(q2_arr)} Q2 Land+Expand unweighted ARR",
                    "Dated customer next step and close date reset.",
                ],
                [
                    "Create future pipe",
                    f"{new_count} new opps / {_fmt_meur(new_arr)}",
                    "Rep-owned coverage list and creation plan.",
                ],
                [
                    "Clean forecast",
                    f"{_forecast_row_count(wb)} forecast rows audited",
                    "Refresh activity, NextStep, stage, and forecast category.",
                ],
                [
                    "Renewal reconciliation",
                    f"{_fmt_meur(raw_renewal_acv)} FY26 renewal ACV",
                    "Keep ACV separate from ARR.",
                ],
            ],
            (24, 28, 34),
        ),
        TableImageSpec(
            "S25_Next14DaysCadence",
            "TC_S25_Next14Days",
            _next_14_days_rows(),
            (12, 18, 34, 40),
        ),
        TableImageSpec(
            "S26_ActionItems",
            "TC_S26_ActionItems",
            _action_item_rows(wb),
            (4, 8, 12, 38, 54, 12),
            render_scale=2.50,
        ),
        TableImageSpec(
            "S27_DecisionChecklist",
            "TC_S27_DecisionChecklist",
            _decision_checklist_rows(),
            (5, 16, 44, 35),
        ),
    ]


def _delete_defined_name(wb: Any, name: str) -> None:
    try:
        del wb.defined_names[name]
    except KeyError:
        pass


def _write_sheet(wb: Any, spec: TableImageSpec) -> None:
    if spec.sheet in wb.sheetnames:
        del wb[spec.sheet]
    ws = wb.create_sheet(spec.sheet)
    for row in spec.rows:
        ws.append(row)

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            cell.font = Font(name="Aptos", size=8 * spec.render_scale, color="1A1D31")
            cell.fill = BODY_FILL
            header = str(ws.cell(1, cell.column).value or "")
            if cell.row > 1:
                if any(token in header for token in ("ARR", "ACV", "Amount")):
                    cell.number_format = '"EUR" #,##0.0,,"M"'
                elif "Date" in header or header == "Close":
                    cell.number_format = "yyyy-mm-dd"
                elif "Prob" in header or "share" in header.lower():
                    if isinstance(cell.value, (int, float)) and abs(float(cell.value)) > 1:
                        cell.value = float(cell.value) / 100
                    cell.number_format = "0%"
            if cell.row == 1:
                cell.fill = HEADER_FILL
                cell.font = Font(
                    name="Aptos",
                    size=8 * spec.render_scale,
                    color="FFFFFF",
                    bold=True,
                )
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = (18 if row_idx == 1 else 28) * spec.render_scale
    ws.freeze_panes = "A2"
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        if spec.widths and col_idx <= len(spec.widths):
            ws.column_dimensions[letter].width = spec.widths[col_idx - 1] * spec.render_scale
        else:
            width = 12
            for cell in ws[letter]:
                value = cell.value or ""
                width = max(width, min(30, len(str(value)) + 2))
            ws.column_dimensions[letter].width = width * spec.render_scale

    _delete_defined_name(wb, spec.name)
    attr_text = f"{quote_sheetname(spec.sheet)}!$A$1:${get_column_letter(ws.max_column)}${ws.max_row}"
    wb.defined_names.add(DefinedName(spec.name, attr_text=attr_text))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--director-slug", default=DEFAULT_DIRECTOR)
    parser.add_argument(
        "--source",
        type=Path,
        help="Connected factory workbook to copy.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="Output workbook with table-image ranges.",
    )
    args = parser.parse_args()

    director_dir = ROOT / "state" / args.period / args.director_slug
    source = args.source or director_dir / "factory" / "connected" / "connected_factory.xlsx"
    output = args.output or director_dir / "factory" / "connected" / "connected_factory_table_images.xlsx"
    if not source.exists():
        raise SystemExit(f"missing connected factory workbook: {source}")

    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)

    wb = load_workbook(output)
    specs = _specs(wb)
    for spec in specs:
        _write_sheet(wb, spec)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(output)
    print(output)
    for spec in specs:
        print(f"{spec.name}: {spec.sheet}!A1:{get_column_letter(len(spec.rows[0]))}{len(spec.rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
