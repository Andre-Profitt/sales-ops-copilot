#!/usr/bin/env python3
"""Audit meeting-spine decks against the slide-surface contract.

Catches regressions in the action layer:
- A slide that should be a table-image becoming a native PPT table.
- A slide that should be a native PPT rebuild silently losing its table.
- Stale think-cell OLE objects re-appearing on any slide.
- Native chart insertion on a slide with no contract entry for it.

This is the regression detector P0 of TEMPLATE_VISUAL_QUALITY_AUDIT_2026-05-03
asks for. Runs after the brand-style gate; it is structure-aware where the
brand-style gate is style-aware.
"""

from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.presentation import Presentation as PresentationType

from period_context import DEFAULT_PERIOD, context_for_period


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_CONTRACT = ROOT / "config" / "meeting_spine_slide_contract.2026-Q2.json"


@dataclass
class SlideObservation:
    index: int
    title: str
    tables: int
    pictures: int
    charts: int
    ole_objects: int
    text_shapes: int
    other_shapes: int
    observed_kind: str


@dataclass
class SlideFinding:
    index: int
    severity: str
    code: str
    message: str


@dataclass
class DeckAuditResult:
    filename: str
    status: str
    slide_count: int
    slides: list[SlideObservation] = field(default_factory=list)
    findings: list[SlideFinding] = field(default_factory=list)


def _classify(s: SlideObservation) -> str:
    if s.ole_objects > 0:
        return "ole_present"
    has_table = s.tables > 0
    has_picture = s.pictures > 0
    has_chart = s.charts > 0
    if has_chart and not (has_table or has_picture):
        return "native_chart"
    if has_table and has_picture:
        return "native_ppt_mixed"
    if has_table and not has_picture:
        return "native_ppt_table"
    if has_picture and not has_table:
        return "table_image"
    return "editorial_native"


def _slide_title(slide) -> str:
    for shape in slide.shapes:
        if hasattr(shape, "is_placeholder") and shape.is_placeholder:
            ph = shape.placeholder_format
            if ph and ph.idx == 0 and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    return text[:80]
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return shape.text_frame.text.strip()[:80]
    return ""


def _observe_slide(index: int, slide) -> SlideObservation:
    counts = {
        "tables": 0,
        "pictures": 0,
        "charts": 0,
        "ole_objects": 0,
        "text_shapes": 0,
        "other_shapes": 0,
    }
    for shape in slide.shapes:
        if getattr(shape, "has_table", False):
            counts["tables"] += 1
        elif "OLE" in str(getattr(shape, "shape_type", "")):
            counts["ole_objects"] += 1
        elif getattr(shape, "name", "") == "Pic" or "Picture" in str(
            getattr(shape, "shape_type", "")
        ):
            counts["pictures"] += 1
        elif getattr(shape, "has_chart", False):
            counts["charts"] += 1
        elif getattr(shape, "has_text_frame", False):
            counts["text_shapes"] += 1
        else:
            counts["other_shapes"] += 1
    obs = SlideObservation(
        index=index,
        title=_slide_title(slide),
        tables=counts["tables"],
        pictures=counts["pictures"],
        charts=counts["charts"],
        ole_objects=counts["ole_objects"],
        text_shapes=counts["text_shapes"],
        other_shapes=counts["other_shapes"],
        observed_kind="",
    )
    obs.observed_kind = _classify(obs)
    return obs


def _audit_slide(slide_obs: SlideObservation, contract_entry: dict) -> list[SlideFinding]:
    findings: list[SlideFinding] = []
    expected = str(contract_entry.get("expected_kind", ""))
    alternates = set(contract_entry.get("alternative_kinds") or [])
    allowed = {expected} | alternates
    if slide_obs.ole_objects > 0:
        findings.append(
            SlideFinding(
                index=slide_obs.index,
                severity="fail",
                code="OLE_REGRESSION",
                message=(
                    f"slide {slide_obs.index} has {slide_obs.ole_objects} OLE object(s); "
                    "stale think-cell metadata likely re-introduced (scrub gate should have caught this)"
                ),
            )
        )
    if slide_obs.observed_kind not in allowed:
        findings.append(
            SlideFinding(
                index=slide_obs.index,
                severity="fail",
                code="SURFACE_REGRESSION",
                message=(
                    f"slide {slide_obs.index} ({contract_entry.get('purpose', '?')}): expected `{expected}`"
                    + (f" or one of {sorted(alternates)}" if alternates else "")
                    + f", observed `{slide_obs.observed_kind}` "
                    f"(t={slide_obs.tables} pic={slide_obs.pictures} ch={slide_obs.charts} ole={slide_obs.ole_objects})"
                ),
            )
        )
    if "expected_tables_min" in contract_entry:
        lo = int(contract_entry["expected_tables_min"])
        hi = int(contract_entry.get("expected_tables_max", lo))
        if not (lo <= slide_obs.tables <= hi):
            findings.append(
                SlideFinding(
                    index=slide_obs.index,
                    severity="fail",
                    code="TABLE_COUNT_DRIFT",
                    message=(
                        f"slide {slide_obs.index}: expected {lo}-{hi} native PPT tables, observed {slide_obs.tables}"
                    ),
                )
            )
    if "expected_pictures_min" in contract_entry:
        lo = int(contract_entry["expected_pictures_min"])
        hi = int(contract_entry.get("expected_pictures_max", lo))
        if not (lo <= slide_obs.pictures <= hi):
            findings.append(
                SlideFinding(
                    index=slide_obs.index,
                    severity="fail",
                    code="TABLE_IMAGE_DRIFT",
                    message=(
                        f"slide {slide_obs.index}: expected {lo}-{hi} table-image picture(s), observed {slide_obs.pictures}"
                    ),
                )
            )
    if (
        expected == "native_chart"
        and slide_obs.charts == 0
        and slide_obs.observed_kind != "native_chart"
    ):
        findings.append(
            SlideFinding(
                index=slide_obs.index,
                severity="fail",
                code="NATIVE_CHART_MISSING",
                message=f"slide {slide_obs.index} ({contract_entry.get('purpose', '?')}): native chart contract entry but 0 charts present",
            )
        )
    if "title_starts_with" in contract_entry:
        prefix = str(contract_entry["title_starts_with"])
        if not slide_obs.title.startswith(prefix):
            findings.append(
                SlideFinding(
                    index=slide_obs.index,
                    severity="warn",
                    code="TITLE_DRIFT",
                    message=f"slide {slide_obs.index} title `{slide_obs.title[:60]}…` does not start with `{prefix}`",
                )
            )
    return findings


def audit_deck(path: Path, contract: dict) -> DeckAuditResult:
    prs: PresentationType = Presentation(str(path))
    slide_count = len(prs.slides)
    expected_count = int(contract.get("slide_count", 0))
    findings: list[SlideFinding] = []
    if expected_count and slide_count != expected_count:
        findings.append(
            SlideFinding(
                index=0,
                severity="fail",
                code="DECK_SLIDE_COUNT_DRIFT",
                message=f"expected {expected_count} slides, observed {slide_count}",
            )
        )
    contract_by_index = {int(s["index"]): s for s in contract.get("slides", [])}
    slides: list[SlideObservation] = []
    for i, slide in enumerate(prs.slides, start=1):
        obs = _observe_slide(i, slide)
        slides.append(obs)
        entry = contract_by_index.get(i)
        if entry is None:
            findings.append(
                SlideFinding(
                    index=i,
                    severity="warn",
                    code="UNCONTRACTED_SLIDE",
                    message=f"slide {i} has no contract entry (observed `{obs.observed_kind}`)",
                )
            )
            continue
        findings.extend(_audit_slide(obs, entry))
    if any(f.severity == "fail" for f in findings):
        status = "fail"
    elif any(f.severity == "warn" for f in findings):
        status = "warn"
    else:
        status = "pass"
    return DeckAuditResult(
        filename=path.name,
        status=status,
        slide_count=slide_count,
        slides=slides,
        findings=findings,
    )


def run_audit(period: str, package_dir: Path, contract_path: Path) -> dict:
    context = context_for_period(period)
    contract = json.loads(contract_path.read_text(encoding="utf-8"))
    decks = sorted(
        path
        for path in package_dir.glob(context.meeting_spine_pattern)
        if not path.name.startswith("~$")
    )
    results = [audit_deck(deck, contract) for deck in decks]
    if not decks:
        status = "fail"
    elif any(r.status == "fail" for r in results):
        status = "fail"
    elif any(r.status == "warn" for r in results):
        status = "warn"
    else:
        status = "pass"
    return {
        "schema": "meeting-spine-audit-gate/v1",
        "status": status,
        "period": context.period,
        "package_dir": str(package_dir),
        "contract": str(contract_path),
        "deck_count": len(decks),
        "results": [_to_dict(r) for r in results],
    }


def _to_dict(result: DeckAuditResult) -> dict:
    return {
        "filename": result.filename,
        "status": result.status,
        "slide_count": result.slide_count,
        "slides": [asdict(s) for s in result.slides],
        "findings": [asdict(f) for f in result.findings],
    }


def write_markdown(payload: dict, path: Path) -> None:
    lines = [
        "# Meeting-Spine Audit Gate",
        "",
        f"- Status: `{payload['status']}`",
        f"- Period: `{payload['period']}`",
        f"- Decks: `{payload['deck_count']}`",
        f"- Contract: `{payload['contract']}`",
        "",
        "| Deck | Status | Slides | Findings |",
        "|---|---:|---:|---|",
    ]
    for r in payload["results"]:
        findings = "; ".join(
            f"[{f['severity']}] {f['code']}: {f['message']}" for f in r["findings"]
        )
        lines.append(f"| {r['filename']} | {r['status']} | {r['slide_count']} | {findings} |")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--package-dir", type=Path, default=None)
    parser.add_argument("--contract", type=Path, default=DEFAULT_CONTRACT)
    parser.add_argument("--json-output", type=Path)
    parser.add_argument("--markdown-output", type=Path)
    args = parser.parse_args()

    context = context_for_period(args.period)
    package_dir = (args.package_dir or context.review_package_dir).expanduser().resolve()
    contract_path = args.contract.expanduser().resolve()
    payload = run_audit(args.period, package_dir, contract_path)

    output_dir = ROOT / "state" / context.period / "__regional__" / "meeting_spine_audit"
    output_dir.mkdir(parents=True, exist_ok=True)
    json_output = args.json_output or output_dir / "meeting_spine_audit.json"
    markdown_output = args.markdown_output or output_dir / "meeting_spine_audit.md"
    json_output.parent.mkdir(parents=True, exist_ok=True)
    json_output.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    write_markdown(payload, markdown_output)
    print(json.dumps(payload, indent=2))
    return 2 if payload["status"] == "fail" else 0


if __name__ == "__main__":
    raise SystemExit(main())
