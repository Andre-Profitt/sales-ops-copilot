#!/usr/bin/env python3
"""Inspect packaged meeting-spine decks for SimCorp brand-style drift."""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from pptx import Presentation

from period_context import DEFAULT_PERIOD, context_for_period


ROOT = Path(__file__).resolve().parent.parent
APPROVED_FONT_NAMES = {"Aptos", "Arial"}
APPROVED_FONT_SIZES_PT = {
    6.0,
    7.0,
    8.0,
    9.0,
    10.0,
    12.0,
    14.0,
    16.0,
    18.0,
    24.0,
}
MIN_FONT_PT = 6.0
MAX_DISTINCT_FONT_SIZES = 11
MAX_SLIDE_FONT_SIZES = 7
FONT_SIZE_TOLERANCE = 0.05
ROUNDED_GEOMETRY_TOKENS = (
    "roundRect",
    "snipRoundRect",
    "round2SameRect",
    "round2DiagRect",
)
STALE_THINKCELL_TOKENS = (
    "think-cell data - do not delete",
    "TCLayout.ActiveDocument.1",
    "THINKCELLSHAPEDONOTDELETE",
    "THINKCELLUNDODONOTDELETE",
)
REQUIRED_PRESENTATION_SUPPORT_RELS = {
    "theme": "ppt/theme/theme1.xml",
    "viewProps": "ppt/viewProps.xml",
    "presProps": "ppt/presProps.xml",
    "tableStyles": "ppt/tableStyles.xml",
}


@dataclass
class DeckBrandStyleResult:
    filename: str
    status: str
    slide_count: int | None
    rounded_geometry_count: int
    stale_thinkcell_ownership_count: int
    missing_package_support: list[str] = field(default_factory=list)
    surface_counts: dict[str, int] = field(default_factory=dict)
    explicit_font_names: dict[str, int] = field(default_factory=dict)
    explicit_font_sizes: dict[str, int] = field(default_factory=dict)
    findings: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    error: str | None = None


def _zip_ok(path: Path) -> bool:
    try:
        with ZipFile(path) as zf:
            return zf.testzip() is None
    except BadZipFile:
        return False


def _rounded_geometry_count(path: Path) -> int:
    count = 0
    with ZipFile(path) as zf:
        for name in zf.namelist():
            if not name.startswith("ppt/slides/slide") or not name.endswith(".xml"):
                continue
            data = zf.read(name).decode("utf-8", errors="ignore")
            count += sum(data.count(token) for token in ROUNDED_GEOMETRY_TOKENS)
    return count


def _stale_thinkcell_ownership_count(path: Path) -> int:
    count = 0
    with ZipFile(path) as zf:
        for name in zf.namelist():
            if not (
                name.startswith("ppt/slides/")
                or name.startswith("ppt/tags/")
                or name.startswith("ppt/embeddings/")
            ):
                continue
            text = zf.read(name).decode("utf-8", errors="ignore")
            count += sum(text.count(token) for token in STALE_THINKCELL_TOKENS)
    return count


def _missing_package_support(path: Path) -> list[str]:
    missing: list[str] = []
    with ZipFile(path) as zf:
        names = set(zf.namelist())
        try:
            rels = zf.read("ppt/_rels/presentation.xml.rels").decode("utf-8", errors="ignore")
        except KeyError:
            return ["ppt/_rels/presentation.xml.rels"]
        for rel_name, part_name in REQUIRED_PRESENTATION_SUPPORT_RELS.items():
            if part_name not in names:
                missing.append(part_name)
            if rel_name not in rels:
                missing.append(f"presentation relationship: {rel_name}")
    return missing


def _iter_text_frames(shape: object) -> list[object]:
    frames: list[object] = []
    if getattr(shape, "has_text_frame", False):
        frames.append(shape.text_frame)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                frames.append(cell.text_frame)
    if hasattr(shape, "shapes"):
        for child in shape.shapes:
            frames.extend(_iter_text_frames(child))
    return frames


def _preview(text: str, limit: int = 48) -> str:
    value = " ".join(text.split())
    if len(value) <= limit:
        return value
    return value[: limit - 1].rstrip() + "..."


def _is_approved_font_size(size: float) -> bool:
    return any(abs(size - allowed) <= FONT_SIZE_TOLERANCE for allowed in APPROVED_FONT_SIZES_PT)


def _scan_text_style(prs: Presentation) -> tuple[Counter[str], Counter[str], list[str], list[str]]:
    font_names: Counter[str] = Counter()
    font_sizes: Counter[str] = Counter()
    findings: list[str] = []
    warnings: list[str] = []
    off_scale: Counter[str] = Counter()
    small_runs: list[str] = []
    non_brand_fonts: Counter[str] = Counter()
    slide_sizes: defaultdict[int, set[float]] = defaultdict(set)

    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            shape_name = getattr(shape, "name", "shape")
            for frame in _iter_text_frames(shape):
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        text = run.text.strip()
                        if not text:
                            continue
                        if run.font.name:
                            font_names[run.font.name] += 1
                            if run.font.name not in APPROVED_FONT_NAMES:
                                non_brand_fonts[run.font.name] += 1
                        if run.font.size is None:
                            continue
                        size = round(run.font.size.pt, 2)
                        size_key = f"{size:.2f}"
                        font_sizes[size_key] += 1
                        slide_sizes[slide_idx].add(size)
                        if size < MIN_FONT_PT:
                            small_runs.append(
                                f"slide {slide_idx} {shape_name}: {size:.2f}pt `{_preview(text)}`"
                            )
                        if not _is_approved_font_size(size):
                            off_scale[size_key] += 1

    if non_brand_fonts:
        fonts = ", ".join(f"{name}={count}" for name, count in sorted(non_brand_fonts.items()))
        findings.append(f"non-brand explicit fonts: {fonts}")
    if small_runs:
        findings.append(f"font sizes below {MIN_FONT_PT:.1f}pt: " + "; ".join(small_runs[:12]))
    if off_scale:
        sizes = ", ".join(f"{size}pt={count}" for size, count in sorted(off_scale.items()))
        findings.append(f"off-scale explicit font sizes: {sizes}")
    if len(font_sizes) > MAX_DISTINCT_FONT_SIZES:
        warnings.append(
            f"deck uses {len(font_sizes)} explicit font sizes; target <= {MAX_DISTINCT_FONT_SIZES}"
        )
    noisy_slides = [
        f"slide {slide_idx}={len(sizes)}"
        for slide_idx, sizes in sorted(slide_sizes.items())
        if len(sizes) > MAX_SLIDE_FONT_SIZES
    ]
    if noisy_slides:
        warnings.append("slides with too many explicit sizes: " + ", ".join(noisy_slides[:12]))
    return font_names, font_sizes, findings, warnings


def _surface_counts(prs: Presentation) -> dict[str, int]:
    counts = {
        "native_powerpoint_tables": 0,
        "table_image_pictures": 0,
        "embedded_ole_objects": 0,
        "native_powerpoint_charts": 0,
    }
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                counts["native_powerpoint_tables"] += 1
            if getattr(shape, "name", "") == "Pic":
                counts["table_image_pictures"] += 1
            if getattr(shape, "has_chart", False):
                counts["native_powerpoint_charts"] += 1
            if "EMBEDDED_OLE_OBJECT" in str(getattr(shape, "shape_type", "")):
                counts["embedded_ole_objects"] += 1
    return counts


def inspect_deck(path: Path) -> DeckBrandStyleResult:
    if not _zip_ok(path):
        return DeckBrandStyleResult(
            filename=path.name,
            status="fail",
            slide_count=None,
            rounded_geometry_count=0,
            stale_thinkcell_ownership_count=0,
            findings=["invalid zip package"],
            error="invalid zip package",
        )
    try:
        rounded_count = _rounded_geometry_count(path)
        stale_tc_count = _stale_thinkcell_ownership_count(path)
        missing_support = _missing_package_support(path)
        prs = Presentation(path)
        surface_counts = _surface_counts(prs)
        font_names, font_sizes, findings, warnings = _scan_text_style(prs)
    except Exception as exc:  # pragma: no cover - defensive CLI path
        return DeckBrandStyleResult(
            filename=path.name,
            status="fail",
            slide_count=None,
            rounded_geometry_count=0,
            stale_thinkcell_ownership_count=0,
            findings=[str(exc)],
            error=str(exc),
        )

    if rounded_count:
        findings.append(f"rounded/curved PowerPoint geometry found: {rounded_count}")
    if stale_tc_count:
        findings.append(f"stale think-cell ownership metadata found: {stale_tc_count}")
    if missing_support:
        findings.append("missing PowerPoint support parts/rels: " + ", ".join(missing_support))
    status = "fail" if findings else ("warn" if warnings else "pass")
    return DeckBrandStyleResult(
        filename=path.name,
        status=status,
        slide_count=len(prs.slides),
        rounded_geometry_count=rounded_count,
        stale_thinkcell_ownership_count=stale_tc_count,
        missing_package_support=missing_support,
        surface_counts=surface_counts,
        explicit_font_names=dict(sorted(font_names.items())),
        explicit_font_sizes=dict(sorted(font_sizes.items(), key=lambda item: float(item[0]))),
        findings=findings,
        warnings=warnings,
    )


def run_brand_style_gate(period: str, package_dir: Path, output_dir: Path) -> dict[str, object]:
    context = context_for_period(period)
    decks = sorted(
        path
        for path in package_dir.glob(context.meeting_spine_pattern)
        if not path.name.startswith("~$")
    )
    results = [inspect_deck(deck) for deck in decks]
    if not decks:
        status = "fail"
    elif any(result.status == "fail" for result in results):
        status = "fail"
    elif any(result.status == "warn" for result in results):
        status = "warn"
    else:
        status = "pass"
    return {
        "schema": "review-package-brand-style-gate/v1",
        "status": status,
        "period": context.period,
        "package_dir": str(package_dir),
        "output_dir": str(output_dir),
        "deck_count": len(decks),
        "rules": {
            "rounded_geometry_tokens": list(ROUNDED_GEOMETRY_TOKENS),
            "stale_thinkcell_tokens": list(STALE_THINKCELL_TOKENS),
            "required_presentation_support_rels": REQUIRED_PRESENTATION_SUPPORT_RELS,
            "approved_font_names": sorted(APPROVED_FONT_NAMES),
            "approved_font_sizes_pt": sorted(APPROVED_FONT_SIZES_PT),
            "min_font_pt": MIN_FONT_PT,
            "max_distinct_font_sizes": MAX_DISTINCT_FONT_SIZES,
            "max_slide_font_sizes": MAX_SLIDE_FONT_SIZES,
        },
        "results": [asdict(result) for result in results],
    }


def write_markdown(payload: dict[str, object], path: Path) -> None:
    lines = [
        "# Review Package Brand Style Gate",
        "",
        f"- Status: `{payload['status']}`",
        f"- Decks: `{payload['deck_count']}`",
        f"- Output: `{payload['output_dir']}`",
        "",
        "| Deck | Status | Slides | Rounded Geometry | Stale TC Ownership | PPT Tables | Table Images | OLE | Font Sizes | Findings | Warnings |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---|---|",
    ]
    for row in payload.get("results", []):
        result = dict(row)
        findings = "; ".join(result.get("findings") or [])
        warnings = "; ".join(result.get("warnings") or [])
        surface_counts = result.get("surface_counts") or {}
        lines.append(
            "| {deck} | {status} | {slides} | {rounded} | {stale_tc} | {ppt_tables} | {table_images} | {ole} | {font_sizes} | {findings} | {warnings} |".format(
                deck=result.get("filename"),
                status=result.get("status"),
                slides=result.get("slide_count"),
                rounded=result.get("rounded_geometry_count"),
                stale_tc=result.get("stale_thinkcell_ownership_count"),
                ppt_tables=surface_counts.get("native_powerpoint_tables", ""),
                table_images=surface_counts.get("table_image_pictures", ""),
                ole=surface_counts.get("embedded_ole_objects", ""),
                font_sizes=len(result.get("explicit_font_sizes") or {}),
                findings=findings,
                warnings=warnings,
            )
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--package-dir", type=Path, default=None)
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--json-output", type=Path)
    parser.add_argument("--markdown-output", type=Path)
    args = parser.parse_args()

    context = context_for_period(args.period)
    package_dir = (args.package_dir or context.review_package_dir).expanduser().resolve()
    output_dir = (
        (
            args.output_dir
            or ROOT
            / "state"
            / context.period
            / "__regional__"
            / "brand_style_gate"
            / "review_package"
        )
        .expanduser()
        .resolve()
    )
    payload = run_brand_style_gate(context.period, package_dir, output_dir)
    json_output = args.json_output or output_dir / "review_package_brand_style_gate.json"
    markdown_output = args.markdown_output or output_dir / "review_package_brand_style_gate.md"
    json_output.parent.mkdir(parents=True, exist_ok=True)
    json_output.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    write_markdown(payload, markdown_output)
    print(json.dumps(payload, indent=2))
    return 2 if payload["status"] == "fail" else 0


if __name__ == "__main__":
    raise SystemExit(main())
