"""Migrate the 4 production LAND director Excel models to embed
chart-binding named ranges (audit finding F-06).

For every binding name in `_BINDING_TO_RANGE_MANIFEST` with
`named_range_safe=true`, ensure a workbook-scoped Excel named range
exists pointing at the source range. The build_ppttc.py refactor reads
those named ranges by name instead of by hard-coded position, so the
pipeline survives Pivots-sheet layout drift.

Usage:

    .venv/bin/python scripts/add_chart_binding_named_ranges.py
    .venv/bin/python scripts/add_chart_binding_named_ranges.py --dry-run
    .venv/bin/python scripts/add_chart_binding_named_ranges.py --director Jesper-Tyrer

The migration is idempotent -- re-running does not duplicate names. The
56 existing data named ranges (`Data_*`, `ClosedCFQ_*`, `Renewals12mo_*`,
`Parameters!*`) are preserved untouched. Pre-mutation backup is written
to `<dir>/land.model.pre-named-ranges.xlsx` (only on first run; existing
backup is not clobbered).
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parent.parent
DIRECTOR_SLUGS = ("Jesper-Tyrer", "Sarah-Pittroff", "Patrick-Gaughan", "Megan-Miceli")
DEFAULT_PERIOD = "2026-Q2"


# Manifest of bindings whose Excel source is FIXED (not driven by
# `_last_nonempty_row` or trends.json). Each entry corresponds to a call
# in `scripts/build_ppttc.py`. `named_range_safe=false` entries are
# documented for completeness but are NOT migrated -- they have variable
# row counts or non-Excel sources.
_BINDING_TO_RANGE_MANIFEST: list[dict[str, Any]] = [
    # --- Matrix-shaped (fixed) ---
    {
        "name": "S04_PipeMovement",
        "function": "_pipe_movement_chart_entry",
        "build_ppttc_line": 459,
        "excel_sheet": "Pipe_Movement",
        "excel_range": "A2:B6",
        "shape": "matrix",
        "python_transforms_after": ["sign-flip negative buckets", "EUR -> mEUR scaling"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S05_PipelineByStage",
        "function": "_pipeline_by_stage_chart_entry",
        "build_ppttc_line": 469,
        "excel_sheet": "Pipeline_By_Stage",
        "excel_range": "A2:B9",
        "shape": "matrix",
        "python_transforms_after": ["EUR -> mEUR scaling", "drop blank stage rows"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S06_PipelineAging",
        "function": "_pipeline_aging_chart_entry",
        "build_ppttc_line": 480,
        "excel_sheet": "Pipeline_Aging",
        "excel_range": "A2:E7",
        "shape": "matrix",
        "python_transforms_after": ["EUR -> mEUR scaling"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S12_GRRProxyTable",
        "function": "_ppttc_entries_from_context",
        "build_ppttc_line": 790,
        "excel_sheet": "Retention",
        "excel_range": "A1:B4",
        "shape": "matrix",
        "python_transforms_after": [],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S13_ForecastCategory",
        "function": "_forecast_category_chart_entry",
        "build_ppttc_line": 366,
        "excel_sheet": "Forecast_Category",
        "excel_range": "A1:C7",
        "shape": "matrix",
        "python_transforms_after": ["EUR -> mEUR scaling"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S16_StageByIndustry",
        "function": "_stage_by_industry_chart_entry",
        "build_ppttc_line": 522,
        "excel_sheet": "Pivots",
        "excel_range": "A5:M13",
        "shape": "matrix",
        "python_transforms_after": ["top-5 industries + Other", "top-3 stages + Other"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S18_WinsLossesQTD",
        "function": "_wins_losses_chart_entry",
        "build_ppttc_line": 587,
        "excel_sheet": "Wins_Losses_QTD",
        "excel_range": "A1:D3",
        "shape": "matrix",
        "python_transforms_after": ["EUR -> mEUR scaling"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S19_Velocity",
        "function": "_velocity_chart_entry",
        "build_ppttc_line": 602,
        "excel_sheet": "Velocity",
        "excel_range": "A3:E10",
        "shape": "matrix",
        "python_transforms_after": ["compact stage labels", "median age in days"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S21_ConcentrationTable",
        "function": "_concentration_chart_entry/_concentration_entries",
        "build_ppttc_line": 616,
        "excel_sheet": "Concentration",
        "excel_range": "A11:C15",
        "shape": "matrix",
        "python_transforms_after": ["share to %"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S22_StaleActivity",
        "function": "_stale_activity_chart_entry",
        "build_ppttc_line": 627,
        "excel_sheet": "Stale_Activity",
        "excel_range": "A1:C5",
        "shape": "matrix",
        "python_transforms_after": ["EUR -> mEUR scaling"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S24_AccountExpansion",
        "function": "_ppttc_entries_from_context",
        "build_ppttc_line": 803,
        "excel_sheet": "Account_Expansion",
        "excel_range": "A1:F16",
        "shape": "matrix",
        "python_transforms_after": [],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S25_PipelineCreationVelocity",
        "function": "_pipeline_creation_velocity_chart_entry",
        "build_ppttc_line": 638,
        "excel_sheet": "Pipeline_Creation_Velocity",
        "excel_range": "A1:C13",
        "shape": "matrix",
        "python_transforms_after": ["week label formatting", "EUR -> mEUR scaling"],
        "named_range_safe": True,
        "reader": "model",
    },
    # --- Single-cell (named_cell) ---
    {
        "name": "S12_GRRProxyFootnote_src",
        "function": "_ppttc_entries_from_context",
        "build_ppttc_line": 791,
        "excel_sheet": "Retention",
        "excel_range": "A5",
        "shape": "cell",
        "python_transforms_after": ["str() coercion"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S22_StaleActivityFootnote_src",
        "function": "_ppttc_entries_from_context",
        "build_ppttc_line": 801,
        "excel_sheet": "Stale_Activity",
        "excel_range": "A7",
        "shape": "cell",
        "python_transforms_after": ["str() coercion"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S21_LargestAccount_src",
        "function": "_concentration_entries",
        "build_ppttc_line": 703,
        "excel_sheet": "Concentration",
        "excel_range": "B5",
        "shape": "cell",
        "python_transforms_after": [],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S21_LargestArr_src",
        "function": "_concentration_entries",
        "build_ppttc_line": 704,
        "excel_sheet": "Concentration",
        "excel_range": "B6",
        "shape": "cell",
        "python_transforms_after": ["currency formatting"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S21_LargestShare_src",
        "function": "_concentration_entries",
        "build_ppttc_line": 705,
        "excel_sheet": "Concentration",
        "excel_range": "B7",
        "shape": "cell",
        "python_transforms_after": ["percentage formatting"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S21_ThresholdFlag_src",
        "function": "_concentration_entries",
        "build_ppttc_line": 706,
        "excel_sheet": "Concentration",
        "excel_range": "B8",
        "shape": "cell",
        "python_transforms_after": [],
        "named_range_safe": True,
        "reader": "model",
    },
    # --- Sales_Velocity 5-row block: each row is its own binding ---
    {
        "name": "S23_OpenOpps_src",
        "function": "_sales_velocity_entries",
        "build_ppttc_line": 680,
        "excel_sheet": "Sales_Velocity",
        "excel_range": "B2",
        "shape": "cell",
        "python_transforms_after": ["int coercion"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S23_WinRate_src",
        "function": "_sales_velocity_entries",
        "build_ppttc_line": 680,
        "excel_sheet": "Sales_Velocity",
        "excel_range": "B3",
        "shape": "cell",
        "python_transforms_after": ["percentage formatting"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S23_AvgDealSize_src",
        "function": "_sales_velocity_entries",
        "build_ppttc_line": 680,
        "excel_sheet": "Sales_Velocity",
        "excel_range": "B4",
        "shape": "cell",
        "python_transforms_after": ["currency formatting"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S23_AvgCycleDays_src",
        "function": "_sales_velocity_entries",
        "build_ppttc_line": 680,
        "excel_sheet": "Sales_Velocity",
        "excel_range": "B5",
        "shape": "cell",
        "python_transforms_after": ["int + ' days' suffix"],
        "named_range_safe": True,
        "reader": "model",
    },
    {
        "name": "S23_Velocity_src",
        "function": "_sales_velocity_entries",
        "build_ppttc_line": 680,
        "excel_sheet": "Sales_Velocity",
        "excel_range": "B6",
        "shape": "cell",
        "python_transforms_after": ["currency formatting + ' / day' suffix"],
        "named_range_safe": True,
        "reader": "model",
    },
    # --- DOCUMENTED-ONLY: variable row counts or non-Excel sources ---
    {
        "name": "S15_ByOwner",
        "function": "_by_owner_chart_entry",
        "build_ppttc_line": 506,
        "excel_sheet": "By_Owner",
        "excel_range": "A2:B{last_row}",
        "shape": "matrix",
        "python_transforms_after": ["_last_nonempty_row dynamic bound", "drop zero-ARR rows"],
        "named_range_safe": False,
        "unsafe_reason": "variable last_row from _last_nonempty_row(model, 'By_Owner', ...)",
        "reader": "model",
    },
    {
        "name": "S17_TerritoryPerformance",
        "function": "_territory_chart_entry/_territory_bar_matrix",
        "build_ppttc_line": 357,
        "excel_sheet": "Territory_Performance",
        "excel_range": "B{r},D{r} per row",
        "shape": "matrix",
        "python_transforms_after": ["per-row cell_value loop", "drop TOTAL/blank/zero"],
        "named_range_safe": False,
        "unsafe_reason": "variable last_row + per-row cell_value (not a contiguous matrix call)",
        "reader": "model",
    },
    {
        "name": "S07_TopDealsLand",
        "function": "_ppttc_entries_from_context",
        "build_ppttc_line": 778,
        "excel_sheet": "Top_Deals_Land",
        "excel_range": "A1:H{last_row}",
        "shape": "matrix",
        "python_transforms_after": [],
        "named_range_safe": False,
        "unsafe_reason": "variable last_row + reads from legacy.xlsx (not land.model.xlsx)",
        "reader": "legacy",
    },
    {
        "name": "S08_TopDealsExpand",
        "function": "_ppttc_entries_from_context",
        "build_ppttc_line": 781,
        "excel_sheet": "Top_Deals_Expand",
        "excel_range": "A1:H{last_row}",
        "shape": "matrix",
        "python_transforms_after": [],
        "named_range_safe": False,
        "unsafe_reason": "variable last_row + reads from legacy.xlsx (not land.model.xlsx)",
        "reader": "legacy",
    },
    {
        "name": "S09_PendingCommercialApproval",
        "function": "_ppttc_entries_from_context",
        "build_ppttc_line": 785,
        "excel_sheet": "Pending_Commercial_Approval",
        "excel_range": "A3:H{last_row}",
        "shape": "matrix",
        "python_transforms_after": [],
        "named_range_safe": False,
        "unsafe_reason": "variable last_row + reads from legacy.xlsx (not land.model.xlsx)",
        "reader": "legacy",
    },
    {
        "name": "S11_RenewalPipeline",
        "function": "_ppttc_entries_from_context",
        "build_ppttc_line": 788,
        "excel_sheet": "At_Risk_Renewals",
        "excel_range": "A1:H{last_row}",
        "shape": "matrix",
        "python_transforms_after": [],
        "named_range_safe": False,
        "unsafe_reason": "variable last_row + reads from legacy.xlsx (not land.model.xlsx)",
        "reader": "legacy",
    },
    {
        "name": "S26_ActionItems",
        "function": "_action_items_table",
        "build_ppttc_line": 332,
        "excel_sheet": "(none)",
        "excel_range": "(none)",
        "shape": "table",
        "python_transforms_after": ["built from trends.json action_items"],
        "named_range_safe": False,
        "unsafe_reason": "no Excel source; built from trends.json",
        "reader": "trends",
    },
]


def _safe_bindings() -> list[dict[str, Any]]:
    return [b for b in _BINDING_TO_RANGE_MANIFEST if b["named_range_safe"]]


def _xlsx_path(period: str, slug: str) -> Path:
    return ROOT / "state" / period / slug / "land.model.xlsx"


def _backup_path(period: str, slug: str) -> Path:
    return ROOT / "state" / period / slug / "land.model.pre-named-ranges.xlsx"


def _split_a1(ref: str) -> tuple[str, str]:
    """Split A1-style ref like 'AB12' into ('AB', '12')."""
    col_chars: list[str] = []
    for ch in ref:
        if ch.isalpha():
            col_chars.append(ch)
        else:
            break
    col = "".join(col_chars)
    return col, ref[len(col) :]


def _format_refers_to(sheet: str, range_str: str) -> str:
    """Format an absolute Excel destination spec for a defined name.

    Quote the sheet name if it contains a space; always anchor cells with
    `$` so the named range stays absolute even if the source is moved.
    """
    sheet_part = f"'{sheet}'" if any(c.isspace() for c in sheet) else sheet
    if ":" in range_str:
        # matrix range like A5:M13
        start, end = range_str.split(":", 1)
        s_col, s_row = _split_a1(start)
        e_col, e_row = _split_a1(end)
        anchored = f"${s_col}${s_row}:${e_col}${e_row}"
    else:
        # single-cell ref like B5
        col, row = _split_a1(range_str)
        anchored = f"${col}${row}"
    return f"{sheet_part}!{anchored}"


def add_named_ranges(
    xlsx_path: Path,
    bindings: list[dict[str, Any]],
    *,
    dry_run: bool = False,
) -> dict[str, Any]:
    """Add chart-binding named ranges to one workbook.

    Returns a per-file report: which names were added vs. skipped (already
    present), plus the resolved `refers_to` strings.
    """
    wb = load_workbook(xlsx_path)
    existing_names = set(wb.defined_names)
    added: list[dict[str, str]] = []
    skipped: list[dict[str, str]] = []

    for binding in bindings:
        name = binding["name"]
        sheet = binding["excel_sheet"]
        range_str = binding["excel_range"]
        refers_to = _format_refers_to(sheet, range_str)
        if name in existing_names:
            skipped.append({"name": name, "refers_to": refers_to, "reason": "already present"})
            continue
        if not dry_run:
            wb.defined_names[name] = DefinedName(name=name, attr_text=refers_to)
        added.append({"name": name, "refers_to": refers_to})

    if added and not dry_run:
        wb.save(xlsx_path)

    return {
        "xlsx_path": str(xlsx_path),
        "before_count": len(existing_names),
        "after_count": len(existing_names) + len(added),
        "added_count": len(added),
        "skipped_count": len(skipped),
        "added": added,
        "skipped": skipped,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Add chart-binding named ranges to LAND director Excel models."
    )
    parser.add_argument(
        "--period",
        default=DEFAULT_PERIOD,
        help=f"Quarter slug (default: {DEFAULT_PERIOD}).",
    )
    parser.add_argument(
        "--director",
        action="append",
        default=None,
        help=(
            "Director slug (Jesper-Tyrer / Sarah-Pittroff / Patrick-Gaughan / "
            "Megan-Miceli). Repeatable. Default: all 4."
        ),
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Print what would be added without writing.",
    )
    parser.add_argument(
        "--manifest-out",
        type=Path,
        default=None,
        help="If given, write the binding-to-range manifest JSON here.",
    )
    args = parser.parse_args()

    targets = args.director or list(DIRECTOR_SLUGS)
    safe = _safe_bindings()

    print(f"Manifest: {len(_BINDING_TO_RANGE_MANIFEST)} bindings total, {len(safe)} safe.")

    reports: list[dict[str, Any]] = []
    for slug in targets:
        xlsx = _xlsx_path(args.period, slug)
        if not xlsx.exists():
            print(f"  SKIP {slug}: missing {xlsx}", file=sys.stderr)
            continue
        backup = _backup_path(args.period, slug)
        if not args.dry_run and not backup.exists():
            shutil.copy2(xlsx, backup)
            print(f"  backup -> {backup}")
        report = add_named_ranges(xlsx, safe, dry_run=args.dry_run)
        reports.append(report)
        print(
            f"  {slug}: {report['before_count']} -> {report['after_count']} "
            f"(+{report['added_count']} added, {report['skipped_count']} already present)"
        )

    if args.manifest_out is not None:
        args.manifest_out.parent.mkdir(parents=True, exist_ok=True)
        manifest_payload = {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "period": args.period,
            "directors": targets,
            "dry_run": args.dry_run,
            "bindings": _BINDING_TO_RANGE_MANIFEST,
            "reports": reports,
        }
        args.manifest_out.write_text(json.dumps(manifest_payload, indent=2))
        print(f"manifest -> {args.manifest_out}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
