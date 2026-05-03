#!/usr/bin/env python3
"""Run a resumable long-running polish loop for a Sales Director LAND deck."""

from __future__ import annotations

import argparse
import datetime as dt
import html
import json
import shutil
import subprocess
import sys
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DIRECTOR = "Jesper Tyrer"
DEFAULT_PERIOD = "2026-Q2"
DEFAULT_HOST = "Windows-VM"
EXPECTED_SLIDES = 28
MIN_NATIVE_TABLES = 12

FORBIDDEN_TEXT = (
    "Keywords:",
    "This slide contains",
    "Insert your desired text",
    "Insert chart title here",
    "[think-cell",
    "Pipe-movement bridge",
    "Stage × Industry",
    "Sales velocity",
    "Pipeline creation velocity",
    "GRR proxy",
    "old APAC",
    "Rebekka",
    "current deck had",
    "ready for rehearsal",
    "publish blockers",
    "Original Apr 20 pack",
    "CreatedDate velocity proxy",
    "sales-velocity proxy",
    "unwtd",
    "ARR unwtd",
    "ARR (unwtd)",
)

JESPER_REQUIRED_TEXT = (
    "Jesper Tyrer",
    "2026-Q2",
    "May 2026",
    "Friday, May 1, 2026",
    "Q1 opened",
    "EUR 17.1M",
    "Q1 lost",
    "EUR 5.5M",
    "Q1 slips",
    "EUR 9.4M",
    "FY26 renewals",
    "EUR 33.5M",
    "reconcile",
    "Current state workbook",
    "Q1 Land losses",
    "Q2-Q3 risk triage",
    "EUR 4.9M",
    "overdue-close",
    "EUR 10.7M",
    "EUR 5.6M",
    "Commit = EUR 2.7M",
    "zero recent activity",
    "top 7 open deals",
    "top 5 = 89%",
    "Amova Asset Management",
    "3 owners carry 50 pushes",
    "EUR 22.0M",
    "11 open deals pushed",
    "Edwina Chow owns 5",
    "Approved 2026",
    "Coolabah",
    "EUR 3.8M",
    "6 -> 12",
    "EUR 4.6M",
    "1W / 16L",
    "EUR 8.6M lost",
    "unweighted",
    "weighted ARR",
    "unweighted ARR",
    "Renewal ACV",
    "Land+Expand",
    "Danantara",
    "LTH",
    "Krungthai",
    "Bank Mandiri",
    "Temasek",
    "HKMA",
)


@dataclass
class GateResult:
    name: str
    status: str
    elapsed_seconds: float = 0.0
    command: list[str] = field(default_factory=list)
    stdout_log: str | None = None
    stderr_log: str | None = None
    notes: list[str] = field(default_factory=list)
    returncode: int | None = None


def _python() -> str:
    venv = ROOT / ".venv" / "bin" / "python"
    return str(venv) if venv.exists() else sys.executable


def _slug(value: str) -> str:
    return value.replace(" ", "-")


def _director_dir(period: str, director: str) -> Path:
    return ROOT / "state" / period / _slug(director)


def _factory_dir(period: str, director: str) -> Path:
    return _director_dir(period, director) / "factory"


def _deck_path(period: str, director: str) -> Path:
    slug = _slug(director)
    return _director_dir(period, director) / f"{slug}-LAND-{period}.pptx"


def _connected_workbook(period: str, director: str) -> Path:
    return _factory_dir(period, director) / "connected" / "connected_factory.xlsx"


def _connected_spec(period: str, director: str) -> Path:
    return _factory_dir(period, director) / "connected" / "connected_factory_spec.json"


def _cmd_text(command: list[str]) -> str:
    return " ".join(command)


def _tail(text: str, lines: int = 8) -> str:
    parts = text.strip().splitlines()
    return "\n".join(parts[-lines:])


def _run_gate(
    name: str,
    command: list[str],
    *,
    iteration_dir: Path,
    timeout_seconds: int | None = None,
) -> GateResult:
    logs_dir = iteration_dir / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)
    stdout_log = logs_dir / f"{name}.stdout.log"
    stderr_log = logs_dir / f"{name}.stderr.log"
    started = time.time()
    try:
        result = subprocess.run(
            command,
            cwd=ROOT,
            text=True,
            capture_output=True,
            check=False,
            timeout=timeout_seconds,
        )
        stdout_log.write_text(result.stdout)
        stderr_log.write_text(result.stderr)
        notes: list[str] = []
        if result.returncode != 0:
            if result.stdout.strip():
                notes.append("stdout tail:\n" + _tail(result.stdout))
            if result.stderr.strip():
                notes.append("stderr tail:\n" + _tail(result.stderr))
        return GateResult(
            name=name,
            status="pass" if result.returncode == 0 else "fail",
            elapsed_seconds=round(time.time() - started, 2),
            command=command,
            stdout_log=str(stdout_log),
            stderr_log=str(stderr_log),
            notes=notes,
            returncode=result.returncode,
        )
    except subprocess.TimeoutExpired as exc:
        stdout_log.write_text(exc.stdout or "")
        stderr_log.write_text(exc.stderr or "")
        return GateResult(
            name=name,
            status="fail",
            elapsed_seconds=round(time.time() - started, 2),
            command=command,
            stdout_log=str(stdout_log),
            stderr_log=str(stderr_log),
            notes=[f"timeout after {timeout_seconds}s"],
            returncode=None,
        )


def _pptx_text_and_counts(path: Path) -> tuple[str, int, int, int]:
    with ZipFile(path) as zf:
        names = zf.namelist()
        slides = [
            name
            for name in names
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        ]
        embeddings = [name for name in names if name.startswith("ppt/embeddings/")]
        xml = "\n".join(zf.read(name).decode("utf-8", "ignore") for name in slides)
    return html.unescape(xml), len(slides), len(embeddings), xml.count("<a:tbl")


def _audit_pptx(period: str, director: str) -> GateResult:
    started = time.time()
    path = _deck_path(period, director)
    errors: list[str] = []
    notes: list[str] = []
    if not path.exists() or path.stat().st_size == 0:
        return GateResult("deck_package_audit", "fail", notes=[f"missing or empty deck: {path}"])
    try:
        text, slide_count, embedding_count, native_tables = _pptx_text_and_counts(path)
    except BadZipFile as exc:
        return GateResult("deck_package_audit", "fail", notes=[f"invalid pptx zip: {exc}"])

    if slide_count != EXPECTED_SLIDES:
        errors.append(f"slide count {slide_count} != {EXPECTED_SLIDES}")
    if embedding_count:
        errors.append(f"native think-cell embeddings present: {embedding_count}")
    if native_tables < MIN_NATIVE_TABLES:
        errors.append(f"native table count {native_tables} < {MIN_NATIVE_TABLES}")

    required = JESPER_REQUIRED_TEXT if director == DEFAULT_DIRECTOR else (director, period)
    missing = [needle for needle in required if needle not in text]
    if missing:
        errors.append("missing required deck text: " + ", ".join(missing))
    leaked = [needle for needle in FORBIDDEN_TEXT if needle in text]
    if leaked:
        errors.append("forbidden text leaked: " + ", ".join(leaked))

    notes.append(
        f"deck={path} slides={slide_count} embeddings={embedding_count} native_tables={native_tables} size={path.stat().st_size}"
    )
    return GateResult(
        "deck_package_audit",
        "pass" if not errors else "fail",
        elapsed_seconds=round(time.time() - started, 2),
        notes=notes + errors,
    )


def _audit_connected_recalc(period: str, director: str, *, iteration_dir: Path) -> GateResult:
    started = time.time()
    workbook = _connected_workbook(period, director)
    soffice = shutil.which("soffice")
    if not soffice:
        return GateResult("connected_recalc_smoke", "fail", notes=["soffice not found on PATH"])
    if not workbook.exists():
        return GateResult("connected_recalc_smoke", "fail", notes=[f"missing workbook: {workbook}"])

    recalc_dir = iteration_dir / "connected_recalc"
    recalc_dir.mkdir(parents=True, exist_ok=True)
    cmd = [soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(recalc_dir), str(workbook)]
    gate = _run_gate("connected_recalc_convert", cmd, iteration_dir=iteration_dir, timeout_seconds=180)
    if gate.status != "pass":
        gate.name = "connected_recalc_smoke"
        return gate

    converted = recalc_dir / workbook.name
    if not converted.exists():
        matches = sorted(recalc_dir.glob("*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
        converted = matches[0] if matches else converted
    if not converted.exists():
        return GateResult("connected_recalc_smoke", "fail", notes=[f"recalculated workbook not produced in {recalc_dir}"])

    wb = load_workbook(converted, data_only=True, read_only=True)
    statuses = [row[3] for row in wb["Audit_Checks"].iter_rows(min_row=2, max_row=11, values_only=True)]
    closed_lost = wb["Out_S05_PipelineMovement"]["E2"].value
    q2_arr = wb["Out_S02_ExecutiveSummary"]["B2"].value
    errors: list[str] = []
    if any(status != "PASS" for status in statuses):
        errors.append(f"audit statuses not all PASS: {statuses}")
    if not isinstance(closed_lost, (int, float)) or closed_lost >= 0:
        errors.append(f"closed lost waterfall value is not negative: {closed_lost!r}")
    if not isinstance(q2_arr, (int, float)) or q2_arr < 0:
        errors.append(f"Q2 ARR output did not calculate: {q2_arr!r}")
    notes = [f"converted={converted}", f"q2_arr={q2_arr}", f"closed_lost={closed_lost}"]
    return GateResult(
        "connected_recalc_smoke",
        "pass" if not errors else "fail",
        elapsed_seconds=round(time.time() - started, 2),
        command=cmd,
        notes=notes + errors,
    )


def _audit_render(period: str, director: str, *, iteration_dir: Path) -> GateResult:
    started = time.time()
    png_dir = iteration_dir / "render" / "png"
    pngs = sorted(png_dir.glob("slide-*.png"))
    errors: list[str] = []
    if len(pngs) != EXPECTED_SLIDES:
        errors.append(f"rendered png count {len(pngs)} != {EXPECTED_SLIDES}")
    too_small = [path.name for path in pngs if path.stat().st_size < 20_000]
    if too_small:
        errors.append("small/blank-looking rendered slides: " + ", ".join(too_small[:8]))
    return GateResult(
        "render_artifact_audit",
        "pass" if not errors else "fail",
        elapsed_seconds=round(time.time() - started, 2),
        notes=[f"png_dir={png_dir}", f"png_count={len(pngs)}"] + errors,
    )


def _write_iteration_report(
    *,
    run_dir: Path,
    iteration_dir: Path,
    iteration: int,
    gates: list[GateResult],
    started_at: str,
    director: str,
    period: str,
) -> None:
    status = "PASS" if all(gate.status in {"pass", "skip"} for gate in gates) else "FAIL"
    payload = {
        "schema": "director-deck-polish-iteration/v1",
        "director": director,
        "period": period,
        "iteration": iteration,
        "started_at": started_at,
        "finished_at": dt.datetime.now().isoformat(timespec="seconds"),
        "status": status,
        "gates": [asdict(gate) for gate in gates],
        "artifacts": {
            "deck": str(_deck_path(period, director)),
            "connected_workbook": str(_connected_workbook(period, director)),
        },
    }
    (iteration_dir / "manifest.json").write_text(json.dumps(payload, indent=2))

    lines = [
        f"# Deck Polish Iteration {iteration} - {director} {period}",
        "",
        f"- Started: {started_at}",
        f"- Finished: {payload['finished_at']}",
        f"- Status: {status}",
        "",
        "| Gate | Status | Seconds | Notes |",
        "|---|---:|---:|---|",
    ]
    for gate in gates:
        notes = "<br>".join(note.replace("\n", "<br>") for note in gate.notes)
        lines.append(f"| {gate.name} | {gate.status.upper()} | {gate.elapsed_seconds:.2f} | {notes} |")
    lines.extend(["", "## Commands", ""])
    for gate in gates:
        if gate.command:
            lines.append(f"- `{_cmd_text(gate.command)}`")
    (iteration_dir / "report.md").write_text("\n".join(lines))

    latest = run_dir / "latest_report.md"
    latest.write_text((iteration_dir / "report.md").read_text())


def _write_run_summary(run_dir: Path, iteration_statuses: list[tuple[int, str, Path]]) -> None:
    lines = [
        "# Longhaul Deck Polish Run",
        "",
        f"- Updated: {dt.datetime.now().isoformat(timespec='seconds')}",
        "",
        "| Iteration | Status | Report |",
        "|---:|---:|---|",
    ]
    for iteration, status, report in iteration_statuses:
        lines.append(f"| {iteration} | {status} | `{report}` |")
    (run_dir / "summary.md").write_text("\n".join(lines))


def _should_refresh_source(iteration: int, refresh_every: int, skip_source: bool) -> bool:
    if skip_source:
        return False
    if refresh_every <= 0:
        return iteration == 1
    return iteration == 1 or (iteration - 1) % refresh_every == 0


def _run_iteration(args: argparse.Namespace, *, run_dir: Path, iteration: int) -> tuple[str, Path]:
    started_at = dt.datetime.now().isoformat(timespec="seconds")
    iteration_dir = run_dir / f"iter-{iteration:03d}-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}"
    iteration_dir.mkdir(parents=True, exist_ok=True)
    director_dir = _director_dir(args.period, args.director)
    gates: list[GateResult] = []

    if _should_refresh_source(iteration, args.refresh_source_every, args.skip_source):
        gates.append(
            _run_gate(
                "source_refresh",
                [_python(), "scripts/land_brief.py", "--director", args.director, "--period", args.period],
                iteration_dir=iteration_dir,
                timeout_seconds=args.source_timeout_seconds,
            )
        )
    else:
        gates.append(GateResult("source_refresh", "skip", notes=["reusing existing source artifacts"]))

    gates.append(
        _run_gate(
            "workbook_contract",
            [_python(), "scripts/validate_workbook_contract.py", str(director_dir)],
            iteration_dir=iteration_dir,
            timeout_seconds=120,
        )
    )
    gates.append(
        _run_gate(
            "excel_analysis_coverage",
            [_python(), "scripts/validate_excel_analysis_coverage.py", str(director_dir)],
            iteration_dir=iteration_dir,
            timeout_seconds=120,
        )
    )

    gates.append(
        _run_gate(
            "connected_factory_build",
            [
                _python(),
                "scripts/build_connected_factory_workbook.py",
                "--director",
                args.director,
                "--period",
                args.period,
            ],
            iteration_dir=iteration_dir,
            timeout_seconds=180,
        )
    )
    gates.append(
        _run_gate(
            "connected_factory_validate",
            [
                _python(),
                "scripts/validate_connected_factory_spec.py",
                "--spec",
                str(_connected_spec(args.period, args.director)),
                "--workbook",
                str(_connected_workbook(args.period, args.director)),
            ],
            iteration_dir=iteration_dir,
            timeout_seconds=120,
        )
    )
    gates.append(_audit_connected_recalc(args.period, args.director, iteration_dir=iteration_dir))

    deck_cmd = [
        _python(),
        "scripts/build_land_presentation_decks.py",
        "--director",
        args.director,
        "--period",
        args.period,
        "--host",
        args.host,
    ]
    if args.skip_seed:
        deck_cmd.append("--skip-seed")
    gates.append(
        _run_gate(
            "deck_build",
            deck_cmd,
            iteration_dir=iteration_dir,
            timeout_seconds=args.deck_timeout_seconds,
        )
    )
    gates.append(_audit_pptx(args.period, args.director))
    if args.director == DEFAULT_DIRECTOR:
        gates.append(
            _run_gate(
                "jesper_apac_intel_coverage",
                [
                    _python(),
                    "scripts/audit_jesper_apac_intel_coverage.py",
                    str(_deck_path(args.period, args.director)),
                    "--json-output",
                    str(iteration_dir / "jesper_apac_intel_coverage.json"),
                    "--markdown-output",
                    str(iteration_dir / "jesper_apac_intel_coverage.md"),
                ],
                iteration_dir=iteration_dir,
                timeout_seconds=120,
            )
        )
    else:
        gates.append(GateResult("jesper_apac_intel_coverage", "skip", notes=["Jesper/APAC only"]))

    if args.skip_powerpoint_open:
        gates.append(GateResult("powerpoint_open_smoke", "skip", notes=["skipped by CLI flag"]))
    else:
        gates.append(
            _run_gate(
                "powerpoint_open_smoke",
                [
                    _python(),
                    "scripts/powerpoint_open_smoke.py",
                    str(_deck_path(args.period, args.director)),
                    "--host",
                    args.host,
                ],
                iteration_dir=iteration_dir,
                timeout_seconds=args.powerpoint_timeout_seconds,
            )
        )

    if args.skip_render:
        gates.append(GateResult("visual_render_smoke", "skip", notes=["skipped by CLI flag"]))
    else:
        render_dir = iteration_dir / "render"
        gates.append(
            _run_gate(
                "visual_render_smoke",
                [
                    _python(),
                    "scripts/render_deck_for_review.py",
                    str(_deck_path(args.period, args.director)),
                    "--output-dir",
                    str(render_dir),
                ],
                iteration_dir=iteration_dir,
                timeout_seconds=args.render_timeout_seconds,
            )
        )
        gates.append(_audit_render(args.period, args.director, iteration_dir=iteration_dir))

    status = "PASS" if all(gate.status in {"pass", "skip"} for gate in gates) else "FAIL"
    _write_iteration_report(
        run_dir=run_dir,
        iteration_dir=iteration_dir,
        iteration=iteration,
        gates=gates,
        started_at=started_at,
        director=args.director,
        period=args.period,
    )
    return status, iteration_dir / "report.md"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--director", default=DEFAULT_DIRECTOR)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--host", default=DEFAULT_HOST)
    parser.add_argument("--iterations", type=int, default=1, help="Number of iterations; 0 means run until stop-file/max-minutes.")
    parser.add_argument("--max-minutes", type=float, default=0, help="Maximum wall-clock minutes; 0 means no wall-clock cap.")
    parser.add_argument("--sleep-seconds", type=int, default=300)
    parser.add_argument("--refresh-source-every", type=int, default=1, help="Refresh Salesforce/source facts every N iterations; 0 means first iteration only.")
    parser.add_argument("--skip-source", action="store_true")
    parser.add_argument("--skip-seed", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--skip-powerpoint-open", action="store_true")
    parser.add_argument("--skip-render", action="store_true")
    parser.add_argument("--until-pass", action="store_true", help="Stop after the first fully passing iteration.")
    parser.add_argument("--run-dir", type=Path, default=None)
    parser.add_argument("--source-timeout-seconds", type=int, default=900)
    parser.add_argument("--deck-timeout-seconds", type=int, default=900)
    parser.add_argument("--powerpoint-timeout-seconds", type=int, default=300)
    parser.add_argument("--render-timeout-seconds", type=int, default=300)
    args = parser.parse_args()

    run_root = args.run_dir or (_factory_dir(args.period, args.director) / "longhaul")
    run_root.mkdir(parents=True, exist_ok=True)
    stop_file = run_root / "STOP"
    started = time.time()
    iteration_statuses: list[tuple[int, str, Path]] = []
    iteration = 0

    print(f"longhaul_run_dir={run_root}")
    print(f"stop_file={stop_file}")

    while True:
        if stop_file.exists():
            print(f"STOP file present: {stop_file}")
            break
        if args.max_minutes > 0 and (time.time() - started) / 60 >= args.max_minutes:
            print(f"max-minutes reached: {args.max_minutes}")
            break
        if args.iterations > 0 and iteration >= args.iterations:
            break

        iteration += 1
        status, report = _run_iteration(args, run_dir=run_root, iteration=iteration)
        iteration_statuses.append((iteration, status, report))
        _write_run_summary(run_root, iteration_statuses)
        print(f"iteration={iteration} status={status} report={report}")

        if status == "PASS" and args.until_pass:
            break
        if args.iterations > 0 and iteration >= args.iterations:
            break
        if args.sleep_seconds > 0:
            time.sleep(args.sleep_seconds)

    _write_run_summary(run_root, iteration_statuses)
    return 0 if iteration_statuses and iteration_statuses[-1][1] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
