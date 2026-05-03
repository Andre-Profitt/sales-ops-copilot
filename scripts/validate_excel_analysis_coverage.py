#!/usr/bin/env python3
"""Validate the factory Excel-analysis coverage spec against workbook support.

This is intentionally independent of deck generation. It verifies that the
named analyses in state/<period>/<director>/factory/excel_analysis_spec.json
are present and grounded in existing sheets, tables, and headers.

CLI:
    python3 scripts/validate_excel_analysis_coverage.py state/2026-Q2/Jesper-Tyrer
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_ANALYSES = {
    "q1_accountability",
    "q2_readiness",
    "approval_governance",
    "owner_push_stale_coaching",
    "renewal_acv_reconciliation",
    "loss_reasons",
}

REQUIRED_WORKBOOK_KEYS = {"current_model", "deck_companion"}


def _load_json(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text())
    except json.JSONDecodeError as exc:
        raise ValueError(f"invalid JSON in {path}: {exc}") from exc


def _check_headers(ws: Any, headers: list[Any], header_row: int, label: str) -> list[str]:
    errors: list[str] = []
    for col_idx, expected in enumerate(headers, start=1):
        if expected is None:
            continue
        actual = ws.cell(row=header_row, column=col_idx).value
        if actual != expected:
            coord = ws.cell(row=header_row, column=col_idx).coordinate
            errors.append(f"{label}!{coord}: expected {expected!r}, got {actual!r}")
    return errors


def validate(director_dir: Path, spec_path: Path | None = None) -> list[str]:
    errors: list[str] = []
    if spec_path is None:
        spec_path = director_dir / "factory" / "excel_analysis_spec.json"
    if not spec_path.exists():
        return [f"missing spec: {spec_path}"]

    spec = _load_json(spec_path)
    if spec.get("schema") != "director-land-excel-analysis-spec/v1":
        errors.append(f"unexpected schema: {spec.get('schema')!r}")

    workbooks = spec.get("workbooks") or {}
    missing_workbook_keys = REQUIRED_WORKBOOK_KEYS - set(workbooks)
    if missing_workbook_keys:
        errors.append(f"missing workbook keys: {sorted(missing_workbook_keys)}")

    loaded: dict[str, Any] = {}
    for key, rel_path in workbooks.items():
        path = director_dir / rel_path
        if not path.exists():
            errors.append(f"workbook missing for {key}: {path}")
            continue
        loaded[key] = load_workbook(path, read_only=False, data_only=False)

    analyses = spec.get("analyses") or []
    analysis_ids = {item.get("id") for item in analyses}
    missing_analyses = REQUIRED_ANALYSES - analysis_ids
    extra_none = None in analysis_ids
    if missing_analyses:
        errors.append(f"missing required analyses: {sorted(missing_analyses)}")
    if extra_none:
        errors.append("one or more analyses missing id")

    for analysis in analyses:
        analysis_id = analysis.get("id", "(missing id)")
        if not analysis.get("deck_tables"):
            errors.append(f"{analysis_id}: deck_tables must be non-empty")
        if not analysis.get("publish_rule"):
            errors.append(f"{analysis_id}: publish_rule must be present")

        sources = analysis.get("required_sources") or []
        if not sources:
            errors.append(f"{analysis_id}: required_sources must be non-empty")
            continue
        source_workbook_keys = {source.get("workbook") for source in sources}
        for required_key in REQUIRED_WORKBOOK_KEYS:
            if required_key not in source_workbook_keys:
                errors.append(f"{analysis_id}: missing {required_key} source")

        for source in sources:
            workbook_key = source.get("workbook")
            sheet_name = source.get("sheet")
            if workbook_key not in loaded:
                errors.append(f"{analysis_id}: unknown or unloaded workbook {workbook_key!r}")
                continue
            wb = loaded[workbook_key]
            if sheet_name not in wb.sheetnames:
                errors.append(f"{analysis_id}: sheet missing in {workbook_key}: {sheet_name}")
                continue
            ws = wb[sheet_name]
            table_name = source.get("table")
            if table_name and table_name not in ws.tables:
                errors.append(
                    f"{analysis_id}: table missing on {workbook_key}.{sheet_name}: {table_name}"
                )
            headers = source.get("headers")
            if headers:
                label = f"{analysis_id}:{workbook_key}.{sheet_name}"
                errors.extend(_check_headers(ws, headers, int(source.get("header_row", 1)), label))

    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "director_dir",
        type=Path,
        help="Path to state/<period>/<director>/ containing land.model.xlsx and land.xlsx",
    )
    parser.add_argument("--spec", type=Path, default=None)
    args = parser.parse_args()

    errors = validate(args.director_dir, args.spec)
    print(f"=== Excel analysis coverage: {args.director_dir.name} ===")
    if errors:
        print("FAIL")
        for error in errors:
            print(f"  - {error}")
        return 1
    print(f"OK: {len(REQUIRED_ANALYSES)} named analyses covered by workbook sources")
    return 0


if __name__ == "__main__":
    sys.exit(main())
