#!/usr/bin/env python3
"""Build a short approval deck for the May 2026 think-cell visual direction."""

from __future__ import annotations

import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
DIRECTOR_DIR = ROOT / "state" / "2026-Q2" / "Jesper-Tyrer"
SOURCE_XLSX = DIRECTOR_DIR / "land.xlsx"
OUT_DIR = DIRECTOR_DIR / "factory" / "design-mockups"
OUT_PPTX = OUT_DIR / "may_2026_template_improvement_sample_v2.pptx"

FONT = "Microsoft Sans Serif"
BLACK = RGBColor(0x00, 0x00, 0x00)
NAVY = RGBColor(0x1A, 0x1D, 0x31)
BLUE = RGBColor(0x08, 0x3E, 0xA7)
CYAN = RGBColor(0x2F, 0xC3, 0xD2)
TEAL = RGBColor(0x00, 0x86, 0x7B)
PURPLE = RGBColor(0x4B, 0x17, 0xB6)
RED = RGBColor(0xE7, 0x00, 0x4C)
AMBER = RGBColor(0xE4, 0xA1, 0x1B)
GREEN = RGBColor(0x1D, 0x9A, 0x5A)
GREY = RGBColor(0x66, 0x66, 0x66)
MID_GREY = RGBColor(0xA3, 0xAA, 0xB7)
LIGHT_GREY = RGBColor(0xF3, 0xF5, 0xF8)
LINE = RGBColor(0xD8, 0xDE, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def inch(value: float) -> Emu:
    return Emu(int(value * 914400))


def delete_all_slides(prs: Presentation) -> None:
    ids = prs.slides._sldIdLst  # noqa: SLF001
    for index in range(len(prs.slides) - 1, -1, -1):
        prs.part.drop_rel(ids[index].rId)
        del ids[index]


def add_text(
    slide: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: float,
    *,
    bold: bool = False,
    color: RGBColor = NAVY,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> Any:
    shape = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_rect(
    slide: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor,
    *,
    line: RGBColor | None = None,
    radius: bool = False,
) -> Any:
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, inch(x), inch(y), inch(w), inch(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.6)
    else:
        shape.line.fill.background()
    return shape


def add_line(slide: Any, x1: float, y1: float, x2: float, y2: float, color: RGBColor = LINE, width: float = 0.8) -> Any:
    line = slide.shapes.add_connector(1, inch(x1), inch(y1), inch(x2), inch(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_header(slide: Any, title: str, subtitle: str, lane: str, slide_no: int) -> None:
    add_text(slide, 0.78, 0.48, 9.7, 0.52, title, 26, bold=True, color=BLACK)
    add_text(slide, 0.80, 1.03, 10.0, 0.25, subtitle, 8.3, color=GREY)
    add_rect(slide, 11.35, 0.56, 1.25, 0.22, LIGHT_GREY, line=LINE)
    add_text(slide, 11.43, 0.595, 1.08, 0.11, lane, 5.8, bold=True, color=GREY, align=PP_ALIGN.CENTER)
    add_footer(slide, slide_no)


def add_footer(slide: Any, slide_no: int) -> None:
    y = 7.03
    add_text(
        slide,
        2.26,
        y + 0.055,
        7.9,
        0.12,
        "Visual direction sample | Salesforce snapshot 2026-04-30 | Land+Expand ARR unweighted unless noted; Renewal ACV separate.",
        5.8,
        color=GREY,
    )


def parse_meur(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value) / 1_000_000
    text = str(value)
    match = re.search(r"(-?\d+(?:\.\d+)?)", text)
    return float(match.group(1)) if match else 0.0


def compact(value: str, limit: int) -> str:
    if len(value) <= limit:
        return value
    return value[: max(0, limit - 3)].rstrip() + "..."


def load_rows(sheet_name: str) -> list[list[Any]]:
    wb = load_workbook(SOURCE_XLSX, data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def add_kpi_slide(prs: Presentation, slide_no: int) -> None:
    rows = load_rows("Pipeline_Total")
    forecast = load_rows("Forecast_Category")
    win_loss = load_rows("Wins_Losses_QTD")
    kpis = {row[0]: row[1] for row in rows[1:] if row and row[0]}
    forecast_map = {row[0]: row for row in forecast[4:] if row and row[0]}
    won = win_loss[1]
    lost = win_loss[2]

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "May APAC operating read",
        "Approval sample for the proposed SimCorp / think-cell visual system.",
        "DASHBOARD STRIP",
        slide_no,
    )
    add_text(
        slide,
        0.82,
        1.55,
        10.8,
        0.38,
        "Closeable Q2 pipe is concentrated in named deals; Commit carries the call and Best Case has no useful cushion.",
        15.8,
        bold=True,
        color=BLACK,
    )
    add_text(
        slide,
        0.82,
        1.96,
        10.5,
        0.24,
        "The tile treatment is intentionally restrained: fewer boxes, stronger hierarchy, exact metric basis, and no ARR/ACV blending.",
        8.2,
        color=GREY,
    )
    add_line(slide, 0.82, 2.42, 12.45, 2.42, NAVY, 1.5)

    items = [
        ("Closeable Q2 L+E ARR", kpis.get("2026-Q2 closeable Land+Expand ARR", "5.1 mEUR"), "26 opps | unweighted ARR", BLUE),
        ("Commit ARR", forecast_map.get("Commit", ["Commit", 11, "2.8 mEUR"])[2], "11 opps | unweighted ARR", TEAL),
        ("Pipeline ARR", forecast_map.get("Pipeline", ["Pipeline", 8, "2.3 mEUR"])[2], "8 opps | unweighted ARR", PURPLE),
        ("QTD closed", f"Won {won[2]} / Lost {lost[2]}", "Land+Expand ARR; lost is negative in movement views", RED),
    ]
    x0, y0, colw = 0.82, 2.66, 2.9
    for idx, (label, value, note, accent) in enumerate(items):
        x = x0 + idx * colw
        add_rect(slide, x, y0, 0.045, 0.95, accent)
        add_text(slide, x + 0.16, y0 + 0.04, colw - 0.28, 0.16, label.upper(), 6.5, bold=True, color=GREY)
        add_text(slide, x + 0.16, y0 + 0.28, colw - 0.28, 0.28, str(value).replace("mEUR", "mEUR"), 20, bold=True, color=NAVY)
        add_text(slide, x + 0.16, y0 + 0.68, colw - 0.28, 0.22, note, 6.8, color=GREY)
        if idx < len(items) - 1:
            add_line(slide, x + colw - 0.15, y0 + 0.06, x + colw - 0.15, y0 + 0.88, LINE, 0.7)

    add_line(slide, 0.82, 4.05, 12.45, 4.05, LINE, 0.8)
    add_text(slide, 0.82, 4.30, 5.4, 0.25, "Readout for the director", 12, bold=True, color=NAVY)
    bullets = [
        "Inspect Danantara, LTH, Krungthai, Mandiri, Temasek, and HKMA as named deal decisions.",
        "Keep QTD losses visible; EUR 6.0M lost is larger than EUR 1.6M won.",
        "Renewal ACV stays separate from Land+Expand ARR in both text and charts.",
    ]
    for i, bullet in enumerate(bullets):
        add_text(slide, 0.82, 4.72 + i * 0.36, 5.7, 0.22, "- " + bullet, 9, color=NAVY)
    add_text(slide, 7.00, 4.30, 4.6, 0.25, "What to approve here", 12, bold=True, color=NAVY)
    notes = [
        "Thin metric strip instead of boxed AI-looking tiles.",
        "Clear labels: unweighted ARR, weighted ARR, ACV.",
        "Enough whitespace to feel designed, not dashboard-dumped.",
    ]
    for i, note in enumerate(notes):
        add_text(slide, 7.00, 4.72 + i * 0.36, 4.8, 0.22, "- " + note, 9, color=NAVY)


def add_waterfall_slide(prs: Presentation, slide_no: int) -> None:
    win_loss = load_rows("Wins_Losses_QTD")
    pipeline = load_rows("Pipeline_Total")
    won = parse_meur(win_loss[1][2])
    lost = -parse_meur(win_loss[2][2])
    remaining = parse_meur(pipeline[1][1])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "QTD closed vs remaining closeable pipe",
        "Waterfall/contribution style: loss bars subtract, current closeable pipe remains a separate stock.",
        "WATERFALL",
        slide_no,
    )
    chart_x, chart_y, chart_w, chart_h = 0.98, 1.70, 7.65, 4.45
    add_line(slide, chart_x, chart_y + chart_h * 0.52, chart_x + chart_w, chart_y + chart_h * 0.52, MID_GREY, 1.0)
    add_text(slide, chart_x, chart_y - 0.03, 1.4, 0.14, "mEUR", 7, color=GREY)
    values = [won, lost, remaining]
    labels = ["QTD won", "QTD lost", "Q2 closeable"]
    colors = [GREEN, RED, BLUE]
    max_abs = 6.5
    base_y = chart_y + chart_h * 0.52
    scale = (chart_h * 0.42) / max_abs
    bar_w = 1.0
    for idx, value in enumerate(values):
        x = chart_x + 1.0 + idx * 2.1
        h = abs(value) * scale
        y = base_y - h if value >= 0 else base_y
        add_rect(slide, x, y, bar_w, h, colors[idx])
        add_text(slide, x - 0.18, y - 0.28 if value >= 0 else y + h + 0.10, 1.35, 0.18, f"{value:+.1f}", 11, bold=True, color=colors[idx], align=PP_ALIGN.CENTER)
        add_text(slide, x - 0.18, chart_y + chart_h - 0.10, 1.35, 0.18, labels[idx], 8, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        if idx == 1:
            add_text(slide, x - 0.35, base_y - 0.34, 1.75, 0.14, "subtracts", 7, bold=True, color=RED, align=PP_ALIGN.CENTER)
    for tick in [-6, -3, 0, 3, 6]:
        y = base_y - tick * scale
        add_line(slide, chart_x + 0.15, y, chart_x + chart_w - 0.45, y, RGBColor(0xEC, 0xEF, 0xF4), 0.5)
        add_text(slide, chart_x - 0.05, y - 0.06, 0.35, 0.10, str(tick), 5.8, color=GREY, align=PP_ALIGN.RIGHT)

    add_rect(slide, 9.05, 1.82, 3.2, 3.65, LIGHT_GREY, line=LINE)
    add_text(slide, 9.28, 2.08, 2.55, 0.22, "Director implication", 11.5, bold=True, color=BLACK)
    text = [
        "The view forces the commercial reality: losses are not a positive segment.",
        "Current closeable ARR should be inspected as named rows, not treated as generic coverage.",
        "Use this only where movement logic is clear; avoid waterfall for unrelated snapshots.",
    ]
    for i, item in enumerate(text):
        add_text(slide, 9.28, 2.55 + i * 0.55, 2.55, 0.34, "- " + item, 8.2, color=NAVY)


def add_scatter_slide(prs: Presentation, slide_no: int) -> None:
    rows = load_rows("Q2_Readiness")[1:9]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Deal risk inspection map",
        "Scatter/Bubble direction: plot value against probability, then use color and label to show inspection need.",
        "SCATTER / BUBBLE",
        slide_no,
    )
    x0, y0, w, h = 1.05, 1.75, 7.3, 4.55
    add_line(slide, x0, y0 + h, x0 + w, y0 + h, NAVY, 1.1)
    add_line(slide, x0, y0, x0, y0 + h, NAVY, 1.1)
    for pct in [0, 25, 50, 75, 100]:
        x = x0 + w * pct / 100
        add_line(slide, x, y0, x, y0 + h, RGBColor(0xEC, 0xEF, 0xF4), 0.5)
        add_text(slide, x - 0.15, y0 + h + 0.12, 0.36, 0.12, str(pct), 5.8, color=GREY, align=PP_ALIGN.CENTER)
    for arr in [0, 0.5, 1.0, 1.5, 2.0]:
        y = y0 + h - h * arr / 2.0
        add_line(slide, x0, y, x0 + w, y, RGBColor(0xEC, 0xEF, 0xF4), 0.5)
        add_text(slide, x0 - 0.50, y - 0.06, 0.38, 0.11, f"{arr:.1f}", 5.8, color=GREY, align=PP_ALIGN.RIGHT)
    add_text(slide, x0 + w * 0.40, y0 + h + 0.38, 1.75, 0.15, "Probability (%)", 7, bold=True, color=GREY, align=PP_ALIGN.CENTER)
    add_text(slide, x0 - 0.70, y0 - 0.25, 0.8, 0.13, "ARR mEUR", 7, bold=True, color=GREY)

    label_offsets = {
        "Danantara": (0.12, -0.25),
        "Lembaga": (0.12, 0.05),
        "Krungthai": (0.40, 0.08),
        "PT Bank": (0.12, -0.05),
        "Temasek": (0.12, -0.20),
        "Hong Kong": (0.12, 0.02),
    }
    for row in rows:
        account = str(row[1])
        arr = parse_meur(row[7])
        prob = float(row[9] or 0)
        readiness = str(row[12] or "")
        push = int(row[10] or 0)
        x = x0 + w * prob / 100
        y = y0 + h - h * min(arr, 2.0) / 2.0
        if "Commercial approval" in readiness:
            color = RED
        elif "Silent" in readiness:
            color = AMBER
        else:
            color = BLUE
        r = 0.08 + min(push, 5) * 0.018 + min(arr, 2.0) * 0.035
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(x - r), inch(y - r), inch(r * 2), inch(r * 2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.fill.transparency = 12
        shape.line.color.rgb = WHITE
        shape.line.width = Pt(1)
        if arr >= 0.25:
            key = next((k for k in label_offsets if account.startswith(k)), "")
            dx, dy = label_offsets.get(key, (0.10, -0.12))
            add_text(slide, x + dx, y + dy, 1.75, 0.18, compact(account, 23), 6.4, bold=True, color=NAVY)
            add_text(slide, x + dx, y + dy + 0.14, 1.75, 0.12, f"{arr:.1f} mEUR | {int(prob)}%", 5.8, color=GREY)

    add_rect(slide, 9.08, 1.80, 3.15, 3.75, WHITE, line=LINE)
    add_text(slide, 9.30, 2.02, 2.7, 0.22, "What this adds", 11.5, bold=True, color=BLACK)
    bullets = [
        "Value and probability are visible at the same time.",
        "Commercial approval gaps are immediately separate from stale activity.",
        "The chart tells the director where to inspect first.",
    ]
    for i, bullet in enumerate(bullets):
        add_text(slide, 9.30, 2.45 + i * 0.46, 2.55, 0.28, "- " + bullet, 8.1, color=NAVY)
    legend = [("Approval gap", RED), ("Silent/stale", AMBER), ("Other watch", BLUE)]
    for i, (label, color) in enumerate(legend):
        add_rect(slide, 9.32, 4.15 + i * 0.26, 0.12, 0.12, color)
        add_text(slide, 9.52, 4.12 + i * 0.26, 1.4, 0.12, label, 6.5, color=GREY)


def _date_x(dt: date, start: date, end: date, x0: float, w: float) -> float:
    return x0 + w * ((dt - start).days / max(1, (end - start).days))


def add_decision_register_slide(prs: Presentation, slide_no: int) -> None:
    rows = load_rows("Action_Items")[1:]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "May decision register",
        "Action-slide direction: use a table unless the workbook has real dated milestones worth plotting.",
        "ACTION TABLE",
        slide_no,
    )
    add_text(
        slide,
        0.86,
        1.48,
        10.9,
        0.24,
        "The prior Gantt-style view overfit the data: every action currently shares the same generic 2026-06-30 due date.",
        10.2,
        bold=True,
        color=BLACK,
    )
    add_text(
        slide,
        0.86,
        1.80,
        10.8,
        0.18,
        "This slide should force operating decisions and evidence requirements, not imply false date precision.",
        8.2,
        color=GREY,
    )

    cols = [
        ("Priority", 0.80),
        ("Evidence gap / rule", 2.45),
        ("Current claim", 3.45),
        ("Decision to force", 3.05),
        ("Owner", 0.95),
        ("Due", 0.80),
    ]
    x0, y0 = 0.72, 2.28
    total_w = sum(w for _, w in cols)
    add_rect(slide, x0, y0, total_w, 0.34, BLUE)
    x = x0
    for label, cw in cols:
        add_text(slide, x + 0.05, y0 + 0.085, cw - 0.10, 0.12, label, 6.2, bold=True, color=WHITE)
        x += cw

    decision_map = {
        "zombie_arr": "Close, disqualify, or attach dated next step.",
        "approval_gap": "Submit approval or remove from Q2 confidence.",
        "coverage_gap": "Assign opener ownership for priority Tier-1 gaps.",
        "activity_drought": "Require logged next step or pull from forecast.",
    }
    for i, row in enumerate(rows[:4]):
        y = y0 + 0.34 + i * 0.72
        add_rect(slide, x0, y, total_w, 0.70, LIGHT_GREY if i % 2 == 0 else WHITE)
        priority, rule, claim, action, due, owner = row[1], row[2], row[3], row[4], row[5], row[6]
        due_text = str(due)[:10]
        priority_color = RED if str(priority) == "HIGH" else AMBER
        values = [
            str(priority),
            str(rule).replace("_", " "),
            compact(str(claim), 88),
            decision_map.get(str(rule), compact(str(action), 70)),
            str(owner).split()[0],
            due_text[5:] if due_text.startswith("2026-") else due_text,
        ]
        x = x0
        for j, ((_, cw), value) in enumerate(zip(cols, values, strict=True)):
            color = priority_color if j == 0 else NAVY
            add_text(slide, x + 0.05, y + 0.16, cw - 0.10, 0.34, value, 6.5, bold=(j in (0, 1)), color=color)
            x += cw

    add_rect(slide, 0.82, 5.68, 5.15, 0.55, LIGHT_GREY, line=LINE)
    add_text(slide, 1.03, 5.84, 4.7, 0.14, "Gantt eligibility rule: use timeline only when due dates are distinct and tied to named deal milestones.", 7.1, bold=True, color=NAVY)
    add_rect(slide, 6.25, 5.68, 5.42, 0.55, LIGHT_GREY, line=LINE)
    add_text(slide, 6.46, 5.84, 4.95, 0.14, "Factory gate: fail timeline charts when all action dates collapse to the same generic month-end date.", 7.1, bold=True, color=NAVY)


def add_table_slide(prs: Presentation, slide_no: int) -> None:
    rows = load_rows("Q2_Readiness")[1:9]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(
        slide,
        "Named deal table treatment",
        "Dense table direction: preserve row-level intelligence, keep headers consistent, and avoid stretched table-image artifacts.",
        "TABLE LANE",
        slide_no,
    )
    cols = [
        ("Account", 2.05),
        ("Opportunity", 2.00),
        ("Owner", 1.25),
        ("Stage", 1.15),
        ("Close", 0.72),
        ("ARR", 0.70),
        ("Forecast", 0.78),
        ("Readiness", 2.25),
    ]
    x0, y0 = 0.67, 1.65
    row_h = 0.47
    add_rect(slide, x0, y0, sum(w for _, w in cols), 0.36, PURPLE)
    x = x0
    for label, cw in cols:
        add_text(slide, x + 0.05, y0 + 0.09, cw - 0.10, 0.12, label, 6.5, bold=True, color=WHITE)
        x += cw
    for i, row in enumerate(rows):
        y = y0 + 0.36 + i * row_h
        add_rect(slide, x0, y, sum(w for _, w in cols), row_h, LIGHT_GREY if i % 2 == 0 else WHITE)
        values = [
            compact(str(row[1] or ""), 31),
            compact(str(row[2] or ""), 30),
            compact(str(row[3] or ""), 18),
            compact(str(row[5] or ""), 16),
            str(row[6] or "")[5:] if str(row[6] or "").startswith("2026-") else str(row[6] or ""),
            str(row[7] or ""),
            str(row[8] or ""),
            compact(str(row[12] or ""), 38),
        ]
        x = x0
        for j, ((_, cw), value) in enumerate(zip(cols, values, strict=True)):
            color = RED if j == 7 and ("approval" in value.lower() or "silent" in value.lower()) else NAVY
            add_text(slide, x + 0.05, y + 0.12, cw - 0.10, 0.18, value, 5.9, bold=(j in (0, 5)), color=color)
            x += cw
    add_line(slide, x0, y0, x0 + sum(w for _, w in cols), y0, NAVY, 1.2)
    add_text(slide, 0.70, 6.10, 10.7, 0.18, "Production implication: keep dense named-deal tables on the proven Excel -> table-image lane until a stable named think-cell table donor exists.", 7.7, color=GREY)


def build() -> Path:
    prs = Presentation(ROOT / "assets" / "LAND_template.pptx")
    delete_all_slides(prs)
    add_kpi_slide(prs, 1)
    add_waterfall_slide(prs, 2)
    add_scatter_slide(prs, 3)
    add_decision_register_slide(prs, 4)
    add_table_slide(prs, 5)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    print(OUT_PPTX)
    return OUT_PPTX


if __name__ == "__main__":
    build()
