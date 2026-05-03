#!/usr/bin/env python3
"""Build an APAC-only Q2 Forward Look v2 pilot deck."""

from __future__ import annotations

import argparse
import datetime as dt
import json
import shutil
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
PERIOD = "2026-Q2"
SLUG = "Jesper-Tyrer"
SNAPSHOT_DATE = dt.date(2026, 4, 30)
SOURCE_DECK = (
    Path.home()
    / "Downloads"
    / "May 2026 Meeting Spine Candidates"
    / f"{SLUG}-LAND-{PERIOD}-meeting-spine.pptx"
)
SOURCE_WORKBOOK = ROOT / "state" / PERIOD / SLUG / "factory" / "connected" / "connected_factory.xlsx"
TRENDS_PATH = ROOT / "state" / PERIOD / SLUG / "trends.json"
DEFAULT_OUTPUT = (
    Path.home()
    / "Downloads"
    / "May 2026 Meeting Spine Candidates"
    / f"{SLUG}-LAND-{PERIOD}-Q2-forward-look-v2-pilot.pptx"
)


NAVY = RGBColor(0x1A, 0x1D, 0x31)
MUTED = RGBColor(0x62, 0x68, 0x72)
RULE = RGBColor(0xD8, 0xDE, 0xEA)
PANEL = RGBColor(0xF7, 0xF8, 0xFB)
PURPLE = RGBColor(0x3B, 0x00, 0xA5)
TEAL = RGBColor(0x00, 0x8D, 0x8D)
RED = RGBColor(0xE8, 0x00, 0x46)
AMBER = RGBColor(0xD9, 0x89, 0x00)
BLUE = RGBColor(0x08, 0x3E, 0xA7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BLUE = RGBColor(0xEC, 0xF2, 0xFF)

SF_BASE = "https://simcorp.my.salesforce.com/lightning/r"
REPORT_LINKS = {
    "zombie": "https://simcorp.my.salesforce.com/lightning/r/Report/00OTb000008nijFMAQ/view",
    "coverage_gap": "https://simcorp.my.salesforce.com/lightning/r/Report/00OTb000008nirJMAQ/view",
    "no_activity_cfq": "https://simcorp.my.salesforce.com/lightning/r/Report/00OTb000008TZgvMAG/view",
    "no_activity_30": "https://simcorp.my.salesforce.com/lightning/r/Report/00OTb000008TaEnMAK/view",
    "pending_approval": "https://simcorp.my.salesforce.com/lightning/r/Report/00OTb000008mvx3MAA/view",
    "approval_candidates": "https://simcorp.my.salesforce.com/lightning/r/Report/00OTb000008ekp7MAA/view",
    "approval_global": "https://simcorp.my.salesforce.com/lightning/r/Report/00OTb000008fBEDMA2/view",
    "no_approval_flow": "https://simcorp.my.salesforce.com/lightning/r/Report/00OTb000008fAlBMAU/view",
}


@dataclass
class Deal:
    id: str
    account: str
    opportunity: str
    owner: str
    stage: str
    close_date: dt.date
    arr_meur: float
    forecast: str
    probability: float
    push: int
    readiness: str
    next_step: str

    @property
    def days_to_close(self) -> int:
        return (self.close_date - SNAPSHOT_DATE).days

    @property
    def risk(self) -> str:
        text = f"{self.readiness} {self.next_step}".casefold()
        if "commercial approval gap" in text or "silent 90d" in text or self.days_to_close <= 0:
            return "high"
        if "no next step" in text or "no activity" in text or self.push >= 3:
            return "watch"
        return "ok"


def _parse_arr(value: Any) -> float:
    text = str(value or "").replace("mEUR", "").replace("EUR", "").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def _parse_date(value: Any) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return dt.date.fromisoformat(str(value)[:10])


def _load_trends() -> dict[str, Any]:
    return json.loads(TRENDS_PATH.read_text(encoding="utf-8"))


def _kpi(trends: dict[str, Any], name: str) -> float:
    for item in trends.get("kpis", []):
        if item.get("name") == name:
            return float(item.get("value") or 0.0)
    return 0.0


def _meur(value: float, decimals: int = 1) -> str:
    return f"EUR {value / 1_000_000:.{decimals}f}M"


def _pct(value: float) -> str:
    return f"{round(value * 100):.0f}%"


def _original_intel() -> dict[str, Any]:
    wb = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
    ws = wb["Raw_Original_Intel"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(value or "").strip() for value in rows[0]]
    intel: dict[str, Any] = {}
    for row in rows[1:]:
        record = {headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))}
        key = str(record.get("Key") or "").strip()
        if key:
            intel[key] = record.get("Value")
    return intel


def _deals() -> list[Deal]:
    trends = _load_trends()
    id_by_name = {
        str(item.get("name") or ""): str(item.get("id") or "")
        for item in trends.get("top_deals_named", [])
    }
    output: list[Deal] = []
    for record in trends.get("q2_deal_readiness", []):
        account = str(record.get("account") or "").strip()
        opportunity = str(record.get("name") or "").strip()
        if not account and not opportunity:
            continue
        try:
            close_date = _parse_date(record.get("close_date"))
        except (TypeError, ValueError):
            continue
        output.append(
            Deal(
                id=id_by_name.get(opportunity, ""),
                account=account,
                opportunity=opportunity,
                owner=str(record.get("owner") or "").strip(),
                stage=str(record.get("stage") or "").strip(),
                close_date=close_date,
                arr_meur=float(record.get("arr_eur") or 0) / 1_000_000,
                forecast=str(record.get("forecast_category") or "").strip(),
                probability=float(record.get("probability") or 0),
                push=int(float(record.get("push_count") or 0)),
                readiness=str(record.get("readiness") or "").strip(),
                next_step=str(record.get("next_step") or "").strip(),
            )
        )
    wanted = [
        "Danantara",
        "LTH",
        "Krungthai",
        "Mandiri",
        "TEM - Additional",
        "HKMA",
    ]
    selected: list[Deal] = []
    for needle in wanted:
        match = next(
            (
                deal
                for deal in output
                if needle.casefold() in f"{deal.account} {deal.opportunity}".casefold()
            ),
            None,
        )
        if match:
            selected.append(match)
    return selected


def _remove_all_shapes(slide: Any) -> None:
    for shape in list(slide.shapes):
        shape.element.getparent().remove(shape.element)


def _text(slide: Any, left: float, top: float, width: float, height: float, text: str, size: float, *, bold: bool = False, color: RGBColor = NAVY, align: PP_ALIGN | None = None) -> Any:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    para = frame.paragraphs[0]
    para.text = text
    if align is not None:
        para.alignment = align
    run = para.runs[0] if para.runs else para.add_run()
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _link_text(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    url: str,
    size: float = 6.2,
) -> Any:
    box = _text(slide, left, top, width, height, text, size, bold=True, color=BLUE)
    para = box.text_frame.paragraphs[0]
    run = para.runs[0] if para.runs else para.add_run()
    run.hyperlink.address = url
    return box


def _source_strip(slide: Any, source: str, *, action_label: str | None = None, action_url: str | None = None) -> None:
    _line(slide, 0.58, 6.17, 11.75, 6.17, RULE, 0.5)
    _text(slide, 0.58, 6.23, 7.6, 0.18, source, 5.6, color=MUTED)
    if action_label and action_url:
        _link_text(slide, 9.10, 6.23, 2.6, 0.18, action_label, action_url, 5.8)


def _line(slide: Any, x1: float, y1: float, x2: float, y2: float, color: RGBColor = RULE, width: float = 0.75) -> Any:
    shape = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def _panel(slide: Any, left: float, top: float, width: float, height: float) -> Any:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = RULE
    shape.line.width = Pt(0.6)
    return shape


def _tile(
    slide: Any,
    left: float,
    top: float,
    width: float,
    title: str,
    value: str,
    note: str,
    accent: RGBColor,
    *,
    value_size: float = 16.0,
) -> None:
    _panel(slide, left, top, width, 0.98)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.035))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    _text(slide, left + 0.16, top + 0.15, width - 0.32, 0.18, title, 6.9, bold=True, color=MUTED)
    _text(slide, left + 0.16, top + 0.39, width - 0.32, 0.32, value, value_size, bold=True, color=NAVY)
    _text(slide, left + 0.16, top + 0.75, width - 0.32, 0.16, note, 6.3, color=MUTED)


def _set_cell(
    cell: Any,
    text: str,
    *,
    header: bool = False,
    fill: RGBColor | None = None,
    size: float = 6.2,
    bold: bool = False,
    url: str | None = None,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    cell.text = text
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    para = cell.text_frame.paragraphs[0]
    para.alignment = align
    para.space_after = Pt(0)
    run = para.runs[0] if para.runs else para.add_run()
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = header or bold
    run.font.color.rgb = WHITE if header else NAVY
    if url:
        run.hyperlink.address = url


def _risk_color(risk: str) -> RGBColor:
    return {"high": RED, "watch": AMBER, "ok": TEAL}.get(risk, BLUE)


def _deal_label(deal: Deal) -> str:
    if "Danantara" in deal.account:
        return "Danantara"
    if deal.account.startswith("Lembaga"):
        return "LTH"
    if "Krungthai" in deal.account:
        return "Krungthai"
    if "Mandiri" in deal.account:
        return "Mandiri"
    if "Temasek" in deal.account:
        return "Temasek"
    if "Hong Kong" in deal.account:
        return "HKMA"
    return deal.account[:12]


def _opportunity_url(deal: Deal) -> str | None:
    if not deal.id:
        return None
    return f"{SF_BASE}/Opportunity/{deal.id}/view"


def _may_ask(deal: Deal) -> str:
    label = _deal_label(deal)
    if label == "Danantara":
        return "Attach buyer proof + dated next step."
    if label == "LTH":
        return "Resolve August decision risk."
    if label == "Krungthai":
        return "Reset overdue close date."
    if label == "Mandiri":
        return "Submit/confirm Commercial Approval."
    return "Attach next-step evidence."


def _draw_inspection_map(slide: Any, deals: list[Deal], left: float, top: float, width: float, height: float) -> None:
    _panel(slide, left, top, width, height)
    _text(slide, left + 0.18, top + 0.16, width - 0.36, 0.22, "Q2 close inspection map", 10.2, bold=True)
    _text(
        slide,
        left + 0.18,
        top + 0.42,
        width - 0.36,
        0.24,
        "X = days to close from Apr 30; Y = unweighted ARR; color = readiness risk.",
        6.6,
        color=MUTED,
    )

    plot_l = left + 0.50
    plot_t = top + 0.88
    plot_w = width - 0.80
    plot_h = height - 1.35
    axis_bottom = plot_t + plot_h
    _line(slide, plot_l, axis_bottom, plot_l + plot_w, axis_bottom, RGBColor(0xA8, 0xB0, 0xC0), 0.8)
    _line(slide, plot_l, plot_t, plot_l, axis_bottom, RGBColor(0xA8, 0xB0, 0xC0), 0.8)
    for tick in [0, 15, 30, 45, 60]:
        x = plot_l + plot_w * (tick / 65)
        _line(slide, x, axis_bottom, x, axis_bottom + 0.04, RGBColor(0xA8, 0xB0, 0xC0), 0.5)
        _text(slide, x - 0.12, axis_bottom + 0.08, 0.32, 0.12, str(tick), 5.8, color=MUTED, align=PP_ALIGN.CENTER)
        if tick in {30, 60}:
            _line(slide, x, plot_t, x, axis_bottom, RGBColor(0xE8, 0xEB, 0xF1), 0.45)
    for value in [0.5, 1.0, 1.5, 2.0]:
        y = axis_bottom - plot_h * (value / 2.0)
        _line(slide, plot_l - 0.04, y, plot_l, y, RGBColor(0xA8, 0xB0, 0xC0), 0.5)
        _text(slide, plot_l - 0.42, y - 0.06, 0.32, 0.12, f"{value:.1f}", 5.8, color=MUTED, align=PP_ALIGN.RIGHT)
        _line(slide, plot_l, y, plot_l + plot_w, y, RGBColor(0xEA, 0xED, 0xF2), 0.35)
    _text(slide, plot_l + plot_w * 0.34, axis_bottom + 0.31, 1.6, 0.14, "days to close", 6.2, color=MUTED, align=PP_ALIGN.CENTER)
    _text(slide, plot_l - 0.46, plot_t - 0.22, 0.6, 0.14, "ARR mEUR", 6.2, color=MUTED)

    label_offsets = {
        "Danantara": (0.08, -0.20),
        "LTH": (0.05, 0.10),
        "Krungthai": (0.05, -0.15),
        "Mandiri": (-0.72, -0.17),
        "Temasek": (0.05, -0.04),
        "HKMA": (0.05, 0.12),
    }
    for deal in deals:
        x = plot_l + plot_w * (max(0, min(65, deal.days_to_close)) / 65)
        y = axis_bottom - plot_h * (max(0.0, min(2.0, deal.arr_meur)) / 2.0)
        diameter = 0.14 + min(0.18, max(0.02, deal.arr_meur * 0.055))
        bubble = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x - diameter / 2),
            Inches(y - diameter / 2),
            Inches(diameter),
            Inches(diameter),
        )
        bubble.fill.solid()
        bubble.fill.fore_color.rgb = _risk_color(deal.risk)
        bubble.line.color.rgb = WHITE
        bubble.line.width = Pt(0.8)
        label = _deal_label(deal)
        dx, dy = label_offsets.get(label, (0.04, -0.12))
        _text(slide, x + dx, y + dy, 0.75, 0.13, label, 5.9, bold=True, color=NAVY)

    legend_y = top + height - 0.29
    for idx, (label, color) in enumerate([("High", RED), ("Watch", AMBER), ("OK", TEAL)]):
        lx = left + 0.20 + idx * 0.78
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(lx), Inches(legend_y), Inches(0.08), Inches(0.08))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()
        _text(slide, lx + 0.11, legend_y - 0.01, 0.45, 0.11, label, 5.9, color=MUTED)


def _draw_deal_table(slide: Any, deals: list[Deal], left: float, top: float, width: float, height: float) -> None:
    rows = len(deals) + 1
    cols = 6
    table_shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = table_shape.table
    widths = [1.22, 1.07, 0.72, 0.70, 0.76, 2.22]
    for idx, col_width in enumerate(widths):
        table.columns[idx].width = Inches(col_width)
    headers = ["Deal", "Owner", "ARR", "Close", "Fcst", "May ask"]
    for col, header in enumerate(headers):
        _set_cell(table.cell(0, col), header, header=True, fill=PURPLE, size=6.5)
    for row_idx, deal in enumerate(deals, start=1):
        values = [
            _deal_label(deal),
            deal.owner.split()[0] if deal.owner else "",
            f"{deal.arr_meur:.1f}",
            deal.close_date.strftime("%b %-d") if hasattr(deal.close_date, "strftime") else str(deal.close_date),
            deal.forecast[:4],
            _may_ask(deal),
        ]
        for col_idx, value in enumerate(values):
            _set_cell(
                table.cell(row_idx, col_idx),
                value,
                fill=RGBColor(0xF2, 0xF4, 0xF8) if row_idx % 2 == 0 else WHITE,
                size=6.0 if col_idx == 5 else 6.3,
                bold=col_idx == 0,
                url=_opportunity_url(deal) if col_idx == 0 else None,
            )


def _build_q1_slide(prs: Presentation, trends: dict[str, Any], original: dict[str, Any]) -> None:
    slide = prs.slides[2]
    _remove_all_shapes(slide)
    _text(slide, 0.58, 0.42, 9.2, 0.48, "Q1 accountability and May reset", 24, bold=True)
    _text(
        slide,
        0.58,
        1.02,
        10.9,
        0.24,
        "Keep the original APAC spine visible, but separate historical Land accountability from current May Q2 ARR.",
        8.4,
        color=MUTED,
    )
    _line(slide, 0.58, 1.36, 11.75, 1.36, RULE, 0.8)
    _tile(slide, 0.58, 1.62, 2.55, "Q1 opened baseline", _meur(float(original["q1_promised_opened_arr_eur"])), "Original APAC commentary.", BLUE)
    _tile(slide, 3.40, 1.62, 2.55, "Q1 Land won", f"{int(original['q1_land_wins'])} / {_meur(float(original['q1_land_wins_arr_eur']))}", "Original sidecar.", TEAL, value_size=14.0)
    _tile(slide, 6.22, 1.62, 2.55, "Q1 Land lost", f"{int(original['q1_land_lost'])} / {_meur(float(original['q1_land_lost_arr_eur']))}", "Loss accountability.", RED, value_size=14.0)
    _tile(slide, 9.04, 1.62, 2.55, "Q1 slipped", f"{int(original['q1_slipped_deals'])} / {_meur(float(original['q1_slipped_arr_eur']))}", "Use as coaching lens.", AMBER, value_size=14.0)

    rows = [
        ("Current Q2 closeable", _meur(_kpi(trends, "total_pipeline_arr")), "Land+Expand unweighted ARR; Renewal ACV excluded."),
        ("QTD L+E result", "Won 1.6M / Lost 6.0M", "Loss pressure is already larger than won ARR."),
        ("Original Q2 baseline", f"{int(original['q2_original_deals'])} deals / {_meur(float(original['q2_original_arr_eur']))}", "The earlier pack flagged no recent activity."),
        ("May review use", "Confirm, reset, or remove", "Every Q2 deal needs dated next step and owner evidence."),
    ]
    table_shape = slide.shapes.add_table(len(rows) + 1, 3, Inches(0.58), Inches(3.05), Inches(11.2), Inches(2.22))
    table = table_shape.table
    for idx, width in enumerate([2.15, 2.60, 6.45]):
        table.columns[idx].width = Inches(width)
    for idx, header in enumerate(["Lens", "Current fact", "May implication"]):
        _set_cell(table.cell(0, idx), header, header=True, fill=BLUE, size=7.0)
    for row_idx, row in enumerate(rows, 1):
        for col_idx, value in enumerate(row):
            _set_cell(
                table.cell(row_idx, col_idx),
                value,
                fill=RGBColor(0xF4, 0xF6, 0xFA) if row_idx % 2 == 0 else WHITE,
                size=6.6,
                bold=col_idx == 0,
            )
    _source_strip(
        slide,
        "Source: connected_factory.xlsx::Raw_Original_Intel + trends.json. ARR = Land+Expand; Renewal ACV separate.",
    )


def _build_forecast_slide(prs: Presentation, trends: dict[str, Any], deals: list[Deal], original: dict[str, Any]) -> None:
    slide = prs.slides[3]
    _remove_all_shapes(slide)
    _text(slide, 0.58, 0.42, 8.2, 0.48, "May forecast quality", 24, bold=True)
    _text(slide, 0.58, 1.02, 10.9, 0.24, "Snapshot as of 2026-04-30. Omitted and out-of-quarter pipeline are context, not coverage.", 8.4, color=MUTED)
    _line(slide, 0.58, 1.36, 11.75, 1.36, RULE, 0.8)
    total = _kpi(trends, "total_pipeline_arr")
    commit = sum(d.arr_meur for d in deals if d.forecast == "Commit") * 1_000_000
    pipeline = sum(d.arr_meur for d in deals if d.forecast == "Pipeline") * 1_000_000
    commit_count = sum(1 for d in deals if d.forecast == "Commit")
    pipeline_count = sum(1 for d in deals if d.forecast == "Pipeline")
    _tile(slide, 0.58, 1.58, 2.62, "Closeable L+E ARR", _meur(total), "Unweighted current-quarter ARR.", BLUE)
    _tile(slide, 3.48, 1.58, 2.62, "Commit ARR", _meur(commit), f"11 opps; {commit_count} material named.", TEAL)
    _tile(slide, 6.38, 1.58, 2.62, "Pipeline ARR", _meur(pipeline), f"8 opps; {pipeline_count} material named.", PURPLE)
    _tile(slide, 9.28, 1.58, 2.62, "Best Case ARR", "EUR 0.0M", "No material cushion.", RED)

    rows = [
        ("Commit proof", "LTH, Krungthai, Temasek, HKMA", "Attach dated next step / decision evidence."),
        ("Pipeline upside", "Danantara, Mandiri", "Danantara needs buyer proof; Mandiri needs approval gate."),
        ("Original forecast spine", f"{int(original['forecast_mix_deals'])} PI deals / {_meur(float(original['forecast_mix_weighted_arr_eur']))}", f"Commit was {_pct(float(original['commit_share_original']))}; keep weighted/unweighted labeled."),
        ("Risk to call", "No Best Case cushion", "Use deal inspection, not forecast optimism."),
    ]
    table_shape = slide.shapes.add_table(len(rows) + 1, 3, Inches(0.58), Inches(3.06), Inches(11.2), Inches(2.18))
    table = table_shape.table
    for idx, width in enumerate([2.15, 3.85, 5.20]):
        table.columns[idx].width = Inches(width)
    for idx, header in enumerate(["Issue", "Named evidence", "May action"]):
        _set_cell(table.cell(0, idx), header, header=True, fill=BLUE, size=7.0)
    for row_idx, row in enumerate(rows, 1):
        for col_idx, value in enumerate(row):
            _set_cell(
                table.cell(row_idx, col_idx),
                value,
                fill=RGBColor(0xF4, 0xF6, 0xFA) if row_idx % 2 == 0 else WHITE,
                size=6.45,
                bold=col_idx == 0,
            )
    _source_strip(
        slide,
        "Source: trends.json::kpis + q2_deal_readiness; original PI context from Raw_Original_Intel.",
        action_label="Open no-activity CFQ report",
        action_url=REPORT_LINKS["no_activity_cfq"],
    )


def _build_q2_slide(prs: Presentation, deals: list[Deal]) -> None:
    slide = prs.slides[4]
    _remove_all_shapes(slide)
    _text(slide, 0.58, 0.34, 8.6, 0.44, "Q2 forward look: May deal inspection", 24, bold=True)
    _text(
        slide,
        0.58,
        0.94,
        10.9,
        0.25,
        "Salesforce snapshot as of 2026-04-30. Land+Expand unweighted ARR only; Renewal ACV remains separate.",
        8.4,
        color=MUTED,
    )
    _line(slide, 0.58, 1.30, 11.75, 1.30, RULE, 0.8)

    _text(slide, 0.58, 1.49, 3.0, 0.22, "Operating read", 10.2, bold=True)
    _text(
        slide,
        0.58,
        1.78,
        4.45,
        0.86,
        "Q2 close call is not a summary-chart problem: Commit depends on LTH, Krungthai, Temasek, and HKMA evidence; Pipeline upside is Danantara and Mandiri.",
        9.0,
        color=NAVY,
    )
    _text(
        slide,
        0.58,
        2.73,
        4.45,
        0.72,
        "Original APAC signal stays live: the prior pack flagged 6 Q2 deals / EUR 5.0M with zero recent activity. Current May review should force dated next steps and approval status.",
        8.1,
        color=MUTED,
    )

    _draw_inspection_map(slide, deals, 5.18, 1.48, 6.14, 3.03)
    _text(slide, 0.58, 3.84, 4.2, 0.22, "Named deal actions", 10.2, bold=True)
    _draw_deal_table(slide, deals, 0.58, 4.14, 10.75, 1.54)
    _text(
        slide,
        0.58,
        5.76,
        10.6,
        0.18,
        "Click deal names for Salesforce records. Factory note: table remains Excel-linked; inspection map becomes the next stable think-cell chart lane after approval.",
        6.2,
        color=MUTED,
    )
    _source_strip(
        slide,
        "Source: trends.json::q2_deal_readiness + connected_factory.xlsx::Raw_Current_Q2_Readiness.",
        action_label="Open no-activity CFQ report",
        action_url=REPORT_LINKS["no_activity_cfq"],
    )


def _build_renewal_slide(prs: Presentation, trends: dict[str, Any], original: dict[str, Any]) -> None:
    slide = prs.slides[6]
    _remove_all_shapes(slide)
    renewals = trends.get("fy26_renewals", [])
    current_total = sum(float(row.get("acv_eur") or 0) for row in renewals)
    q2_total = sum(
        float(row.get("acv_eur") or 0)
        for row in renewals
        if str(row.get("close_date") or "") <= "2026-06-30"
    )
    _text(slide, 0.58, 0.42, 8.2, 0.48, "FY26 renewal watchlist", 24, bold=True)
    _text(slide, 0.58, 1.02, 10.9, 0.24, "ACV only. Land+Expand ARR is intentionally excluded from this slide.", 8.4, color=MUTED)
    _line(slide, 0.58, 1.36, 11.75, 1.36, RULE, 0.8)
    _tile(slide, 0.58, 1.58, 2.75, "Original APAC baseline", _meur(float(original["fy26_renewals_acv_eur_original"])), "3 FY26 renewals from Apr 20 ETL.", TEAL)
    _tile(slide, 3.60, 1.58, 2.75, "Current FY26 open ACV", _meur(current_total), f"{len(renewals)} open renewal rows.", BLUE)
    _tile(slide, 6.62, 1.58, 2.75, "Current Q2 ACV", _meur(q2_total), "Fullerton only in Q2.", AMBER)
    _tile(slide, 9.64, 1.58, 2.18, "Basis status", "Reconcile", "Do not blend baselines.", RED, value_size=14.0)

    table_shape = slide.shapes.add_table(len(renewals) + 1, 8, Inches(0.58), Inches(3.04), Inches(11.2), Inches(1.72))
    table = table_shape.table
    widths = [0.45, 1.02, 2.25, 2.36, 1.42, 1.20, 1.00, 1.50]
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
    for idx, header in enumerate(["#", "Close", "Account", "Opportunity", "Owner", "Stage", "ACV", "Risk"]):
        _set_cell(table.cell(0, idx), header, header=True, fill=BLUE, size=6.7)
    for row_idx, row in enumerate(renewals, 1):
        values = [
            str(row_idx),
            str(row.get("close_date") or ""),
            str(row.get("account") or ""),
            str(row.get("name") or ""),
            str(row.get("owner") or ""),
            str(row.get("stage") or ""),
            _meur(float(row.get("acv_eur") or 0)),
            str(row.get("risk_level") or ""),
        ]
        for col_idx, value in enumerate(values):
            _set_cell(
                table.cell(row_idx, col_idx),
                value,
                fill=RGBColor(0xF4, 0xF6, 0xFA) if row_idx % 2 == 0 else WHITE,
                size=5.65 if col_idx in {2, 3} else 5.9,
                bold=col_idx == 2,
            )
    _text(
        slide,
        0.58,
        5.02,
        10.8,
        0.36,
        "Monthly review use: Fullerton is the only current-quarter renewal action; BSP and Temasek are Q3 FY26 renewal ACV. Keep the Apr 20 EUR 33.5M baseline as historical context until the source basis is reconciled.",
        8.0,
        color=NAVY,
    )
    _source_strip(
        slide,
        "Source: trends.json::fy26_renewals + connected_factory.xlsx::Raw_Current_FY26_Renewals. Renewal values are ACV.",
    )


def _build_action_slide(prs: Presentation, trends: dict[str, Any]) -> None:
    slide = prs.slides[13]
    _remove_all_shapes(slide)
    action_items = {item.get("rule_id"): item for item in trends.get("action_items", [])}
    rows = [
        (
            "Q2 activity drought",
            "Danantara, LTH, Krungthai, Mandiri, BOCI, Temasek, HKMA",
            action_items.get("activity_drought", {}).get("claim", "26 Q2 opps with no activity in 30d"),
            "Every named owner logs next step/task before forecast review.",
            "CFQ no-activity report",
            REPORT_LINKS["no_activity_cfq"],
        ),
        (
            "Commercial approval",
            "Mandiri current-quarter gap; all-open gaps need report support",
            "Q2: Mandiri EUR 0.5M. All-open: 6 gaps / EUR 10.2M exposure.",
            "Submit/confirm Stage 3 approval gate; do not mix Q2 and all-open claims.",
            "Pending approval report",
            REPORT_LINKS["pending_approval"],
        ),
        (
            "Zombie pipeline",
            "EPF, SSO, Coolabah, Temasek, Mainstream BPO",
            action_items.get("zombie_arr", {}).get("claim", "Stale pipeline needs cleanup"),
            "Rep-by-rep close/disqualify review by EOM.",
            "Zombie report",
            REPORT_LINKS["zombie"],
        ),
        (
            "Tier-1 coverage",
            "Japan Trustee Services Bank; Janus Henderson Singapore; Asset Management One",
            action_items.get("coverage_gap", {}).get("claim", "194 Tier-1 accounts without open opps"),
            "Assign opener ownership and create/confirm coverage plan.",
            "Coverage gap report",
            REPORT_LINKS["coverage_gap"],
        ),
    ]
    _text(slide, 0.58, 0.42, 8.8, 0.48, "May action register", 24, bold=True)
    _text(slide, 0.58, 1.02, 10.9, 0.24, "Named operating queues with direct Salesforce action links.", 8.4, color=MUTED)
    _line(slide, 0.58, 1.36, 11.75, 1.36, RULE, 0.8)
    table_shape = slide.shapes.add_table(len(rows) + 1, 5, Inches(0.58), Inches(1.66), Inches(11.2), Inches(3.92))
    table = table_shape.table
    for idx, width in enumerate([1.55, 3.15, 3.00, 2.35, 1.15]):
        table.columns[idx].width = Inches(width)
    table.rows[0].height = Inches(0.38)
    for row_idx in range(1, len(rows) + 1):
        table.rows[row_idx].height = Inches(0.78)
    for idx, header in enumerate(["Queue", "Named evidence", "Fact", "May output", "Action"]):
        _set_cell(table.cell(0, idx), header, header=True, fill=BLUE, size=6.9)
    for row_idx, row in enumerate(rows, 1):
        for col_idx, value in enumerate(row[:5]):
            _set_cell(
                table.cell(row_idx, col_idx),
                value,
                fill=RGBColor(0xF4, 0xF6, 0xFA) if row_idx % 2 == 0 else WHITE,
                size=6.0 if col_idx in {1, 2, 3} else 6.15,
                bold=col_idx == 0,
                url=row[5] if col_idx == 4 else None,
            )
    _source_strip(
        slide,
        "Source: trends.json::action_items plus verified Salesforce reports. Counts remain horizon-labeled.",
        action_label="Open all report links from table",
        action_url=REPORT_LINKS["no_activity_cfq"],
    )


def build(output: Path) -> dict[str, Any]:
    deals = _deals()
    if not deals:
        raise SystemExit("no APAC Q2 readiness deals found")
    trends = _load_trends()
    original = _original_intel()
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE_DECK, output)
    prs = Presentation(output)
    _build_q1_slide(prs, trends, original)
    _build_forecast_slide(prs, trends, deals, original)
    _build_q2_slide(prs, deals)
    _build_renewal_slide(prs, trends, original)
    _build_action_slide(prs, trends)
    prs.save(output)
    payload = {
        "schema": "apac-forward-look-v2-pilot/v2",
        "source_deck": str(SOURCE_DECK),
        "source_workbook": str(SOURCE_WORKBOOK),
        "source_trends": str(TRENDS_PATH),
        "output_deck": str(output),
        "rebuilt_slides": [3, 4, 5, 7, 14],
        "deals": [asdict(deal) for deal in deals],
        "report_links": REPORT_LINKS,
    }
    trace = output.with_suffix(".json")
    trace.write_text(json.dumps(payload, indent=2, default=str) + "\n", encoding="utf-8")
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    payload = build(args.output)
    print(f"output={payload['output_deck']}")
    print(f"trace={Path(payload['output_deck']).with_suffix('.json')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
