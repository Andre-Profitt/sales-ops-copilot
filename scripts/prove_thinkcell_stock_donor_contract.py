#!/usr/bin/env python3
"""Build and prove stock think-cell donor contracts.

This is for chart families that are present in the installed think-cell POTX
corpus but are not yet authored as named objects in the SimCorp LAND seed.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import shutil
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Callable
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

from ppttc_template import _patch_cfb_stream, template_named_elements
from prove_thinkcell_native_chart_contract import (
    Check,
    NativeChartContract,
    _bridge_ppttc,
    _check_bound_package,
    _proof_render,
    _render_deck,
)
from thinkcell_cfb import CfbStreamEdit, replace_cfb_stream_data


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_PERIOD = "2026-Q2"
DEFAULT_DIRECTOR_SLUG = "Jesper-Tyrer"
DEFAULT_CONTRACT = "QTR04_DealRisk_Scatter"
SCATTER_DONOR = Path(
    "/Library/Application Support/Microsoft/think-cell/templates/think-cell Charts/Scatter, Bubble/Scatter, Bubble.potx"
)
GANTT_DONOR = Path(
    "/Library/Application Support/Microsoft/think-cell/templates/think-cell Charts/Timeline, Gantt/Timeline, Gantt.potx"
)
PPTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
POTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

ET.register_namespace("p", P_NS)
ET.register_namespace("a", A_NS)
ET.register_namespace("r", R_NS)
ET.register_namespace("", PKG_REL_NS)


Cell = dict[str, Any] | None
Table = list[list[Cell]]
DataBuilder = Callable[[str, str], tuple[Table, list[str]]]


@dataclass(frozen=True)
class StockDonorContract:
    contract: str
    target_name: str
    donor_template: Path
    ole_parts: tuple[str, ...]
    slide: int
    data_builder: DataBuilder
    required_terms: tuple[str, ...]
    note: str
    bound_required_terms: tuple[str, ...] = ()
    bound_dynamic_term_count: int = 0


@dataclass
class StockProofPaths:
    work_dir: Path
    seed: Path
    ppttc: Path
    bound: Path
    render_dir: Path


def _cell(value: Any) -> Cell:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return {"number": float(value)}
    return {"string": str(value)}


def _date_cell(value: Any) -> Cell:
    if value is None:
        return None
    text = str(value)
    if " " in text:
        text = text.split(" ", 1)[0]
    return {"date": text}


def _parse_meur(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value) / 1_000_000 if abs(float(value)) > 1000 else float(value)
    text = str(value)
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text.replace(",", ""))
    if not match:
        return 0.0
    return float(match.group(0))


def _connected_workbook(period: str, director_slug: str) -> Path:
    return ROOT / "state" / period / director_slug / "factory" / "connected" / "connected_factory_table_images.xlsx"


def _rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    out: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        out.append({str(headers[index]): value for index, value in enumerate(row) if index < len(headers)})
    return out


def _deal_risk_table(period: str, director_slug: str) -> tuple[Table, list[str]]:
    path = _connected_workbook(period, director_slug)
    source_rows = _rows(path, "Raw_Current_Q2_Readiness")[:8]
    table: Table = [[_cell("Deal"), _cell("Probability (%)"), _cell("ARR (mEUR)"), _cell("Push count"), _cell("Readiness")]]
    for index, row in enumerate(source_rows, 1):
        label = str(row.get("Account") or row.get("Opportunity") or f"Deal {index}")[:28]
        table.append(
            [
                _cell(label),
                _cell(float(row.get("Prob") or 0)),
                _cell(_parse_meur(row.get("ARR"))),
                _cell(max(1.0, float(row.get("Push") or 0) + 1.0)),
                _cell(str(row.get("Readiness") or "review")[:24]),
            ]
        )
    terms = [str(row.get("Account") or row.get("Opportunity") or "")[:28] for row in source_rows[:2]]
    return table, terms


def _renewal_gantt_table(period: str, director_slug: str, *, q2_only: bool) -> tuple[Table, list[str]]:
    path = _connected_workbook(period, director_slug)
    source_rows = _rows(path, "Raw_Current_FY26_Renewals")
    if q2_only:
        source_rows = [row for row in source_rows if str(row.get("Close Date") or "") <= "2026-06-30"]
    source_rows = source_rows[:8]
    table: Table = [
        [_cell("Renewal"), _cell("Owner"), _cell("Stage"), _cell("Start"), _cell("Close Date"), _cell("ACV")]
    ]
    for row in source_rows:
        close_date = row.get("Close Date")
        table.append(
            [
                _cell(row.get("Account")),
                _cell(row.get("Owner")),
                _cell(row.get("Stage")),
                _date_cell("2026-04-01"),
                _date_cell(close_date),
                _cell(row.get("ACV")),
            ]
        )
    terms = [str(row.get("Account")) for row in source_rows[:2] if row.get("Account")]
    return table, terms


def _fy26_renewal_gantt_table(period: str, director_slug: str) -> tuple[Table, list[str]]:
    return _renewal_gantt_table(period, director_slug, q2_only=False)


def _q2_renewal_gantt_table(period: str, director_slug: str) -> tuple[Table, list[str]]:
    return _renewal_gantt_table(period, director_slug, q2_only=True)


CONTRACTS: dict[str, StockDonorContract] = {
    DEFAULT_CONTRACT: StockDonorContract(
        contract=DEFAULT_CONTRACT,
        target_name=DEFAULT_CONTRACT,
        donor_template=SCATTER_DONOR,
        ole_parts=("ppt/embeddings/oleObject3.bin",),
        slide=1,
        data_builder=_deal_risk_table,
        required_terms=("ARR (mEUR)", "Probability (%)"),
        bound_required_terms=("ARR (mEUR)", "Probability (%)"),
        bound_dynamic_term_count=0,
        note="Stock scatter donor slide 1 patched from empty m_strName.",
    ),
    "QTR05_FY26RenewalTimeline_Gantt": StockDonorContract(
        contract="QTR05_FY26RenewalTimeline_Gantt",
        target_name="QTR05_FY26RenewalTimeline_Gantt",
        donor_template=GANTT_DONOR,
        ole_parts=("ppt/embeddings/oleObject3.bin",),
        slide=1,
        data_builder=_fy26_renewal_gantt_table,
        required_terms=("Renewal", "Close Date", "ACV"),
        bound_required_terms=("Renewal",),
        bound_dynamic_term_count=2,
        note="Stock Gantt donor slide 1 patched from empty m_strName values.",
    ),
    "QTR06_Q2RenewalTimeline_Gantt": StockDonorContract(
        contract="QTR06_Q2RenewalTimeline_Gantt",
        target_name="QTR06_Q2RenewalTimeline_Gantt",
        donor_template=GANTT_DONOR,
        ole_parts=("ppt/embeddings/oleObject3.bin",),
        slide=1,
        data_builder=_q2_renewal_gantt_table,
        required_terms=("Renewal", "Close Date", "ACV"),
        bound_required_terms=("Renewal",),
        bound_dynamic_term_count=2,
        note="Stock Gantt donor slide 1 patched from empty m_strName values.",
    ),
}


def _work_paths(period: str, contract: str) -> StockProofPaths:
    work_dir = ROOT / "state" / "thinkcell_bridge" / "build_scaffold" / period / "work" / contract
    return StockProofPaths(
        work_dir=work_dir,
        seed=work_dir / f"{contract}-stock-donor-seed.pptx",
        ppttc=work_dir / f"{contract}-stock-donor-{period}.ppttc",
        bound=work_dir / f"{contract}-stock-donor-{period}-bound.pptx",
        render_dir=work_dir / "rendered_stock_donor",
    )


def _patch_empty_names(data: bytes, target_name: str) -> tuple[bytes, int]:
    replacement = f"<m_strName>{target_name}</m_strName>".encode("utf-8")
    try:
        result = replace_cfb_stream_data(
            data,
            (CfbStreamEdit(name="think-cellXML", replacements=((b"<m_strName></m_strName>", replacement),)),),
        )
        if result.replacements_made:
            return result.data, result.replacements_made
    except ValueError:
        pass
    patched = _patch_cfb_stream(
        data,
        "think-cellXML",
        {rb"<m_strName>\s*</m_strName>": replacement},
        regex=True,
    )
    return patched, data.count(b"<m_strName></m_strName>")


def _q(tag: str, ns: str = P_NS) -> str:
    return f"{{{ns}}}{tag}"


def _textbox(
    *,
    shape_id: int,
    name: str,
    x: int,
    y: int,
    cx: int,
    cy: int,
    text: str,
    size: int,
    bold: bool = False,
    color: str = "1A1D31",
) -> ET.Element:
    shape = ET.Element(_q("sp"))
    nv = ET.SubElement(shape, _q("nvSpPr"))
    ET.SubElement(nv, _q("cNvPr")).attrib.update({"id": str(shape_id), "name": name})
    ET.SubElement(nv, _q("cNvSpPr")).attrib.update({"txBox": "1"})
    ET.SubElement(nv, _q("nvPr"))
    sp_pr = ET.SubElement(shape, _q("spPr"))
    xfrm = ET.SubElement(sp_pr, _q("xfrm", A_NS))
    ET.SubElement(xfrm, _q("off", A_NS)).attrib.update({"x": str(x), "y": str(y)})
    ET.SubElement(xfrm, _q("ext", A_NS)).attrib.update({"cx": str(cx), "cy": str(cy)})
    geom = ET.SubElement(sp_pr, _q("prstGeom", A_NS))
    geom.attrib["prst"] = "rect"
    ET.SubElement(geom, _q("avLst", A_NS))
    ET.SubElement(sp_pr, _q("noFill", A_NS))
    line = ET.SubElement(sp_pr, _q("ln", A_NS))
    ET.SubElement(line, _q("noFill", A_NS))
    body = ET.SubElement(shape, _q("txBody"))
    body_pr = ET.SubElement(body, _q("bodyPr", A_NS))
    body_pr.attrib.update({"wrap": "square", "rtlCol": "0"})
    ET.SubElement(body_pr, _q("spAutoFit", A_NS))
    ET.SubElement(body, _q("lstStyle", A_NS))
    para = ET.SubElement(body, _q("p", A_NS))
    run = ET.SubElement(para, _q("r", A_NS))
    rpr = ET.SubElement(run, _q("rPr", A_NS))
    rpr.attrib.update({"lang": "en-US", "sz": str(size)})
    if bold:
        rpr.attrib["b"] = "1"
    fill = ET.SubElement(rpr, _q("solidFill", A_NS))
    ET.SubElement(fill, _q("srgbClr", A_NS)).attrib["val"] = color
    ET.SubElement(rpr, _q("latin", A_NS)).attrib["typeface"] = "Aptos"
    ET.SubElement(run, _q("t", A_NS)).text = text
    ET.SubElement(para, _q("endParaRPr", A_NS)).attrib.update({"lang": "en-US", "sz": str(size)})
    return shape


def _clean_qtr04_seed(path: Path) -> dict[str, Any]:
    """Trim the stock Scatter/Bubble POTX to one clean SimCorp pilot slide.

    The official donor carries instructional comments, placeholder titles, and
    a second unrelated bubble example. Those prove donor availability, but they
    are not acceptable insertion-pilot evidence. This keeps the named think-cell
    chart object and removes the stock presentation residue before binding.
    """
    with ZipFile(path) as zin:
        parts = {item.filename: zin.read(item.filename) for item in zin.infolist()}

    removed_parts = [
        "ppt/slides/slide2.xml",
        "ppt/slides/_rels/slide2.xml.rels",
        "ppt/notesSlides/notesSlide2.xml",
        "ppt/notesSlides/_rels/notesSlide2.xml.rels",
    ]
    for part in removed_parts:
        parts.pop(part, None)

    slide_root = ET.fromstring(parts["ppt/slides/slide1.xml"])
    sp_tree = slide_root.find(f".//{{{P_NS}}}spTree")
    removed_shapes = 0
    if sp_tree is not None:
        for child in list(sp_tree):
            if child.tag in {_q("nvGrpSpPr"), _q("grpSpPr")}:
                continue
            if child.tag == _q("graphicFrame"):
                cnv = child.find(f".//{{{P_NS}}}cNvPr")
                name = cnv.attrib.get("name", "") if cnv is not None else ""
                xfrm = child.find(f"{{{P_NS}}}xfrm")
                if name.startswith("Chart") and xfrm is not None:
                    off = xfrm.find(f"{{{A_NS}}}off")
                    ext = xfrm.find(f"{{{A_NS}}}ext")
                    if off is not None:
                        off.attrib.update({"x": "548640", "y": "1371600"})
                    if ext is not None:
                        ext.attrib.update({"cx": "10058400", "cy": "4023360"})
                continue
            sp_tree.remove(child)
            removed_shapes += 1
        sp_tree.append(
            _textbox(
                shape_id=9001,
                name="SimCorp QTR04 title",
                x=548640,
                y=365760,
                cx=10058400,
                cy=365760,
                text="Q2 deal-risk scatter: probability vs ARR",
                size=2400,
                bold=True,
            )
        )
        sp_tree.append(
            _textbox(
                shape_id=9002,
                name="SimCorp QTR04 basis",
                x=548640,
                y=822960,
                cx=10058400,
                cy=228600,
                text=(
                    "Land+Expand ARR is unweighted unless explicitly labeled weighted; "
                    "Type IN ('Land','Expand'); Renewal ACV excluded."
                ),
                size=900,
                color="626872",
            )
        )
        sp_tree.append(
            _textbox(
                shape_id=9003,
                name="SimCorp QTR04 axis note",
                x=548640,
                y=5852160,
                cx=10058400,
                cy=182880,
                text="X = Probability (%); Y = ARR (mEUR); bubble size = push pressure. Source: connected_factory_table_images.xlsx::Raw_Current_Q2_Readiness.",
                size=720,
                color="626872",
            )
        )
    slide_xml = ET.tostring(slide_root, encoding="utf-8", xml_declaration=True)
    if b'xmlns:v="' not in slide_xml:
        slide_xml = slide_xml.replace(b"<p:sld ", b'<p:sld xmlns:v="urn:schemas-microsoft-com:vml" ', 1)
    parts["ppt/slides/slide1.xml"] = slide_xml

    slide_rels = ET.fromstring(parts["ppt/slides/_rels/slide1.xml.rels"])
    removed_note_rels = 0
    for rel in list(slide_rels):
        if rel.attrib.get("Type", "").endswith("/notesSlide"):
            slide_rels.remove(rel)
            removed_note_rels += 1
    parts["ppt/slides/_rels/slide1.xml.rels"] = ET.tostring(slide_rels, encoding="utf-8", xml_declaration=True)

    pres_root = ET.fromstring(parts["ppt/presentation.xml"])
    sld_id_lst = pres_root.find(f"{{{P_NS}}}sldIdLst")
    removed_slide_ids = 0
    if sld_id_lst is not None:
        for idx, child in enumerate(list(sld_id_lst)):
            if idx > 0:
                sld_id_lst.remove(child)
                removed_slide_ids += 1
    parts["ppt/presentation.xml"] = ET.tostring(pres_root, encoding="utf-8", xml_declaration=True)

    pres_rels = ET.fromstring(parts["ppt/_rels/presentation.xml.rels"])
    removed_presentation_rels = 0
    for rel in list(pres_rels):
        if rel.attrib.get("Target") == "slides/slide2.xml":
            pres_rels.remove(rel)
            removed_presentation_rels += 1
    parts["ppt/_rels/presentation.xml.rels"] = ET.tostring(pres_rels, encoding="utf-8", xml_declaration=True)

    ct_root = ET.fromstring(parts["[Content_Types].xml"])
    removed_content_types = 0
    for override in list(ct_root):
        part_name = override.attrib.get("PartName", "")
        if part_name in {"/ppt/slides/slide2.xml", "/ppt/notesSlides/notesSlide2.xml"}:
            ct_root.remove(override)
            removed_content_types += 1
    parts["[Content_Types].xml"] = ET.tostring(ct_root, encoding="utf-8", xml_declaration=True)

    with ZipFile(path, "w", ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)

    return {
        "removed_parts": [part for part in removed_parts if part not in parts],
        "removed_shapes": removed_shapes,
        "removed_note_rels": removed_note_rels,
        "removed_slide_ids": removed_slide_ids,
        "removed_presentation_rels": removed_presentation_rels,
        "removed_content_types": removed_content_types,
    }


def _build_seed(contract: StockDonorContract, output: Path) -> Check:
    if not contract.donor_template.exists():
        return Check("stock_donor_seed", "fail", {"donor_template": str(contract.donor_template), "error": "missing"})
    output.parent.mkdir(parents=True, exist_ok=True)
    replacements = 0
    touched: list[str] = []
    with ZipFile(contract.donor_template) as zin, ZipFile(output, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(POTX_CONTENT_TYPE.encode("utf-8"), PPTX_CONTENT_TYPE.encode("utf-8"))
            if item.filename in contract.ole_parts:
                data, count = _patch_empty_names(data, contract.target_name)
                replacements += count
                touched.append(item.filename)
            zout.writestr(item, data)
    cleanup: dict[str, Any] | None = None
    if contract.contract == DEFAULT_CONTRACT:
        cleanup = _clean_qtr04_seed(output)
    names = set(template_named_elements(output))
    ok = contract.target_name in names and replacements >= 1
    return Check(
        "stock_donor_seed",
        "pass" if ok else "fail",
        {
            "donor_template": str(contract.donor_template),
            "seed_pptx": str(output),
            "target_name": contract.target_name,
            "ole_parts": list(contract.ole_parts),
            "parts_touched": touched,
            "replacements": replacements,
            "target_name_present": contract.target_name in names,
            "name_count": len(names),
            "note": contract.note,
            "cleanup": cleanup,
        },
    )


def _write_ppttc(contract: StockDonorContract, output: Path, seed: Path, period: str, director_slug: str) -> Check:
    try:
        table, dynamic_terms = contract.data_builder(period, director_slug)
    except Exception as exc:  # noqa: BLE001 - proof artifact should preserve exact failure.
        return Check("stock_donor_ppttc_payload", "fail", {"error": str(exc)})
    payload = [{"template": str(seed.resolve()), "data": [{"name": contract.target_name, "table": table}]}]
    output.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    text = output.read_text(encoding="utf-8")
    terms = tuple(dict.fromkeys((*contract.required_terms, *dynamic_terms)))
    term_checks = [{"term": term, "present": term in text} for term in terms]
    ok = bool(table) and all(item["present"] for item in term_checks)
    return Check(
        "stock_donor_ppttc_payload",
        "pass" if ok else "fail",
        {
            "contract_ppttc": str(output),
            "template": str(seed),
            "target_name": contract.target_name,
            "rows": len(table),
            "cols": max((len(row) for row in table), default=0),
            "required_terms": term_checks,
            "dynamic_terms": dynamic_terms,
            "preview_rows": table[:4],
        },
    )


def _write_markdown(path: Path, payload: dict[str, Any]) -> None:
    lines = [
        f"# {payload['contract']} Stock Donor Proof",
        "",
        f"- Status: `{payload['status']}`",
        f"- Period: `{payload['period']}`",
        f"- Director: `{payload['director_slug']}`",
        f"- Donor template: `{payload['donor_template']}`",
        f"- Seed: `{payload['seed_pptx']}`",
        f"- Bound deck: `{payload['deck']}`",
        f"- Render dir: `{payload['render_dir']}`",
        "",
        "## Checks",
        "",
        "| Check | Status |",
        "|---|---|",
    ]
    for check in payload["checks"]:
        lines.append(f"| `{check['name']}` | `{check['status']}` |")
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def prove_contract(period: str, director_slug: str, contract: StockDonorContract) -> dict[str, Any]:
    paths = _work_paths(period, contract.contract)
    if paths.render_dir.exists():
        shutil.rmtree(paths.render_dir)
    paths.work_dir.mkdir(parents=True, exist_ok=True)
    checks = [
        _build_seed(contract, paths.seed),
        _write_ppttc(contract, paths.ppttc, paths.seed, period, director_slug),
    ]
    dynamic_terms: tuple[str, ...] = ()
    if len(checks) > 1 and checks[1].status == "pass":
        dynamic_terms = tuple(str(term) for term in checks[1].details.get("dynamic_terms", []) if term)
    bound_terms = contract.bound_required_terms or contract.required_terms
    if contract.bound_dynamic_term_count:
        bound_terms = tuple(dict.fromkeys((*bound_terms, *dynamic_terms[: contract.bound_dynamic_term_count])))
    native_contract = NativeChartContract(
        contract=contract.contract,
        source_name="",
        target_name=contract.target_name,
        slide=contract.slide,
        required_terms=bound_terms,
        source_seed=contract.donor_template,
    )
    if all(check.status == "pass" for check in checks):
        checks.append(_bridge_ppttc(paths.ppttc, paths.seed, paths.bound))
    if all(check.status == "pass" for check in checks):
        checks.append(_check_bound_package(paths.bound, native_contract))
    if all(check.status == "pass" for check in checks):
        checks.append(_render_deck(paths.bound, paths.render_dir))
    if all(check.status == "pass" for check in checks):
        checks.append(_proof_render(paths.render_dir, native_contract))
    status = "pass" if all(check.status == "pass" for check in checks) else "fail"
    return {
        "schema": "thinkcell-stock-donor-proof/v1",
        "created_at_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "status": status,
        "contract": contract.contract,
        "period": period,
        "director_slug": director_slug,
        "deck": str(paths.bound),
        "workbook": str(_connected_workbook(period, director_slug)),
        "donor_template": str(contract.donor_template),
        "seed_pptx": str(paths.seed),
        "ppttc": str(paths.ppttc),
        "render_dir": str(paths.render_dir),
        "targets": [
            {
                "name": contract.target_name,
                "slide": contract.slide,
                "source_template": str(contract.donor_template),
                "ole_parts": list(contract.ole_parts),
                "payload_required_terms": list(contract.required_terms),
                "bound_required_terms": list(bound_terms),
            }
        ],
        "checks": [asdict(check) for check in checks],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--director-slug", default=DEFAULT_DIRECTOR_SLUG)
    parser.add_argument("--contract", default=DEFAULT_CONTRACT)
    args = parser.parse_args()

    contract = CONTRACTS.get(args.contract)
    if not contract:
        raise SystemExit(f"unsupported stock donor contract: {args.contract}")

    payload = prove_contract(args.period, args.director_slug, contract)
    paths = _work_paths(args.period, contract.contract)
    proof_json = paths.work_dir / f"{contract.contract}-proof.json"
    proof_md = paths.work_dir / f"{contract.contract}-proof.md"
    proof_json.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    _write_markdown(proof_md, payload)
    print(f"status={payload['status']}")
    print(f"proof_json={proof_json}")
    print(f"proof_md={proof_md}")
    return 0 if payload["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
