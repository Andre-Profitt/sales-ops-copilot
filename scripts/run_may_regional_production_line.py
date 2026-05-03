#!/usr/bin/env python3
"""Run the May 2026 regional Sales Director production line.

This is the top-level, repeatable QA/package command for the current proven
lane. It intentionally defaults to the fast finalization path: source-aware text
polish, meeting-spine rebuild, publish gates, strict APAC intel audit, regional
goal audit, and local review packaging.

The heavier Excel/PowerPoint/think-cell table-image refresh is available behind
explicit flags because it uses the Windows desktop COM bridge.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import shutil
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from pptx import Presentation

from _directors import canonical_directors
from period_context import DEFAULT_PERIOD, context_for_period


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DOWNLOADS_PACKAGE = context_for_period(DEFAULT_PERIOD).review_package_dir


@dataclass
class StepResult:
    name: str
    status: str
    command: list[str]
    elapsed_seconds: float | None = None
    returncode: int | None = None
    stdout_log: str | None = None
    stderr_log: str | None = None
    notes: list[str] | None = None


def _python() -> str:
    venv = ROOT / ".venv" / "bin" / "python"
    return str(venv) if venv.exists() else sys.executable


def _slug(value: str) -> str:
    return value.replace(" ", "-")


def _selected_directors(director_slug: str | None) -> list[dict[str, Any]]:
    directors = canonical_directors()
    if not director_slug:
        return directors
    return [director for director in directors if _slug(str(director["name"])) == director_slug]


def _run_step(name: str, command: list[str], *, run_dir: Path, plan_only: bool) -> StepResult:
    if plan_only:
        return StepResult(name=name, status="planned", command=command)

    logs = run_dir / "logs"
    logs.mkdir(parents=True, exist_ok=True)
    started = time.time()
    result = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, check=False)
    stdout_log = logs / f"{name}.stdout.log"
    stderr_log = logs / f"{name}.stderr.log"
    stdout_log.write_text(result.stdout, encoding="utf-8")
    stderr_log.write_text(result.stderr, encoding="utf-8")
    return StepResult(
        name=name,
        status="pass" if result.returncode == 0 else "fail",
        command=command,
        elapsed_seconds=round(time.time() - started, 2),
        returncode=result.returncode,
        stdout_log=str(stdout_log),
        stderr_log=str(stderr_log),
    )


def _can_continue(steps: list[StepResult]) -> bool:
    return all(step.status in {"pass", "planned", "skip"} for step in steps)


def _linked_deck(period: str, slug: str) -> Path:
    return ROOT / "state" / period / slug / f"{slug}-LAND-{period}-table-image-linked.pptx"


def _meeting_spine(period: str, slug: str) -> Path:
    return ROOT / "state" / period / slug / "factory" / "meeting-spine" / f"{slug}-LAND-{period}-meeting-spine.pptx"


def _zip_ok(path: Path) -> bool:
    try:
        with ZipFile(path) as zf:
            return zf.testzip() is None
    except BadZipFile:
        return False


def _deck_text(path: Path) -> str:
    prs = Presentation(path)
    return "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text
    )


def _meeting_smoke_one(period: str, director: dict[str, Any], required: list[str], forbidden: list[str]) -> dict[str, Any]:
    slug = _slug(str(director["name"]))
    deck = _meeting_spine(period, slug)
    zip_ok = deck.exists() and _zip_ok(deck)
    slide_count = 0
    text = ""
    if deck.exists():
        prs = Presentation(deck)
        slide_count = len(prs.slides)
        text = _deck_text(deck)
    missing = [needle for needle in required if needle not in text]
    found_forbidden = [needle for needle in forbidden if needle in text]
    status = "pass" if zip_ok and slide_count == 16 and not missing and not found_forbidden else "fail"
    return {
        "director": str(director["name"]),
        "slug": slug,
        "territory": str(director["scope_label"]),
        "deck": str(deck),
        "status": status,
        "zip_ok": zip_ok,
        "slide_count": slide_count,
        "required_missing": missing,
        "forbidden_found": found_forbidden,
    }


def _write_meeting_smoke(
    period: str,
    output_dir: Path,
    *,
    directors: list[dict[str, Any]],
    jobs: int,
    required_text: list[str],
) -> tuple[Path, Path, bool]:
    required = required_text
    forbidden = [
        "SC Test",
        "Test Account",
        "Test MASB",
        "5000%",
        "9000%",
        "#NAME",
        "#NULL",
        "Click to add subtitle",
        "Title of the section",
        "no director-specific prior-review target pack",
        "prior territory concentration spine is not applied",
        "prior territory risk spine is not applied",
    ]
    if jobs > 1 and len(directors) > 1:
        with ThreadPoolExecutor(max_workers=min(jobs, len(directors))) as executor:
            results = list(
                executor.map(
                    lambda director: _meeting_smoke_one(period, director, required, forbidden),
                    directors,
                )
            )
    else:
        results = [_meeting_smoke_one(period, director, required, forbidden) for director in directors]

    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "meeting_spine_smoke_report.json"
    md_path = output_dir / "meeting_spine_smoke_report.md"
    payload = {
        "schema": "meeting-spine-smoke/v1",
        "period": period,
        "status": "pass" if all(row["status"] == "pass" for row in results) else "fail",
        "results": results,
    }
    json_path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    lines = ["# Meeting Spine Smoke Report", ""]
    for row in results:
        lines.append(
            f"- {row['status'].upper()} {row['slug']}: "
            f"slides={row['slide_count']} zip_ok={row['zip_ok']} "
            f"forbidden={len(row['forbidden_found'])} missing={len(row['required_missing'])}"
        )
    md_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return json_path, md_path, payload["status"] == "pass"


def _copy_review_package(period: str, package_dir: Path, *, directors: list[dict[str, Any]]) -> list[str]:
    package_dir.mkdir(parents=True, exist_ok=True)
    copied: list[str] = []
    for lock_file in package_dir.glob("~$*"):
        lock_file.unlink(missing_ok=True)
    for existing in package_dir.glob("*.pptx"):
        existing.unlink()

    for director in directors:
        slug = _slug(str(director["name"]))
        source = _meeting_spine(period, slug)
        target = package_dir / source.name
        shutil.copy2(source, target)
        copied.append(str(target))

    report_paths = [
        ROOT / "state" / period / "__regional__" / "meeting_spine" / "meeting_spine_manifest.json",
        ROOT / "state" / period / "__regional__" / "meeting_spine" / "meeting_spine_smoke_report.md",
        ROOT / "state" / period / "__regional__" / "publish_gate" / "regional_publish_gate.md",
        ROOT / "state" / period / "__regional__" / "goal_audit" / "regional_deck_goal_audit.md",
        ROOT / "state" / period / "__regional__" / "template_contract_gate.md",
        ROOT / "state" / period / "__regional__" / "template_contract_gate.json",
        ROOT / "state" / period / "__regional__" / "thinkcell_deck_blueprint" / "thinkcell_deck_factory_blueprint.md",
        ROOT / "state" / period / "__regional__" / "thinkcell_visual_plan" / "thinkcell_visual_contract_plan.md",
        ROOT / "docs" / "thinkcell-corpus" / "deck-factory-blueprint.md",
        ROOT / "docs" / "thinkcell-corpus" / "deck-factory-visual-plan.md",
    ]
    for source in report_paths:
        if source.exists():
            target = package_dir / source.name
            shutil.copy2(source, target)
            copied.append(str(target))
    return copied


def _write_runbook_summary(run_dir: Path, manifest: dict[str, Any]) -> Path:
    path = run_dir / "production_summary.md"
    lines = [
        f"# {manifest['production_summary_title']}",
        "",
        f"- Status: `{manifest['status']}`",
        f"- Period: `{manifest['period']}`",
        f"- Review package: `{manifest.get('review_package', '')}`",
        "",
        "## Steps",
        "",
        "| Step | Status | Seconds |",
        "|---|---:|---:|",
    ]
    for step in manifest["steps"]:
        lines.append(f"| {step['name']} | {step['status']} | {step.get('elapsed_seconds') or ''} |")
    if manifest.get("residual_risks"):
        lines.extend(["", "## Residual Risks", ""])
        lines.extend(f"- {risk}" for risk in manifest["residual_risks"])
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return path


def _copy_final_package_reports(period: str, package_dir: Path, run_dir: Path, summary_path: Path) -> None:
    status_dir = ROOT / "state" / period / "__regional__" / "production_status"
    report_sources = [
        status_dir / "regional_production_status.md",
        status_dir / "regional_production_status.json",
        summary_path,
        run_dir / "review_package_validation.json",
        ROOT / "state" / period / "__regional__" / "template_contract_gate.md",
        ROOT / "state" / period / "__regional__" / "template_contract_gate.json",
        ROOT / "state" / period / "__regional__" / "visual_gate" / "review_package" / "review_package_visual_gate.md",
        ROOT / "state" / period / "__regional__" / "visual_gate" / "review_package" / "review_package_visual_gate.json",
        ROOT
        / "state"
        / period
        / "__regional__"
        / "brand_style_gate"
        / "review_package"
        / "review_package_brand_style_gate.md",
        ROOT
        / "state"
        / period
        / "__regional__"
        / "brand_style_gate"
        / "review_package"
        / "review_package_brand_style_gate.json",
    ]
    for source in report_sources:
        if source.exists():
            shutil.copy2(source, package_dir / source.name)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--host", default="Windows-VM")
    parser.add_argument("--director-slug")
    parser.add_argument("--plan-only", action="store_true")
    parser.add_argument(
        "--refresh-source",
        action="store_true",
        help="Refresh Salesforce-derived trends.json, brief.md, land.xlsx, and land.model.xlsx before deck gates.",
    )
    parser.add_argument(
        "--source-only",
        action="store_true",
        help="Stop after source/envelope/connected-factory refresh. Does not run deck gates or package decks.",
    )
    parser.add_argument(
        "--rebuild-connected-factories",
        action="store_true",
        help="Rebuild connected Excel factory workbooks from the refreshed source workbooks.",
    )
    parser.add_argument(
        "--refresh-sparse-tables",
        action="store_true",
        help="Run the Windows COM sparse table-image refresh before gates.",
    )
    parser.add_argument(
        "--full-table-image-refresh",
        action="store_true",
        help="Run the heavier regional table-image factory refresh before gates.",
    )
    parser.add_argument("--skip-package", action="store_true")
    parser.add_argument(
        "--sharepoint-publish",
        action="store_true",
        help="After local gates pass, quarantine prior top-level May assets, upload final assets, and validate SharePoint.",
    )
    parser.add_argument(
        "--sharepoint-validate",
        action="store_true",
        help="After local gates pass, validate the May SharePoint folder without moving or uploading files.",
    )
    parser.add_argument("--package-dir", type=Path, default=DEFAULT_DOWNLOADS_PACKAGE)
    parser.add_argument("--jobs", type=int, default=1, help="Parallelize local per-director deck/audit stages.")
    args = parser.parse_args()
    try:
        period_context = context_for_period(args.period)
    except ValueError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    if args.source_only and not args.refresh_source:
        print("error: --source-only requires --refresh-source", file=sys.stderr)
        return 2
    if args.refresh_source and not args.full_table_image_refresh and not args.source_only:
        print(
            "error: --refresh-source changes the Excel layer; use --full-table-image-refresh "
            "to refresh PowerPoint links, or --source-only to stop after Excel artifacts.",
            file=sys.stderr,
        )
        return 2
    if args.sharepoint_publish and args.skip_package:
        print("error: --sharepoint-publish requires the local package gates", file=sys.stderr)
        return 2
    if args.sharepoint_publish and args.director_slug:
        print("error: --sharepoint-publish is all-director only", file=sys.stderr)
        return 2

    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    run_dir = ROOT / "state" / args.period / "__regional__" / "production_runs" / stamp
    if not args.plan_only:
        run_dir.mkdir(parents=True, exist_ok=True)
    steps: list[StepResult] = []
    directors = _selected_directors(args.director_slug)
    if not directors:
        raise SystemExit(f"unknown director slug: {args.director_slug}")
    rebuild_connected_factories = args.rebuild_connected_factories or (
        args.refresh_source and (args.full_table_image_refresh or args.source_only)
    )

    if args.refresh_source:
        cmd = [
            _python(),
            "scripts/land_brief.py",
            "--period",
            args.period,
            "--snapshot-date",
            period_context.snapshot_date,
        ]
        if args.director_slug:
            cmd.extend(["--director", str(directors[0]["name"])])
        else:
            cmd.append("--all-directors")
            if args.jobs > 1:
                cmd.extend(["--jobs", str(args.jobs)])
        steps.append(_run_step("source_refresh", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if args.refresh_source and _can_continue(steps):
        cmd = [
            _python(),
            "scripts/run_land_to_deck.py",
            "--period",
            args.period,
            "--snapshot-date",
            period_context.snapshot_date,
            "--skip-regen",
            "--validate-only",
        ]
        if args.director_slug:
            cmd.extend(["--director", str(directors[0]["name"])])
        else:
            cmd.append("--all-directors")
        steps.append(_run_step("source_envelope_validation", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if rebuild_connected_factories and _can_continue(steps):
        cmd = [_python(), "scripts/build_connected_factory_workbook.py", "--period", args.period]
        if args.director_slug:
            cmd.extend(["--director", str(directors[0]["name"])])
        else:
            cmd.append("--all-directors")
        steps.append(_run_step("connected_factory_rebuild", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if args.source_only:
        status = "planned" if args.plan_only else ("pass" if _can_continue(steps) else "fail")
        manifest = {
            "schema": "may-regional-production-line/v1",
            "status": status,
            "period": args.period,
            "month_label": period_context.month_label,
            "snapshot_date": period_context.snapshot_date,
            "kickoff_date": period_context.kickoff_date,
            "production_summary_title": period_context.production_summary_title,
            "refresh_source": args.refresh_source,
            "source_only": True,
            "rebuild_connected_factories": rebuild_connected_factories,
            "full_table_image_refresh": False,
            "created_at_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
            "run_dir": str(run_dir),
            "review_package": None,
            "copied": [],
            "residual_risks": [
                "Source-only runs intentionally leave PowerPoint table-image links stale until --full-table-image-refresh is run.",
            ],
            "steps": [asdict(step) for step in steps],
        }
        if args.plan_only:
            print("plan_only_no_write=true")
            print(json.dumps(manifest, indent=2))
            for step in steps:
                print(f"{step.status.upper():>7} {step.name}")
            return 0
        manifest_path = run_dir / "manifest.json"
        manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
        summary_path = _write_runbook_summary(run_dir, manifest)
        print(f"manifest={manifest_path}")
        print(f"summary={summary_path}")
        for step in steps:
            print(f"{step.status.upper():>7} {step.name}")
        return 0 if status == "pass" else 2

    if _can_continue(steps):
        steps.append(
            _run_step(
                "template_contract_gate",
                [
                    _python(),
                    "scripts/run_template_contract_gate.py",
                    "--json-output",
                    str(ROOT / "state" / args.period / "__regional__" / "template_contract_gate.json"),
                    "--markdown-output",
                    str(ROOT / "state" / args.period / "__regional__" / "template_contract_gate.md"),
                ],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if _can_continue(steps):
        steps.append(
            _run_step(
                "thinkcell_ppttc_validation",
                [_python(), "scripts/run_ppttc_factory_validation.py", "--period", args.period, "--strict"],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if _can_continue(steps):
        steps.append(
            _run_step(
                "thinkcell_deck_blueprint",
                [_python(), "scripts/build_thinkcell_deck_factory_blueprint.py", "--period", args.period],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if args.full_table_image_refresh:
        cmd = [
            _python(),
            "scripts/run_regional_table_image_factory.py",
            "--period",
            args.period,
            "--host",
            args.host,
            "--refresh-existing",
            "--finalize-close",
        ]
        if args.director_slug:
            cmd.extend(["--director-slug", args.director_slug])
        steps.append(_run_step("full_table_image_refresh", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if args.refresh_sparse_tables and _can_continue(steps):
        cmd = [_python(), "scripts/refresh_sparse_table_images.py", "--period", args.period, "--host", args.host]
        if args.director_slug:
            cmd.extend(["--director-slug", args.director_slug])
        steps.append(_run_step("sparse_table_refresh", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if _can_continue(steps):
        cmd = [_python(), "scripts/build_regional_intelligence_specs.py", "--period", args.period]
        if args.director_slug:
            cmd.extend(["--director-slug", args.director_slug])
        steps.append(_run_step("regional_intelligence_specs", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if _can_continue(steps):
        cmd = [_python(), "scripts/build_thinkcell_visual_contract_plan.py", "--period", args.period]
        if args.director_slug:
            cmd.extend(["--director-slug", args.director_slug])
        if args.jobs > 1:
            cmd.extend(["--jobs", str(args.jobs)])
        steps.append(_run_step("thinkcell_visual_contract_plan", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if _can_continue(steps):
        cmd = [_python(), "scripts/build_ai_deck_builder_review_workbook.py", "--period", args.period]
        if args.director_slug:
            cmd.extend(["--director-slug", args.director_slug])
        steps.append(_run_step("ai_deck_builder_review_workbook", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if _can_continue(steps):
        cmd = [_python(), "scripts/polish_regional_linked_deck_text.py", "--period", args.period]
        if args.director_slug:
            cmd.extend(["--director-slug", args.director_slug])
        steps.append(_run_step("source_aware_text_polish", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if _can_continue(steps):
        cmd = [_python(), "scripts/fix_table_image_aspect_ratios.py", "--period", args.period, "--linked-decks"]
        if args.director_slug:
            cmd.extend(["--director-slug", args.director_slug])
        if args.jobs > 1:
            cmd.extend(["--jobs", str(args.jobs)])
        steps.append(_run_step("table_image_aspect_polish", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if _can_continue(steps):
        cmd = [_python(), "scripts/build_regional_meeting_spine_decks.py", "--period", args.period]
        if args.director_slug:
            cmd.extend(["--director-slug", args.director_slug])
        if args.jobs > 1:
            cmd.extend(["--jobs", str(args.jobs)])
        steps.append(_run_step("meeting_spine_build", cmd, run_dir=run_dir, plan_only=args.plan_only))

    if _can_continue(steps):
        if args.plan_only:
            steps.append(StepResult("meeting_spine_smoke", "planned", ["internal"]))
        else:
            smoke_dir = ROOT / "state" / args.period / "__regional__" / "meeting_spine"
            json_path, md_path, ok = _write_meeting_smoke(
                args.period,
                smoke_dir,
                directors=directors,
                jobs=args.jobs,
                required_text=list(period_context.required_deck_text),
            )
            steps.append(
                StepResult(
                    "meeting_spine_smoke",
                    "pass" if ok else "fail",
                    ["internal"],
                    notes=[str(json_path), str(md_path)],
                )
            )

    if _can_continue(steps):
        publish_cmd = [
            _python(),
            "scripts/run_regional_deck_publish_gate.py",
            "--period",
            args.period,
            "--jobs",
            str(args.jobs),
        ]
        if args.director_slug:
            publish_cmd.extend(["--director-slug", args.director_slug])
        steps.append(
            _run_step(
                "regional_publish_gate",
                publish_cmd,
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if _can_continue(steps):
        apac_slug = period_context.apac_strict_slug
        full = _linked_deck(args.period, apac_slug)
        full_json = ROOT / "state" / args.period / apac_slug / "factory" / "jesper_apac_intel_coverage_current_full.json"
        full_md = ROOT / "state" / args.period / apac_slug / "factory" / "jesper_apac_intel_coverage_current_full.md"
        steps.append(
            _run_step(
                "apac_strict_intel_full",
                [
                    _python(),
                    "scripts/audit_jesper_apac_intel_coverage.py",
                    str(full),
                    "--json-output",
                    str(full_json),
                    "--markdown-output",
                    str(full_md),
                ],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if _can_continue(steps):
        apac_slug = period_context.apac_strict_slug
        spine = _meeting_spine(args.period, apac_slug)
        spine_json = ROOT / "state" / args.period / apac_slug / "factory" / "jesper_apac_intel_coverage_current_meeting_spine.json"
        spine_md = ROOT / "state" / args.period / apac_slug / "factory" / "jesper_apac_intel_coverage_current_meeting_spine.md"
        steps.append(
            _run_step(
                "apac_strict_intel_meeting_spine",
                [
                    _python(),
                    "scripts/audit_jesper_apac_intel_coverage.py",
                    str(spine),
                    "--json-output",
                    str(spine_json),
                    "--markdown-output",
                    str(spine_md),
                ],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if _can_continue(steps):
        goal_cmd = [
            _python(),
            "scripts/audit_regional_decks_against_goals.py",
            "--period",
            args.period,
            "--jobs",
            str(args.jobs),
        ]
        if args.director_slug:
            goal_cmd.extend(["--director-slug", args.director_slug])
        steps.append(_run_step("regional_goal_audit", goal_cmd, run_dir=run_dir, plan_only=args.plan_only))

    copied: list[str] = []
    if (not args.skip_package) and _can_continue(steps):
        if args.plan_only:
            steps.append(StepResult("downloads_review_package", "planned", ["internal"]))
        else:
            copied = _copy_review_package(args.period, args.package_dir, directors=directors)
            steps.append(
                StepResult(
                    "downloads_review_package",
                    "pass",
                    ["internal"],
                    notes=[f"copied={len(copied)}", str(args.package_dir)],
                )
            )

    if (not args.skip_package) and _can_continue(steps):
        validation_json = run_dir / "review_package_validation.json"
        steps.append(
            _run_step(
                "review_package_validation",
                [
                    _python(),
                    "scripts/validate_may_review_package.py",
                    "--period",
                    args.period,
                    "--package-dir",
                    str(args.package_dir),
                    "--json-output",
                    str(validation_json),
                ],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if (not args.skip_package) and _can_continue(steps):
        visual_dir = ROOT / "state" / args.period / "__regional__" / "visual_gate" / "review_package"
        steps.append(
            _run_step(
                "review_package_visual_gate",
                [
                    _python(),
                    "scripts/run_review_package_visual_gate.py",
                    "--period",
                    args.period,
                    "--package-dir",
                    str(args.package_dir),
                    "--output-dir",
                    str(visual_dir),
                    "--jobs",
                    str(args.jobs),
                ],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if (not args.skip_package) and _can_continue(steps):
        style_dir = ROOT / "state" / args.period / "__regional__" / "brand_style_gate" / "review_package"
        steps.append(
            _run_step(
                "review_package_brand_style_gate",
                [
                    _python(),
                    "scripts/run_review_package_brand_style_gate.py",
                    "--period",
                    args.period,
                    "--package-dir",
                    str(args.package_dir),
                    "--output-dir",
                    str(style_dir),
                ],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if args.sharepoint_publish and _can_continue(steps):
        steps.append(
            _run_step(
                "sharepoint_containment",
                [
                    _python(),
                    "scripts/contain_may_sharepoint_uploads.py",
                    "--period",
                    args.period,
                    "--execute",
                ],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if args.sharepoint_publish and _can_continue(steps):
        steps.append(
            _run_step(
                "sharepoint_upload",
                [_python(), "scripts/upload_may_regional_assets_sharepoint.py", "--period", args.period],
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    if (args.sharepoint_publish or args.sharepoint_validate) and _can_continue(steps):
        validate_cmd = [_python(), "scripts/validate_may_sharepoint_upload.py", "--period", args.period]
        if args.sharepoint_publish:
            validate_cmd.append("--upload-manifest")
        steps.append(
            _run_step(
                "sharepoint_validation",
                validate_cmd,
                run_dir=run_dir,
                plan_only=args.plan_only,
            )
        )

    residual_risks = [
        "Native think-cell tables are still bridged by linked table images; this is reliable but not the final native table donor lane.",
        "Desktop COM refresh remains available behind explicit flags and should stay gated by PowerPoint open/publish checks.",
        "Action cadence should remain a decision register unless real differentiated due dates exist.",
    ]
    status = "planned" if args.plan_only else ("pass" if _can_continue(steps) else "fail")
    manifest = {
        "schema": "may-regional-production-line/v1",
        "status": status,
        "period": args.period,
        "month_label": period_context.month_label,
        "snapshot_date": period_context.snapshot_date,
        "kickoff_date": period_context.kickoff_date,
        "production_summary_title": period_context.production_summary_title,
        "refresh_source": args.refresh_source,
        "rebuild_connected_factories": rebuild_connected_factories,
        "full_table_image_refresh": args.full_table_image_refresh,
        "sharepoint_publish": args.sharepoint_publish,
        "sharepoint_validate": args.sharepoint_validate,
        "created_at_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "run_dir": str(run_dir),
        "review_package": str(args.package_dir) if not args.skip_package else None,
        "copied": copied,
        "residual_risks": residual_risks,
        "steps": [asdict(step) for step in steps],
    }
    if args.plan_only:
        if (not args.skip_package) and _can_continue(steps):
            steps.append(
                StepResult(
                    name="production_status_report",
                    status="planned",
                    command=[_python(), "scripts/report_regional_production_status.py", "--period", args.period],
                )
            )
            if args.sharepoint_publish:
                steps.append(
                    StepResult(
                        name="sharepoint_evidence_upload",
                        status="planned",
                        command=[_python(), "scripts/upload_may_sharepoint_evidence.py", "--period", args.period],
                    )
                )
                steps.append(
                    StepResult(
                        name="sharepoint_final_validation",
                        status="planned",
                        command=[
                            _python(),
                            "scripts/validate_may_sharepoint_upload.py",
                            "--period",
                            args.period,
                            "--upload-manifest",
                        ],
                    )
                )
        manifest["steps"] = [asdict(step) for step in steps]
        print("plan_only_no_write=true")
        print(json.dumps(manifest, indent=2))
        for step in steps:
            print(f"{step.status.upper():>7} {step.name}")
        return 0

    manifest_path = run_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    summary_path = _write_runbook_summary(run_dir, manifest)

    if (not args.skip_package) and status == "pass":
        logs = run_dir / "logs"
        logs.mkdir(parents=True, exist_ok=True)
        status_command = [_python(), "scripts/report_regional_production_status.py", "--period", args.period]
        started = time.time()
        result = subprocess.run(status_command, cwd=ROOT, text=True, capture_output=True, check=False)
        stdout_log = logs / "production_status_report.stdout.log"
        stderr_log = logs / "production_status_report.stderr.log"
        stdout_log.write_text(result.stdout, encoding="utf-8")
        stderr_log.write_text(result.stderr, encoding="utf-8")
        steps.append(
            StepResult(
                name="production_status_report",
                status="pass" if result.returncode == 0 else "fail",
                command=status_command,
                elapsed_seconds=round(time.time() - started, 2),
                returncode=result.returncode,
                stdout_log=str(stdout_log),
                stderr_log=str(stderr_log),
            )
        )
        status = "pass" if _can_continue(steps) else "fail"
        manifest["status"] = status
        manifest["steps"] = [asdict(step) for step in steps]
        manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
        summary_path = _write_runbook_summary(run_dir, manifest)
        if status == "pass":
            subprocess.run(status_command, cwd=ROOT, text=True, capture_output=True, check=False)
            _copy_final_package_reports(args.period, args.package_dir, run_dir, summary_path)
            if args.sharepoint_publish:
                steps.append(
                    _run_step(
                        "sharepoint_evidence_upload",
                        [_python(), "scripts/upload_may_sharepoint_evidence.py", "--period", args.period],
                        run_dir=run_dir,
                        plan_only=False,
                    )
                )
                if _can_continue(steps):
                    steps.append(
                        _run_step(
                            "sharepoint_final_validation",
                            [
                                _python(),
                                "scripts/validate_may_sharepoint_upload.py",
                                "--period",
                                args.period,
                                "--upload-manifest",
                            ],
                            run_dir=run_dir,
                            plan_only=False,
                        )
                    )
                status = "pass" if _can_continue(steps) else "fail"
                manifest["status"] = status
                manifest["steps"] = [asdict(step) for step in steps]
                manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
                summary_path = _write_runbook_summary(run_dir, manifest)
                if status == "pass":
                    _copy_final_package_reports(args.period, args.package_dir, run_dir, summary_path)

    print(f"manifest={manifest_path}")
    print(f"summary={summary_path}")
    for step in steps:
        print(f"{step.status.upper():>7} {step.name}")
    return 0 if status in {"pass", "planned"} else 2


if __name__ == "__main__":
    raise SystemExit(main())
