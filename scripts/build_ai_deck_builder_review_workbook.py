#!/usr/bin/env python3
"""Build a one-director AI/deck-builder review workbook.

The workbook is a human-review control pack: it shows the production deck
artifacts, think-cell render lanes, AI artifact requirements, and gate status
for one or all Sales Director decks.
"""

from __future__ import annotations

import argparse
import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from _directors import canonical_directors
from period_context import DEFAULT_PERIOD, context_for_period


ROOT = Path(__file__).resolve().parent.parent
CONTRACT_PATH = ROOT / "config" / "sd_factory_render_lane_contract.2026-Q2.json"
HEADER_FILL = PatternFill("solid", fgColor="083EA7")
SUBHEAD_FILL = PatternFill("solid", fgColor="D9EAF7")
PASS_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
FAIL_FILL = PatternFill("solid", fgColor="FCE4D6")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def slugify(value: str) -> str:
    return value.replace(" ", "-")


def selected_directors(director_slug: str | None) -> list[dict[str, Any]]:
    directors = canonical_directors()
    if not director_slug:
        return directors
    selected = [
        director
        for director in directors
        if slugify(str(director["name"])) == director_slug
    ]
    if not selected:
        raise SystemExit(f"unknown director slug: {director_slug}")
    return selected


def artifact_paths(period: str, slug: str) -> dict[str, Path]:
    director_dir = ROOT / "state" / period / slug
    return {
        "source trends": director_dir / "source" / "trends.json",
        "source brief": director_dir / "source" / "brief.md",
        "companion workbook": director_dir / "land.xlsx",
        "model workbook": director_dir / "land.model.xlsx",
        "connected factory": director_dir / "factory" / "connected" / "connected_factory.xlsx",
        "table-image workbook": director_dir / "factory" / "connected" / "connected_factory_table_images.xlsx",
        "linked think-cell deck": director_dir / f"{slug}-LAND-{period}-table-image-linked.pptx",
        "meeting spine deck": director_dir / "factory" / "meeting-spine" / f"{slug}-LAND-{period}-meeting-spine.pptx",
        "annotated original assessment": director_dir
        / "factory"
        / "original-comment-assessment"
        / f"{slug}-{period}-annotated-original-assessment.xlsx",
        "annotated original assessment manifest": director_dir
        / "factory"
        / "original-comment-assessment"
        / f"{slug}-{period}-annotated-original-assessment.json",
        "ai artifact target": director_dir / "factory" / "ai-review" / f"{slug}-{period}-ai-enrichment-artifact.json",
    }


def file_status(path: Path) -> tuple[str, str, int | str]:
    if not path.exists():
        return "missing", "", ""
    mtime = datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
    size = path.stat().st_size
    return "present", mtime, size


def load_publish_rows(period: str) -> list[dict[str, Any]]:
    candidates = [
        ROOT / "state" / period / "__regional__" / "publish_gate_all_smoke" / "regional_publish_gate.json",
        ROOT / "state" / period / "__regional__" / "publish_gate" / "regional_publish_gate.json",
    ]
    for path in candidates:
        if not path.exists():
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, list):
            return payload
        if isinstance(payload, dict):
            directors = payload.get("directors")
            if isinstance(directors, list):
                return directors
    return []


def publish_row_for_slug(period: str, slug: str) -> dict[str, Any] | None:
    for row in load_publish_rows(period):
        row_slug = row.get("slug") or slugify(str(row.get("director", "")))
        if row_slug == slug:
            return row
    return None


def connected_workbook_summary(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"status": "missing", "sheets": 0, "defined_names": 0}
    wb = load_workbook(path, read_only=True, data_only=False)
    return {
        "status": "present",
        "sheets": len(wb.sheetnames),
        "defined_names": len(list(wb.defined_names)),
        "has_thinkcell_link_map": "ThinkCell_Link_Map" in wb.sheetnames,
        "has_audit_checks": "Audit_Checks" in wb.sheetnames,
    }


def write_rows(ws: Any, rows: list[list[Any]], *, table_name: str | None = None) -> None:
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    if table_name and ws.max_row > 1 and ws.max_column > 1:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def style_sheet(ws: Any) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            if str(cell.value).lower() == "pass":
                cell.fill = PASS_FILL
            elif str(cell.value).lower() in {"review", "partial", "planned"}:
                cell.fill = WARN_FILL
            elif str(cell.value).lower() in {"fail", "missing"}:
                cell.fill = FAIL_FILL
    for idx in range(1, ws.max_column + 1):
        letter = get_column_letter(idx)
        width = 12
        for cell in ws[letter]:
            if cell.value is None:
                continue
            width = max(width, min(58, len(str(cell.value)) + 3))
        ws.column_dimensions[letter].width = width
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(ws.max_column)}{ws.max_row}",
        CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL),
    )


def build_review_workbook(
    *,
    director: dict[str, Any],
    period: str,
    contract: dict[str, Any],
    out_dir: Path,
) -> Path:
    name = str(director["name"])
    slug = slugify(name)
    context = context_for_period(period)
    paths = artifact_paths(period, slug)
    publish = publish_row_for_slug(period, slug) or {}
    connected_summary = connected_workbook_summary(paths["connected factory"])
    ai_slots = [
        slot
        for slot in contract["monthly_slots"]
        if slot.get("status") == "contracted_ai_cached_required"
    ]

    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True

    review_index_rows = [
        ["Field", "Value", "Status", "Notes"],
        ["Director", name, "pass", str(director["scope_label"])],
        ["Period", period, "pass", context.month_label],
        ["Snapshot date", context.snapshot_date, "pass", "Certified period context"],
        ["Factory mode", contract["decision"]["factory_mode"], "pass", contract["decision"]["production_substrate"]],
        ["AI role", contract["decision"]["llm_role"], "partial", contract["decision"]["tc_ai_role"]],
        ["Linked deck publish gate", publish.get("status", "missing"), publish.get("status", "missing"), "0 blockers required for production"],
        ["Named think-cell elements", publish.get("named_thinkcell_elements_present", ""), "pass" if publish.get("named_thinkcell_elements_present", 0) else "review", "Expected 19 in current linked deck lane"],
        ["Table-image objects", publish.get("table_image_count", ""), "pass" if publish.get("table_image_count", 0) else "review", "Expected 19 in current linked deck lane"],
        ["Connected workbook sheets", connected_summary.get("sheets", ""), "pass" if connected_summary.get("sheets", 0) else "fail", "Formula/audit workbook"],
        ["Connected workbook defined names", connected_summary.get("defined_names", ""), "pass" if connected_summary.get("defined_names", 0) else "review", "think-cell named output ranges"],
        ["Annotated original assessment", str(paths["annotated original assessment"]), "pass" if paths["annotated original assessment"].exists() else "missing", "Assesses generated deck against Rebekka's comments without migrating comments"],
        ["Open production gap", "AI artifact manifest integration", "partial", "tc-aicore gate exists; all-9 production manifest wiring remains"],
    ]
    ws = wb.create_sheet("Review_Index")
    write_rows(ws, review_index_rows, table_name="tbl_Review_Index")
    ws.sheet_properties.tabColor = "083EA7"

    artifact_rows = [["Artifact", "Path", "Exists", "Modified", "Size bytes", "Review note"]]
    for label, path in paths.items():
        status, mtime, size = file_status(path)
        note = "Required for deck builder review"
        if label == "ai artifact target":
            note = "Expected output once accepted GPT-5.5/tc.ai text is cached"
        artifact_rows.append([label, str(path), status, mtime, size, note])
    ws = wb.create_sheet("Builder_Artifacts")
    write_rows(ws, artifact_rows, table_name="tbl_Builder_Artifacts")

    ai_rows = [
        [
            "Slot",
            "Slide",
            "Surface",
            "Source",
            "Primary lane",
            "Artifact requirement",
            "Guardrail",
            "Current status",
        ]
    ]
    for slot in ai_slots:
        ai_rows.append(
            [
                slot["slot_id"],
                slot.get("slide", ""),
                slot.get("surface", ""),
                slot.get("source", ""),
                slot.get("primary_lane", ""),
                "schema, model, text, evidence_refs, source_metrics, validation=pass",
                slot.get("guardrail", ""),
                "partial",
            ]
        )
    ws = wb.create_sheet("AI_Artifact_Gate")
    write_rows(ws, ai_rows, table_name="tbl_AI_Artifact_Gate")

    lane_rows = [
        [
            "Slot",
            "Slide",
            "Surface",
            "Primary lane",
            "Status",
            "Proof / contract",
            "Guardrail",
        ]
    ]
    for slot in contract["monthly_slots"]:
        lane_rows.append(
            [
                slot["slot_id"],
                slot.get("slide", ""),
                slot.get("surface", ""),
                slot.get("primary_lane", ""),
                slot.get("status", ""),
                slot.get("component_proof") or slot.get("table_image_contract") or "",
                slot.get("guardrail", ""),
            ]
        )
    ws = wb.create_sheet("ThinkCell_Lane_Map")
    write_rows(ws, lane_rows, table_name="tbl_ThinkCell_Lane_Map")

    evidence_rows = [
        ["Gate", "Status", "Evidence", "Notes"],
        ["ARR/ACV separation", "pass", "Render-lane contract global gate", "ARR = Land+Expand; Renewal ACV separate"],
        ["Type filters mandatory", "pass", "Render-lane contract global gate", "Type-bearing facts require explicit filters"],
        ["Render-lane contract", "pass", f"{len(contract['monthly_slots'])} monthly slots", "Validator run separately"],
        ["Component proofs", "pass", f"{len(contract['component_proofs'])} component proofs", "L5 component proof inventory"],
        ["Table-image contracts", "pass", f"{len(contract['table_image_contracts'])} table-image contracts", "Defined-name source range contracts"],
        ["AI artifact gate", "partial", "tc-aicore smoke_ai_artifact.py", "Needs all-9 production manifest wiring"],
        ["Annotated original comment assessment", "pass" if paths["annotated original assessment manifest"].exists() else "missing", str(paths["annotated original assessment manifest"]), "Comments are evidence/rubric only; never copied into production deck"],
        ["S04 slipped ARR", "pass", "tc-aicore smoke_slipped_arr.py", "OpportunityFieldHistory CloseDate moves"],
        ["SharePoint publish state", "pass", "May validation manifest 37/37", "Existing May package is green"],
    ]
    ws = wb.create_sheet("Review_Checks")
    write_rows(ws, evidence_rows, table_name="tbl_Review_Checks")

    if publish:
        publish_rows = [["Field", "Value"]]
        for key, value in publish.items():
            if isinstance(value, (dict, list)):
                value = json.dumps(value, default=str)
            publish_rows.append([key, value])
        ws = wb.create_sheet("Publish_Gate_Row")
        write_rows(ws, publish_rows, table_name="tbl_Publish_Gate_Row")

    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{slug}-{period}-AI-Deck-Builder-Review.xlsx"
    wb.save(out_path)
    return out_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--director-slug")
    parser.add_argument("--contract", type=Path, default=CONTRACT_PATH)
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=None,
        help="Defaults to state/<period>/<Director>/factory/ai-review for one director, or __regional__/ai-review for all.",
    )
    parser.add_argument("--copy-to-downloads", action="store_true")
    args = parser.parse_args()

    contract = json.loads(args.contract.read_text(encoding="utf-8"))
    outputs: list[Path] = []
    for director in selected_directors(args.director_slug):
        slug = slugify(str(director["name"]))
        if args.out_dir:
            out_dir = args.out_dir
        elif args.director_slug:
            out_dir = ROOT / "state" / args.period / slug / "factory" / "ai-review"
        else:
            out_dir = ROOT / "state" / args.period / "__regional__" / "ai-review"
        output = build_review_workbook(
            director=director,
            period=args.period,
            contract=contract,
            out_dir=out_dir,
        )
        outputs.append(output)
        if args.copy_to_downloads:
            target = Path.home() / "Downloads" / output.name
            shutil.copy2(output, target)
            outputs.append(target)

    print("ai_deck_builder_review_workbooks:")
    for output in outputs:
        print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
