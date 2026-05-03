#!/usr/bin/env python3
"""Validate the local May 2026 meeting-spine review package."""

from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from pptx import Presentation

from _directors import canonical_directors
from period_context import DEFAULT_PERIOD, context_for_period


DEFAULT_PACKAGE = context_for_period(DEFAULT_PERIOD).review_package_dir
REQUIRED_REPORTS = {
    "meeting_spine_manifest.json",
    "meeting_spine_smoke_report.md",
    "regional_publish_gate.md",
    "regional_deck_goal_audit.md",
}
FORBIDDEN_TEXT = {
    "SC Test",
    "Test Account",
    "Test MASB",
    "SimCorp Test",
    "Demo Account",
    "Sample Account",
    "Sample Co",
    "Lorem ipsum",
    "5000%",
    "9000%",
    "#NAME",
    "#NULL",
    "#REF",
    "#VALUE",
    "#DIV/0",
    "#N/A",
    "Click to add subtitle",
    "Title of the section",
    "no director-specific prior-review target pack",
    "prior territory concentration spine is not applied",
    "prior territory risk spine is not applied",
}
MIN_SALESFORCE_REPORT_LINKS = 5
MIN_SALESFORCE_OPPORTUNITY_LINKS = 3


@dataclass
class DeckCheck:
    filename: str
    status: str
    slide_count: int | None
    zip_ok: bool
    missing_text: list[str]
    forbidden_text: list[str]
    stretched_table_images: list[str]
    salesforce_link_count: int
    salesforce_report_links: int
    salesforce_opportunity_links: int
    error: str | None = None


def _table_image_aspect_findings(prs: Presentation) -> list[str]:
    findings: list[str] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.name != "Pic":
                continue
            try:
                image_width, image_height = shape.image.size
            except Exception:
                continue
            if not image_width or not image_height or not int(shape.height):
                continue
            image_ratio = image_width / image_height
            shape_ratio = int(shape.width) / int(shape.height)
            distortion = shape_ratio / image_ratio if image_ratio else 1.0
            if distortion < 0.88 or distortion > 1.14:
                findings.append(
                    f"slide {slide_idx}: shape_ratio={shape_ratio:.2f} image_ratio={image_ratio:.2f}"
                )
    return findings


def _slug(value: str) -> str:
    return value.replace(" ", "-")


def _zip_ok(path: Path) -> bool:
    try:
        with ZipFile(path) as zf:
            return zf.testzip() is None
    except BadZipFile:
        return False


def _salesforce_link_counts(path: Path) -> tuple[int, int, int]:
    link_count = 0
    report_links = 0
    opportunity_links = 0
    with ZipFile(path) as zf:
        rel_names = [
            name
            for name in zf.namelist()
            if name.startswith("ppt/slides/_rels/") and name.endswith(".rels")
        ]
        for name in rel_names:
            data = zf.read(name).decode("utf-8", errors="ignore")
            for target in data.split('Target="')[1:]:
                url = target.split('"', 1)[0]
                if "simcorp.my.salesforce.com" not in url:
                    continue
                link_count += 1
                if "/Report/" in url:
                    report_links += 1
                if "/Opportunity/" in url:
                    opportunity_links += 1
    return link_count, report_links, opportunity_links


def _deck_check(path: Path, *, required_text: set[str]) -> DeckCheck:
    zip_ok = _zip_ok(path)
    if not zip_ok:
        return DeckCheck(path.name, "fail", None, False, [], [], [], 0, 0, 0, "invalid zip package")
    try:
        prs = Presentation(path)
        text = "\n".join(
            shape.text
            for slide in prs.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text
        )
    except Exception as exc:  # pragma: no cover - defensive CLI gate
        return DeckCheck(path.name, "fail", None, zip_ok, [], [], [], 0, 0, 0, str(exc))

    text_lower = text.lower()
    missing = sorted(needle for needle in required_text if needle.lower() not in text_lower)
    forbidden = sorted(needle for needle in FORBIDDEN_TEXT if needle.lower() in text_lower)
    stretched = _table_image_aspect_findings(prs)
    link_count, report_links, opportunity_links = _salesforce_link_counts(path)
    links_ok = (
        report_links >= MIN_SALESFORCE_REPORT_LINKS
        and opportunity_links >= MIN_SALESFORCE_OPPORTUNITY_LINKS
    )
    status = (
        "pass"
        if len(prs.slides) == 16 and not missing and not forbidden and not stretched and links_ok
        else "fail"
    )
    return DeckCheck(
        path.name,
        status,
        len(prs.slides),
        zip_ok,
        missing,
        forbidden,
        stretched,
        link_count,
        report_links,
        opportunity_links,
    )


def validate_package(package_dir: Path, *, period: str = DEFAULT_PERIOD) -> dict[str, object]:
    period_context = context_for_period(period)
    lock_files = sorted(path.name for path in package_dir.glob("~$*"))
    expected_decks = {
        f"{_slug(str(director['name']))}-LAND-{period_context.period}-meeting-spine.pptx"
        for director in canonical_directors()
    }
    present_decks = {
        path.name
        for path in package_dir.glob(period_context.meeting_spine_pattern)
        if not path.name.startswith("~$")
    }
    missing_decks = sorted(expected_decks - present_decks)
    extra_decks = sorted(present_decks - expected_decks)
    missing_reports = sorted(name for name in REQUIRED_REPORTS if not (package_dir / name).exists())
    deck_checks = [
        _deck_check(package_dir / name, required_text=set(period_context.required_deck_text))
        for name in sorted(present_decks & expected_decks)
    ]
    status = "pass"
    if lock_files or missing_decks or extra_decks or missing_reports:
        status = "fail"
    if any(check.status != "pass" for check in deck_checks):
        status = "fail"
    return {
        "schema": "may-review-package-validation/v1",
        "status": status,
        "period": period_context.period,
        "month_label": period_context.month_label,
        "snapshot_date": period_context.snapshot_date,
        "package_dir": str(package_dir),
        "expected_deck_count": len(expected_decks),
        "present_deck_count": len(present_decks),
        "lock_files": lock_files,
        "missing_decks": missing_decks,
        "extra_decks": extra_decks,
        "missing_reports": missing_reports,
        "deck_checks": [asdict(check) for check in deck_checks],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--package-dir", type=Path, default=DEFAULT_PACKAGE)
    parser.add_argument("--json-output", type=Path)
    args = parser.parse_args()

    try:
        payload = validate_package(args.package_dir, period=args.period)
    except ValueError as exc:
        print(f"error: {exc}")
        return 2
    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, indent=2))
    return 0 if payload["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
