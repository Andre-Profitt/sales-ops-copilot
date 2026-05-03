#!/usr/bin/env python3
"""Build a reversible APAC-only meeting-spine test deck from the linked deck."""

from __future__ import annotations

import json
import shutil
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
PERIOD = "2026-Q2"
SLUG = "Jesper-Tyrer"
DIRECTOR_DIR = ROOT / "state" / PERIOD / SLUG
SOURCE_DECK = DIRECTOR_DIR / f"{SLUG}-LAND-{PERIOD}-table-image-linked.pptx"
OUT_DIR = DIRECTOR_DIR / "factory" / "meeting-spine-test"
OUT_DECK = OUT_DIR / f"{SLUG}-LAND-{PERIOD}-meeting-spine-test.pptx"
MANIFEST = OUT_DIR / "meeting_spine_manifest.json"

KEEP_SLIDES = [
    1,   # Cover
    2,   # May operating summary
    4,   # Original targets / Q1 accountability
    5,   # Forecast quality
    7,   # May deal readiness
    9,   # Commercial approval gaps
    11,  # FY26 renewal watchlist
    15,  # Owner pipeline coverage
    16,  # Owner coaching focus
    18,  # QTD wins and losses
    21,  # Concentration risk
    22,  # Q2-Q3 risk triage
    24,  # May territory plan
    26,  # May action register
    27,  # May decision checklist
    28,  # Thank you
]


def _slide_title(slide: object) -> str:
    texts: list[str] = []
    for shape in slide.shapes:  # type: ignore[attr-defined]
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            texts.append(" ".join(shape.text.split()))
    return texts[0] if texts else ""


def delete_slides_except(path: Path, keep_slides: list[int]) -> dict[str, object]:
    prs = Presentation(path)
    original_count = len(prs.slides)
    keep_indexes = {slide_no - 1 for slide_no in keep_slides}
    kept_titles = {
        str(slide_no): _slide_title(prs.slides[slide_no - 1])
        for slide_no in keep_slides
        if 1 <= slide_no <= original_count
    }
    slide_id_list = prs.slides._sldIdLst  # noqa: SLF001
    for index in range(original_count - 1, -1, -1):
        if index in keep_indexes:
            continue
        rel_id = slide_id_list[index].rId
        prs.part.drop_rel(rel_id)
        del slide_id_list[index]
    prs.save(path)
    return {
        "source_deck": str(SOURCE_DECK),
        "output_deck": str(path),
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "original_slide_count": original_count,
        "output_slide_count": len(keep_slides),
        "kept_source_slides": keep_slides,
        "kept_titles": kept_titles,
        "purpose": "APAC-only reduced meeting spine for approval before changing the regional factory.",
    }


def main() -> int:
    if not SOURCE_DECK.exists():
        raise SystemExit(f"missing source deck: {SOURCE_DECK}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE_DECK, OUT_DECK)
    manifest = delete_slides_except(OUT_DECK, KEEP_SLIDES)
    MANIFEST.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(OUT_DECK)
    print(MANIFEST)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
