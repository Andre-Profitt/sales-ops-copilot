#!/usr/bin/env python3
"""Build connected Excel factory workbooks for Sales Director LAND decks."""

from __future__ import annotations

import json
import re
import argparse
from collections import defaultdict
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

from _directors import canonical_directors
from sales_director_row_filters import is_internal_sales_record

ROOT = Path(__file__).resolve().parent.parent
BASE_SPEC_PATH = ROOT / "config" / "connected_thinkcell_factory.jesper_apac.json"
ORIGINAL_WORKBOOK_DIR = Path("/Users/test/crm-analytics/output/director_live_workbooks/2026-04-20")
ORIGINAL_DECK_DIR = Path("/Users/test/crm-analytics/output/simcorp_director_decks/2026-04-20/land-only")
ORIGINAL_SLUG_ALIASES = {
    "Adam Steinhouse": "adam-steinhaus",
    "Mourad": "mourad-essofi",
}

Q2_START = "DATE(2026,4,1)"
Q2_END = "DATE(2026,7,1)"
Q3_END = "DATE(2026,10,1)"

HEADER_FILL = PatternFill("solid", fgColor="083EA7")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBHEAD_FILL = PatternFill("solid", fgColor="D9EAF7")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
THIN = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass(frozen=True)
class FactoryContext:
    director: str
    territory: str
    period: str
    slug: str
    original_slug: str
    original_workbook: Path
    original_sidecar: Path
    original_deck: Path
    current_companion_workbook: Path
    current_model_workbook: Path
    current_ppttc: Path
    output_workbook: Path
    generated_spec: Path


def slugify(value: str) -> str:
    return value.replace(" ", "-")


def original_slug(director_name: str) -> str:
    return ORIGINAL_SLUG_ALIASES.get(director_name, director_name.lower().replace(" ", "-"))


def director_index() -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}
    for director in canonical_directors():
        index[director["name"].casefold()] = director
        index[slugify(director["name"]).casefold()] = director
    return index


def resolve_director(value: str) -> dict[str, Any]:
    try:
        return director_index()[value.casefold()]
    except KeyError as exc:
        raise SystemExit(f"Unknown director: {value}") from exc


def factory_context(director: dict[str, Any], period: str) -> FactoryContext:
    name = str(director["name"])
    slug = slugify(name)
    legacy_slug = original_slug(name)
    director_dir = ROOT / "state" / period / slug
    connected_dir = director_dir / "factory" / "connected"
    return FactoryContext(
        director=name,
        territory=str(director["scope_label"]),
        period=period,
        slug=slug,
        original_slug=legacy_slug,
        original_workbook=ORIGINAL_WORKBOOK_DIR / f"{legacy_slug}.xlsx",
        original_sidecar=ORIGINAL_DECK_DIR / f"{legacy_slug}-LAND.json",
        original_deck=ORIGINAL_DECK_DIR / f"{legacy_slug}-LAND.pptx",
        current_companion_workbook=director_dir / "land.xlsx",
        current_model_workbook=director_dir / "land.model.xlsx",
        current_ppttc=director_dir / f"{slug}-LAND-{period}.ppttc",
        output_workbook=connected_dir / "connected_factory.xlsx",
        generated_spec=connected_dir / "connected_factory_spec.json",
    )


def _sidecar_number(sidecar: dict[str, Any], key: str) -> float:
    value = sidecar.get(key)
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _director_intel_lines(ctx: FactoryContext, sidecar: dict[str, Any]) -> list[str]:
    open_arr = _sidecar_number(sidecar, "open_land_arr")
    weighted_arr = _sidecar_number(sidecar, "open_land_arr_wtd")
    q1_lost_arr = _sidecar_number(sidecar, "q1_land_lost_arr")
    renewal_acv = _sidecar_number(sidecar, "q2_renewals_acv") + _sidecar_number(sidecar, "q3_renewals_acv")
    return [
        "LAND scope: Type Land in the original ETL workbook; current factory adds Expand ARR while keeping Renewal ACV separate.",
        f"Original sidecar baseline: {int(_sidecar_number(sidecar, 'open_land_deals'))} open Land deals, EUR {open_arr / 1_000_000:.2f}M ARR unweighted, EUR {weighted_arr / 1_000_000:.2f}M weighted.",
        f"Q1 accountability: {int(_sidecar_number(sidecar, 'q1_land_wins'))} Land wins / EUR {_sidecar_number(sidecar, 'q1_land_wins_arr') / 1_000_000:.2f}M ARR vs {int(_sidecar_number(sidecar, 'q1_land_lost'))} Land losses / EUR {q1_lost_arr / 1_000_000:.2f}M ARR.",
        f"Renewals: Q2-Q3 renewal ACV is EUR {renewal_acv / 1_000_000:.2f}M and must not be blended into ARR.",
        f"Commercial approvals: {int(_sidecar_number(sidecar, 'approved_2026'))} approved, {int(_sidecar_number(sidecar, 'conditionally_approved'))} conditionally approved, {int(_sidecar_number(sidecar, 'missing_stage3'))} missing Stage 3.",
        "ARR and Renewal ACV must never be blended.",
    ]


def generated_spec(base_spec: dict[str, Any], ctx: FactoryContext, sidecar: dict[str, Any]) -> dict[str, Any]:
    spec = deepcopy(base_spec)
    spec["schema"] = "connected-thinkcell-factory/v1"
    spec["director"] = ctx.director
    spec["territory"] = ctx.territory
    spec["period"] = ctx.period
    spec["generated_from_base_spec"] = str(BASE_SPEC_PATH)
    spec["purpose"] = (
        "Factory contract for rebuilding a Sales Director territory deck as Salesforce raw data -> "
        "formula/audit Excel -> named Excel output ranges -> think-cell objects -> validated PowerPoint."
    )
    spec["principles"] = [
        "Original director ETL intelligence is the narrative baseline; current facts update it.",
        "Raw extract sheets are immutable inputs; model/output sheets are formula-driven and auditable.",
        "ARR is Land + Expand only; Renewal ACV stays separate.",
        "Every PowerPoint object must trace to source tabs, model logic, an output range, and a validation rule.",
        "think-cell is the connected visualization layer, not the source of business logic.",
    ]
    spec["artifacts"] = {
        "original_etl_workbook": str(ctx.original_workbook),
        "original_etl_deck": str(ctx.original_deck),
        "original_etl_sidecar": str(ctx.original_sidecar),
        "current_formula_workbook": str(ctx.current_model_workbook.relative_to(ROOT)),
        "current_companion_workbook": str(ctx.current_companion_workbook.relative_to(ROOT)),
        "current_ppttc": str(ctx.current_ppttc.relative_to(ROOT)),
        "thinkcell_seed": "assets/LAND_thinkcell_seed.pptx",
    }
    if ctx.director == "Jesper Tyrer":
        return spec

    if "original_apac_slide_order" in spec:
        spec["original_director_slide_order"] = spec.pop("original_apac_slide_order")
    spec.pop("original_apac_intel_to_preserve", None)
    spec["original_director_intel_to_preserve"] = _director_intel_lines(ctx, sidecar)
    for obj in spec.get("objects", []):
        obj["slide_title"] = str(obj.get("slide_title", "")).replace("APAC", ctx.territory)
        obj["intelligence_role"] = str(obj.get("intelligence_role", "")).replace("APAC", ctx.territory)
        current_tab_map = {
            "Q2_Readiness": "Top_Deals_Land",
            "FY26_Renewals": "At_Risk_Renewals",
        }
        obj["current_tabs"] = [
            current_tab_map.get(str(tab), str(tab))
            for tab in obj.get("current_tabs", [])
        ]
        obj["validation_rules"] = [
            "director_scope_present",
            "arr_acv_separated",
            "output_range_formula_driven",
        ]
        if str(obj.get("object_type")) == "chart":
            obj["validation_rules"].append("chart_source_range_named")
        if "table" in str(obj.get("object_type")) or "table" in str(obj.get("visualization_lane")):
            obj["validation_rules"].append("table_link_blocker_stated")
    return spec


def resolve(path_text: str) -> Path:
    path = Path(path_text)
    return path if path.is_absolute() else ROOT / path


def safe_table_name(title: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_]", "_", title)
    if not cleaned or cleaned[0].isdigit():
        cleaned = f"T_{cleaned}"
    return f"tbl_{cleaned}"[:240]


def normalize_headers(headers: Iterable[Any]) -> list[str]:
    seen: dict[str, int] = {}
    normalized: list[str] = []
    for idx, header in enumerate(headers, start=1):
        text = str(header).strip() if header not in (None, "") else f"Column_{idx}"
        text = text.replace("\n", " ")
        if text in seen:
            seen[text] += 1
            text = f"{text}_{seen[text]}"
        else:
            seen[text] = 1
        normalized.append(text)
    return normalized


def style_range(ws: Any, max_col: int | None = None) -> None:
    max_col = max_col or ws.max_column
    for cell in ws[1][:max_col]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=max_col):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
    ws.freeze_panes = "A2"
    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        width = 14
        for cell in ws[letter]:
            if cell.value is None:
                continue
            width = max(width, min(38, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def add_table(ws: Any, name: str | None = None) -> str | None:
    if ws.max_row < 2 or ws.max_column < 1:
        return None
    table_name = name or safe_table_name(ws.title)
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    return table_name


def write_rows(ws: Any, rows: list[list[Any]], table: bool = True) -> str | None:
    if not rows:
        return None
    headers = normalize_headers(rows[0])
    ws.append(headers)
    for row in rows[1:]:
        record = {headers[idx]: row[idx] for idx in range(min(len(headers), len(row)))}
        if is_internal_sales_record(record):
            continue
        ws.append(list(row))
    style_range(ws)
    return add_table(ws) if table else None


def copy_sheet(
    wb: Workbook,
    src_wb: Any,
    source_sheet: str,
    target_sheet: str,
    *,
    header_row: int = 1,
    table: bool = True,
) -> str | None:
    src = src_wb[source_sheet]
    ws = wb.create_sheet(target_sheet)
    rows: list[list[Any]] = []
    for row in src.iter_rows(min_row=header_row, values_only=True):
        rows.append(list(row))
    while rows and all(value is None for value in rows[-1]):
        rows.pop()
    if not rows:
        return None
    return write_rows(ws, rows, table=table)


def add_named_range(wb: Workbook, name: str, target: str) -> None:
    sheet, cell_range = target.split("!", 1)
    abs_range = _absolute_range(cell_range)
    attr_text = f"{quote_sheetname(sheet)}!{abs_range}"
    wb.defined_names.add(DefinedName(name, attr_text=attr_text))


def _absolute_range(cell_range: str) -> str:
    parts = cell_range.split(":")
    return ":".join(_absolute_cell(part) for part in parts)


def _absolute_cell(cell_ref: str) -> str:
    match = re.fullmatch(r"\$?([A-Z]+)\$?([0-9]+)", cell_ref)
    if not match:
        return cell_ref
    return f"${match.group(1)}${match.group(2)}"


def lookup_formula(key: str, value_col: int = 2) -> str:
    return f'=IFERROR(VLOOKUP("{key}",tbl_Raw_Original_Intel,{value_col},FALSE),"")'


def eur_formula(value_expr: str) -> str:
    return f'=IFERROR({value_expr}/1000000,0)'


def set_money_format(ws: Any, cells: Iterable[str]) -> None:
    for coord in cells:
        ws[coord].number_format = '#,##0.0,," mEUR"'


def table_sumifs(table: str, sum_col: str, filters: list[tuple[str, str]]) -> str:
    pieces = [f"{table}[{sum_col}]"]
    for col, criterion in filters:
        pieces.append(f"{table}[{col}]")
        pieces.append(criterion)
    return f"SUMIFS({','.join(pieces)})"


def table_countifs(table: str, filters: list[tuple[str, str]]) -> str:
    pieces: list[str] = []
    for col, criterion in filters:
        pieces.append(f"{table}[{col}]")
        pieces.append(criterion)
    return f"COUNTIFS({','.join(pieces)})"


def plus(parts: Iterable[str]) -> str:
    return " + ".join(parts)


def original_commentary_rows(spec: dict[str, Any], sidecar: dict[str, Any]) -> list[list[Any]]:
    if spec.get("director") == "Jesper Tyrer":
        return [
            ["q1_promised_opened_arr_eur", 17100000, "EUR", "connected_factory_spec.original_apac_intel", "Rounded original deck commentary."],
            ["q1_slipped_deals", 9, "count", "connected_factory_spec.original_apac_intel", "Rounded original deck commentary."],
            ["q1_slipped_arr_eur", 9400000, "EUR", "connected_factory_spec.original_apac_intel", "Rounded original deck commentary."],
            ["q2_original_deals", 6, "count", "connected_factory_spec.original_apac_intel", "Original Q2 close focus."],
            ["q2_original_arr_eur", 5000000, "EUR", "connected_factory_spec.original_apac_intel", "Original Q2 close focus, rounded."],
            ["concentration_top7_arr_eur", 9600000, "EUR", "connected_factory_spec.original_apac_intel", "Rounded original deck commentary."],
            ["concentration_top5_share", 0.89, "share", "connected_factory_spec.original_apac_intel", "Top five open deals share."],
            ["largest_original_deal_account", "Amova Asset Management", "text", "connected_factory_spec.original_apac_intel", ""],
            ["largest_original_deal_arr_eur", 2600000, "EUR", "connected_factory_spec.original_apac_intel", "Rounded original deck commentary."],
            ["risk_top4_q2_q3_arr_eur", 4900000, "EUR", "connected_factory_spec.original_apac_intel", "Top four Q2-Q3 risk deals."],
            ["owner_push_owners", 3, "count", "connected_factory_spec.original_apac_intel", "Owner coaching baseline."],
            ["owner_push_count", 50, "count", "connected_factory_spec.original_apac_intel", "Owner coaching baseline."],
            ["owner_push_arr_eur", 22000000, "EUR", "connected_factory_spec.original_apac_intel", "Owner coaching baseline, rounded."],
            ["forecast_mix_weighted_arr_eur", 5600000, "EUR", "connected_factory_spec.original_apac_intel", "Weighted across Pipeline Inspection deals."],
            ["forecast_mix_deals", 12, "count", "connected_factory_spec.original_apac_intel", "Original deck commentary."],
            ["commit_arr_eur_original", 2700000, "EUR", "connected_factory_spec.original_apac_intel", "Rounded original deck commentary."],
            ["commit_share_original", 0.48, "share", "connected_factory_spec.original_apac_intel", "Original deck commentary."],
            ["fy26_renewals_count_original", 3, "count", "connected_factory_spec.original_apac_intel", ""],
            ["fy26_renewals_acv_eur_original", 33500000, "EUR", "connected_factory_spec.original_apac_intel", "Rounded original deck commentary."],
        ]

    source = "original_etl_sidecar.generic_director_baseline"
    open_land_deals = _sidecar_number(sidecar, "open_land_deals")
    open_land_arr = _sidecar_number(sidecar, "open_land_arr")
    weighted_arr = _sidecar_number(sidecar, "open_land_arr_wtd")
    renewal_count = _sidecar_number(sidecar, "q2_renewals") + _sidecar_number(sidecar, "q3_renewals")
    renewal_acv = _sidecar_number(sidecar, "q2_renewals_acv") + _sidecar_number(sidecar, "q3_renewals_acv")
    return [
        ["q1_promised_opened_arr_eur", open_land_arr, "EUR", source, "Generic baseline uses original sidecar open Land ARR; no separate promise commentary found."],
        ["q1_slipped_deals", 0, "count", source, "No director-specific slipped-deal commentary in sidecar."],
        ["q1_slipped_arr_eur", 0, "EUR", source, "No director-specific slipped ARR commentary in sidecar."],
        ["q2_original_deals", open_land_deals, "count", source, "Original open Land deal count from sidecar."],
        ["q2_original_arr_eur", open_land_arr, "EUR", source, "Original open Land ARR from sidecar."],
        ["concentration_top7_arr_eur", open_land_arr, "EUR", source, "Detailed concentration model is computed from current raw data."],
        ["concentration_top5_share", 0, "share", source, "Not available in sidecar."],
        ["largest_original_deal_account", "", "text", source, "Computed in current workbook where available."],
        ["largest_original_deal_arr_eur", 0, "EUR", source, "Computed in current workbook where available."],
        ["risk_top4_q2_q3_arr_eur", 0, "EUR", source, "Computed in current workbook where available."],
        ["owner_push_owners", 0, "count", source, "No owner-push summary in sidecar."],
        ["owner_push_count", 0, "count", source, "No owner-push summary in sidecar."],
        ["owner_push_arr_eur", 0, "EUR", source, "No owner-push summary in sidecar."],
        ["forecast_mix_weighted_arr_eur", weighted_arr, "EUR", source, "Original sidecar weighted ARR."],
        ["forecast_mix_deals", open_land_deals, "count", source, "Original sidecar open Land deal count."],
        ["commit_arr_eur_original", 0, "EUR", source, "Commit split comes from current workbook."],
        ["commit_share_original", 0, "share", source, "Commit share comes from current workbook."],
        ["fy26_renewals_count_original", renewal_count, "count", source, "Q2+Q3 renewal count from original sidecar."],
        ["fy26_renewals_acv_eur_original", renewal_acv, "EUR", source, "Q2+Q3 renewal ACV from original sidecar."],
    ]


def build_raw_original_intel(wb: Workbook, spec: dict[str, Any], sidecar: dict[str, Any]) -> None:
    ws = wb.create_sheet("Raw_Original_Intel")
    rows = [
        ["Key", "Value", "Unit", "Source", "Notes"],
        ["director", spec["director"], "text", "factory_spec", ""],
        ["territory", spec["territory"], "text", "factory_spec", ""],
        ["period", spec["period"], "text", "factory_spec", ""],
        ["orientation_date", "2026-05-01", "date", "factory_spec", "Use exact date in deck copy."],
        ["scope_original", "Land-only", "text", "original_etl_deck", "Original director deck was Land-only."],
        [
            "scope_current",
            "Land+Expand ARR; Renewal ACV separate",
            "text",
            "sales_ops_contract",
            "Do not blend ARR and renewal ACV.",
        ],
        ["open_land_deals_sidecar", sidecar["open_land_deals"], "count", "original_etl_sidecar", ""],
        ["open_land_arr_eur_sidecar", sidecar["open_land_arr"], "EUR", "original_etl_sidecar", ""],
        ["open_land_arr_wtd_eur_sidecar", sidecar["open_land_arr_wtd"], "EUR", "original_etl_sidecar", ""],
        ["q1_land_wins", sidecar["q1_land_wins"], "count", "original_etl_sidecar", ""],
        ["q1_land_wins_arr_eur", sidecar["q1_land_wins_arr"], "EUR", "original_etl_sidecar", ""],
        ["q1_land_lost", sidecar["q1_land_lost"], "count", "original_etl_sidecar", ""],
        ["q1_land_lost_arr_eur", sidecar["q1_land_lost_arr"], "EUR", "original_etl_sidecar", ""],
        ["q2_renewals", sidecar["q2_renewals"], "count", "original_etl_sidecar", ""],
        ["q2_renewals_acv_eur", sidecar["q2_renewals_acv"], "EUR", "original_etl_sidecar", ""],
        ["q3_renewals", sidecar["q3_renewals"], "count", "original_etl_sidecar", ""],
        ["q3_renewals_acv_eur", sidecar["q3_renewals_acv"], "EUR", "original_etl_sidecar", ""],
        ["approved_2026", sidecar["approved_2026"], "count", "original_etl_sidecar", ""],
        ["conditionally_approved", sidecar["conditionally_approved"], "count", "original_etl_sidecar", ""],
        ["missing_stage3", sidecar["missing_stage3"], "count", "original_etl_sidecar", ""],
        *original_commentary_rows(spec, sidecar),
    ]
    write_rows(ws, rows)


def build_raw_accounts(wb: Workbook, source_rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("Raw_Accounts")
    accounts = sorted({row["Account"] for row in source_rows if row.get("Account")})
    rows: list[list[Any]] = [["Account", "Source", "Notes"]]
    for account in accounts:
        rows.append([account, "Raw_Pipeline_Open", "Unique account from open pipeline extract"])
    write_rows(ws, rows)


def load_sheet_rows(path: Path, sheet: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = normalize_headers([cell.value for cell in ws[1]])
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in values):
            continue
        record = {headers[idx]: values[idx] for idx in range(min(len(headers), len(values)))}
        if is_internal_sales_record(record):
            continue
        rows.append(record)
    return rows


def _date_in_q2(value: Any) -> bool:
    if not isinstance(value, datetime):
        return False
    return datetime(2026, 4, 1) <= value < datetime(2026, 7, 1)


def build_raw_current_q2_readiness(wb: Workbook, current_data_rows: list[dict[str, Any]]) -> None:
    rows = [
        ["#", "Account", "Opportunity", "Owner", "Type", "Stage", "Close Date", "ARR", "Forecast", "Prob", "Push", "Last Activity", "Readiness", "Next Step"]
    ]
    candidates: list[dict[str, Any]] = []
    for row in current_data_rows:
        if row.get("Type") not in {"Land", "Expand"} or not _date_in_q2(row.get("CloseDate")):
            continue
        try:
            arr = float(row.get("ARR_EUR") or 0)
        except (TypeError, ValueError):
            arr = 0
        candidates.append({**row, "_arr": arr})
    for idx, row in enumerate(sorted(candidates, key=lambda item: item["_arr"], reverse=True)[:12], start=1):
        rows.append(
            [
                idx,
                row.get("AccountName", ""),
                row.get("Id", ""),
                row.get("OwnerName", ""),
                row.get("Type", ""),
                row.get("StageName", ""),
                row.get("CloseDate", ""),
                row.get("ARR_EUR", 0),
                row.get("ForecastCategoryName", ""),
                "",
                "",
                "",
                "Fallback from formula model; enrich next-step evidence in companion workbook.",
                "",
            ]
        )
    ws = wb.create_sheet("Raw_Current_Q2_Readiness")
    write_rows(ws, rows, table=True)


def build_raw_current_fy26_renewals(wb: Workbook, current_data_rows: list[dict[str, Any]]) -> None:
    rows = [["#", "Close Date", "Account", "Opportunity", "Owner", "Stage", "ACV", "Prob", "Risk"]]
    candidates = [row for row in current_data_rows if row.get("Type") == "Renewal"]
    for idx, row in enumerate(sorted(candidates, key=lambda item: item.get("CloseDate") or datetime.max)[:12], start=1):
        rows.append(
            [
                idx,
                row.get("CloseDate", ""),
                row.get("AccountName", ""),
                row.get("Id", ""),
                row.get("OwnerName", ""),
                row.get("StageName", ""),
                row.get("ACV_EUR", 0),
                "",
                "Renewal ACV only",
            ]
        )
    ws = wb.create_sheet("Raw_Current_FY26_Renewals")
    write_rows(ws, rows, table=True)


def distinct_top(rows: list[dict[str, Any]], key: str, amount_key: str, type_filter: set[str]) -> list[str]:
    totals: defaultdict[str, float] = defaultdict(float)
    for row in rows:
        if row.get("Type") not in type_filter:
            continue
        name = row.get(key)
        if not name:
            continue
        amount = row.get(amount_key) or 0
        try:
            totals[str(name)] += float(amount)
        except (TypeError, ValueError):
            continue
    return [name for name, _ in sorted(totals.items(), key=lambda item: item[1], reverse=True)]


def build_model_original_intel(wb: Workbook) -> None:
    ws = wb.create_sheet("Model_OriginalIntel")
    rows = [
        ["Key", "Value", "Unit", "Source"],
        ["q1_promised_opened_arr_eur", lookup_formula("q1_promised_opened_arr_eur"), "EUR", lookup_formula("q1_promised_opened_arr_eur", 4)],
        ["open_land_deals_sidecar", lookup_formula("open_land_deals_sidecar"), "count", lookup_formula("open_land_deals_sidecar", 4)],
        ["open_land_arr_eur_sidecar", lookup_formula("open_land_arr_eur_sidecar"), "EUR", lookup_formula("open_land_arr_eur_sidecar", 4)],
        ["q1_land_wins", lookup_formula("q1_land_wins"), "count", lookup_formula("q1_land_wins", 4)],
        ["q1_land_wins_arr_eur", lookup_formula("q1_land_wins_arr_eur"), "EUR", lookup_formula("q1_land_wins_arr_eur", 4)],
        ["q1_land_lost", lookup_formula("q1_land_lost"), "count", lookup_formula("q1_land_lost", 4)],
        ["q1_land_lost_arr_eur", lookup_formula("q1_land_lost_arr_eur"), "EUR", lookup_formula("q1_land_lost_arr_eur", 4)],
        ["q1_slipped_deals", lookup_formula("q1_slipped_deals"), "count", lookup_formula("q1_slipped_deals", 4)],
        ["q1_slipped_arr_eur", lookup_formula("q1_slipped_arr_eur"), "EUR", lookup_formula("q1_slipped_arr_eur", 4)],
        ["q2_original_arr_eur", lookup_formula("q2_original_arr_eur"), "EUR", lookup_formula("q2_original_arr_eur", 4)],
        ["fy26_renewals_acv_eur_original", lookup_formula("fy26_renewals_acv_eur_original"), "EUR", lookup_formula("fy26_renewals_acv_eur_original", 4)],
    ]
    write_rows(ws, rows)
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 2).number_format = '#,##0.0'


def build_pipeline_total(wb: Workbook) -> None:
    table = "tbl_Raw_Current_Model_Data"
    land_q2 = table_sumifs(table, "ARR_EUR", [("Type", '"Land"'), ("CloseDate", f'">="&{Q2_START}'), ("CloseDate", f'"<"&{Q2_END}')])
    expand_q2 = table_sumifs(table, "ARR_EUR", [("Type", '"Expand"'), ("CloseDate", f'">="&{Q2_START}'), ("CloseDate", f'"<"&{Q2_END}')])
    land_later = table_sumifs(table, "ARR_EUR", [("Type", '"Land"'), ("CloseDate", f'">="&{Q2_END}')])
    expand_later = table_sumifs(table, "ARR_EUR", [("Type", '"Expand"'), ("CloseDate", f'">="&{Q2_END}')])
    renewal_q2 = table_sumifs(table, "ACV_EUR", [("Type", '"Renewal"'), ("CloseDate", f'">="&{Q2_START}'), ("CloseDate", f'"<"&{Q2_END}')])
    rows = [
        ["KPI", "Value (EUR)", "Value (mEUR)", "Source formula"],
        ["2026-Q2 closeable Land+Expand ARR", f"={land_q2}+{expand_q2}", "=B2/1000000", "Type IN (Land, Expand); CloseDate 2026-04-01 to 2026-06-30"],
        ["Open Land+Expand beyond 2026-Q2", f"={land_later}+{expand_later}", "=B3/1000000", "Out-of-quarter context only; not in Q2 forecast"],
        ["2026-Q2 renewal ACV", f"={renewal_q2}", "=B4/1000000", "Type=Renewal only; never added to ARR"],
    ]
    ws = wb.create_sheet("Pipeline_Total")
    write_rows(ws, rows, table=False)
    set_money_format(ws, ["B2", "B3", "B4"])


def build_wins_losses_qtd(wb: Workbook) -> None:
    table = "tbl_Raw_Current_Closed_CFQ"
    won_count = plus([
        table_countifs(table, [("IsWon", "TRUE"), ("Type", '"Land"')]),
        table_countifs(table, [("IsWon", "TRUE"), ("Type", '"Expand"')]),
    ])
    lost_count = plus([
        table_countifs(table, [("IsWon", "FALSE"), ("Type", '"Land"')]),
        table_countifs(table, [("IsWon", "FALSE"), ("Type", '"Expand"')]),
    ])
    won_arr = plus([
        table_sumifs(table, "ARR_EUR", [("IsWon", "TRUE"), ("Type", '"Land"')]),
        table_sumifs(table, "ARR_EUR", [("IsWon", "TRUE"), ("Type", '"Expand"')]),
    ])
    lost_arr = plus([
        table_sumifs(table, "ARR_EUR", [("IsWon", "FALSE"), ("Type", '"Land"')]),
        table_sumifs(table, "ARR_EUR", [("IsWon", "FALSE"), ("Type", '"Expand"')]),
    ])
    won_acv = table_sumifs(table, "ACV_EUR", [("IsWon", "TRUE"), ("Type", '"Renewal"')])
    lost_acv = table_sumifs(table, "ACV_EUR", [("IsWon", "FALSE"), ("Type", '"Renewal"')])
    ws = wb.create_sheet("Wins_Losses_QTD")
    write_rows(
        ws,
        [
            ["Outcome", "Count", "ARR (Land+Expand, EUR)", "ACV (Renewal, EUR)"],
            ["Won", f"={won_count}", f"={won_arr}", f"={won_acv}"],
            ["Lost", f"={lost_count}", f"={lost_arr}", f"={lost_acv}"],
        ],
        table=False,
    )
    set_money_format(ws, ["C2", "C3", "D2", "D3"])


def build_forecast_category(wb: Workbook) -> None:
    table = "tbl_Raw_Current_Model_Data"
    ws = wb.create_sheet("Forecast_Category")
    rows = [["Category", "# Opps", "ARR (EUR)", "Readout"]]
    for category in ["Pipeline", "Best Case", "Commit", "Omitted"]:
        count_formula = plus([
            table_countifs(table, [("Type", '"Land"'), ("ForecastCategoryName", f'"{category}"'), ("CloseDate", f'">="&{Q2_START}'), ("CloseDate", f'"<"&{Q2_END}')]),
            table_countifs(table, [("Type", '"Expand"'), ("ForecastCategoryName", f'"{category}"'), ("CloseDate", f'">="&{Q2_START}'), ("CloseDate", f'"<"&{Q2_END}')]),
        ])
        arr_formula = plus([
            table_sumifs(table, "ARR_EUR", [("Type", '"Land"'), ("ForecastCategoryName", f'"{category}"'), ("CloseDate", f'">="&{Q2_START}'), ("CloseDate", f'"<"&{Q2_END}')]),
            table_sumifs(table, "ARR_EUR", [("Type", '"Expand"'), ("ForecastCategoryName", f'"{category}"'), ("CloseDate", f'">="&{Q2_START}'), ("CloseDate", f'"<"&{Q2_END}')]),
        ])
        rows.append([category, f"={count_formula}", f"={arr_formula}", "Omitted is shown separately, not counted as coverage." if category == "Omitted" else "ARR only."])
    rows.append(["TOTAL", "=SUM(B2:B5)", "=SUM(C2:C5)", "Renewal ACV excluded."])
    write_rows(ws, rows, table=False)
    set_money_format(ws, ["C2", "C3", "C4", "C5", "C6"])


def build_pipe_movement(wb: Workbook) -> None:
    ws = wb.create_sheet("Pipe_Movement")
    write_rows(
        ws,
        [
            ["Bucket", "ARR (EUR)", "Note"],
            ["Opening pipe (start of 2026-Q2)", 0, "No trusted prior Q2 snapshot in this artifact; residual is shown explicitly."],
            ["(+) Creation / advancement residual", "=B6-B2-B4-B5", "Residual needed to reconcile to closing pipe."],
            ["(-) Won this Q (Land+Expand)", "=-Wins_Losses_QTD!C2", "Closed-won is subtracted from open pipeline."],
            ["(-) Closed lost / no opportunity this Q (Land+Expand)", "=-Wins_Losses_QTD!C3", "Closed lost is negative in the bridge."],
            ["Closing pipe (2026-Q2 CFQ)", "=Pipeline_Total!B2", "Closeable Land+Expand ARR only."],
        ],
        table=False,
    )
    set_money_format(ws, ["B2", "B3", "B4", "B5", "B6"])


def build_retention_and_renewals(wb: Workbook) -> None:
    table = "tbl_Raw_Renewals"
    q2_acv = table_sumifs(
        table,
        "ACV Unweighted (EUR)",
        [("Close Date", f'">="&{Q2_START}'), ("Close Date", f'"<"&{Q2_END}')],
    )
    q3_acv = table_sumifs(
        table,
        "ACV Unweighted (EUR)",
        [("Close Date", f'">="&{Q2_END}'), ("Close Date", f'"<"&{Q3_END}')],
    )
    rows = [
        ["Metric", "Value", "Unit", "Source"],
        ["FY26 renewal ACV", f'=SUM({table}[ACV Unweighted (EUR)])', "EUR ACV", "Raw_Renewals"],
        ["Q2 renewal ACV", f"={q2_acv}", "EUR ACV", "Raw_Renewals"],
        ["Q3 renewal ACV", f"={q3_acv}", "EUR ACV", "Raw_Renewals"],
        ["Renewal count", f'=COUNTA({table}[Opportunity])', "count", "Raw_Renewals"],
    ]
    ws = wb.create_sheet("Retention")
    write_rows(ws, rows, table=False)
    ws2 = wb.create_sheet("Model_Renewals")
    write_rows(ws2, rows, table=False)
    for sheet in [ws, ws2]:
        set_money_format(sheet, ["B2", "B3", "B4"])


def build_by_owner(wb: Workbook, current_data_rows: list[dict[str, Any]]) -> None:
    owners = distinct_top(current_data_rows, "OwnerName", "ARR_EUR", {"Land", "Expand"})[:12]
    table = "tbl_Raw_Current_Model_Data"
    rows = [["Owner", "# Open opps", "Open ARR (EUR)", "Action cue"]]
    for owner in owners:
        owner_q = f'"{owner}"'
        count_formula = plus([
            table_countifs(table, [("Type", '"Land"'), ("OwnerName", owner_q)]),
            table_countifs(table, [("Type", '"Expand"'), ("OwnerName", owner_q)]),
        ])
        arr_formula = plus([
            table_sumifs(table, "ARR_EUR", [("Type", '"Land"'), ("OwnerName", owner_q)]),
            table_sumifs(table, "ARR_EUR", [("Type", '"Expand"'), ("OwnerName", owner_q)]),
        ])
        rows.append([owner, f"={count_formula}", f"={arr_formula}", '=IF(C2>3000000,"Capacity/risk review","Coach active Q2 actions")'])
    ws = wb.create_sheet("By_Owner")
    write_rows(ws, rows, table=False)
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 4).value = f'=IF(C{row_idx}>3000000,"Capacity/risk review","Coach active Q2 actions")'
        ws.cell(row_idx, 3).number_format = '#,##0.0,," mEUR"'


def build_weighted_forecast(wb: Workbook) -> None:
    ws = wb.create_sheet("Weighted_Forecast")
    rows = [
        ["Stage", "Open ARR (EUR)", "Forward rate", "Weighted ARR (EUR)", "Rate source"],
        ["3 - Engagement", '=SUMIFS(tbl_Raw_Current_Model_Data[ARR_EUR],tbl_Raw_Current_Model_Data[StageName],"3 - Engagement")', 0.2, "=B2*C2", "Commercial handbook proxy"],
        ["4 - Shortlisted", '=SUMIFS(tbl_Raw_Current_Model_Data[ARR_EUR],tbl_Raw_Current_Model_Data[StageName],"4 - Shortlisted")', 0.4, "=B3*C3", "Commercial handbook proxy"],
        ["5 - Preferred", '=SUMIFS(tbl_Raw_Current_Model_Data[ARR_EUR],tbl_Raw_Current_Model_Data[StageName],"5 - Preferred")', 0.8, "=B4*C4", "Commercial handbook proxy"],
        ["6 - Contracting", '=SUMIFS(tbl_Raw_Current_Model_Data[ARR_EUR],tbl_Raw_Current_Model_Data[StageName],"6 - Contracting")', 0.9, "=B5*C5", "Commercial handbook proxy"],
        ["TOTAL", "=SUM(B2:B5)", "", "=SUM(D2:D5)", "ARR only"],
    ]
    write_rows(ws, rows, table=False)
    set_money_format(ws, ["B2", "B3", "B4", "B5", "B6", "D2", "D3", "D4", "D5", "D6"])


def build_pipeline_creation(wb: Workbook) -> None:
    ws = wb.create_sheet("Pipeline_Creation_Velocity")
    rows: list[list[Any]] = [["Week starting", "# New opps", "New ARR (EUR)", "Source"]]
    start = datetime(2026, 2, 9)
    for idx in range(14):
        week = start + timedelta(days=7 * idx)
        next_week = week + timedelta(days=7)
        date_expr = f"DATE({week.year},{week.month},{week.day})"
        next_expr = f"DATE({next_week.year},{next_week.month},{next_week.day})"
        count_formula = plus([
            table_countifs("tbl_Raw_Current_Model_Data", [("Type", '"Land"'), ("CreatedDate", f'">="&{date_expr}'), ("CreatedDate", f'"<"&{next_expr}')]),
            table_countifs("tbl_Raw_Current_Model_Data", [("Type", '"Expand"'), ("CreatedDate", f'">="&{date_expr}'), ("CreatedDate", f'"<"&{next_expr}')]),
        ])
        arr_formula = plus([
            table_sumifs("tbl_Raw_Current_Model_Data", "ARR_EUR", [("Type", '"Land"'), ("CreatedDate", f'">="&{date_expr}'), ("CreatedDate", f'"<"&{next_expr}')]),
            table_sumifs("tbl_Raw_Current_Model_Data", "ARR_EUR", [("Type", '"Expand"'), ("CreatedDate", f'">="&{date_expr}'), ("CreatedDate", f'"<"&{next_expr}')]),
        ])
        rows.append([week, f"={count_formula}", f"={arr_formula}", "CreatedDate from current Salesforce model extract"])
    write_rows(ws, rows, table=False)
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 1).number_format = "yyyy-mm-dd"
        ws.cell(row_idx, 3).number_format = '#,##0.0,," mEUR"'


def build_account_expansion(wb: Workbook, current_data_rows: list[dict[str, Any]]) -> None:
    accounts = distinct_top(current_data_rows, "AccountName", "ARR_EUR", {"Land", "Expand"})[:15]
    rows = [["#", "Account", "Land ARR (EUR)", "Expand ARR (EUR)", "Renewal ACV (EUR)", "# Motions"]]
    table = "tbl_Raw_Current_Model_Data"
    for idx, account in enumerate(accounts, start=1):
        account_q = f'"{account}"'
        land_arr = table_sumifs(table, "ARR_EUR", [("AccountName", account_q), ("Type", '"Land"')])
        expand_arr = table_sumifs(table, "ARR_EUR", [("AccountName", account_q), ("Type", '"Expand"')])
        renewal_acv = table_sumifs(table, "ACV_EUR", [("AccountName", account_q), ("Type", '"Renewal"')])
        rows.append(
            [
                idx,
                account,
                f"={land_arr}",
                f"={expand_arr}",
                f"={renewal_acv}",
                f'=IF(C{idx+1}>0,1,0)+IF(D{idx+1}>0,1,0)+IF(E{idx+1}>0,1,0)',
            ]
        )
    ws = wb.create_sheet("Account_Expansion")
    write_rows(ws, rows, table=False)
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in [3, 4, 5]:
            ws.cell(row_idx, col_idx).number_format = '#,##0.0,," mEUR"'


def build_custom_models(wb: Workbook) -> None:
    model_rows = {
        "Model_Q1_Accountability": [
            ["Metric", "Value", "Unit", "Source", "Action"],
            ["Q1 opened / promised", lookup_formula("q1_promised_opened_arr_eur"), "EUR ARR", lookup_formula("q1_promised_opened_arr_eur", 4), "Use as original target baseline."],
            ["Q1 Land wins", lookup_formula("q1_land_wins"), "count", lookup_formula("q1_land_wins", 4), "Show delivered count."],
            ["Q1 Land won ARR", lookup_formula("q1_land_wins_arr_eur"), "EUR ARR", lookup_formula("q1_land_wins_arr_eur", 4), "Delivered."],
            ["Q1 Land losses", lookup_formula("q1_land_lost"), "count", lookup_formula("q1_land_lost", 4), "Loss count must stay visible."],
            ["Q1 Land lost ARR", lookup_formula("q1_land_lost_arr_eur"), "EUR ARR", lookup_formula("q1_land_lost_arr_eur", 4), "Do not add; this is negative accountability."],
            ["Q1 slipped deals", lookup_formula("q1_slipped_deals"), "count", lookup_formula("q1_slipped_deals", 4), "Slippage, not coverage."],
            ["Q1 slipped ARR", lookup_formula("q1_slipped_arr_eur"), "EUR ARR", lookup_formula("q1_slipped_arr_eur", 4), "Slippage, not creation."],
        ],
        "Model_LossDrivers": [
            ["Reason / bucket", "Count", "ARR (EUR)", "Evidence", "Director action"],
            ["Q1 Land losses", lookup_formula("q1_land_lost"), lookup_formula("q1_land_lost_arr_eur"), "Original sidecar", "Separate from closed-won."],
            ["Missing reason code", '=COUNTIFS(tbl_Raw_Won_Lost[Type],"Land",tbl_Raw_Won_Lost[Reason],"")', '=SUMIFS(tbl_Raw_Won_Lost[ARR Unweighted (EUR)],tbl_Raw_Won_Lost[Type],"Land",tbl_Raw_Won_Lost[Reason],"")', "Raw_Won_Lost Reason blank", "Repair Salesforce hygiene."],
            ["No Opportunity / Lost stage", '=COUNTIFS(tbl_Raw_Won_Lost[Type],"Land",tbl_Raw_Won_Lost[Stage],"0 - No Opportunity")', '=SUMIFS(tbl_Raw_Won_Lost[ARR Unweighted (EUR)],tbl_Raw_Won_Lost[Type],"Land",tbl_Raw_Won_Lost[Stage],"0 - No Opportunity")', "Raw_Won_Lost Stage", "Review qualification exits."],
            ["Buying process stopped", '=COUNTIFS(tbl_Raw_Won_Lost[Reason],"Buying process stopped by prospect/customer")', '=SUMIFS(tbl_Raw_Won_Lost[ARR Unweighted (EUR)],tbl_Raw_Won_Lost[Reason],"Buying process stopped by prospect/customer")', "Raw_Won_Lost Reason", "Confirm if reopenable."],
        ],
        "Model_Q2_Readiness": [
            ["#", "Account", "Opportunity", "Owner", "Type", "Stage", "Close Date", "ARR", "Forecast", "Readiness"],
        ],
        "Model_DealRisk": [
            ["Account", "Opportunity", "Owner", "Stage", "Forecast", "ARR", "Risk reason"],
        ],
        "Model_OwnerCoaching": [
            ["Owner", "# Open opps", "Open ARR", "Action cue", "Original director risk", "Source"],
        ],
        "Model_PushedDeals": [
            ["Account", "Opportunity", "Owner", "Stage", "Movement", "Old Close", "New Close", "ARR (EUR)"],
        ],
        "Model_CommercialApproval": [
            ["#", "Account", "Opportunity", "Owner", "Stage", "Close Date", "Type", "ARR (EUR)", "Action"],
        ],
        "Model_ActionContract": [
            ["#", "Priority", "Rule", "Claim", "Suggested action", "Due date", "Owner"],
        ],
        "Trend_QoQ": [
            ["Quarter", "# Won", "Booked ARR (EUR)", "Delta QoQ", "Source"],
            ["2026-Q2", "=Wins_Losses_QTD!B2", "=Wins_Losses_QTD!C2", "", "Current CFQ closed-won"],
            ["Q1 original wins", lookup_formula("q1_land_wins"), lookup_formula("q1_land_wins_arr_eur"), "", "Original sidecar"],
        ],
    }
    for title, rows in model_rows.items():
        ws = wb.create_sheet(title)
        write_rows(ws, rows, table=False)

    # Formula-populate repeated model tables.
    for row_idx in range(2, 11):
        n = row_idx - 1
        ws = wb["Model_Q2_Readiness"]
        formulas = [
            n,
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Account],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Opportunity],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Owner],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Type],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Stage],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Close Date],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[ARR],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Forecast],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Readiness],{n}),"")',
        ]
        for col_idx, value in enumerate(formulas, start=1):
            ws.cell(row_idx, col_idx).value = value

    for row_idx in range(2, 9):
        n = row_idx - 1
        ws = wb["Model_DealRisk"]
        formulas = [
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Account],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Opportunity],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Owner],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Stage],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Forecast],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[ARR],{n}),"")',
            f'=IFERROR(INDEX(tbl_Raw_Current_Q2_Readiness[Readiness],{n}),"")',
        ]
        for col_idx, value in enumerate(formulas, start=1):
            ws.cell(row_idx, col_idx).value = value

    for row_idx in range(2, 9):
        n = row_idx - 1
        ws = wb["Model_PushedDeals"]
        for col_idx, col_letter in enumerate(["A", "B", "C", "D", "E", "F", "G", "I"], start=1):
            ws.cell(row_idx, col_idx).value = f'=IFERROR(Raw_CloseDate_History!{col_letter}{row_idx},"")'

    for row_idx in range(2, 11):
        n = row_idx - 1
        ws = wb["Model_CommercialApproval"]
        columns = ["#", "Account", "Opportunity", "Owner", "Stage", "Close Date", "Type", "ARR (EUR)"]
        for col_idx, col_name in enumerate(columns, start=1):
            if col_name == "#":
                ws.cell(row_idx, col_idx).value = f'=IFERROR(Raw_Current_Pending_Approval!A{row_idx},"")'
            else:
                ws.cell(row_idx, col_idx).value = f'=IFERROR(INDEX(tbl_Raw_Current_Pending_Approval[{col_name}],{n}),"")'
        ws.cell(row_idx, 9).value = '=IF(B{0}<>"","Submit/clear approval gate","")'.format(row_idx)

    for row_idx in range(2, 9):
        n = row_idx - 1
        ws = wb["Model_ActionContract"]
        for col_idx, col_name in enumerate(["#", "Priority", "Rule", "Claim", "Suggested action", "Due date", "Owner"], start=1):
            if col_name == "#":
                ws.cell(row_idx, col_idx).value = f'=IFERROR(Raw_Current_Action_Items!A{row_idx},"")'
            else:
                ws.cell(row_idx, col_idx).value = f'=IFERROR(INDEX(tbl_Raw_Current_Action_Items[{col_name}],{n}),"")'

    owner_ws = wb["Model_OwnerCoaching"]
    for row_idx in range(2, 8):
        owner_ws.cell(row_idx, 1).value = f'=IFERROR(By_Owner!A{row_idx},"")'
        owner_ws.cell(row_idx, 2).value = f'=IFERROR(By_Owner!B{row_idx},"")'
        owner_ws.cell(row_idx, 3).value = f'=IFERROR(By_Owner!C{row_idx},"")'
        owner_ws.cell(row_idx, 4).value = f'=IFERROR(By_Owner!D{row_idx},"")'
        owner_ws.cell(row_idx, 5).value = ""
        owner_ws.cell(row_idx, 6).value = "By_Owner + original director intel"
    owner_ws["E2"] = '=TEXT(VLOOKUP("owner_push_count",tbl_Raw_Original_Intel,2,FALSE),"0")&" pushes / "&TEXT(VLOOKUP("owner_push_arr_eur",tbl_Raw_Original_Intel,2,FALSE)/1000000,"0.0")&"m EUR"'

    for ws in wb.worksheets:
        if ws.title.startswith("Model_") or ws.title in {"Trend_QoQ"}:
            style_range(ws)


def build_output_sheets(wb: Workbook, spec: dict[str, Any]) -> None:
    outputs: dict[str, list[list[Any]]] = {
        "Out_S02_ExecutiveSummary": [
            ["Metric", "Value", "Unit", "Source"],
            ["2026-Q2 closeable Land+Expand unweighted ARR", "=Pipeline_Total!B2", "EUR unweighted ARR", "Current model formula"],
            ["QTD won / closed-lost unweighted ARR", '=TEXT(Wins_Losses_QTD!C2/1000000,"0.0")&" / "&TEXT(Wins_Losses_QTD!C3/1000000,"0.0")', "mEUR unweighted ARR", "Wins_Losses_QTD"],
            ["FY26 renewal ACV", "=Model_Renewals!B2", "EUR ACV", "Renewal ACV separate"],
            ["May 1 action posture", "=Model_ActionContract!D2", "text", "Action_Items"],
        ],
        "Out_S04_Q1Accountability": [
            ["Q1 accountability lens", "Original director baseline", "Current proof", "Director action"],
            ["Opened / promised", "=Model_Q1_Accountability!B2", "Original deck commentary", "Use as baseline, not current pipeline."],
            ["Won", '=Model_Q1_Accountability!B3&" / "&TEXT(Model_Q1_Accountability!B4/1000000,"0.0")&"m"', "Original sidecar", "Show delivered."],
            ["Lost", '=Model_Q1_Accountability!B5&" / "&TEXT(Model_Q1_Accountability!B6/1000000,"0.0")&"m"', "Original sidecar", "Keep loss accountability visible."],
            ["Slipped", '=Model_Q1_Accountability!B7&" / "&TEXT(Model_Q1_Accountability!B8/1000000,"0.0")&"m"', "Original deck commentary", "Separate slippage from creation."],
            ["Q2 close focus", '=VLOOKUP("q2_original_deals",tbl_Raw_Original_Intel,2,FALSE)&" / "&TEXT(VLOOKUP("q2_original_arr_eur",tbl_Raw_Original_Intel,2,FALSE)/1000000,"0.0")&"m"', "Original director deck", "Orient May execution."],
            ["Concentration", '=TEXT(VLOOKUP("concentration_top5_share",tbl_Raw_Original_Intel,2,FALSE),"0%")&" top-5 share"', "Original director deck", "Name account exposure."],
            ["Renewals", '=VLOOKUP("fy26_renewals_count_original",tbl_Raw_Original_Intel,2,FALSE)&" / "&TEXT(VLOOKUP("fy26_renewals_acv_eur_original",tbl_Raw_Original_Intel,2,FALSE)/1000000,"0.0")&"m ACV"', "Original director deck", "Do not blend into ARR."],
        ],
        "Out_S05_PipelineMovement": [
            ["Bucket", "Opening pipe", "+ Creation / advancement", "Won", "Closed lost", "Closing pipe"],
            ["Unweighted ARR (EUR)", "=Pipe_Movement!B2", "=Pipe_Movement!B3", "=-Wins_Losses_QTD!C2", "=-Wins_Losses_QTD!C3", "=Pipe_Movement!B6"],
            ["Sign rule", "start", "positive residual", "negative", "negative", "end"],
        ],
        "Out_S06_LossDrivers": [
            ["Reason / bucket", "Count", "Unweighted ARR (EUR)", "Evidence", "Director action"],
            ["Q1 Land losses", "=Model_LossDrivers!B2", "=Model_LossDrivers!C2", "=Model_LossDrivers!D2", "=Model_LossDrivers!E2"],
            ["Missing reason code", "=Model_LossDrivers!B3", "=Model_LossDrivers!C3", "=Model_LossDrivers!D3", "=Model_LossDrivers!E3"],
            ["No Opportunity / Lost stage", "=Model_LossDrivers!B4", "=Model_LossDrivers!C4", "=Model_LossDrivers!D4", "=Model_LossDrivers!E4"],
            ["Buying process stopped", "=Model_LossDrivers!B5", "=Model_LossDrivers!C5", "=Model_LossDrivers!D5", "=Model_LossDrivers!E5"],
            ["", "", "", "", ""],
            ["", "", "", "", ""],
            ["", "", "", "", ""],
        ],
        "Out_S07_Q2Readiness": [["#", "Account", "Opportunity", "Owner", "Type", "Stage", "Close Date", "Unweighted ARR (EUR)", "Forecast", "Readiness"]],
        "Out_S08_ForecastCategory": [
            ["Category", "# Opps", "Unweighted ARR (EUR)", "Readout"],
            ["Pipeline", "=Forecast_Category!B2", "=Forecast_Category!C2", "=Forecast_Category!D2"],
            ["Best Case", "=Forecast_Category!B3", "=Forecast_Category!C3", "=Forecast_Category!D3"],
            ["Commit", "=Forecast_Category!B4", "=Forecast_Category!C4", "=Forecast_Category!D4"],
            ["Omitted", "=Forecast_Category!B5", "=Forecast_Category!C5", "=Forecast_Category!D5"],
            ["TOTAL", "=Forecast_Category!B6", "=Forecast_Category!C6", "=Forecast_Category!D6"],
        ],
        "Out_S09_DealRisk": [["Account", "Opportunity", "Owner", "Stage", "Forecast", "Unweighted ARR (EUR)", "Risk reason"]],
        "Out_S10_OwnerCoaching": [["Owner", "# Open opps", "Open unweighted ARR (EUR)", "Action cue", "Original director risk", "Source"]],
        "Out_S11_PushedDeals": [["Account", "Opportunity", "Owner", "Stage", "Movement", "Old Close", "New Close", "Unweighted ARR (EUR)"]],
        "Out_S12_CommercialApprovals": [["#", "Account", "Opportunity", "Owner", "Stage", "Close Date", "Type", "Unweighted ARR (EUR)"]],
        "Out_S13_Renewals": [["#", "Close Date", "Account", "Opportunity", "Owner", "Stage", "ACV", "Prob", "Risk / note"]],
        "Out_S14_PipelineCreation": [["Week starting", "# New opps", "New unweighted ARR (EUR)", "Source", "Director readout", "Action"]],
        "Out_S15_ActionContract": [["#", "Priority", "Rule", "Claim", "Suggested action", "Owner"]],
    }

    for row_idx in range(2, 11):
        n = row_idx - 1
        outputs["Out_S07_Q2Readiness"].append([f"=Model_Q2_Readiness!A{row_idx}", f"=Model_Q2_Readiness!B{row_idx}", f"=Model_Q2_Readiness!C{row_idx}", f"=Model_Q2_Readiness!D{row_idx}", f"=Model_Q2_Readiness!E{row_idx}", f"=Model_Q2_Readiness!F{row_idx}", f"=Model_Q2_Readiness!G{row_idx}", f"=Model_Q2_Readiness!H{row_idx}", f"=Model_Q2_Readiness!I{row_idx}", f"=Model_Q2_Readiness!J{row_idx}"])
        if n <= 7:
            outputs["Out_S09_DealRisk"].append([f"=Model_DealRisk!A{row_idx}", f"=Model_DealRisk!B{row_idx}", f"=Model_DealRisk!C{row_idx}", f"=Model_DealRisk!D{row_idx}", f"=Model_DealRisk!E{row_idx}", f"=Model_DealRisk!F{row_idx}", f"=Model_DealRisk!G{row_idx}"])
            outputs["Out_S10_OwnerCoaching"].append([f"=Model_OwnerCoaching!A{row_idx}", f"=Model_OwnerCoaching!B{row_idx}", f"=Model_OwnerCoaching!C{row_idx}", f"=Model_OwnerCoaching!D{row_idx}", f"=Model_OwnerCoaching!E{row_idx}", f"=Model_OwnerCoaching!F{row_idx}"])
            outputs["Out_S11_PushedDeals"].append([f"=Model_PushedDeals!A{row_idx}", f"=Model_PushedDeals!B{row_idx}", f"=Model_PushedDeals!C{row_idx}", f"=Model_PushedDeals!D{row_idx}", f"=Model_PushedDeals!E{row_idx}", f"=Model_PushedDeals!F{row_idx}", f"=Model_PushedDeals!G{row_idx}", f"=Model_PushedDeals!H{row_idx}"])
            outputs["Out_S12_CommercialApprovals"].append([f"=Model_CommercialApproval!A{row_idx}", f"=Model_CommercialApproval!B{row_idx}", f"=Model_CommercialApproval!C{row_idx}", f"=Model_CommercialApproval!D{row_idx}", f"=Model_CommercialApproval!E{row_idx}", f"=Model_CommercialApproval!F{row_idx}", f"=Model_CommercialApproval!G{row_idx}", f"=Model_CommercialApproval!H{row_idx}"])

    for row_idx in range(2, 7):
        outputs["Out_S13_Renewals"].append([row_idx - 1, f'=IFERROR(INDEX(tbl_Raw_Renewals[Close Date],{row_idx-1}),"")', f'=IFERROR(INDEX(tbl_Raw_Renewals[Account],{row_idx-1}),"")', f'=IFERROR(INDEX(tbl_Raw_Renewals[Opportunity],{row_idx-1}),"")', f'=IFERROR(INDEX(tbl_Raw_Renewals[Owner],{row_idx-1}),"")', f'=IFERROR(INDEX(tbl_Raw_Renewals[Stage],{row_idx-1}),"")', f'=IFERROR(INDEX(tbl_Raw_Renewals[ACV Unweighted (EUR)],{row_idx-1}),"")', f'=IFERROR(INDEX(tbl_Raw_Renewals[Probability %],{row_idx-1}),"")', "Renewal ACV only"])

    for row_idx in range(2, 9):
        outputs["Out_S14_PipelineCreation"].append([f"=Pipeline_Creation_Velocity!A{row_idx}", f"=Pipeline_Creation_Velocity!B{row_idx}", f"=Pipeline_Creation_Velocity!C{row_idx}", f"=Pipeline_Creation_Velocity!D{row_idx}", '=IF(C{0}=0,"No new unweighted ARR that week","Created unweighted ARR")'.format(row_idx), "Review creation pace"])
        outputs["Out_S15_ActionContract"].append([f"=Model_ActionContract!A{row_idx}", f"=Model_ActionContract!B{row_idx}", f"=Model_ActionContract!C{row_idx}", f"=Model_ActionContract!D{row_idx}", f"=Model_ActionContract!E{row_idx}", f"=Model_ActionContract!G{row_idx}"])

    for obj in spec["objects"]:
        title = obj["output_sheet_target"]
        ws = wb.create_sheet(title)
        write_rows(ws, outputs[title], table=False)
        add_named_range(wb, obj["excel_name_target"], obj["output_range_target"])
        # A named output table is useful for human audit; think-cell still uses the named range.
        ws["A1"].comment = None
        if "ARR" in str(ws["A1"].value) or "ACV" in str(ws["A1"].value):
            pass

    for sheet_name in ["Out_S05_PipelineMovement"]:
        ws = wb[sheet_name]
        ws["E2"].fill = OK_FILL
        ws["E3"].fill = OK_FILL


def build_factory_index(wb: Workbook, spec: dict[str, Any]) -> None:
    ws = wb.create_sheet("Factory_Index", 0)
    rows = [["Object ID", "Output sheet", "Output range", "Excel name", "PPT name", "think-cell status", "Validation rules"]]
    for obj in spec["objects"]:
        rows.append(
            [
                obj["object_id"],
                obj["output_sheet_target"],
                obj["output_range_target"],
                obj["excel_name_target"],
                obj["ppt_name"],
                obj["thinkcell_status"],
                "; ".join(obj["validation_rules"]),
            ]
        )
    write_rows(ws, rows, table=True)
    for row_idx in range(2, ws.max_row + 1):
        status = str(ws.cell(row_idx, 6).value)
        ws.cell(row_idx, 6).fill = WARN_FILL if "blocked" in status else OK_FILL


def metric_basis_for_object(obj: dict[str, Any]) -> str:
    text = " ".join(
        str(obj.get(key, ""))
        for key in (
            "object_id",
            "object_type",
            "intelligence_role",
            "visualization_lane",
            "output_range_target",
        )
    ).lower()
    if "renewal" in text or "acv" in text:
        return "Renewal ACV only; never blended with Land+Expand unweighted ARR."
    if "weighted" in text:
        return "Explicit weighted ARR where labeled; otherwise Land+Expand unweighted ARR."
    if "arr" in text or "eur" in text or "pipeline" in text or "deal" in text or "owner" in text or "loss" in text:
        return "Land+Expand unweighted ARR in EUR unless a field explicitly says weighted ARR."
    return "Non-currency / text control."


def build_thinkcell_link_map(wb: Workbook, spec: dict[str, Any]) -> None:
    ws = wb.create_sheet("ThinkCell_Link_Map", 2)
    rows = [
        [
            "Slide",
            "Slide title",
            "Object ID",
            "PowerPoint / think-cell name",
            "Excel defined name",
            "Excel output range",
            "Output sheet",
            "Object type",
            "Metric basis",
            "Raw source tabs",
            "Model tabs",
            "think-cell status",
            "Current publish path",
            "Validation rules",
        ]
    ]
    for obj in spec["objects"]:
        rows.append(
            [
                obj.get("slide_id", ""),
                obj.get("slide_title", ""),
                obj.get("object_id", ""),
                obj.get("ppt_name", ""),
                obj.get("excel_name_target", ""),
                obj.get("output_range_target", ""),
                obj.get("output_sheet_target", ""),
                obj.get("object_type", ""),
                metric_basis_for_object(obj),
                "; ".join(obj.get("raw_tabs", [])),
                "; ".join(obj.get("model_tabs", [])),
                obj.get("thinkcell_status", ""),
                obj.get("current_status", ""),
                "; ".join(obj.get("validation_rules", [])),
            ]
        )
    write_rows(ws, rows, table=True)
    for row_idx in range(2, ws.max_row + 1):
        status = str(ws.cell(row_idx, 12).value)
        ws.cell(row_idx, 12).fill = WARN_FILL if "blocked" in status else OK_FILL
    add_named_range(wb, "rng_ThinkCell_Link_Map", f"ThinkCell_Link_Map!A1:N{ws.max_row}")


def build_audit_checks(wb: Workbook, spec: dict[str, Any], sidecar: dict[str, Any]) -> None:
    ws = wb.create_sheet("Audit_Checks", 1)
    rows = [
        ["Check", "Expected", "Workbook formula / source", "Status"],
        ["Original sidecar open Land deals", _sidecar_number(sidecar, "open_land_deals"), lookup_formula("open_land_deals_sidecar"), '=IF(B2=C2,"PASS","REVIEW")'],
        ["Original sidecar open Land ARR", _sidecar_number(sidecar, "open_land_arr"), lookup_formula("open_land_arr_eur_sidecar"), '=IF(ABS(B3-C3)<1,"PASS","REVIEW")'],
        ["Original Q1 Land wins", _sidecar_number(sidecar, "q1_land_wins"), lookup_formula("q1_land_wins"), '=IF(B4=C4,"PASS","REVIEW")'],
        ["Original Q1 Land won ARR", _sidecar_number(sidecar, "q1_land_wins_arr"), lookup_formula("q1_land_wins_arr_eur"), '=IF(ABS(B5-C5)<1,"PASS","REVIEW")'],
        ["Original Q1 Land lost", _sidecar_number(sidecar, "q1_land_lost"), lookup_formula("q1_land_lost"), '=IF(B6=C6,"PASS","REVIEW")'],
        ["Original Q1 Land lost ARR", _sidecar_number(sidecar, "q1_land_lost_arr"), lookup_formula("q1_land_lost_arr_eur"), '=IF(ABS(B7-C7)<1,"PASS","REVIEW")'],
        ["Original Q2 renewal ACV", _sidecar_number(sidecar, "q2_renewals_acv"), lookup_formula("q2_renewals_acv_eur"), '=IF(ABS(B8-C8)<1,"PASS","REVIEW")'],
        ["Original Q3 renewal ACV", _sidecar_number(sidecar, "q3_renewals_acv"), lookup_formula("q3_renewals_acv_eur"), '=IF(ABS(B9-C9)<1,"PASS","REVIEW")'],
        ["Bridge closed-lost sign", "negative", '=IF(Out_S05_PipelineMovement!E2<0,"negative","not negative")', '=IF(B10=C10,"PASS","FAIL")'],
        ["ARR/ACV separation", "separate", '=IF(ISNUMBER(SEARCH("ACV",Out_S13_Renewals!G1)),"separate","mixed")', '=IF(B11=C11,"PASS","FAIL")'],
        ["think-cell link map rows", len(spec["objects"]), '=COUNTA(ThinkCell_Link_Map!C:C)-1', '=IF(B12=C12,"PASS","FAIL")'],
        ["ARR basis is explicit in output headers", "explicit", '=IF(COUNTIF(Out_S07_Q2Readiness!1:1,"*Unweighted ARR*")>0,"explicit","missing")', '=IF(B13=C13,"PASS","FAIL")'],
    ]
    write_rows(ws, rows, table=True)
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 4).fill = OK_FILL


def build_connected_factory(ctx: FactoryContext, spec: dict[str, Any], sidecar: dict[str, Any]) -> Path:
    missing = [
        path
        for path in (
            ctx.original_workbook,
            ctx.original_sidecar,
            ctx.current_companion_workbook,
            ctx.current_model_workbook,
        )
        if not path.exists()
    ]
    if missing:
        raise FileNotFoundError("missing factory input(s): " + ", ".join(str(path) for path in missing))

    original = load_workbook(ctx.original_workbook, data_only=False)
    current_companion = load_workbook(ctx.current_companion_workbook, data_only=False)
    current_model = load_workbook(ctx.current_model_workbook, data_only=False)

    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    current_data_rows = load_sheet_rows(ctx.current_model_workbook, "Data")

    # Raw/source layer.
    source_map = {
        "Raw_Pipeline_Open": (original, "Pipeline Open FY26", 1),
        "Raw_Won_Lost": (original, "Won Lost FY26", 1),
        "Raw_Commercial_Approval": (original, "Commercial Approval", 1),
        "Raw_Approvals": (original, "Commercial Approval", 1),
        "Raw_Renewals": (original, "Renewals FY26", 1),
        "Raw_Pipeline_Inspection": (original, "Pipeline Inspection", 1),
        "Raw_Activity": (original, "Activity Volume", 1),
        "Raw_CloseDate_History": (original, "Q1 Movement", 1),
        "Raw_Q1_Movement": (original, "Q1 Movement", 1),
        "Raw_Stage_History": (original, "Stage History", 1),
        "Raw_Forecast_Category_History": (original, "Forecast Category History", 1),
        "Raw_Forecast_Items": (original, "Commit Items", 1),
        "Raw_Current_Action_Items": (current_companion, "Action_Items", 1),
        "Raw_Current_Pending_Approval": (current_companion, "Pending_Commercial_Approval", 3),
        "Raw_Current_Model_Data": (current_model, "Data", 1),
        "Raw_Current_Closed_CFQ": (current_model, "Data_Closed_CFQ", 1),
        "Raw_Current_Closed_Won_6mo": (current_model, "Data_Closed_Won_6mo", 1),
        "Raw_Current_Renewals_12mo": (current_model, "Data_Renewals_12mo", 1),
    }
    for target, (src_wb, source, header_row) in source_map.items():
        copy_sheet(wb, src_wb, source, target, header_row=header_row, table=True)
    if "Q2_Readiness" in current_companion.sheetnames:
        copy_sheet(wb, current_companion, "Q2_Readiness", "Raw_Current_Q2_Readiness", table=True)
    else:
        build_raw_current_q2_readiness(wb, current_data_rows)
    if "FY26_Renewals" in current_companion.sheetnames:
        copy_sheet(wb, current_companion, "FY26_Renewals", "Raw_Current_FY26_Renewals", table=True)
    else:
        build_raw_current_fy26_renewals(wb, current_data_rows)

    build_raw_original_intel(wb, spec, sidecar)
    build_raw_accounts(wb, load_sheet_rows(ctx.original_workbook, "Pipeline Open FY26"))

    # Formula/model layer.
    build_model_original_intel(wb)
    build_pipeline_total(wb)
    build_wins_losses_qtd(wb)
    build_forecast_category(wb)
    build_pipe_movement(wb)
    build_retention_and_renewals(wb)
    build_by_owner(wb, current_data_rows)
    build_weighted_forecast(wb)
    build_pipeline_creation(wb)
    build_account_expansion(wb, current_data_rows)
    build_custom_models(wb)

    # Output/named-range layer.
    build_output_sheets(wb, spec)
    build_factory_index(wb, spec)
    build_thinkcell_link_map(wb, spec)
    build_audit_checks(wb, spec, sidecar)

    # Put audit/index up front after all sheets are created.
    wb._sheets.sort(key=lambda ws: {"Factory_Index": 0, "Audit_Checks": 1, "ThinkCell_Link_Map": 2}.get(ws.title, 10))

    ctx.output_workbook.parent.mkdir(parents=True, exist_ok=True)
    ctx.generated_spec.write_text(json.dumps(spec, indent=2) + "\n", encoding="utf-8")
    wb.save(ctx.output_workbook)
    print(f"Wrote {ctx.output_workbook}")
    print(f"Wrote {ctx.generated_spec}")
    print(f"Defined output names: {len(spec['objects'])}")
    print("Formula recalc mode: auto/fullCalcOnLoad")
    return ctx.output_workbook


def selected_directors(args: argparse.Namespace) -> list[dict[str, Any]]:
    if args.all_directors:
        return canonical_directors()
    return [resolve_director(args.director)]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--director", default="Jesper Tyrer", help="Director display name or slug.")
    group.add_argument("--all-directors", action="store_true", help="Build connected factories for all canonical directors.")
    parser.add_argument("--period", default="2026-Q2")
    parser.add_argument("--base-spec", type=Path, default=BASE_SPEC_PATH)
    args = parser.parse_args()

    base_spec = json.loads(args.base_spec.read_text(encoding="utf-8"))
    built: list[Path] = []
    failures: list[str] = []
    for director in selected_directors(args):
        ctx = factory_context(director, args.period)
        try:
            sidecar = json.loads(ctx.original_sidecar.read_text(encoding="utf-8"))
            spec = generated_spec(base_spec, ctx, sidecar)
            built.append(build_connected_factory(ctx, spec, sidecar))
        except Exception as exc:
            failures.append(f"{ctx.director}: {type(exc).__name__}: {exc}")

    if failures:
        print("FAIL")
        for failure in failures:
            print(f"  - {failure}")
        return 1
    print("connected_factory_outputs:")
    for path in built:
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
