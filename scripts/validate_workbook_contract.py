"""Workbook contract validator — assert that land.model.xlsx and land.xlsx
have the sheet/header/range structure that the deck template binds to.

Codex review 2026-04-30 finding #5: template/runbook drifted from the
live workbook on slides 12 / 17 / 18 / 22 / 23 / 24. This script is the
CI gate that catches that drift early.

CLI:
    python3 scripts/validate_workbook_contract.py state/2026-Q2/Jesper-Tyrer
    # exits 0 if all contracts hold; 1 if any drift is detected.

The CONTRACTS dict below is the canonical truth: when the model.xlsx or
the deck slide-by-slide table changes, update both sides AND this file
in lockstep.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# (sheet_name, expected_header_row, expected_columns_left_to_right) — model.xlsx
MODEL_CONTRACTS: list[tuple[str, int, list[str]]] = [
    ("Pipeline_Total", 1, ["KPI", "Value (EUR)", "Value (mEUR)", "Source formula"]),
    (
        "Pipeline_By_Stage",
        1,
        [None, "ARR (EUR)", "ARR (mEUR)", "# Opps"],
    ),  # col A is dynamic period
    (
        "Pipeline_Aging",
        1,
        ["Age bucket (days since CreatedDate)", "Min", "Max", "ARR (EUR)", "# Opps"],
    ),
    ("By_Owner", 1, ["Owner", "Open ARR (Land+Expand, all CloseDates)", "# Opps"]),
    (
        "Velocity",
        1,
        ["Stage", "# Open opps", "Avg age (days)", "# > 60 days old", "# > 180 days old"],
    ),
    ("Concentration", 11, ["Slice", "ARR (EUR)", "% of total"]),  # header on row 11 (after intro)
    (
        "Weighted_Forecast",
        1,
        ["Stage", "Open ARR (EUR)", "Forward rate", "Weighted ARR (EUR)", "Rate source"],
    ),
    ("Pipe_Movement", 1, ["Bucket", "ARR (EUR)", "Note"]),
    ("ARR_Roll", 1, ["Month", "Booked ARR (EUR)", "# Won"]),
    ("Trend_MoM", 1, ["Month", "# Won", "Booked ARR (EUR)", "Δ MoM (Booked ARR)"]),
    ("Trend_QoQ", 1, ["Quarter", "# Won", "Booked ARR (EUR)", "Δ QoQ (Booked ARR)"]),
    ("Retention", 1, ["Metric", "Value"]),
    ("Wins_Losses_QTD", 1, ["Outcome", "#", "ARR (Land+Expand, EUR)", "ACV (Renewal, EUR)"]),
    ("Forecast_Category", 1, ["Category", "# Opps", "ARR (EUR)"]),
    ("Top_Accounts", 1, ["#", "Account", "# Opps", "ARR (EUR)"]),
    ("Territory_Performance", 1, ["#", "Country", "# Opps", "Open ARR (EUR)"]),
    ("Sales_Velocity", 1, ["Metric", "Value", "Note"]),
    (
        "Account_Expansion",
        1,
        ["#", "Account", "Land ARR (EUR)", "Expand ARR (EUR)", "Renewal ACV (EUR)", "# Motions"],
    ),
    ("Pipeline_Creation_Velocity", 1, ["Week starting", "# New opps", "New ARR (EUR)"]),
    ("Stale_Activity", 1, ["Stage", "# Stale opps", "ARR (EUR)"]),
]

# Excel tables that downstream pivots / SUMIFS depend on.
REQUIRED_TABLES = ["tblData", "tblClosedCFQ", "tblClosedWon6mo", "tblRenewals12mo"]

# Named ranges that downstream formulas reference.
REQUIRED_NAMED_RANGES = [
    "period_start",
    "period_end",
    "today",
    "eur_to_meur",
    "Data_ARR_EUR",
    "Data_ACV_EUR",
    "Data_Type",
    "Data_StageName",
    "Data_AccountName",
    "Data_OwnerName",
    "Data_BillingCountry",
    "Data_Industry",
    "Data_CloseDate",
    "Data_CreatedDate",
    "ClosedCFQ_IsWon",
    "ClosedCFQ_Type",
    "ClosedCFQ_ARR_EUR",
    "ClosedCFQ_ACV_EUR",
    "ClosedWon6mo_ARR_EUR",
    "ClosedWon6mo_CloseDate",
    "ClosedWon6mo_CreatedDate",
    "Renewals12mo_ACV_EUR",
    "Renewals12mo_IsWon",
]

# Legacy land.xlsx contracts (named-account list sheets).
LEGACY_CONTRACTS: list[tuple[str, int, list[str]]] = [
    (
        "Top_Deals_Land",
        1,
        ["#", "Account", "Opportunity", "Owner", "Stage", "Close Date", "Age (days)", "ARR (EUR)"],
    ),
    (
        "Top_Deals_Expand",
        1,
        ["#", "Account", "Opportunity", "Owner", "Stage", "Close Date", "Age (days)", "ARR (EUR)"],
    ),
    (
        "Pending_Commercial_Approval",
        3,  # row 1 = title, row 2 = blank, row 3 = headers
        ["#", "Account", "Opportunity", "Owner", "Stage", "Close Date", "Type", "ARR (EUR)"],
    ),
    (
        "At_Risk_Renewals",
        1,
        ["#", "Account", "Owner", "Stage", "Close Date", "ACV", "Risk", "Risk score (0-4)"],
    ),
]


def _check_headers(
    ws: Any, expected: list[str | None], header_row: int, sheet_name: str
) -> list[str]:
    """Return a list of error strings (empty if pass)."""
    errors = []
    for col_idx, expected_col in enumerate(expected, start=1):
        if expected_col is None:
            continue  # wildcard
        actual = ws.cell(row=header_row, column=col_idx).value
        if actual != expected_col:
            errors.append(
                f"  ✗ {sheet_name}!{ws.cell(row=header_row, column=col_idx).coordinate}: "
                f"expected '{expected_col}', got '{actual}'"
            )
    return errors


def validate_model(path: Path) -> list[str]:
    errors: list[str] = []
    if not path.exists():
        return [f"  ✗ model file missing: {path}"]
    wb = load_workbook(path)

    for sheet_name, header_row, expected in MODEL_CONTRACTS:
        if sheet_name not in wb.sheetnames:
            errors.append(f"  ✗ sheet missing: {sheet_name}")
            continue
        errors.extend(_check_headers(wb[sheet_name], expected, header_row, sheet_name))

    # Tables
    if "Data" in wb.sheetnames:
        existing_tables = list(wb["Data"].tables.keys())
        for tbl in REQUIRED_TABLES[:1]:  # tblData lives on Data sheet
            if tbl not in existing_tables:
                errors.append(f"  ✗ Excel Table missing on Data sheet: {tbl}")
    for tbl_name in REQUIRED_TABLES[1:]:
        # Closed-history tables live on dedicated sheets — find any sheet with it
        found = False
        for ws in wb.worksheets:
            if tbl_name in ws.tables:
                found = True
                break
        if not found:
            errors.append(f"  ✗ Excel Table missing: {tbl_name}")

    # Named ranges
    defined = set(wb.defined_names)
    for name in REQUIRED_NAMED_RANGES:
        if name not in defined:
            errors.append(f"  ✗ named range missing: {name}")

    return errors


def validate_legacy(path: Path) -> list[str]:
    errors: list[str] = []
    if not path.exists():
        return [f"  ✗ legacy file missing: {path}"]
    wb = load_workbook(path)

    for sheet_name, header_row, expected in LEGACY_CONTRACTS:
        if sheet_name not in wb.sheetnames:
            errors.append(f"  ✗ sheet missing: {sheet_name}")
            continue
        errors.extend(_check_headers(wb[sheet_name], expected, header_row, sheet_name))

    return errors


def main() -> int:
    ap = argparse.ArgumentParser(description="Validate workbook contracts.")
    ap.add_argument(
        "director_dir",
        type=Path,
        help="Path to state/<period>/<director>/ — must contain land.model.xlsx + land.xlsx",
    )
    args = ap.parse_args()

    model_errors = validate_model(args.director_dir / "land.model.xlsx")
    legacy_errors = validate_legacy(args.director_dir / "land.xlsx")

    print(f"=== Workbook contract validation: {args.director_dir.name} ===")
    print()
    print("Model (land.model.xlsx):")
    if not model_errors:
        print(f"  ✓ all {len(MODEL_CONTRACTS)} sheet contracts pass + tables + named ranges")
    else:
        for e in model_errors:
            print(e)
    print()
    print("Legacy (land.xlsx):")
    if not legacy_errors:
        print(f"  ✓ all {len(LEGACY_CONTRACTS)} sheet contracts pass")
    else:
        for e in legacy_errors:
            print(e)

    return 1 if (model_errors or legacy_errors) else 0


if __name__ == "__main__":
    sys.exit(main())
