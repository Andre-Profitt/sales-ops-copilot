#!/usr/bin/env python3
"""Build a deep PowerPoint corpus for think-cell templates and SimCorp use."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import tempfile
import textwrap
import xml.etree.ElementTree as ET
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_ROOT = Path("/Library/Application Support/Microsoft/think-cell/templates")
DEFAULT_WORKDIR = ROOT / "state" / "thinkcell_bridge" / "deep_corpus"
SF_FIT_JSON = ROOT / "state" / "2026-Q2" / "__regional__" / "thinkcell_sf_fit" / "thinkcell_quarter_salesforce_fit.json"
OUT_PPTX = DEFAULT_WORKDIR / "SimCorp_ThinkCell_Template_Corpus_Deep_Dive.pptx"

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
C_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

NAVY = RGBColor(8, 62, 167)
INK = RGBColor(26, 29, 49)
MUTED = RGBColor(92, 99, 112)
LIGHT = RGBColor(244, 247, 251)
GRID = RGBColor(214, 222, 234)
CORAL = RGBColor(239, 62, 74)
GREEN = RGBColor(26, 151, 99)
TEAL = RGBColor(0, 163, 173)
GOLD = RGBColor(244, 172, 45)
WHITE = RGBColor(255, 255, 255)


HIGH_FIT = {
    "think-cell Charts/Bar, Column/Bar, Column.potx",
    "think-cell Charts/Waterfall/Waterfall.potx",
    "think-cell Charts/Line, Area/Line, Area.potx",
    "think-cell Charts/Scatter, Bubble/Scatter, Bubble.potx",
    "think-cell Charts/Timeline, Gantt/Timeline, Gantt.potx",
    "think-cell Charts/Mekko/Mekko.potx",
    "think-cell Charts/Annotations/Annotations.potx",
    "Dashboards, Statistics/Dashboards, Statistics.potx",
    "Tables/Tables.potx",
    "Useful Elements/Tables/Tables.potx",
    "Processes, Flow Charts, Phases/Processes, Flow Charts, Phases.potx",
    "Matrices, SWOT Analyses/Matrices, SWOT Analyses.potx",
    "Maps/Maps.potx",
    "Agendas, Schedules, Timetables/Agendas, Schedules, Timetables.potx",
    "Timelines, Milestones, Project Planning/Timelines, Milestones, Project Planning.potx",
}


FAMILY_NOTES = {
    "Bar, Column": {
        "use": "Default analytical workhorse: rankings, stage, forecast, owner, geography, age buckets.",
        "simcorp": "Universal for Q2: all nine directors have enough category spread.",
        "avoid": "Do not hide ARR/ACV separation inside stacks.",
    },
    "Waterfall": {
        "use": "Movement logic: opening, additions, wins, losses, slips, closing.",
        "simcorp": "Use from workbook movement bridge; live Salesforce alone is current state, not movement.",
        "avoid": "Never bridge ARR and renewal ACV together.",
    },
    "Line, Area": {
        "use": "Time series: pipeline creation, activity trend, close-date movement.",
        "simcorp": "Good only when workbook exposes real weekly/monthly points.",
        "avoid": "No single-period trend theatre.",
    },
    "Scatter, Bubble": {
        "use": "Deal inspection: probability, ARR, age, push count, activity recency.",
        "simcorp": "Eligible for 8/9 Q2 directors; Patrick needs a different y-axis or fallback.",
        "avoid": "No anonymous or decorative quadrants.",
    },
    "Timeline, Gantt": {
        "use": "Milestones: approval dates, renewal decision windows, close-plan commitments.",
        "simcorp": "FY26 renewals qualify for 8/9; Q2-only renewal Gantt qualifies for only 3/9.",
        "avoid": "No action Gantt from free-text NextStep.",
    },
    "Mekko": {
        "use": "Two-dimensional mix when width and stack both matter.",
        "simcorp": "Technically eligible in 8/9, but should remain rare and claim-led.",
        "avoid": "Sparse matrix or exact-value comparisons.",
    },
    "Tables": {
        "use": "Named evidence: deals, approvals, renewals, owner actions.",
        "simcorp": "Production-safe through table-image lane, not native think-cell table automation.",
        "avoid": "Native think-cell tables until a real named data-backed donor exists.",
    },
    "Dashboards, Statistics": {
        "use": "KPI strip and operating summary patterns.",
        "simcorp": "Use restraint: three to five metrics with basis labels.",
        "avoid": "Full dashboard clones and traffic-light financials.",
    },
    "Processes, Flow Charts, Phases": {
        "use": "Governance, approval flow, operating cadence, handoffs.",
        "simcorp": "Useful for Commercial Approval workflow, not ARR by stage.",
        "avoid": "Replacing the 8-stage metric view with process art.",
    },
    "Matrices, SWOT Analyses": {
        "use": "Risk heat map, triage, segmentation, decision rights.",
        "simcorp": "Use only when two axes create a decision.",
        "avoid": "Generic SWOT in a pipeline review.",
    },
    "Maps": {
        "use": "Geography when geography is the decision variable.",
        "simcorp": "6/9 have country spread, but ranked bars are still clearer.",
        "avoid": "Map-as-decoration.",
    },
    "Agendas, Schedules, Timetables": {
        "use": "Meeting spine, cadence, operating schedule.",
        "simcorp": "Useful for director meeting flow and decision cadence.",
        "avoid": "Generic page furniture.",
    },
    "Timelines, Milestones, Project Planning": {
        "use": "Roadmaps and milestone narratives.",
        "simcorp": "Use for strategic or enablement timelines; use think-cell Gantt for row-level dated data.",
        "avoid": "Undated process explanation.",
    },
}


@dataclass
class SlideInfo:
    number: int
    title: str
    text_sample: list[str]
    chart_refs: int
    graphic_frames: int
    pictures: int


@dataclass
class TemplateInfo:
    relative_path: str
    group: str
    family: str
    slide_count: int
    chart_refs: int
    ole_objects: int
    tags: int
    named_payloads: int
    slides: list[SlideInfo]


def _slug(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", value).strip("-")


def _slide_sort_key(name: str) -> int:
    match = re.search(r"slide(\d+)\.xml$", name)
    return int(match.group(1)) if match else 0


def _extract_text(root: ET.Element) -> list[str]:
    texts: list[str] = []
    for node in root.iter(f"{A_NS}t"):
        value = " ".join((node.text or "").split())
        if value and value not in texts:
            texts.append(value)
    return texts


def _family_from_path(relative: str) -> tuple[str, str]:
    parts = relative.split("/")
    group = parts[0]
    family = parts[-2] if len(parts) > 1 else parts[-1].replace(".potx", "")
    if group == "think-cell Charts":
        family = parts[-2]
    if group == "Useful Elements" and len(parts) > 2:
        family = parts[-2]
    return group, family


def analyze_template(path: Path) -> TemplateInfo:
    relative = path.relative_to(TEMPLATE_ROOT).as_posix()
    group, family = _family_from_path(relative)
    slides: list[SlideInfo] = []
    chart_refs = 0
    with ZipFile(path) as zf:
        names = zf.namelist()
        slide_names = sorted(
            [name for name in names if re.match(r"ppt/slides/slide\d+\.xml$", name)],
            key=_slide_sort_key,
        )
        for idx, slide_name in enumerate(slide_names, start=1):
            root = ET.fromstring(zf.read(slide_name))
            texts = _extract_text(root)
            title = texts[0] if texts else f"Slide {idx}"
            charts = sum(1 for _ in root.iter(f"{C_NS}chart"))
            graphics = sum(1 for _ in root.iter(f"{P_NS}graphicFrame"))
            pictures = sum(1 for _ in root.iter(f"{P_NS}pic"))
            chart_refs += charts
            slides.append(
                SlideInfo(
                    number=idx,
                    title=title,
                    text_sample=texts[:8],
                    chart_refs=charts,
                    graphic_frames=graphics,
                    pictures=pictures,
                )
            )
        ole_objects = len([name for name in names if name.startswith("ppt/embeddings/")])
        tags = len([name for name in names if name.startswith("ppt/tags/tag")])
        named_payloads = 0
        for name in names:
            if name.startswith("ppt/slides/") and name.endswith(".xml"):
                named_payloads += zf.read(name).decode("utf-8", "ignore").count("m_strName")
    return TemplateInfo(
        relative_path=relative,
        group=group,
        family=family,
        slide_count=len(slides),
        chart_refs=chart_refs,
        ole_objects=ole_objects,
        tags=tags,
        named_payloads=named_payloads,
        slides=slides,
    )


def _require_binary(name: str) -> str:
    path = shutil.which(name)
    if not path:
        raise RuntimeError(f"required binary missing: {name}")
    return path


def render_template(path: Path, render_root: Path) -> list[Path]:
    relative = path.relative_to(TEMPLATE_ROOT).as_posix()
    out_dir = render_root / _slug(relative)
    out_dir.mkdir(parents=True, exist_ok=True)
    existing = sorted(out_dir.glob("page-*.png"))
    if existing:
        return existing
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        result = subprocess.run(
            [
                _require_binary("soffice"),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(tmp_dir),
                str(path),
            ],
            text=True,
            capture_output=True,
            check=False,
        )
        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or result.stdout.strip())
        pdfs = sorted(tmp_dir.glob("*.pdf"), key=lambda p: p.stat().st_mtime, reverse=True)
        if not pdfs:
            raise RuntimeError(f"no PDF produced for {path}")
        pdf_path = out_dir / f"{_slug(relative)}.pdf"
        shutil.copy2(pdfs[0], pdf_path)
    result = subprocess.run(
        [_require_binary("pdftoppm"), "-png", "-r", "105", str(pdf_path), str(out_dir / "page")],
        text=True,
        capture_output=True,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip())
    return sorted(out_dir.glob("page-*.png"))


def fmt_eur(value: float) -> str:
    if abs(value) >= 1_000_000:
        return f"EUR {value / 1_000_000:.1f}M"
    if abs(value) >= 1_000:
        return f"EUR {value / 1_000:.0f}K"
    return f"EUR {value:.0f}"


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)

    def blank(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = WHITE
        return slide

    def text(self, slide, text: str, x: float, y: float, w: float, h: float, *, size=12, bold=False, color=INK, align=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.02)
        tf.margin_right = Inches(0.02)
        p = tf.paragraphs[0]
        if align:
            p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return box

    def header(self, slide, title: str, subtitle: str = "") -> None:
        self.text(slide, title, 0.55, 0.33, 10.6, 0.42, size=22, bold=True)
        if subtitle:
            self.text(slide, subtitle, 0.57, 0.78, 11.4, 0.28, size=8.7, color=MUTED)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.12), Inches(12.22), Inches(0.02))
        rule.fill.solid()
        rule.fill.fore_color.rgb = GRID
        rule.line.fill.background()

    def footer(self, slide) -> None:
        self.text(
            slide,
            "SimCorp think-cell corpus | Salesforce-backed Q2 2026 lens | ARR and Renewal ACV separated",
            0.55,
            7.08,
            9.2,
            0.18,
            size=7.2,
            color=MUTED,
        )

    def rect(self, slide, x: float, y: float, w: float, h: float, color: RGBColor, line: RGBColor | None = None):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        if line:
            shape.line.color.rgb = line
        else:
            shape.line.fill.background()
        return shape

    def pill(self, slide, label: str, value: str, x: float, y: float, w: float = 2.35, color: RGBColor = NAVY):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.73))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT
        shape.line.color.rgb = GRID
        self.text(slide, label, x + 0.12, y + 0.1, w - 0.24, 0.15, size=7.2, color=MUTED)
        self.text(slide, value, x + 0.12, y + 0.31, w - 0.24, 0.28, size=17, bold=True, color=color)

    def add_img(self, slide, path: Path, x: float, y: float, w: float, h: float | None = None):
        if h is None:
            return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_thumbnail_grid(deck: Deck, slide, images: list[Path], labels: list[str], *, x=0.6, y=1.35, cols=3, thumb_w=3.65, thumb_h=2.05, gap_x=0.28, gap_y=0.62):
    for idx, path in enumerate(images):
        row = idx // cols
        col = idx % cols
        px = x + col * (thumb_w + gap_x)
        py = y + row * (thumb_h + gap_y)
        deck.rect(slide, px - 0.02, py - 0.02, thumb_w + 0.04, thumb_h + 0.04, WHITE, GRID)
        deck.add_img(slide, path, px, py, thumb_w, thumb_h)
        label = labels[idx] if idx < len(labels) else ""
        deck.text(slide, label[:78], px, py + thumb_h + 0.08, thumb_w, 0.25, size=7.6, color=MUTED)


def add_family_slide(deck: Deck, info: TemplateInfo, pages: list[Path], title: str | None = None) -> None:
    slide = deck.blank()
    fam_title = title or info.family
    note = FAMILY_NOTES.get(info.family, FAMILY_NOTES.get(fam_title, {}))
    deck.header(
        slide,
        fam_title,
        f"{info.relative_path} | {info.slide_count} slides | charts {info.chart_refs} | OLE {info.ole_objects} | tags {info.tags}",
    )
    deck.footer(slide)
    selected = pages[: min(6, len(pages))]
    labels = [f"S{item.number:02d}: {item.title}" for item in info.slides[: len(selected)]]
    add_thumbnail_grid(deck, slide, selected, labels, x=0.55, y=1.35, cols=3, thumb_w=3.35, thumb_h=1.88)
    deck.rect(slide, 10.95, 1.35, 1.85, 4.75, LIGHT, GRID)
    deck.text(slide, "Use", 11.1, 1.52, 1.55, 0.18, size=9, bold=True, color=NAVY)
    deck.text(slide, note.get("use", ""), 11.1, 1.78, 1.5, 0.95, size=8.5, color=INK)
    deck.text(slide, "SimCorp", 11.1, 3.0, 1.55, 0.18, size=9, bold=True, color=GREEN)
    deck.text(slide, note.get("simcorp", ""), 11.1, 3.25, 1.5, 0.95, size=8.5, color=INK)
    deck.text(slide, "Avoid", 11.1, 4.5, 1.55, 0.18, size=9, bold=True, color=CORAL)
    deck.text(slide, note.get("avoid", ""), 11.1, 4.75, 1.5, 0.95, size=8.5, color=INK)


def add_table(deck: Deck, slide, rows: list[list[str]], x: float, y: float, w: float, h: float, col_widths: list[float] | None = None, font_size: float = 7.5):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    if col_widths:
        for idx, width in enumerate(col_widths):
            table.columns[idx].width = Inches(width)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
            for paragraph in cell.text_frame.paragraphs:
                if paragraph.runs:
                    paragraph.runs[0].font.size = Pt(font_size)
                    paragraph.runs[0].font.bold = i == 0
                    paragraph.runs[0].font.color.rgb = WHITE if i == 0 else INK
    return table


def build_deck(inventory: list[TemplateInfo], renders: dict[str, list[Path]], sf_fit: dict[str, object], out_path: Path) -> None:
    deck = Deck()
    prs = deck.prs

    total_slides = sum(info.slide_count for info in inventory)
    total_charts = sum(info.chart_refs for info in inventory)
    total_ole = sum(info.ole_objects for info in inventory)
    total_tags = sum(info.tags for info in inventory)
    named_payloads = sum(info.named_payloads for info in inventory)

    # 1 Cover
    s = deck.blank()
    deck.rect(s, 0, 0, 13.333, 0.24, NAVY)
    deck.text(s, "SimCorp think-cell Template Corpus", 0.68, 0.95, 9.8, 0.55, size=32, bold=True)
    deck.text(s, "Deep dive across the pre-populated template library and what it can do in the Q2 Sales Director decks", 0.7, 1.68, 10.8, 0.42, size=14.5, color=MUTED)
    deck.text(s, "Evidence-led, Salesforce-backed, and constrained by SimCorp ARR/ACV rules", 0.7, 2.35, 8.2, 0.25, size=10.5, bold=True, color=NAVY)
    deck.pill(s, "Stock templates", str(len(inventory)), 0.72, 3.25, 2.15)
    deck.pill(s, "Template slides", str(total_slides), 3.05, 3.25, 2.15)
    deck.pill(s, "Chart refs", str(total_charts), 5.38, 3.25, 2.15)
    deck.pill(s, "OLE parts", str(total_ole), 7.71, 3.25, 2.15)
    deck.pill(s, "Named payloads", str(named_payloads), 10.04, 3.25, 2.15, CORAL)
    deck.text(s, "The stock gallery is broad, but automation readiness is narrow: use the gallery as visual and donor intelligence, then bind only through SimCorp-owned named objects and gates.", 0.73, 5.15, 11.0, 0.7, size=18, bold=True)
    deck.footer(s)

    # 2 Method
    s = deck.blank()
    deck.header(s, "How this corpus was built", "Local think-cell install + OpenXML inspection + rendered template thumbnails + live Salesforce Q2 fit.")
    deck.footer(s)
    method = [
        ("Inventory", f"{len(inventory)} .potx files, {total_slides} slides, {total_charts} chart refs, {total_ole} OLE parts, {total_tags} tag files."),
        ("Automation check", "Stock templates contain zero named m_strName payloads, so they are not strict .ppttc templates."),
        ("Visual check", "High-fit families rendered into page thumbnails so we can inspect actual layout grammar, not just filenames."),
        ("Salesforce check", "Q2 open Land/Expand/Renewal rows queried from preprod and filtered by director scope plus internal/test rules."),
        ("Deck check", "Existing regional publish gate and packaged visual render gate still pass all nine current Q2 meeting-spine decks."),
    ]
    y = 1.45
    for label, body in method:
        deck.rect(s, 0.75, y + 0.05, 0.14, 0.5, NAVY)
        deck.text(s, label, 1.08, y, 2.1, 0.25, size=13, bold=True, color=NAVY)
        deck.text(s, body, 3.0, y, 9.2, 0.42, size=13, color=INK)
        y += 0.92

    # 3 Library topology
    s = deck.blank()
    deck.header(s, "What is actually in the template library", "The gallery is not just charts. It spans chart families, tables, management pages, process models, maps, symbols, and presentation furniture.")
    deck.footer(s)
    by_group = defaultdict(lambda: {"templates": 0, "slides": 0, "charts": 0})
    for info in inventory:
        by_group[info.group]["templates"] += 1
        by_group[info.group]["slides"] += info.slide_count
        by_group[info.group]["charts"] += info.chart_refs
    rows = [["Group", "Templates", "Slides", "Chart refs", "SimCorp read"]]
    for group, stats in sorted(by_group.items(), key=lambda kv: (-kv[1]["slides"], kv[0])):
        read = {
            "Useful Elements": "Useful selectively; most are support objects, not core analytics.",
            "Processes, Flow Charts, Phases": "Governance/cadence, not metrics.",
            "think-cell Charts": "Core chart donor families.",
            "Maps": "Rare, geography-led decisions only.",
            "Mental Models, Frameworks, Concepts": "Strategy decks, not monthly pipeline default.",
            "Timelines, Milestones, Project Planning": "Roadmap/cadence; use Gantt for dated row data.",
            "Matrices, SWOT Analyses": "Triage/segmentation when two axes matter.",
        }.get(group, "Conditional or appendix use.")
        rows.append([group, str(stats["templates"]), str(stats["slides"]), str(stats["charts"]), read])
    add_table(deck, s, rows, 0.55, 1.35, 12.25, 4.95, [3.2, 0.85, 0.75, 0.85, 6.2], 7.5)

    # 4 Decision map
    s = deck.blank()
    deck.header(s, "Decision-first selection map", "Choose the visual from the operating question, then check whether the data shape qualifies.")
    deck.footer(s)
    decision_rows = [
        ["Question", "Best template family", "Q2 deck use", "Fallback"],
        ["What changed?", "Waterfall", "Pipeline movement / QTD movement", "Movement table"],
        ["Where is value concentrated?", "Bar/Column or Scatter", "Owner, country, named deal risk", "Ranked table"],
        ["Which deals need action?", "Table", "Readiness, approvals, renewals", "Decision register"],
        ["What happens when?", "Timeline/Gantt", "FY26 renewal milestones", "Renewal table"],
        ["What is the mix?", "Stacked Bar or Mekko", "Stage x industry only if dense", "Stacked bar / heat table"],
        ["Who owns the risk?", "Bar/Column + table", "Owner coaching", "Named action table"],
        ["Where geographically?", "Ranked Bar, rarely Map", "Country exposure", "Ranked table"],
    ]
    add_table(deck, s, decision_rows, 0.55, 1.35, 12.25, 4.95, [2.55, 2.2, 4.1, 2.6], 8.2)

    # 5 Salesforce gate
    s = deck.blank()
    deck.header(s, "Live Salesforce says what qualifies this quarter", "This is the check the first summary deck did not show deeply enough.")
    deck.footer(s)
    totals = sf_fit["totals"]
    deck.pill(s, "Raw Q2 rows", str(totals["q2_raw_rows"]), 0.65, 1.42, 2.3)
    deck.pill(s, "Publishable rows", str(totals["q2_publishable_rows"]), 3.12, 1.42, 2.3, GREEN)
    deck.pill(s, "Internal/test removed", str(totals["q2_internal_rows_removed"]), 5.59, 1.42, 2.3, CORAL)
    deck.pill(s, "Deck package", "9/9 pass", 8.06, 1.42, 2.3, GREEN)
    counts = sf_fit["family_eligibility_counts"]
    chart_data = CategoryChartData()
    order = [
        ("bar_column", "Bar/Column"),
        ("scatter_bubble", "Scatter/Bubble"),
        ("timeline_gantt_fy26_renewals", "FY26 Renewal TL"),
        ("timeline_gantt_q2_renewals", "Q2 Renewal TL"),
        ("mekko", "Mekko"),
        ("map", "Map"),
    ]
    chart_data.categories = [label for _, label in order]
    chart_data.add_series("Eligible directors", [counts.get(key, 0) for key, _ in order])
    chart = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.85), Inches(2.55), Inches(6.4), Inches(3.65), chart_data).chart
    chart.has_legend = False
    chart.value_axis.maximum_scale = 9
    chart.value_axis.major_unit = 1
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = NAVY
    chart.series[0].has_data_labels = True
    chart.series[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    chart.series[0].data_labels.font.size = Pt(8)
    deck.text(s, "Interpretation", 8.05, 2.58, 2.0, 0.22, size=14, bold=True)
    deck.text(
        s,
        "Bar/Column is the production default. Scatter and FY26 renewal Timeline are conditional pilots. Q2-only renewal Timeline is not a standard slide. Mekko and maps require a real claim, not merely available data.",
        8.05,
        2.98,
        4.1,
        1.45,
        size=13,
    )

    # Family slides
    info_by_rel = {info.relative_path: info for info in inventory}
    family_order = [
        ("think-cell Charts/Bar, Column/Bar, Column.potx", "Bar/Column: the operating workhorse"),
        ("think-cell Charts/Waterfall/Waterfall.potx", "Waterfall: movement, not snapshots"),
        ("think-cell Charts/Line, Area/Line, Area.potx", "Line/Area: only for real time series"),
        ("think-cell Charts/Scatter, Bubble/Scatter, Bubble.potx", "Scatter/Bubble: named deal inspection"),
        ("think-cell Charts/Timeline, Gantt/Timeline, Gantt.potx", "Timeline/Gantt: real milestones only"),
        ("think-cell Charts/Mekko/Mekko.potx", "Mekko: rare mix analysis"),
        ("Tables/Tables.potx", "Tables: named evidence grammar"),
        ("Useful Elements/Tables/Tables.potx", "Useful Elements/Tables: donor inspiration"),
        ("Dashboards, Statistics/Dashboards, Statistics.potx", "Dashboards/Statistics: KPI strip restraint"),
        ("think-cell Charts/Annotations/Annotations.potx", "Annotations: one or two proof labels"),
        ("Processes, Flow Charts, Phases/Processes, Flow Charts, Phases.potx", "Processes: governance and cadence"),
        ("Matrices, SWOT Analyses/Matrices, SWOT Analyses.potx", "Matrices: triage and decision rights"),
        ("Maps/Maps.potx", "Maps: geography only when it matters"),
        ("Agendas, Schedules, Timetables/Agendas, Schedules, Timetables.potx", "Agendas/Schedules: meeting spine"),
        ("Timelines, Milestones, Project Planning/Timelines, Milestones, Project Planning.potx", "Timelines/Milestones: roadmap context"),
    ]
    for rel, title in family_order:
        if rel in info_by_rel and rel in renders:
            add_family_slide(deck, info_by_rel[rel], renders[rel], title)

    # Director table
    s = deck.blank()
    deck.header(s, "Director-by-director visual eligibility", "Live Salesforce Q2 rows after internal/test filtering. ARR and renewal ACV stay separate.")
    deck.footer(s)
    rows = [["Director", "Territory", "Q2 L+E", "Q2 ARR", "FY26 Ren.", "FY26 ACV", "Scatter", "Renew TL", "Removed"]]
    for director in sf_fit["directors"]:
        fit = director["visual_fit"]
        rows.append(
            [
                director["director"],
                director["territory"],
                str(director["q2_open_arr_count"]),
                fmt_eur(director["q2_open_arr_eur"]),
                str(director["fy26_open_renewal_count"]),
                fmt_eur(director["fy26_open_renewal_acv_eur"]),
                "yes" if fit["scatter_bubble"]["eligible"] else "no",
                "yes" if fit["timeline_gantt"]["eligible_for_fy26_renewals"] else "no",
                str(director["internal_rows_removed"]),
            ]
        )
    add_table(deck, s, rows, 0.38, 1.3, 12.65, 5.55, [1.65, 1.72, 0.55, 0.9, 0.7, 0.9, 0.58, 0.65, 0.62], 6.9)

    # Q2 deck slide map
    s = deck.blank()
    deck.header(s, "Current Q2 meeting spine: where templates apply", "This is the practical upgrade map for this quarter, not a generic gallery translation.")
    deck.footer(s)
    rows = [
        ["Spine area", "Current role", "Template family", "This-quarter action"],
        ["May operating summary", "Executive orientation", "Dashboard/KPI strip", "Keep restrained metric strip; no dashboard clone."],
        ["Pipeline / forecast", "Stage and forecast read", "Bar/Column", "Production default; explicit Land+Expand ARR basis."],
        ["May deal readiness", "Named close evidence", "Table + Scatter pilot", "Keep table; pilot Scatter for eligible directors."],
        ["Commercial approval", "Stage 3+ governance", "Table / Process", "Named exception table; process only in appendix/governance explainer."],
        ["Renewals", "ACV watchlist", "Table + Timeline pilot", "FY26 Timeline for 8/9; Adam table-only."],
        ["Owner coaching", "Rep-level focus", "Bar/Column + Table", "Keep owner ranking plus action cues."],
        ["Concentration/risk", "Where to inspect", "Scatter / Bar / Table", "Conditional Scatter; table fallback."],
        ["Decision register", "Actions", "Table", "Do not make a Gantt from text fields."],
    ]
    add_table(deck, s, rows, 0.48, 1.32, 12.35, 5.45, [2.0, 2.25, 2.0, 5.25], 7.8)

    # Exceptions
    s = deck.blank()
    deck.header(s, "Exceptions are product requirements", "The build layer needs fallbacks so the deck stays honest by director.")
    deck.footer(s)
    exceptions = [
        ("Patrick Gaughan", "Scatter fails", "Positive ARR has one distinct value. Use table/ranked bar unless age, push count, risk score, or activity recency becomes the y-axis."),
        ("Adam Steinhouse", "Renewal timeline fails", "FY26 renewal close dates collapse to 2026-12-31. Use renewal ACV table only."),
        ("Christian Ebbesen", "Internal-row exposure", "31 Q2 rows filtered out as internal/test. Filter is a blocker before visual refresh."),
        ("Sarah Pittroff", "Internal-row exposure", "2 Q2 rows filtered out as internal/test. Same control, lower volume."),
    ]
    y = 1.55
    for who, issue, detail in exceptions:
        deck.rect(s, 0.75, y - 0.03, 11.9, 0.65, LIGHT, GRID)
        deck.text(s, who, 0.95, y + 0.08, 2.4, 0.2, size=12.5, bold=True, color=NAVY)
        deck.text(s, issue, 3.45, y + 0.08, 2.1, 0.2, size=12.5, bold=True, color=CORAL)
        deck.text(s, detail, 5.65, y + 0.05, 6.65, 0.35, size=11, color=INK)
        y += 0.95

    # Automation architecture
    s = deck.blank()
    deck.header(s, "Automation architecture: what is safe", "The template library is a donor and design source; SimCorp-owned gates decide production.")
    deck.footer(s)
    architecture = [
        ("Stock .potx", "Visual/reference donor", "0 named payloads; not strict .ppttc-ready."),
        ("LAND seed", "Named chart/text contract", "42 named elements; charts/text credible."),
        ("Table-image donors", "Production table lane", "Excel COM AddRangeImage; auditable named ranges."),
        ("Native think-cell tables", "Blocked", "No real named data-backed table donor yet."),
        ("Publish gates", "Protection layer", "ARR/ACV, internal/test filters, visual eligibility, render checks."),
    ]
    x = 0.65
    for title, role, detail in architecture:
        deck.rect(s, x, 1.65, 2.35, 3.35, LIGHT, GRID)
        deck.text(s, title, x + 0.18, 1.88, 1.95, 0.24, size=13, bold=True, color=NAVY if role != "Blocked" else CORAL)
        deck.text(s, role, x + 0.18, 2.38, 1.95, 0.24, size=10.5, bold=True, color=INK)
        deck.text(s, detail, x + 0.18, 2.82, 1.95, 1.25, size=10, color=MUTED)
        x += 2.48

    # Pilot plan
    s = deck.blank()
    deck.header(s, "Build plan: deeper, but still controlled", "Pilot two high-value think-cell upgrades before touching the whole regional package.")
    deck.footer(s)
    steps = [
        ("1", "Pilot director", "Use Jesper or Sarah because both have probability/ARR spread and renewal date spread."),
        ("2", "Deal-risk Scatter", "Probability x converted ARR; named labels for material deals; fallback table where axes collapse."),
        ("3", "Renewal Timeline", "FY26 Renewal ACV rows by close date; table-only for Adam."),
        ("4", "Generalize", "Only after render, publish, and fact gates pass on the pilot deck."),
    ]
    x = 0.65
    for num, title, body in steps:
        deck.rect(s, x, 1.55, 2.75, 3.1, WHITE, GRID)
        shape = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.15), Inches(1.78), Inches(0.48), Inches(0.48))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY
        shape.line.fill.background()
        deck.text(s, num, x + 0.31, 1.89, 0.16, 0.14, size=10.5, bold=True, color=WHITE)
        deck.text(s, title, x + 0.18, 2.48, 2.3, 0.24, size=13, bold=True)
        deck.text(s, body, x + 0.18, 2.9, 2.3, 0.92, size=10.5, color=MUTED)
        x += 3.05
    deck.text(s, "Recommendation: build this as a pilot insertion, not a wholesale template replacement.", 0.72, 5.55, 10.4, 0.35, size=18, bold=True)

    # Guardrails
    s = deck.blank()
    deck.header(s, "Guardrails that must survive every visual upgrade", "The corpus is only useful if it protects the business truth.")
    deck.footer(s)
    guards = [
        "ARR = Land + Expand only; Renewal ACV stays separate.",
        "Type-bearing ARR visuals require Type IN ('Land','Expand').",
        "Type-bearing ACV visuals require Type = 'Renewal'.",
        "Headline currency must use converted/reporting-currency EUR basis.",
        "The SimCorp 8-stage process stays stage-explicit; no generic funnel.",
        "Commercial Approval gaps must trace to Stage_20_Approval__c.",
        "Internal/test rows are blocker-level publish failures.",
        "Native think-cell tables stay blocked until a named data-backed donor is proven.",
    ]
    for idx, guard in enumerate(guards):
        y = 1.35 + idx * 0.62
        deck.rect(s, 0.85, y + 0.08, 0.12, 0.12, NAVY)
        deck.text(s, guard, 1.08, y, 11.2, 0.25, size=13.2)

    # Appendix inventory 1
    sorted_inventory = sorted(inventory, key=lambda item: (item.group, item.family, item.relative_path))
    for page_idx in range(0, len(sorted_inventory), 16):
        s = deck.blank()
        deck.header(s, f"Appendix: full template inventory {page_idx // 16 + 1}", "All stock templates are retained in the corpus even when they are not production defaults.")
        deck.footer(s)
        chunk = sorted_inventory[page_idx : page_idx + 16]
        rows = [["Template", "Slides", "Charts", "OLE", "Use posture"]]
        for info in chunk:
            posture = "high-fit" if info.relative_path in HIGH_FIT else ("conditional" if info.group in {"Useful Elements", "Mental Models, Frameworks, Concepts"} else "rare")
            rows.append([info.relative_path.replace(".potx", ""), str(info.slide_count), str(info.chart_refs), str(info.ole_objects), posture])
        add_table(deck, s, rows, 0.35, 1.25, 12.75, 5.8, [7.2, 0.55, 0.55, 0.55, 1.1], 5.9)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workdir", type=Path, default=DEFAULT_WORKDIR)
    parser.add_argument("--output", type=Path, default=OUT_PPTX)
    parser.add_argument("--skip-render", action="store_true")
    args = parser.parse_args()

    args.workdir.mkdir(parents=True, exist_ok=True)
    render_root = args.workdir / "renders"
    templates = sorted(TEMPLATE_ROOT.rglob("*.potx"))
    inventory = [analyze_template(path) for path in templates]
    renders: dict[str, list[Path]] = {}
    if not args.skip_render:
        for path in templates:
            rel = path.relative_to(TEMPLATE_ROOT).as_posix()
            if rel in HIGH_FIT:
                renders[rel] = render_template(path, render_root)
    else:
        for path in templates:
            rel = path.relative_to(TEMPLATE_ROOT).as_posix()
            render_dir = render_root / _slug(rel)
            pages = sorted(render_dir.glob("page-*.png"))
            if pages:
                renders[rel] = pages
    sf_fit = json.loads(SF_FIT_JSON.read_text(encoding="utf-8"))
    corpus = {
        "schema": "thinkcell-deep-corpus/v1",
        "template_count": len(inventory),
        "slide_count": sum(item.slide_count for item in inventory),
        "chart_refs": sum(item.chart_refs for item in inventory),
        "ole_objects": sum(item.ole_objects for item in inventory),
        "tags": sum(item.tags for item in inventory),
        "named_payloads": sum(item.named_payloads for item in inventory),
        "templates": [
            {
                "relative_path": item.relative_path,
                "group": item.group,
                "family": item.family,
                "slide_count": item.slide_count,
                "chart_refs": item.chart_refs,
                "ole_objects": item.ole_objects,
                "tags": item.tags,
                "named_payloads": item.named_payloads,
                "slides": [slide.__dict__ for slide in item.slides],
            }
            for item in inventory
        ],
    }
    (args.workdir / "thinkcell_deep_corpus.json").write_text(json.dumps(corpus, indent=2) + "\n", encoding="utf-8")
    build_deck(inventory, renders, sf_fit, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
