#!/usr/bin/env python3
"""Validate the connected Salesforce -> Excel -> think-cell factory spec."""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


ROOT = Path(__file__).resolve().parent.parent
SUPPORTED_THINKCELL_STATUSES = {
    "native_chart_supported_seeded",
    "text_supported",
    "blocked_no_named_table_donor",
    "chart_supported_table_blocked",
    "table_blocked_chart_possible",
}
SUPPORTED_VISUAL_GATE_RULES = {
    "waterfall_requires_negative_closed_lost",
    "timeline_requires_distinct_dates",
    "scatter_requires_two_noncollapsed_axes",
    "action_register_rejects_generic_gantt",
    "mekko_requires_dense_matrix",
}
REQUIRED_OBJECT_FIELDS = {
    "slide_id",
    "object_id",
    "object_type",
    "intelligence_role",
    "raw_tabs",
    "model_tabs",
    "output_sheet_target",
    "output_range_target",
    "excel_name_target",
    "ppt_name",
    "visualization_lane",
    "thinkcell_status",
    "validation_rules",
}


def _load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _resolve(path_text: str) -> Path:
    path = Path(path_text)
    if path.is_absolute():
        return path
    return ROOT / path


def _load_workbooks(spec: dict[str, Any]) -> dict[str, Any]:
    loaded: dict[str, Any] = {}
    for key, value in (spec.get("artifacts") or {}).items():
        if not str(value).endswith(".xlsx"):
            continue
        path = _resolve(str(value))
        if path.exists():
            loaded[key] = load_workbook(path, read_only=False, data_only=False)
    return loaded


def _defined_names(wb: Any) -> set[str]:
    return {defined_name.name for defined_name in wb.defined_names.values()}


def _split_target(target: str) -> tuple[str, str]:
    sheet, cell_range = target.split("!", 1)
    return sheet.strip("'"), cell_range.replace("$", "")


def _same_range(left: str, right: str) -> bool:
    return range_boundaries(left.replace("$", "")) == range_boundaries(right.replace("$", ""))


def _has_formula(ws: Any, cell_range: str) -> bool:
    for row in ws[cell_range.replace("$", "")]:
        cells = row if isinstance(row, tuple) else (row,)
        for cell in cells:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                return True
    return False


def _sheet_headers(ws: Any) -> dict[str, int]:
    return {
        str(cell.value).strip(): idx
        for idx, cell in enumerate(ws[1], start=1)
        if cell.value not in (None, "")
    }


def _has_header(ws: Any, header: str) -> bool:
    return header in _sheet_headers(ws)


def _column_values_by_header(ws: Any, header: str) -> list[Any]:
    headers = _sheet_headers(ws)
    col_idx = headers.get(header)
    if not col_idx:
        return []
    return [
        ws.cell(row_idx, col_idx).value
        for row_idx in range(2, ws.max_row + 1)
        if ws.cell(row_idx, col_idx).value not in (None, "")
    ]


def _normalize_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip() if value not in (None, "") else ""
    if not text:
        return None
    return text[:10]


def _to_number(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    if not text:
        return None
    import re

    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    number = float(match.group(0))
    if "meur" in text.lower():
        number *= 1_000_000
    return number


def _distinct_nonblank(values: list[Any]) -> set[str]:
    normalized: set[str] = set()
    for value in values:
        if isinstance(value, (date, datetime)):
            date_value = _normalize_date(value)
            if date_value:
                normalized.add(date_value)
        else:
            text = str(value).strip()
            if text:
                normalized.add(text)
    return normalized


def _distinct_numeric(values: list[Any]) -> set[float]:
    return {round(number, 6) for value in values if (number := _to_number(value)) is not None}


def _cell_is_negative(ws: Any, cell_ref: str) -> bool:
    value = ws[cell_ref].value
    if isinstance(value, str):
        return value.replace(" ", "").startswith("=-")
    number = _to_number(value)
    return number is not None and number < 0


def _visual_gate_workbook_errors(spec: dict[str, Any], wb: Any) -> list[str]:
    errors: list[str] = []
    for obj in spec.get("objects") or []:
        object_id = str(obj.get("object_id", "(missing)"))
        gate = obj.get("visual_gate")
        if not isinstance(gate, dict):
            continue
        rule = str(gate.get("rule") or "")
        source_sheet = str(gate.get("source_sheet") or "")
        if source_sheet not in wb.sheetnames:
            continue
        ws = wb[source_sheet]
        primary_visual = str(obj.get("primary_visual") or obj.get("visualization_lane") or "")
        if rule == "waterfall_requires_negative_closed_lost":
            cell_ref = str(gate.get("value_cell") or "")
            if cell_ref and not _cell_is_negative(ws, cell_ref):
                errors.append(f"{object_id}: waterfall closed-lost cell {source_sheet}!{cell_ref} is not negative")
        elif rule == "timeline_requires_distinct_dates" and "timeline" in primary_visual:
            values = [_normalize_date(value) for value in _column_values_by_header(ws, str(gate.get("date_header") or ""))]
            distinct = {value for value in values if value}
            if len(distinct) < int(gate.get("min_distinct_dates") or 2):
                errors.append(f"{object_id}: timeline visual has only {len(distinct)} distinct date(s)")
        elif rule == "scatter_requires_two_noncollapsed_axes" and "scatter" in primary_visual:
            x_values = _distinct_numeric(_column_values_by_header(ws, str(gate.get("x_header") or "")))
            y_values = _distinct_numeric(_column_values_by_header(ws, str(gate.get("y_header") or "")))
            if len(x_values) < int(gate.get("min_distinct_x") or 2) or len(y_values) < int(gate.get("min_distinct_y") or 2):
                errors.append(
                    f"{object_id}: scatter visual axes are collapsed "
                    f"(distinct x={len(x_values)}, y={len(y_values)})"
                )
        elif rule == "action_register_rejects_generic_gantt" and "timeline" in primary_visual:
            values = [_normalize_date(value) for value in _column_values_by_header(ws, str(gate.get("date_header") or ""))]
            values = [value for value in values if value]
            if values:
                most_common = max(values.count(value) for value in set(values)) / len(values)
                if most_common > float(gate.get("max_same_date_ratio") or 0.8):
                    errors.append(f"{object_id}: action timeline uses generic repeated due dates ({most_common:.0%})")
    return errors


def _raw_intel_values(wb: Any) -> dict[str, Any]:
    if "Raw_Original_Intel" not in wb.sheetnames:
        return {}
    ws = wb["Raw_Original_Intel"]
    values: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        values[str(row[0])] = row[1]
    return values


def _almost_equal(actual: Any, expected: float, tolerance: float = 1.0) -> bool:
    try:
        return abs(float(actual) - expected) <= tolerance
    except (TypeError, ValueError):
        return False


def _expected_raw_intel_from_spec(spec: dict[str, Any]) -> dict[str, float]:
    sidecar_text = str((spec.get("artifacts") or {}).get("original_etl_sidecar") or "")
    if not sidecar_text:
        return {}
    sidecar_path = _resolve(sidecar_text)
    if not sidecar_path.exists() or not sidecar_path.is_file():
        return {}
    sidecar = _load_json(sidecar_path)
    return {
        "open_land_deals_sidecar": float(sidecar.get("open_land_deals") or 0),
        "open_land_arr_eur_sidecar": float(sidecar.get("open_land_arr") or 0),
        "q1_land_wins": float(sidecar.get("q1_land_wins") or 0),
        "q1_land_wins_arr_eur": float(sidecar.get("q1_land_wins_arr") or 0),
        "q1_land_lost": float(sidecar.get("q1_land_lost") or 0),
        "q1_land_lost_arr_eur": float(sidecar.get("q1_land_lost_arr") or 0),
        "q2_renewals_acv_eur": float(sidecar.get("q2_renewals_acv") or 0),
        "q3_renewals_acv_eur": float(sidecar.get("q3_renewals_acv") or 0),
    }


def validate_workbook(spec: dict[str, Any], workbook_path: Path) -> list[str]:
    errors: list[str] = []
    if not workbook_path.exists():
        return [f"workbook does not exist: {workbook_path}"]

    wb = load_workbook(workbook_path, data_only=False)
    sheet_names = set(wb.sheetnames)
    defined_names = _defined_names(wb)
    objects = spec.get("objects") or []

    if "ThinkCell_Link_Map" not in sheet_names:
        errors.append("workbook missing ThinkCell_Link_Map sheet")
    else:
        link_ws = wb["ThinkCell_Link_Map"]
        expected_headers = [
            "Slide",
            "Slide title",
            "Object ID",
            "PowerPoint / think-cell name",
            "Excel defined name",
            "Excel output range",
            "Output sheet",
            "Object type",
            "Metric basis",
            "Raw source tabs",
            "Model tabs",
            "think-cell status",
            "Current publish path",
            "Validation rules",
        ]
        actual_headers = [link_ws.cell(1, col_idx).value for col_idx in range(1, len(expected_headers) + 1)]
        if actual_headers != expected_headers:
            errors.append("ThinkCell_Link_Map headers do not match factory contract")
        link_object_ids = {
            str(link_ws.cell(row_idx, 3).value)
            for row_idx in range(2, link_ws.max_row + 1)
            if link_ws.cell(row_idx, 3).value
        }
        expected_object_ids = {str(obj["object_id"]) for obj in objects}
        if link_object_ids != expected_object_ids:
            errors.append(
                "ThinkCell_Link_Map object IDs do not match spec objects: "
                f"missing={sorted(expected_object_ids - link_object_ids)} extra={sorted(link_object_ids - expected_object_ids)}"
            )
        for row_idx in range(2, link_ws.max_row + 1):
            basis = str(link_ws.cell(row_idx, 9).value or "")
            if "ARR" in basis and "unweighted" not in basis and "weighted" not in basis:
                errors.append(f"ThinkCell_Link_Map row {row_idx} has ambiguous ARR metric basis: {basis}")
        if "rng_ThinkCell_Link_Map" not in defined_names:
            errors.append("workbook missing rng_ThinkCell_Link_Map defined name")

    required_raw_tabs = sorted({tab for obj in objects for tab in obj.get("raw_tabs", [])})
    missing_raw = [tab for tab in required_raw_tabs if tab not in sheet_names]
    if missing_raw:
        errors.append(f"workbook missing raw tabs referenced by spec: {missing_raw}")

    required_model_tabs = sorted({tab for obj in objects for tab in obj.get("model_tabs", [])})
    missing_model = [tab for tab in required_model_tabs if tab not in sheet_names]
    if missing_model:
        errors.append(f"workbook missing model tabs referenced by spec: {missing_model}")

    for obj in objects:
        object_id = str(obj["object_id"])
        output_sheet, output_range = _split_target(str(obj["output_range_target"]))
        excel_name = str(obj["excel_name_target"])
        if output_sheet not in sheet_names:
            errors.append(f"{object_id}: output sheet missing: {output_sheet}")
            continue
        if excel_name not in defined_names:
            errors.append(f"{object_id}: defined name missing: {excel_name}")
        else:
            defined_name = wb.defined_names[excel_name]
            destinations = list(defined_name.destinations)
            if len(destinations) != 1:
                errors.append(f"{object_id}: defined name {excel_name} has {len(destinations)} destinations")
            else:
                dest_sheet, dest_range = destinations[0]
                if dest_sheet != output_sheet or not _same_range(dest_range, output_range):
                    errors.append(
                        f"{object_id}: defined name {excel_name} points to {dest_sheet}!{dest_range}, expected {output_sheet}!{output_range}"
                    )
        if not _has_formula(wb[output_sheet], output_range):
            errors.append(f"{object_id}: output range {output_sheet}!{output_range} has no formulas")

    raw_intel = _raw_intel_values(wb)
    expected_values = _expected_raw_intel_from_spec(spec)
    for key, expected in expected_values.items():
        if key not in raw_intel:
            errors.append(f"Raw_Original_Intel missing key: {key}")
        elif not _almost_equal(raw_intel[key], expected):
            errors.append(f"Raw_Original_Intel {key}={raw_intel[key]!r}, expected {expected}")

    if "Out_S05_PipelineMovement" in sheet_names:
        closed_lost_formula = wb["Out_S05_PipelineMovement"]["E2"].value
        if not (isinstance(closed_lost_formula, str) and closed_lost_formula.startswith("=-")):
            errors.append(
                "Out_S05_PipelineMovement!E2 must be a negative closed-lost formula for the waterfall"
            )

    if "Out_S13_Renewals" in sheet_names:
        renewal_header = str(wb["Out_S13_Renewals"]["G1"].value)
        if "ACV" not in renewal_header:
            errors.append("Out_S13_Renewals must label renewal amount as ACV")

    errors.extend(_visual_gate_workbook_errors(spec, wb))

    return errors


def validate(spec_path: Path) -> list[str]:
    errors: list[str] = []
    spec = _load_json(spec_path)
    if spec.get("schema") != "connected-thinkcell-factory/v1":
        errors.append(f"unexpected schema: {spec.get('schema')!r}")

    objects = spec.get("objects") or []
    if not objects:
        errors.append("objects must be non-empty")

    object_ids: set[str] = set()
    excel_names: set[str] = set()
    ppt_names: set[str] = set()
    for idx, obj in enumerate(objects, start=1):
        missing = REQUIRED_OBJECT_FIELDS - set(obj)
        if missing:
            errors.append(f"objects[{idx}] missing fields: {sorted(missing)}")
            continue
        object_id = str(obj["object_id"])
        if object_id in object_ids:
            errors.append(f"duplicate object_id: {object_id}")
        object_ids.add(object_id)
        excel_name = str(obj["excel_name_target"])
        if excel_name in excel_names:
            errors.append(f"duplicate excel_name_target: {excel_name}")
        excel_names.add(excel_name)
        ppt_name = str(obj["ppt_name"])
        if ppt_name in ppt_names:
            errors.append(f"duplicate ppt_name: {ppt_name}")
        ppt_names.add(ppt_name)
        if obj["thinkcell_status"] not in SUPPORTED_THINKCELL_STATUSES:
            errors.append(f"{object_id}: unsupported thinkcell_status {obj['thinkcell_status']!r}")
        if not obj.get("validation_rules"):
            errors.append(f"{object_id}: validation_rules must be non-empty")
        if "table" in str(obj.get("object_type")) and obj["thinkcell_status"] == "native_chart_supported_seeded":
            errors.append(f"{object_id}: table object cannot be marked native_chart_supported_seeded")
        visual_gate = obj.get("visual_gate")
        if visual_gate is not None:
            if not isinstance(visual_gate, dict):
                errors.append(f"{object_id}: visual_gate must be an object")
            else:
                rule = str(visual_gate.get("rule", ""))
                if rule not in SUPPORTED_VISUAL_GATE_RULES:
                    errors.append(f"{object_id}: unsupported visual_gate rule {rule!r}")
                if not visual_gate.get("source_sheet"):
                    errors.append(f"{object_id}: visual_gate missing source_sheet")
                if rule == "waterfall_requires_negative_closed_lost" and not visual_gate.get("value_cell"):
                    errors.append(f"{object_id}: waterfall visual_gate missing value_cell")
                if rule in {"timeline_requires_distinct_dates", "action_register_rejects_generic_gantt"} and not visual_gate.get("date_header"):
                    errors.append(f"{object_id}: {rule} visual_gate missing date_header")
                if rule == "scatter_requires_two_noncollapsed_axes":
                    if not visual_gate.get("x_header") or not visual_gate.get("y_header"):
                        errors.append(f"{object_id}: scatter visual_gate missing x_header or y_header")

    loaded = _load_workbooks(spec)
    if not loaded:
        errors.append("no referenced workbook artifacts could be loaded")
        return errors

    workbook_defined_names = {
        key: _defined_names(wb)
        for key, wb in loaded.items()
    }
    any_existing_output_names = set().union(*workbook_defined_names.values()) if workbook_defined_names else set()
    missing_existing_output_names = [
        str(obj["excel_name_target"])
        for obj in objects
        if str(obj["excel_name_target"]) in any_existing_output_names
    ]
    if missing_existing_output_names:
        errors.append(
            "factory output names already exist unexpectedly before rebuild: "
            + ", ".join(sorted(missing_existing_output_names))
        )

    current_sheet_names = set()
    for wb in loaded.values():
        current_sheet_names.update(wb.sheetnames)

    for obj in objects:
        object_id = str(obj.get("object_id", "(missing)"))
        missing_current_tabs = [
            sheet for sheet in obj.get("current_tabs", []) if sheet not in current_sheet_names
        ]
        if missing_current_tabs:
            errors.append(f"{object_id}: current source tab(s) not found in loaded workbooks: {missing_current_tabs}")
        output_sheet = str(obj.get("output_sheet_target") or "")
        if output_sheet in current_sheet_names:
            errors.append(
                f"{object_id}: output sheet {output_sheet!r} already exists; factory rebuild should create it intentionally"
            )
        visual_gate = obj.get("visual_gate") or {}
        source_sheet = str(visual_gate.get("source_sheet") or "")
        if source_sheet.startswith(("Out_", "Raw_")):
            continue
        if source_sheet and source_sheet not in current_sheet_names:
            errors.append(f"{object_id}: visual_gate source_sheet not found in loaded workbooks: {source_sheet}")
        elif source_sheet:
            source_ws = next((wb[source_sheet] for wb in loaded.values() if source_sheet in wb.sheetnames), None)
            if source_ws is not None:
                rule = str(visual_gate.get("rule") or "")
                if rule in {"timeline_requires_distinct_dates", "action_register_rejects_generic_gantt"}:
                    date_header = str(visual_gate.get("date_header") or "")
                    if date_header and not _has_header(source_ws, date_header):
                        errors.append(f"{object_id}: visual_gate date_header {date_header!r} missing from {source_sheet}")
                if rule == "scatter_requires_two_noncollapsed_axes":
                    for header_key in ("x_header", "y_header"):
                        header = str(visual_gate.get(header_key) or "")
                        if header and not _has_header(source_ws, header):
                            errors.append(f"{object_id}: visual_gate {header_key} {header!r} missing from {source_sheet}")

    table_targets = [
        obj for obj in objects if "table" in str(obj.get("object_type", "")) or "table" in str(obj.get("visualization_lane", ""))
    ]
    table_supported = [
        obj for obj in table_targets if obj.get("thinkcell_status") not in {"blocked_no_named_table_donor", "chart_supported_table_blocked", "table_blocked_chart_possible"}
    ]
    if table_supported:
        errors.append(
            "table objects marked as supported before real table donor proof: "
            + ", ".join(str(obj["object_id"]) for obj in table_supported)
        )

    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--spec",
        type=Path,
        default=ROOT / "config" / "connected_thinkcell_factory.jesper_apac.json",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        help="Optional connected factory workbook to validate against the spec.",
    )
    args = parser.parse_args()

    errors = validate(args.spec)
    spec = _load_json(args.spec)
    if args.workbook:
        errors.extend(validate_workbook(spec, args.workbook))

    print(f"=== Connected factory spec: {args.spec} ===")
    if args.workbook:
        print(f"=== Connected factory workbook: {args.workbook} ===")
    if errors:
        print("FAIL")
        for error in errors:
            print(f"  - {error}")
        return 1
    if args.workbook:
        print("OK: connected factory spec and workbook are valid")
    else:
        print("OK: connected factory spec is structurally valid and honestly gated")
    return 0


if __name__ == "__main__":
    sys.exit(main())
