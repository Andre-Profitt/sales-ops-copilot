#!/usr/bin/env python3
"""Validate the LAND template assets by role."""

from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass, field
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from pptx import Presentation

from ppttc_template import template_named_elements
from run_review_package_brand_style_gate import STALE_THINKCELL_TOKENS


ROOT = Path(__file__).resolve().parent.parent


@dataclass(frozen=True)
class TemplateContract:
    path: str
    role: str
    required: bool = True
    expected_slides: int | None = None
    expected_names: int | None = None
    expected_exact_names: tuple[str, ...] = ()
    min_embedded_ole: int | None = None
    max_embedded_ole: int | None = None
    min_native_charts: int | None = None
    max_stale_tokens: int | None = None


@dataclass
class TemplateContractResult:
    path: str
    role: str
    status: str
    exists: bool
    slide_count: int | None = None
    named_element_count: int | None = None
    native_powerpoint_tables: int = 0
    table_image_pictures: int = 0
    embedded_ole_objects: int = 0
    native_powerpoint_charts: int = 0
    stale_thinkcell_token_count: int = 0
    findings: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


CONTRACTS = (
    TemplateContract(
        path="assets/LAND_template.pptx",
        role="clean_simcorp_shell",
        expected_slides=28,
        expected_names=0,
        max_embedded_ole=0,
        max_stale_tokens=0,
    ),
    TemplateContract(
        path="assets/LAND_thinkcell_seed.pptx",
        role="native_chart_text_seed",
        expected_slides=28,
        expected_names=42,
        min_embedded_ole=12,
        min_native_charts=12,
    ),
    TemplateContract(
        path="assets/LAND_thinkcell_seed_charts.pptx",
        role="native_chart_text_seed_mirror",
        expected_slides=28,
        expected_names=42,
        min_embedded_ole=12,
        min_native_charts=12,
    ),
    TemplateContract(
        path="assets/LAND_thinkcell_table_image_donor.pptx",
        role="single_table_image_donor",
        expected_slides=1,
        expected_names=1,
        expected_exact_names=("ProbeTableImage",),
        min_embedded_ole=1,
    ),
    TemplateContract(
        path="assets/LAND_seed_thinkcell.pptx",
        role="legacy_experimental_seed_not_production",
        required=False,
    ),
)


def _zip_ok(path: Path) -> bool:
    try:
        with ZipFile(path) as zf:
            return zf.testzip() is None
    except BadZipFile:
        return False


def _stale_token_count(path: Path) -> int:
    total = 0
    with ZipFile(path) as zf:
        for name in zf.namelist():
            if not (
                name.startswith("ppt/")
                or name.startswith("docProps/")
                or name.startswith("customXml/")
            ):
                continue
            data = zf.read(name).decode("utf-8", errors="ignore")
            total += sum(data.count(token) for token in STALE_THINKCELL_TOKENS)
    return total


def _surface_counts(path: Path) -> tuple[int, int, int, int, int]:
    prs = Presentation(path)
    tables = 0
    pics = 0
    ole = 0
    charts = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                tables += 1
            if getattr(shape, "name", "") == "Pic":
                pics += 1
            if "EMBEDDED_OLE_OBJECT" in str(getattr(shape, "shape_type", "")):
                ole += 1
            if getattr(shape, "has_chart", False):
                charts += 1
    return len(prs.slides), tables, pics, ole, charts


def inspect_template(contract: TemplateContract, *, root: Path = ROOT) -> TemplateContractResult:
    path = (root / contract.path).resolve()
    result = TemplateContractResult(
        path=str(path),
        role=contract.role,
        status="pass",
        exists=path.exists(),
    )
    if not path.exists():
        if contract.required:
            result.findings.append("missing required template asset")
            result.status = "fail"
        else:
            result.warnings.append("optional legacy asset absent")
            result.status = "warn"
        return result
    if not _zip_ok(path):
        result.findings.append("invalid PPTX zip package")
        result.status = "fail"
        return result

    try:
        slide_count, tables, pics, ole, charts = _surface_counts(path)
        names = sorted(template_named_elements(path))
        stale_count = _stale_token_count(path)
    except Exception as exc:  # pragma: no cover - CLI defensive path
        result.findings.append(str(exc))
        result.status = "fail"
        return result

    result.slide_count = slide_count
    result.native_powerpoint_tables = tables
    result.table_image_pictures = pics
    result.embedded_ole_objects = ole
    result.native_powerpoint_charts = charts
    result.named_element_count = len(names)
    result.stale_thinkcell_token_count = stale_count

    if contract.expected_slides is not None and slide_count != contract.expected_slides:
        result.findings.append(f"slide_count={slide_count}; expected {contract.expected_slides}")
    if contract.expected_names is not None and len(names) != contract.expected_names:
        result.findings.append(f"named_element_count={len(names)}; expected {contract.expected_names}")
    if contract.expected_exact_names and tuple(names) != contract.expected_exact_names:
        result.findings.append(f"named_elements={names}; expected {list(contract.expected_exact_names)}")
    if contract.min_embedded_ole is not None and ole < contract.min_embedded_ole:
        result.findings.append(f"embedded_ole_objects={ole}; expected >= {contract.min_embedded_ole}")
    if contract.max_embedded_ole is not None and ole > contract.max_embedded_ole:
        result.findings.append(f"embedded_ole_objects={ole}; expected <= {contract.max_embedded_ole}")
    if contract.min_native_charts is not None and charts < contract.min_native_charts:
        result.findings.append(f"native_powerpoint_charts={charts}; expected >= {contract.min_native_charts}")
    if contract.max_stale_tokens is not None and stale_count > contract.max_stale_tokens:
        result.findings.append(f"stale_thinkcell_token_count={stale_count}; expected <= {contract.max_stale_tokens}")
    if not contract.required and path.exists():
        result.warnings.append("legacy experimental asset exists; do not use as production template")
    result.status = "fail" if result.findings else ("warn" if result.warnings else "pass")
    return result


def build_report(*, root: Path = ROOT) -> dict[str, object]:
    results = [inspect_template(contract, root=root) for contract in CONTRACTS]
    required_results = [result for result, contract in zip(results, CONTRACTS, strict=True) if contract.required]
    status = "pass" if all(result.status == "pass" for result in required_results) else "fail"
    return {
        "schema": "land-template-contract-gate/v1",
        "status": status,
        "root": str(root),
        "results": [asdict(result) for result in results],
    }


def write_markdown(payload: dict[str, object], path: Path) -> None:
    lines = [
        "# LAND Template Contract Gate",
        "",
        f"- Status: `{payload['status']}`",
        "",
        "| Asset | Role | Status | Slides | Names | OLE | Charts | Stale TC | Findings | Warnings |",
        "|---|---|---:|---:|---:|---:|---:|---:|---|---|",
    ]
    for row in payload.get("results", []):
        result = dict(row)
        lines.append(
            "| {asset} | {role} | {status} | {slides} | {names} | {ole} | {charts} | {stale} | {findings} | {warnings} |".format(
                asset=Path(str(result["path"])).name,
                role=result["role"],
                status=result["status"],
                slides=result.get("slide_count"),
                names=result.get("named_element_count"),
                ole=result.get("embedded_ole_objects"),
                charts=result.get("native_powerpoint_charts"),
                stale=result.get("stale_thinkcell_token_count"),
                findings="; ".join(result.get("findings") or []),
                warnings="; ".join(result.get("warnings") or []),
            )
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--json-output", type=Path, default=ROOT / "state" / "__template_contract_gate.json")
    parser.add_argument("--markdown-output", type=Path, default=ROOT / "state" / "__template_contract_gate.md")
    args = parser.parse_args()
    payload = build_report()
    args.json_output.parent.mkdir(parents=True, exist_ok=True)
    args.json_output.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    write_markdown(payload, args.markdown_output)
    print(json.dumps(payload, indent=2))
    return 0 if payload["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
