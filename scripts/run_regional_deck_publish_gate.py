#!/usr/bin/env python3
"""Run regional Sales Director deck publish gates for the May 2026 batch.

This gate is intentionally stricter than "file exists": it separates the
factory/link status from leadership readiness. A deck can be technically
rebuilt and still fail publish if it carries donor-chart placeholder text,
ambiguous ARR/ACV labeling, weak original-intel coverage, or chart/table
overlap.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pptx import Presentation

from _directors import canonical_directors
from period_context import DEFAULT_PERIOD, context_for_period
from ppttc_template import template_named_elements
from sales_director_row_filters import is_internal_sales_record


ROOT = Path(__file__).resolve().parent.parent
TARGET_TABLE_IMAGE_NAMES = [
    "S04_ReviewDeltaTargets",
    "S05_ForecastQualityTable",
    "S06_HygieneSignals",
    "S07_TopDealsLand",
    "S08_TopDealsExpand",
    "S09_PendingCommercialApproval",
    "S11_RenewalPipeline",
    "S12_GRRProxyTable",
    "S13_ForecastCategoryDetail",
    "S16_OwnerCoaching",
    "S18_QTDLossSpine",
    "S19_DealHygieneSignals",
    "S21_ConcentrationTable",
    "S22_NamedRiskTriage",
    "S23_OperatingRhythm",
    "S24_AccountExpansion",
    "S25_Next14DaysCadence",
    "S26_ActionItems",
    "S27_DecisionChecklist",
]
REQUIRED_TABLE_IMAGE_SLIDES = {4, 5, 6, 7, 8, 9, 11, 12, 13, 16, 18, 19, 21, 22, 23, 24, 25, 26, 27}
# Slots that may legitimately ship as a native think-cell chart instead of a
# table-image picture once promote_native_charts_to_linked_deck.py has run
# against the linked deck. The publish gate accepts either form on these
# slots — table-image for safety, native chart when L5 promotion has landed.
NATIVE_CHART_PROMOTABLE_SLIDES = {4, 5, 6, 13, 16, 17, 18, 19, 21, 22, 25}
DONOR_PLACEHOLDER_PATTERNS = [
    "Revenues, costs, totals",
    "[USD m]",
    "User count [K]",
    "Product A",
    "BU1",
    "BU2",
]
WEAK_TITLE_PATTERNS = [
    r"\bBy owner\b",
    r"\bExec summary\b",
    r"\bPer-territory pipeline mix\b",
]
FORBIDDEN_TEXT_NEEDLES = [
    "SC Test",
    "Test Account",
    "Test MASB",
    "SimCorp Test",
    "Demo Account",
    "Sample Account",
    "Sample Co",
    "Lorem ipsum",
    "5000%",
    "9000%",
    "#NAME",
    "#NULL",
    "#REF",
    "#VALUE",
    "#DIV/0",
    "#N/A",
    "Click to add subtitle",
    "Title of the section",
]
EXCEL_ERROR_TOKENS: tuple[str, ...] = (
    "#NAME?",
    "#NULL!",
    "#REF!",
    "#VALUE!",
    "#DIV/0!",
    "#N/A",
)
WORKBOOK_ERROR_FINDINGS_CAP = 12
_MONTH_NAMES: tuple[str, ...] = (
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
)
MIN_TABLE_IMAGE_WIDTH_RATIO = 0.25
MIN_TABLE_IMAGE_HEIGHT_RATIO = 0.07
MIN_TABLE_IMAGE_ASPECT_DISTORTION = 0.88
MAX_TABLE_IMAGE_ASPECT_DISTORTION = 1.14


@dataclass(frozen=True)
class Box:
    left: int
    top: int
    right: int
    bottom: int

    @property
    def area(self) -> int:
        return max(0, self.right - self.left) * max(0, self.bottom - self.top)

    def overlap_ratio(self, other: "Box") -> float:
        x1 = max(self.left, other.left)
        y1 = max(self.top, other.top)
        x2 = min(self.right, other.right)
        y2 = min(self.bottom, other.bottom)
        overlap = max(0, x2 - x1) * max(0, y2 - y1)
        return 0.0 if self.area == 0 else overlap / self.area


@dataclass
class DirectorGate:
    director: str
    territory: str
    slug: str
    status: str
    factory_status: str
    readiness_status: str
    blockers: list[str]
    polish_items: list[str]
    artifacts: dict[str, str]
    metrics: dict[str, Any]


def slugify(value: str) -> str:
    return value.replace(" ", "-")


def _zip_ok(path: Path) -> bool:
    try:
        with ZipFile(path) as zf:
            return zf.testzip() is None
    except BadZipFile:
        return False


def _deck_text(prs: Presentation) -> str:
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "") or ""
            if text:
                chunks.append(text)
    return re.sub(r"\s+", " ", "\n".join(chunks)).strip()


def _shape_box(shape: Any) -> Box:
    return Box(
        left=int(shape.left),
        top=int(shape.top),
        right=int(shape.left + shape.width),
        bottom=int(shape.top + shape.height),
    )


def _overlap_findings(prs: Presentation) -> list[str]:
    findings: list[str] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        chart_boxes = [
            (shape.name, _shape_box(shape))
            for shape in slide.shapes
            if str(shape.shape_type) == "CHART (3)" or shape.name.startswith("Chart ")
        ]
        table_boxes = [
            (shape.name, _shape_box(shape))
            for shape in slide.shapes
            if str(shape.shape_type) == "TABLE (19)" or shape.name == "Pic"
        ]
        for chart_name, chart_box in chart_boxes:
            for table_name, table_box in table_boxes:
                ratio = chart_box.overlap_ratio(table_box)
                if ratio > 0.03:
                    findings.append(
                        f"slide {slide_idx}: chart {chart_name!r} overlaps table/image {table_name!r} ({ratio:.0%})"
                    )
    return findings


def _defined_names(path: Path) -> set[str]:
    wb = load_workbook(path, read_only=False, data_only=False)
    return {defined_name.name for defined_name in wb.defined_names.values()}


def _workbook_internal_row_findings(path: Path) -> list[str]:
    findings: list[str] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            continue
        header_values = [str(value or "") for value in headers]
        interesting = any(
            header in {"Account", "AccountName", "Opportunity", "OpportunityName", "Name"}
            for header in header_values
        )
        if not interesting:
            continue
        for row_idx, values in enumerate(rows, start=2):
            record = {header: value for header, value in zip(header_values, values) if header}
            if is_internal_sales_record(record):
                account = record.get("Account") or record.get("AccountName") or ""
                opportunity = (
                    record.get("Opportunity")
                    or record.get("OpportunityName")
                    or record.get("Name")
                    or ""
                )
                findings.append(
                    f"{path.name}:{ws.title}!{row_idx} account={account!r} opportunity={opportunity!r}"
                )
                if len(findings) >= 12:
                    return findings
    return findings


def workbook_error_findings(path: Path) -> list[str]:
    """Scan an xlsx for cached cell values that are exact Excel error tokens.

    Whole-cell match (after `.strip()`) against `EXCEL_ERROR_TOKENS` so that
    prose mentioning "#N/A" inside a sentence does not flip the gate. Capped
    at `WORKBOOK_ERROR_FINDINGS_CAP` so a wholly broken workbook does not
    produce a megabyte-sized blocker list.
    """

    findings: list[str] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str):
                        continue
                    stripped = value.strip()
                    if stripped in EXCEL_ERROR_TOKENS:
                        findings.append(f"{ws.title}!{cell.coordinate}={stripped}")
                        if len(findings) >= WORKBOOK_ERROR_FINDINGS_CAP:
                            return findings
    finally:
        wb.close()
    return findings


def forbidden_period_label_findings(text: str, *, period: str = DEFAULT_PERIOD) -> list[str]:
    """Return month/quarter labels in `text` that contradict `period`.

    Allowed for the validated 2026-Q2 lane: `May 2026` and `Q2 2026`.
    Forbidden: any other 2026 month label and any off-quarter `Q1/Q3/Q4 2026`.

    The function delegates period validation to `period_context.context_for_period`,
    so an unsupported period raises `ValueError` instead of silently passing.
    """

    context = context_for_period(period)
    period_year, quarter_str = context.period.split("-Q", 1)
    period_quarter = int(quarter_str)
    allowed_month = context.month_label.split()[0]
    other_months = [name for name in _MONTH_NAMES if name != allowed_month]
    other_quarters = [f"Q{q}" for q in (1, 2, 3, 4) if q != period_quarter]

    findings: list[str] = []
    month_pattern = r"\b(?:" + "|".join(other_months) + r")\s+" + re.escape(period_year) + r"\b"
    for match in re.finditer(month_pattern, text, flags=re.IGNORECASE):
        findings.append(match.group(0))
    quarter_pattern = (
        r"\b(?:"
        + "|".join(re.escape(q) for q in other_quarters)
        + r")\s+"
        + re.escape(period_year)
        + r"\b"
    )
    for match in re.finditer(quarter_pattern, text, flags=re.IGNORECASE):
        findings.append(match.group(0))
    return findings


def _sheet_headers(ws: Any) -> dict[str, int]:
    return {
        str(cell.value).strip(): idx
        for idx, cell in enumerate(ws[1], start=1)
        if cell.value not in (None, "")
    }


def _column_values_by_header(ws: Any, header: str) -> list[Any]:
    col_idx = _sheet_headers(ws).get(header)
    if not col_idx:
        return []
    return [
        ws.cell(row_idx, col_idx).value
        for row_idx in range(2, ws.max_row + 1)
        if ws.cell(row_idx, col_idx).value not in (None, "")
    ]


def _normalize_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip() if value not in (None, "") else ""
    return text[:10] if text else None


def _to_number(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    if not text:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    number = float(match.group(0))
    if "meur" in text.lower():
        number *= 1_000_000
    return number


def _numeric_distinct(values: list[Any]) -> set[float]:
    return {round(number, 6) for value in values if (number := _to_number(value)) is not None}


def _cell_is_negative(ws: Any, cell_ref: str) -> bool:
    value = ws[cell_ref].value
    if isinstance(value, str):
        return value.replace(" ", "").startswith("=-")
    number = _to_number(value)
    return number is not None and number < 0


def _visual_gate_findings(
    spec_path: Path, workbook_path: Path
) -> tuple[list[str], list[str], dict[str, Any]]:
    blockers: list[str] = []
    polish: list[str] = []
    metrics: dict[str, Any] = {}
    if not spec_path.exists() or not workbook_path.exists():
        return blockers, polish, metrics
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    wb = load_workbook(workbook_path, read_only=False, data_only=False)
    for obj in spec.get("objects") or []:
        object_id = str(obj.get("object_id", "(missing)"))
        gate = obj.get("visual_gate")
        if not isinstance(gate, dict):
            continue
        rule = str(gate.get("rule") or "")
        source_sheet = str(gate.get("source_sheet") or "")
        if source_sheet not in wb.sheetnames:
            blockers.append(f"{object_id}: visual gate source sheet missing: {source_sheet}")
            continue
        ws = wb[source_sheet]
        primary_visual = str(obj.get("primary_visual") or obj.get("visualization_lane") or "")
        fallback_visual = str(obj.get("fallback_visual") or "")
        metric_key = f"{object_id}:{rule}"
        if rule == "waterfall_requires_negative_closed_lost":
            cell_ref = str(gate.get("value_cell") or "")
            is_negative = bool(cell_ref and _cell_is_negative(ws, cell_ref))
            metrics[metric_key] = {"source": f"{source_sheet}!{cell_ref}", "pass": is_negative}
            if not is_negative:
                blockers.append(
                    f"{object_id}: waterfall closed-lost value is not negative at {source_sheet}!{cell_ref}"
                )
        elif rule == "timeline_requires_distinct_dates":
            header = str(gate.get("date_header") or "")
            values = [_normalize_date(value) for value in _column_values_by_header(ws, header)]
            distinct = sorted({value for value in values if value})
            minimum = int(gate.get("min_distinct_dates") or 2)
            metrics[metric_key] = {
                "source": f"{source_sheet}.{header}",
                "distinct_dates": distinct,
                "min": minimum,
            }
            if len(distinct) < minimum and "timeline" in primary_visual:
                blockers.append(
                    f"{object_id}: timeline visual is not eligible; only {len(distinct)} distinct date(s), fallback={fallback_visual}"
                )
        elif rule == "scatter_requires_two_noncollapsed_axes":
            x_header = str(gate.get("x_header") or "")
            y_header = str(gate.get("y_header") or "")
            x_values = _numeric_distinct(_column_values_by_header(ws, x_header))
            y_values = _numeric_distinct(_column_values_by_header(ws, y_header))
            min_x = int(gate.get("min_distinct_x") or 2)
            min_y = int(gate.get("min_distinct_y") or 2)
            metrics[metric_key] = {
                "source": f"{source_sheet}.{x_header}/{y_header}",
                "distinct_x": len(x_values),
                "distinct_y": len(y_values),
                "min_x": min_x,
                "min_y": min_y,
            }
            if "scatter" in primary_visual and (len(x_values) < min_x or len(y_values) < min_y):
                blockers.append(
                    f"{object_id}: scatter visual axes collapsed (distinct x={len(x_values)}, y={len(y_values)})"
                )
        elif rule == "action_register_rejects_generic_gantt":
            header = str(gate.get("date_header") or "")
            dates = [_normalize_date(value) for value in _column_values_by_header(ws, header)]
            dates = [value for value in dates if value]
            ratio = 0.0
            if dates:
                ratio = max(Counter(dates).values()) / len(dates)
            metrics[metric_key] = {
                "source": f"{source_sheet}.{header}",
                "rows": len(dates),
                "most_common_date_ratio": round(ratio, 4),
                "max_same_date_ratio": float(gate.get("max_same_date_ratio") or 0.8),
            }
            if "timeline" in primary_visual and ratio > float(
                gate.get("max_same_date_ratio") or 0.8
            ):
                blockers.append(
                    f"{object_id}: action timeline would imply false precision; due-date repeat ratio={ratio:.0%}"
                )
    return blockers, polish, metrics


def _run_connected_factory_gate(spec: Path, workbook: Path) -> tuple[bool, str]:
    result = subprocess.run(
        [
            sys.executable,
            "scripts/validate_connected_factory_spec.py",
            "--spec",
            str(spec),
            "--workbook",
            str(workbook),
        ],
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    output = "\n".join(part for part in (result.stdout.strip(), result.stderr.strip()) if part)
    return result.returncode == 0, output


def _mtime_iso(path: Path) -> str | None:
    if not path.exists():
        return None
    return datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")


def _freshness_findings(paths: dict[str, Path]) -> tuple[list[str], dict[str, str | None]]:
    """Validate freshness through the data -> Excel -> PPT chain."""

    metrics = {label: _mtime_iso(path) for label, path in paths.items()}
    blockers: list[str] = []

    connected = paths["connected_workbook"]
    table_image = paths["table_image_workbook"]
    deck = paths["linked_deck"]

    for label in ("trends", "brief", "land_xlsx", "land_model_xlsx"):
        source = paths[label]
        if (
            source.exists()
            and connected.exists()
            and source.stat().st_mtime > connected.stat().st_mtime
        ):
            blockers.append(f"connected workbook is older than {label}; rebuild connected factory")

    if (
        connected.exists()
        and table_image.exists()
        and connected.stat().st_mtime > table_image.stat().st_mtime
    ):
        blockers.append(
            "table-image workbook is older than connected workbook; rebuild table-image ranges"
        )

    if (
        table_image.exists()
        and deck.exists()
        and table_image.stat().st_mtime > deck.stat().st_mtime
    ):
        blockers.append(
            "linked deck is older than table-image workbook; refresh PowerPoint table images"
        )

    return blockers, metrics


def _gate_director(
    period: str, director: dict[str, Any], *, required_text: list[str]
) -> DirectorGate:
    name = str(director["name"])
    territory = str(director["scope_label"])
    slug = slugify(name)
    director_dir = ROOT / "state" / period / slug
    deck = director_dir / f"{slug}-LAND-{period}-table-image-linked.pptx"
    workbook = director_dir / "factory" / "connected" / "connected_factory_table_images.xlsx"
    base_workbook = director_dir / "factory" / "connected" / "connected_factory.xlsx"
    spec = director_dir / "factory" / "connected" / "connected_factory_spec.json"
    base_deck = director_dir / f"{slug}-LAND-{period}.pptx"
    brief = director_dir / "brief.md"
    trends = director_dir / "trends.json"
    land_xlsx = director_dir / "land.xlsx"
    land_model_xlsx = director_dir / "land.model.xlsx"
    regional_intel = director_dir / "factory" / "regional_intelligence_spec.json"

    blockers: list[str] = []
    polish: list[str] = []
    metrics: dict[str, Any] = {}

    for label, path in {
        "linked_deck": deck,
        "table_image_workbook": workbook,
        "connected_workbook": base_workbook,
        "connected_spec": spec,
        "brief": brief,
        "trends": trends,
        "land_xlsx": land_xlsx,
        "land_model_xlsx": land_model_xlsx,
        "regional_intelligence_spec": regional_intel,
    }.items():
        if not path.exists() or path.stat().st_size == 0:
            blockers.append(f"missing {label}: {path}")

    factory_ok = not blockers
    if factory_ok:
        freshness_blockers, freshness_metrics = _freshness_findings(
            {
                "linked_deck": deck,
                "table_image_workbook": workbook,
                "connected_workbook": base_workbook,
                "trends": trends,
                "brief": brief,
                "land_xlsx": land_xlsx,
                "land_model_xlsx": land_model_xlsx,
            }
        )
        metrics["freshness"] = freshness_metrics
        blockers.extend(freshness_blockers)
        if not _zip_ok(deck):
            blockers.append("linked PPTX zip integrity failed")
        spec_ok, spec_output = _run_connected_factory_gate(spec, base_workbook)
        metrics["connected_factory_gate"] = "pass" if spec_ok else spec_output[:1000]
        if not spec_ok:
            blockers.append("connected factory spec/workbook validation failed")
        visual_blockers, visual_polish, visual_metrics = _visual_gate_findings(spec, base_workbook)
        metrics["visual_gate"] = visual_metrics
        blockers.extend(visual_blockers)
        polish.extend(visual_polish)
        names = _defined_names(workbook)
        missing_names = sorted(set(TARGET_TABLE_IMAGE_NAMES) - names)
        metrics["table_image_defined_names"] = sorted(set(TARGET_TABLE_IMAGE_NAMES) & names)
        if missing_names:
            blockers.append(
                f"table-image workbook missing defined names: {', '.join(missing_names)}"
            )
        for workbook_label, workbook_path in {
            "table-image workbook": workbook,
            "connected workbook": base_workbook,
        }.items():
            internal_rows = _workbook_internal_row_findings(workbook_path)
            if internal_rows:
                blockers.append(
                    f"{workbook_label} contains internal/test account or opportunity rows: "
                    + "; ".join(internal_rows[:6])
                )
            error_findings = workbook_error_findings(workbook_path)
            if error_findings:
                blockers.append(
                    f"{workbook_label} contains cached Excel error cells: "
                    + "; ".join(error_findings[:6])
                )

    if deck.exists():
        prs = Presentation(deck)
        metrics["slide_count"] = len(prs.slides)
        if len(prs.slides) != 28:
            blockers.append(f"expected 28 slides, found {len(prs.slides)}")
        text = _deck_text(prs)
        for required in required_text:
            if required.lower() not in text.lower():
                blockers.append(f"deck text missing required basis label: {required}")
        placeholders = [
            needle for needle in DONOR_PLACEHOLDER_PATTERNS if needle.lower() in text.lower()
        ]
        if placeholders:
            blockers.append(f"donor/chart placeholder text remains: {', '.join(placeholders)}")
        forbidden_text = [
            needle for needle in FORBIDDEN_TEXT_NEEDLES if needle.lower() in text.lower()
        ]
        if forbidden_text:
            blockers.append(f"publish-forbidden text remains: {', '.join(forbidden_text)}")
        period_label_findings = forbidden_period_label_findings(text, period=period)
        if period_label_findings:
            blockers.append(
                "stale/off-period labels in deck text: " + ", ".join(period_label_findings[:8])
            )
        weak_titles = [
            pattern
            for pattern in WEAK_TITLE_PATTERNS
            if re.search(pattern, text, flags=re.IGNORECASE)
        ]
        if weak_titles:
            polish.append(f"weak/generic slide language remains: {', '.join(weak_titles)}")
        overlaps = _overlap_findings(prs)
        if overlaps:
            blockers.extend(overlaps[:8])
        pic_slides = sorted(
            {
                slide_idx
                for slide_idx, slide in enumerate(prs.slides, start=1)
                if any(shape.name == "Pic" for shape in slide.shapes)
            }
        )
        native_chart_slides = sorted(
            {
                slide_idx
                for slide_idx, slide in enumerate(prs.slides, start=1)
                if any(getattr(shape, "has_chart", False) for shape in slide.shapes)
            }
        )
        metrics["table_image_pic_slides"] = pic_slides
        metrics["native_chart_slides"] = native_chart_slides
        # Promotable slots may ship as either a table-image picture OR a
        # native think-cell chart (after promote_native_charts_to_linked_deck.py).
        # Required-but-not-promotable slots must still be pictures.
        strict_required = REQUIRED_TABLE_IMAGE_SLIDES - NATIVE_CHART_PROMOTABLE_SLIDES
        promotable_required = REQUIRED_TABLE_IMAGE_SLIDES & NATIVE_CHART_PROMOTABLE_SLIDES
        accepted_promotable = set(pic_slides) | set(native_chart_slides)
        missing_strict = sorted(strict_required - set(pic_slides))
        missing_promotable = sorted(promotable_required - accepted_promotable)
        missing_pic_slides = sorted(set(missing_strict) | set(missing_promotable))
        if missing_pic_slides:
            blockers.append(
                f"missing refreshed table-image picture on slide(s): {missing_pic_slides}"
            )
        undersized: list[str] = []
        for slide_idx, slide in enumerate(prs.slides, start=1):
            pics = [shape for shape in slide.shapes if shape.name == "Pic"]
            if not pics:
                continue
            largest = max(pics, key=lambda shape: int(shape.width) * int(shape.height))
            width_ratio = int(largest.width) / int(prs.slide_width)
            height_ratio = int(largest.height) / int(prs.slide_height)
            if (
                width_ratio < MIN_TABLE_IMAGE_WIDTH_RATIO
                or height_ratio < MIN_TABLE_IMAGE_HEIGHT_RATIO
            ):
                undersized.append(f"slide {slide_idx}: {width_ratio:.0%}w x {height_ratio:.0%}h")
        if undersized:
            blockers.append(
                "undersized linked table-image object(s): " + "; ".join(undersized[:10])
            )
        stretched: list[str] = []
        for slide_idx, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.name != "Pic":
                    continue
                try:
                    image_width, image_height = shape.image.size
                except Exception:
                    continue
                if not image_width or not image_height or not int(shape.height):
                    continue
                image_ratio = image_width / image_height
                shape_ratio = int(shape.width) / int(shape.height)
                distortion = shape_ratio / image_ratio if image_ratio else 1.0
                if (
                    distortion < MIN_TABLE_IMAGE_ASPECT_DISTORTION
                    or distortion > MAX_TABLE_IMAGE_ASPECT_DISTORTION
                ):
                    stretched.append(
                        f"slide {slide_idx}: shape_ratio={shape_ratio:.2f} image_ratio={image_ratio:.2f}"
                    )
        metrics["table_image_aspect_findings"] = stretched
        if stretched:
            blockers.append("stretched linked table-image object(s): " + "; ".join(stretched[:10]))
        native_table_slides = sorted(
            {
                slide_idx
                for slide_idx, slide in enumerate(prs.slides, start=1)
                if any(str(shape.shape_type) == "TABLE (19)" for shape in slide.shapes)
            }
        )
        metrics["native_table_slides"] = native_table_slides
        if native_table_slides:
            blockers.append(f"native PowerPoint tables remain on slide(s): {native_table_slides}")
        named_elements = template_named_elements(deck)
        metrics["named_thinkcell_elements_present"] = len(
            [name for name in TARGET_TABLE_IMAGE_NAMES if name in named_elements]
        )
        missing_named_elements = [
            name for name in TARGET_TABLE_IMAGE_NAMES if name not in named_elements
        ]
        if missing_named_elements:
            blockers.append(
                f"linked deck missing think-cell table-image names: {', '.join(missing_named_elements)}"
            )

    if base_deck.exists() and deck.exists() and base_deck.stat().st_mtime > deck.stat().st_mtime:
        blockers.append("base deck is newer than linked deck; linked deck may be stale")

    if regional_intel.exists():
        try:
            intel = json.loads(regional_intel.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            blockers.append(f"regional intelligence spec is invalid JSON: {exc}")
            intel = {}
        checks = intel.get("readiness_checks", {}) if isinstance(intel, dict) else {}
        metrics["regional_intelligence_checks"] = checks
        for key in [
            "has_original_sidecar",
            "has_original_intelligence_rows",
            "has_gold_or_audit_context",
            "has_first_two_week_actions",
            "has_named_q2_deals",
            "arr_acv_separated",
            "no_generic_apac_label_for_non_apac",
        ]:
            if checks.get(key) is not True:
                polish.append(f"regional intelligence spec check not proven: {key}")
    elif slug != "Jesper-Tyrer":
        polish.append(
            "needs region-specific original ETL deck intelligence audit; APAC-level manual coverage not yet proven"
        )
        polish.append(
            "needs director-specific first-two-week action language, not generic chart commentary"
        )

    factory_status = (
        "pass"
        if not [b for b in blockers if "placeholder" not in b and "language" not in b]
        else "fail"
    )
    readiness_status = "pass" if not blockers and not polish else "needs_work"
    status = "pass" if readiness_status == "pass" else ("fail" if blockers else "needs_work")

    return DirectorGate(
        director=name,
        territory=territory,
        slug=slug,
        status=status,
        factory_status=factory_status,
        readiness_status=readiness_status,
        blockers=blockers,
        polish_items=polish,
        artifacts={
            "linked_deck": str(deck),
            "table_image_workbook": str(workbook),
            "connected_spec": str(spec),
            "brief": str(brief),
            "trends": str(trends),
            "regional_intelligence_spec": str(regional_intel),
        },
        metrics=metrics,
    )


def _write_markdown(gates: list[DirectorGate], output: Path, *, title: str) -> None:
    lines: list[str] = [
        f"# {title} Regional Deck Publish Gate",
        "",
        "This report separates rebuild/link readiness from leadership publish readiness.",
        "",
        "## Summary",
        "",
        "| Director | Territory | Status | Factory | Readiness | Blockers | Polish |",
        "|---|---|---:|---:|---:|---:|---:|",
    ]
    for gate in gates:
        lines.append(
            f"| {gate.director} | {gate.territory} | {gate.status} | {gate.factory_status} | "
            f"{gate.readiness_status} | {len(gate.blockers)} | {len(gate.polish_items)} |"
        )
    lines.extend(["", "## Detail", ""])
    for gate in gates:
        lines.extend(
            [
                f"### {gate.director} - {gate.territory}",
                "",
                f"- Status: `{gate.status}`",
                f"- Linked deck: `{gate.artifacts['linked_deck']}`",
                f"- Connected Excel: `{gate.artifacts['table_image_workbook']}`",
                f"- Table-image slides: `{gate.metrics.get('table_image_pic_slides', [])}`",
                "",
            ]
        )
        if gate.blockers:
            lines.append("Blockers:")
            lines.extend(f"- {item}" for item in gate.blockers)
            lines.append("")
        if gate.polish_items:
            lines.append("Polish / next work:")
            lines.extend(f"- {item}" for item in gate.polish_items)
            lines.append("")
        visual_gate = gate.metrics.get("visual_gate") or {}
        if visual_gate:
            lines.append("Visual gate metrics:")
            for key, value in visual_gate.items():
                lines.append(f"- `{key}`: `{json.dumps(value, sort_keys=True)}`")
            lines.append("")
    output.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument(
        "--jobs",
        type=int,
        default=1,
        help="Parallel director gates for local file/workbook checks.",
    )
    parser.add_argument("--director-slug", help="Run the publish gate for one director only.")
    args = parser.parse_args()
    try:
        period_context = context_for_period(args.period)
    except ValueError as exc:
        print(f"error: {exc}")
        return 2

    output_dir = args.output_dir or (ROOT / "state" / args.period / "__regional__" / "publish_gate")
    output_dir.mkdir(parents=True, exist_ok=True)
    directors = canonical_directors()
    if args.director_slug:
        directors = [
            director
            for director in directors
            if slugify(str(director["name"])) == args.director_slug
        ]
        if not directors:
            raise SystemExit(f"unknown director slug: {args.director_slug}")
    if args.jobs > 1 and len(directors) > 1:
        with ThreadPoolExecutor(max_workers=min(args.jobs, len(directors))) as executor:
            gates = list(
                executor.map(
                    lambda director: _gate_director(
                        args.period,
                        director,
                        required_text=list(period_context.required_deck_text),
                    ),
                    directors,
                )
            )
    else:
        gates = [
            _gate_director(
                args.period, director, required_text=list(period_context.required_deck_text)
            )
            for director in directors
        ]
    json_path = output_dir / "regional_publish_gate.json"
    md_path = output_dir / "regional_publish_gate.md"
    json_path.write_text(
        json.dumps([asdict(gate) for gate in gates], indent=2) + "\n", encoding="utf-8"
    )
    _write_markdown(gates, md_path, title=period_context.month_label)

    print(f"json={json_path}")
    print(f"markdown={md_path}")
    for gate in gates:
        print(
            f"{gate.status.upper():>10} {gate.director:<22} "
            f"factory={gate.factory_status} readiness={gate.readiness_status} "
            f"blockers={len(gate.blockers)} polish={len(gate.polish_items)}"
        )
    return 0 if all(gate.status == "pass" for gate in gates) else 2


if __name__ == "__main__":
    raise SystemExit(main())
