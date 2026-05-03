"""Build per-director think-cell-ready PowerPoint templates.

This module takes the branded placeholder deck at `assets/LAND_template.pptx`,
bakes the non-think-cell slides with native PowerPoint text/tables, then
injects named think-cell chart donors into the analytical slides at the OpenXML
package layer. The generated template is the `.pptx` that `.ppttc` should
reference.
"""

from __future__ import annotations

import copy
import math
import numbers
import posixpath
import re
import struct
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from sales_director_row_filters import is_internal_sales_record
from thinkcell_cfb import CfbStreamEdit, replace_cfb_stream_data

ROOT = Path(__file__).resolve().parent.parent

BAR_COLUMN_DONOR = Path(
    "/Library/Application Support/Microsoft/think-cell/templates/think-cell Charts/Bar, Column/Bar, Column.potx"
)
MEKKO_DONOR = Path(
    "/Library/Application Support/Microsoft/think-cell/templates/think-cell Charts/Mekko/Mekko.potx"
)
WATERFALL_DONOR = Path(
    "/Library/Application Support/Microsoft/think-cell/templates/think-cell Charts/Waterfall/Waterfall.potx"
)

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

CFB_FREE = 0xFFFFFFFF
CFB_END_OF_CHAIN = 0xFFFFFFFE
CFB_FAT_SECTOR = 0xFFFFFFFD
CFB_DIFAT_SECTOR = 0xFFFFFFFC

ET.register_namespace("", P_NS)
ET.register_namespace("a", A_NS)
ET.register_namespace("r", R_NS)

COLOR_HEADER_FILL = RGBColor(0x08, 0x3E, 0xA7)
COLOR_HEADER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_ROW_ALT = RGBColor(0xF2, 0xF4, 0xF7)
COLOR_BODY_TEXT = RGBColor(0x1A, 0x1D, 0x31)
COLOR_MUTED_TEXT = RGBColor(0x66, 0x66, 0x66)
COLOR_RULE = RGBColor(0xD6, 0xDE, 0xEA)
COLOR_ACCENT = RGBColor(0xEF, 0x3E, 0x4A)


@dataclass(frozen=True)
class DonorSlide:
    template_path: Path
    slide_number: int
    exclude_names: tuple[str, ...] = ()
    exclude_text_prefixes: tuple[str, ...] = ()


@dataclass(frozen=True)
class ChartInjection:
    slide_number: int
    chart_name: str
    donor: DonorSlide
    text_replacements: tuple[tuple[str, str], ...] = ()


COMMON_DONOR_EXCLUDES = (
    "Title 4",
    "tc_columnhead",
    "tc_columnheadline",
    "Rectangle 18",
    "Text Placeholder 15",
    "Rechteck 137",
)

BAR_DONOR = DonorSlide(
    template_path=BAR_COLUMN_DONOR,
    slide_number=6,
    exclude_names=COMMON_DONOR_EXCLUDES
    + (
        "Arrow: Right 20",
        "Arrow: Right 54",
        "Straight Connector 355",
        "Straight Connector 361",
        "Straight Connector 55",
        "Rectangle 124",
        "Rectangle 123",
    ),
    exclude_text_prefixes=(
        "Margin",
        "Market avg.",
        "Our target",
        "Our brand",
        "Competitors",
    ),
)

COLUMN_DONOR = DonorSlide(
    template_path=BAR_COLUMN_DONOR,
    slide_number=3,
    exclude_names=COMMON_DONOR_EXCLUDES
    + (
        "Straight Connector 113",
        "Straight Connector 114",
        "Straight Connector 115",
        "Oval 33",
        "Rectangle 19",
        "Rectangle 22",
        "Rectangle 2",
        "Rectangle 10",
    ),
)

COMBO_DONOR = DonorSlide(
    template_path=BAR_COLUMN_DONOR,
    slide_number=4,
    exclude_names=COMMON_DONOR_EXCLUDES,
)

STACKED_COLUMN_DONOR = DonorSlide(
    template_path=BAR_COLUMN_DONOR,
    slide_number=8,
    exclude_names=COMMON_DONOR_EXCLUDES,
)

MEKKO_DONOR_SPEC = DonorSlide(
    template_path=MEKKO_DONOR,
    slide_number=1,
    exclude_names=COMMON_DONOR_EXCLUDES,
)

WATERFALL_DONOR_SPEC = DonorSlide(
    template_path=WATERFALL_DONOR,
    slide_number=1,
    exclude_names=COMMON_DONOR_EXCLUDES,
)

CHART_INJECTIONS: tuple[ChartInjection, ...] = (
    ChartInjection(
        4,
        "S04_PipeMovement",
        WATERFALL_DONOR_SPEC,
        (
            ("Revenues, costs, totals [USD m]", "ARR (mEUR)"),
            ("BU1", "Additions"),
            ("BU2", "Outflows"),
            ("Totals", "Total"),
        ),
    ),
    ChartInjection(5, "S05_PipelineByStage", BAR_DONOR),
    ChartInjection(6, "S06_PipelineAging", BAR_DONOR),
    ChartInjection(
        13,
        "S13_ForecastCategory",
        COLUMN_DONOR,
        (("User count [K]", "ARR (mEUR)"),),
    ),
    ChartInjection(15, "S15_ByOwner", BAR_DONOR),
    ChartInjection(
        16,
        "S16_StageByIndustry",
        STACKED_COLUMN_DONOR,
        (("Product A", "Stage mix"),),
    ),
    ChartInjection(17, "S17_TerritoryPerformance", BAR_DONOR),
    ChartInjection(
        18,
        "S18_WinsLossesQTD",
        COLUMN_DONOR,
        (("User count [K]", "Value (mEUR)"),),
    ),
    ChartInjection(
        19,
        "S19_Velocity",
        COLUMN_DONOR,
        (("User count [K]", "Median age (days)"),),
    ),
    ChartInjection(
        21,
        "S21_ConcentrationRiskChart",
        COLUMN_DONOR,
        (("User count [K]", "Share (%)"),),
    ),
    ChartInjection(22, "S22_StaleActivity", BAR_DONOR),
    ChartInjection(
        25,
        "S25_PipelineCreationVelocity",
        COLUMN_DONOR,
        (("User count [K]", "New ARR (mEUR)"),),
    ),
)

DECK_MONTH_LABEL = "May 2026"
DECK_DATE_LABEL = "Friday, May 1, 2026"
SNAPSHOT_LABEL = "Salesforce snapshot as of 2026-04-30"

JESPER_ORIGINAL_Q1_TARGET_ROWS = [
    ["Q1 opened", "27 deals / EUR 17.1M unweighted ARR", "Baseline for pipeline creation and inspection volume."],
    ["Q1 committed", "EUR 2.7M unweighted ARR / EUR 2.3M weighted ARR", "Commit quality should be tested against final Q1 delivery."],
    ["Q1 won", "1 deal / EUR 1.9M unweighted ARR", "Sony Life was the only Q1 Land win called out in the prior APAC review."],
    ["Q1 lost", "14 deals / EUR 5.5M unweighted ARR", "Loss discipline and reason-code capture stay in the director review."],
    ["Q1 slips", "9 deals / EUR 9.4M unweighted ARR", "Seven slipped again after Q1; push discipline remains a May review theme."],
    ["FY26 renewals", "3 deals / EUR 33.5M ACV", "Renewal motion is ACV and stays separate from Land unweighted ARR."],
]

JESPER_ORIGINAL_Q1_SLIP_ROWS = [
    ["Temasek Customer Care", "EUR 3.0M", "Slipped from Q1; subsequently lost in April."],
    ["RBI", "EUR 1.9M", "Slipped from Q1; subsequently lost."],
    ["Danantara", "EUR 1.8M", "Slipped into Q2 and remains a top readiness risk."],
    ["KWAP", "EUR 1.3M", "Q1 slip requiring close-date realism review."],
    ["LTH", "EUR 1.1M", "Q1 slip and post-Q1 push; current next step points beyond Q2."],
]

JESPER_ORIGINAL_RENEWAL_BASELINE = {
    "fy26_count": 3,
    "fy26_acv_eur": 33_500_000,
    "q2_count": 1,
    "q2_acv_eur": 455_293,
    "q3_count": 2,
    "q3_acv_eur": 33_093_198,
}

JESPER_ORIGINAL_Q1_LOSS_ROWS = [
    ["Q1 Land losses", "14", "EUR 5.5M", "Prior APAC sidecar; keep separate from Q2 QTD."],
    ["Missing reason code", "9", "EUR 0.0M", "No Opportunity rows need hygiene cleanup even when unweighted ARR is zero."],
    ["Lost with reason", "5", "EUR 5.5M", "Buying process stopped, own solution, competitor, SimCorp stop."],
    ["Stage-at-loss mix", "9 No Opp / 5 Lost", "EUR 5.5M", "Use for qualification and stage-exit coaching."],
    ["Forecast accuracy lens", "1W / 16L", "EUR 8.6M lost", "Broader original closed-out view; win rate was 6%."],
]

JESPER_Q2_Q3_RISK_ROWS = [
    ["Coolabah", "Coolabah capital - F2M - ILF", "3 - Engagement", "2026-09-30", "EUR 2.0M", "3 pushes; Q3 slip risk"],
    ["Danantara", "Danantara - F2B", "3 - Engagement", "2026-06-30", "EUR 1.8M", "Silent 90d; no next step; Q2 close risk"],
    ["BOCI", "BOCI Prudential - M2B", "5 - Preferred", "2026-06-30", "EUR 0.5M", "8 pushes; procurement timing risk"],
    ["Mandiri", "Bank Mandiri - Fund Acctg / TA", "3 - Engagement", "2026-06-30", "EUR 0.5M", "4 pushes; approval gap"],
]

JESPER_ORIGINAL_CONCENTRATION_ROWS = [
    ["Top 7 open deals", "EUR 9.6M", "Original APAC open-book concentration baseline."],
    ["Top 5 = 89%", "Amova Asset Management EUR 2.6M unweighted ARR", "Use as a prior concentration spine, not current-state unweighted ARR."],
    ["Original open book", "12 Land deals", "Compare current Q2 readiness to the Apr 20 Land-only book."],
]

JESPER_ORIGINAL_PUSH_DISCIPLINE_ROWS = [
    ["Owner push spine", "3 owners carry 50 pushes", "EUR 22.0M exposed; coach by push intensity and qualification quality."],
    ["Pushed deals", "11 open deals pushed", "Edwina Chow owns 5; pattern review stays in the director discussion."],
    ["High-push tail", "2 deals at 5+ pushes", "EUR 1.0M unweighted ARR at highest push risk."],
    ["Watch band", "5 deals at 3-4 pushes", "EUR 3.0M unweighted ARR needs close-date discipline."],
]

JESPER_ORIGINAL_FORECAST_MIX_NOTE = (
    "Original APAC forecast mix: EUR 5.6M unweighted ARR across 12 open Pipeline Inspection deals; "
    "Commit = EUR 2.7M unweighted ARR / EUR 2.3M weighted ARR (48%)."
)

JESPER_ORIGINAL_COMMERCIAL_APPROVAL_ROWS = [
    ["Approved 2026", "Coolabah capital - F2M - ILF", "EUR 2.0M unweighted ARR", "2026-03-19", "Enzo Cotroneo"],
    ["Approved 2026", "Danantara - F2B", "EUR 1.8M unweighted ARR", "2026-03-11", "Jesper Tyrer"],
    ["Prior-year approved", "LTH; Krungthai; BOCI", "EUR 2.7M unweighted ARR", "2024-25", "Q2 close governance context"],
    ["Pending / candidate", "Amova AM - IBOR; SSO - F2B", "EUR 5.0M unweighted ARR", "2026-03-11", "Amova surfaced as candidate"],
]


def build_director_template(
    *,
    artifacts: Any,
    base_template_path: Path,
    trends: dict[str, Any],
    brief_sections: dict[str, list[str]],
    model: Any,
    legacy: Any,
    inject_charts: bool = True,
) -> Path:
    """Create the director-specific `.pptx` consumed by `.ppttc`."""

    out_path = artifacts.director_dir / f"{artifacts.slug}-LAND-{artifacts.period}-template.pptx"
    is_apac = _is_apac_artifacts(artifacts)

    prs = Presentation(str(base_template_path))
    _replace_tokens_in_deck(
        prs,
        {
            "{director_name}": artifacts.name,
            "{period}": DECK_MONTH_LABEL,
            "{scope_label}": artifacts.scope_label,
        },
    )
    _fill_section_dividers(prs)

    _fill_cover_slide(prs.slides[0], artifacts=artifacts)
    _fill_exec_summary(prs.slides[1], artifacts=artifacts, trends=trends, model=model, legacy=legacy, is_apac=is_apac)
    _fill_review_delta_slide(prs.slides[3], artifacts=artifacts, trends=trends, legacy=legacy, is_apac=is_apac)
    _fill_forecast_quality_slide(prs.slides[4], trends=trends, legacy=legacy, is_apac=is_apac)
    _fill_pipeline_aging_slide(prs.slides[5], legacy)
    _fill_q2_readiness_slide(prs.slides[6], legacy)
    _fill_close_plan_slide(prs.slides[7], legacy, is_apac=is_apac)
    _fill_pending_approval_slide(prs.slides[8], legacy, is_apac=is_apac)
    _fill_fy26_renewals_slide(prs.slides[10], legacy, is_apac=is_apac)
    _fill_renewal_guardrail_slide(prs.slides[11], legacy, is_apac=is_apac)
    _fill_forecast_category_detail_slide(prs.slides[12], legacy)
    _fill_by_owner_slide(prs.slides[14], legacy, is_apac=is_apac)
    _fill_owner_coaching_slide(prs.slides[15], legacy, is_apac=is_apac)
    _fill_territory_pipeline_slide(prs.slides[16], legacy, is_apac=is_apac)
    _fill_qtd_wins_losses_slide(prs.slides[17], legacy, is_apac=is_apac)
    _fill_deal_hygiene_slide(prs.slides[18], legacy, is_apac=is_apac)
    _fill_concentration_slide(prs.slides[20], model, is_apac=is_apac)
    _fill_named_stale_slide(prs.slides[21], legacy, is_apac=is_apac)
    _fill_sales_velocity_slide(prs.slides[22], legacy, is_apac=is_apac)
    _fill_account_expansion_slide(prs.slides[23], trends, legacy, is_apac=is_apac)
    _fill_next_14_days_slide(prs.slides[24], trends, legacy, is_apac=is_apac)
    _fill_action_items_slide(prs.slides[25], trends, legacy, is_apac=is_apac)
    _fill_risks_outlook_slide(prs.slides[26], trends=trends, legacy=legacy, is_apac=is_apac)
    _remove_empty_prompt_placeholders(prs)

    prs.save(str(out_path))
    if inject_charts:
        _inject_donor_charts(out_path)
    return out_path


def _is_apac_artifacts(artifacts: Any) -> bool:
    return getattr(artifacts, "slug", "") == "Jesper-Tyrer"


def _fill_section_dividers(prs: Presentation) -> None:
    sections = {
        3: ("Pipeline", "Q2 closeable Land+Expand ARR"),
        10: ("Renewals", "FY26 renewal ACV watchlist and guardrails"),
        14: ("Territory", "Owner, forecast, and country mix"),
        20: ("Risks", "Concentration, deal hygiene, and operating cadence"),
    }
    for slide_number, (title, subtitle) in sections.items():
        slide = prs.slides[slide_number - 1]
        _set_shape_text(_find_shape(slide, "Text Placeholder 1"), title)
        _set_shape_text(_find_shape(slide, "Text Placeholder 2"), subtitle)


def _remove_empty_prompt_placeholders(prs: Presentation) -> None:
    for slide in prs.slides:
        for shape in list(slide.shapes):
            if shape.name not in {"Text Placeholder 4", "Text Placeholder 8"}:
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            if shape.text_frame.text.strip():
                continue
            _remove_shape(shape)


def template_has_named_elements(template_path: Path) -> bool:
    """Detect generated templates by looking for non-empty think-cell OLE names."""

    return bool(template_named_elements(template_path))


def template_named_elements(template_path: Path) -> set[str]:
    """Return all automation names discoverable in a think-cell template."""

    binary_pattern = re.compile(rb"<m_strName>([^<]+)</m_strName>")
    escaped_xml_pattern = re.compile(r"&lt;m_strName&gt;(.*?)&lt;/m_strName&gt;")

    names: set[str] = set()
    with ZipFile(template_path) as zf:
        for name in zf.namelist():
            if name.startswith("ppt/embeddings/") and name.endswith(".bin"):
                for stream in _cfb_streams(zf.read(name)):
                    if stream.name == "think-cellXML":
                        for match in binary_pattern.finditer(stream.data):
                            value = match.group(1).decode("utf-8")
                            if value:
                                names.add(value)
                continue
            if name.startswith("ppt/slides/") and name.endswith(".xml"):
                text = zf.read(name).decode("utf-8", "ignore")
                names.update(escaped_xml_pattern.findall(text))
    return names


def _replace_tokens_in_deck(prs: Presentation, tokens: dict[str, str]) -> None:
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text
            if not text:
                continue
            replaced = text
            for needle, value in tokens.items():
                replaced = replaced.replace(needle, value)
            if replaced != text:
                shape.text_frame.text = replaced


def _find_shape(slide: Any, name: str) -> Any | None:
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def _remove_shape(shape: Any | None) -> None:
    if shape is None:
        return
    element = shape._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _set_shape_text(shape: Any | None, text: str) -> None:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    shape.text_frame.text = text


def _set_slide_header(slide: Any, *, title: str, subtitle: str) -> None:
    _set_shape_text(_find_shape(slide, "Text Placeholder 3") or _find_shape(slide, "Text Placeholder 7"), title)
    _set_shape_text(_find_shape(slide, "Text Placeholder 2"), subtitle)


def _content_bounds(slide: Any) -> tuple[Emu, Emu, Emu, Emu]:
    box = _find_shape(slide, "Rectangle 5")
    if box is not None:
        return box.left, box.top, box.width, box.height
    placeholder = _find_shape(slide, "Content Placeholder 1")
    if placeholder is None:
        raise RuntimeError(f"Slide {slide.slide_id} does not contain the expected content area")
    return placeholder.left, placeholder.top, placeholder.width, placeholder.height


def _clear_standard_placeholder(slide: Any) -> tuple[Emu, Emu, Emu, Emu]:
    bounds = _content_bounds(slide)
    _remove_shape(_find_shape(slide, "Content Placeholder 1"))
    _remove_shape(_find_shape(slide, "Rectangle 5"))
    return bounds


def _clear_exec_placeholder(slide: Any) -> None:
    _remove_shape(_find_shape(slide, "Rectangle 9"))


def _shape_text_frame(shape: Any) -> Any:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    return tf


def _set_paragraph_text(
    paragraph: Any,
    text: str,
    *,
    font_size: Pt,
    bold: bool = False,
    color: RGBColor = COLOR_BODY_TEXT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    italic: bool = False,
) -> None:
    paragraph.alignment = align
    paragraph.space_after = Pt(4)
    run = paragraph.add_run()
    run.text = text
    run.font.size = font_size
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = "Aptos"


def _add_textbox(
    slide: Any,
    left: Emu,
    top: Emu,
    width: Emu,
    height: Emu,
    text: str,
    *,
    font_size: Pt,
    bold: bool = False,
    color: RGBColor = COLOR_BODY_TEXT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    italic: bool = False,
) -> Any:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = _shape_text_frame(box)
    paragraph = tf.paragraphs[0]
    _set_paragraph_text(
        paragraph,
        text,
        font_size=font_size,
        bold=bold,
        color=color,
        align=align,
        italic=italic,
    )
    return box


def _add_rule(slide: Any, left: Emu, top: Emu, width: Emu, *, color: RGBColor = COLOR_RULE) -> None:
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Emu(12700))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()


def _format_currency(value: Any) -> str:
    if not isinstance(value, numbers.Real) or isinstance(value, bool):
        return str(value or "")
    numeric = float(value)
    if not math.isfinite(numeric):
        return ""
    absolute = abs(numeric)
    if absolute >= 1_000_000:
        return f"EUR {numeric / 1_000_000:.1f}M"
    if absolute >= 1_000:
        return f"EUR {numeric:,.0f}"
    return f"EUR {numeric:.0f}"


def _format_percentage(value: Any) -> str:
    if not isinstance(value, numbers.Real) or isinstance(value, bool):
        return str(value or "")
    numeric = float(value)
    if not math.isfinite(numeric):
        return ""
    # Salesforce snapshots store probability as either 50 or 0.50 depending
    # on the source tab. Normalize both to a human percent once.
    if abs(numeric) > 1:
        numeric /= 100
    percent = numeric * 100
    return f"{percent:.0f}%" if percent % 1 == 0 else f"{percent:.1f}%"


def _format_table_cell(header: str, value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")

    header_lc = header.lower()
    if isinstance(value, numbers.Real) and not isinstance(value, bool):
        numeric = float(value)
        if not math.isfinite(numeric):
            return ""
        if "%" in header or "proxy" in header_lc or "prob" in header_lc:
            return _format_percentage(numeric)
        if "arr" in header_lc or "acv" in header_lc or "eur" in header_lc or "value" == header_lc:
            return _format_currency(numeric)
        if header.strip() == "#" or "score" in header_lc or "motions" in header_lc or "days" in header_lc:
            return str(int(round(numeric)))
        return f"{numeric:,.1f}" if numeric % 1 else str(int(numeric))

    return str(value)


def _inches(value: float) -> Emu:
    return Emu(int(value * 914400))


def _number_from_display(value: Any) -> float | None:
    if isinstance(value, numbers.Real) and not isinstance(value, bool):
        numeric = float(value)
        return numeric if math.isfinite(numeric) else None
    text = str(value or "").strip()
    if not text:
        return None
    match = re.search(r"(-?\d+(?:\.\d+)?)\s*mEUR", text, flags=re.IGNORECASE)
    if match:
        return float(match.group(1)) * 1_000_000
    match = re.search(r"(-?\d[\d,]*(?:\.\d+)?)", text)
    if not match:
        return None
    try:
        return float(match.group(1).replace(",", ""))
    except ValueError:
        return None


def _format_meur(value: Any) -> str:
    numeric = _number_from_display(value)
    if numeric is None:
        return str(value or "")
    return f"EUR {numeric / 1_000_000:.1f}M"


def _format_count(value: Any) -> str:
    numeric = _number_from_display(value)
    if numeric is None:
        return str(value or "")
    return str(int(round(numeric)))


def _kpi_value(trends: dict[str, Any], name: str) -> Any:
    for kpi in trends.get("kpis", []):
        if kpi.get("name") == name:
            return kpi.get("value")
    return None


def _add_kpi_tile(
    slide: Any,
    *,
    left: Emu,
    top: Emu,
    width: Emu,
    height: Emu,
    label: str,
    value: str,
    note: str = "",
    accent: RGBColor = COLOR_HEADER_FILL,
) -> None:
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    panel.line.color.rgb = RGBColor(0xE3, 0xE7, 0xEF)
    panel.line.width = Pt(0.45)
    accent_rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Emu(36000))
    accent_rule.fill.solid()
    accent_rule.fill.fore_color.rgb = accent
    accent_rule.line.fill.background()
    bottom_rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left,
        top + height - Emu(38000),
        width,
        Emu(9000),
    )
    bottom_rule.fill.solid()
    bottom_rule.fill.fore_color.rgb = RGBColor(0xE8, 0xEC, 0xF2)
    bottom_rule.line.fill.background()
    _add_textbox(
        slide,
        left + Emu(155000),
        top + Emu(120000),
        width - Emu(280000),
        Emu(170000),
        label,
        font_size=Pt(8.6),
        bold=True,
        color=COLOR_MUTED_TEXT,
    )
    _add_textbox(
        slide,
        left + Emu(155000),
        top + Emu(300000),
        width - Emu(280000),
        Emu(300000),
        value,
        font_size=Pt(18 if len(value) <= 18 else 15),
        bold=True,
        color=COLOR_BODY_TEXT,
    )
    if note:
        _add_textbox(
            slide,
            left + Emu(155000),
            top + Emu(625000),
            width - Emu(280000),
            height - Emu(680000),
            note,
            font_size=Pt(8.2),
            color=COLOR_MUTED_TEXT,
        )


def _add_bullet_block(
    slide: Any,
    *,
    left: Emu,
    top: Emu,
    width: Emu,
    height: Emu,
    title: str,
    bullets: list[str],
    font_size: Pt = Pt(12),
) -> None:
    _add_textbox(slide, left, top, width, Emu(240000), title, font_size=Pt(12), bold=True)
    _add_rule(slide, left, top + Emu(285000), width)
    body = slide.shapes.add_textbox(left, top + Emu(350000), width, height - Emu(350000))
    tf = _shape_text_frame(body)
    for idx, bullet in enumerate(bullets):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.level = 0
        _set_paragraph_text(
            paragraph,
            f"- {bullet}",
            font_size=font_size,
            color=COLOR_BODY_TEXT,
        )


def _filtered_body_rows(matrix: list[list[Any]]) -> list[list[Any]]:
    return [row for row in matrix[1:] if any(cell not in (None, "") for cell in row)]


def _publishable_record(record: dict[str, str]) -> bool:
    return not is_internal_sales_record(record)


def _readiness_records(legacy: Any, *, limit: int = 8) -> list[dict[str, str]]:
    last_row = _last_nonempty_row(
        legacy,
        "Q2_Readiness",
        start_row=1,
        columns=["A", "B", "C", "M", "N"],
    )
    matrix = _matrix_rows(legacy, "Q2_Readiness", f"A1:N{last_row}")
    headers = [str(cell or "") for cell in matrix[0]]
    records: list[dict[str, str]] = []
    for row in _filtered_body_rows(matrix):
        record = {
            header: _format_table_cell(header, value)
            for header, value in zip(headers, row)
        }
        if _publishable_record(record):
            records.append(record)
    return records[:limit]


def _forecast_category_records(legacy: Any) -> list[dict[str, str]]:
    matrix = _matrix_rows(legacy, "Forecast_Category", "A4:C9")
    headers = [str(cell or "") for cell in matrix[0]]
    return [
        {header: _format_table_cell(header, value) for header, value in zip(headers, row)}
        for row in _filtered_body_rows(matrix)
    ]


def _pending_approval_records(legacy: Any) -> list[dict[str, str]]:
    last_row = _last_nonempty_row(
        legacy,
        "Pending_Commercial_Approval",
        start_row=3,
        columns=["A", "B", "C", "D", "E", "F", "G", "H"],
    )
    matrix = _matrix_rows(legacy, "Pending_Commercial_Approval", f"A3:H{last_row}")
    headers = [str(cell or "") for cell in matrix[0]]
    records = [
        {header: _format_table_cell(header, value) for header, value in zip(headers, row)}
        for row in _filtered_body_rows(matrix)
    ]
    return [record for record in records if _publishable_record(record)]


def _sheet_records(workbook: Any, sheet_name: str, cell_range: str) -> list[dict[str, str]]:
    matrix = _matrix_rows(workbook, sheet_name, cell_range)
    headers = [str(cell or "") for cell in matrix[0]]
    records = [
        {header: _format_table_cell(header, value) for header, value in zip(headers, row)}
        for row in _filtered_body_rows(matrix)
    ]
    return [record for record in records if _publishable_record(record)]


def _add_horizontal_bar_panel(
    slide: Any,
    *,
    left: Emu,
    top: Emu,
    width: Emu,
    height: Emu,
    rows: list[tuple[str, float, str, str]],
    color: RGBColor = COLOR_HEADER_FILL,
) -> None:
    if not rows:
        _add_textbox(slide, left, top, width, height, "No rows available.", font_size=Pt(12))
        return
    label_w = Emu(int(int(width) * 0.32))
    value_w = Emu(int(int(width) * 0.16))
    bar_w = width - label_w - value_w - Emu(320000)
    row_h = Emu(int(int(height) / max(1, len(rows))))
    max_value = max(value for _, value, _, _ in rows) or 1.0
    for idx, (label, value, value_text, note) in enumerate(rows):
        y = top + Emu(idx * int(row_h))
        _add_textbox(
            slide,
            left,
            y + Emu(35000),
            label_w,
            row_h - Emu(45000),
            _clip_cell(label, 32),
            font_size=Pt(10.5),
            bold=idx == 0,
        )
        track = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left + label_w + Emu(80000),
            y + Emu(95000),
            bar_w,
            Emu(max(70000, int(int(row_h) * 0.38))),
        )
        track.fill.solid()
        track.fill.fore_color.rgb = RGBColor(0xEA, 0xEF, 0xF6)
        track.line.fill.background()
        fill_width = Emu(max(45000, int(int(bar_w) * value / max_value)))
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left + label_w + Emu(80000),
            y + Emu(95000),
            fill_width,
            Emu(max(70000, int(int(row_h) * 0.38))),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()
        _add_textbox(
            slide,
            left + label_w + bar_w + Emu(190000),
            y + Emu(32000),
            value_w,
            row_h - Emu(45000),
            value_text,
            font_size=Pt(10.5),
            bold=True,
            align=PP_ALIGN.RIGHT,
        )
        if note:
            _add_textbox(
                slide,
                left + label_w + Emu(90000),
                y + Emu(max(175000, int(int(row_h) * 0.52))),
                bar_w,
                row_h - Emu(max(175000, int(int(row_h) * 0.52))),
                note,
                font_size=Pt(8),
                color=COLOR_MUTED_TEXT,
            )


def _renewal_records(legacy: Any) -> list[dict[str, str]]:
    last_row = _last_nonempty_row(
        legacy,
        "FY26_Renewals",
        start_row=1,
        columns=["A", "B", "C", "D", "E", "F", "G", "H", "I"],
    )
    matrix = _matrix_rows(legacy, "FY26_Renewals", f"A1:I{last_row}")
    headers = [str(cell or "") for cell in matrix[0]]
    records = [
        {header: _format_table_cell(header, value) for header, value in zip(headers, row)}
        for row in _filtered_body_rows(matrix)
    ]
    return [record for record in records if _publishable_record(record)]


def _normalize_arr_basis_text(value: Any) -> str:
    text = str(value or "")
    text = re.sub(r"\bunwtd\b", "unweighted", text, flags=re.IGNORECASE)

    def replace_arr(match: re.Match[str]) -> str:
        prefix = text[max(0, match.start() - 18) : match.start()].lower()
        if prefix.endswith("weighted ") or prefix.endswith("unweighted "):
            return "ARR"
        return "unweighted ARR"

    return re.sub(r"\bARR\b", replace_arr, text)


def _action_item_claim(trends: dict[str, Any], rule_id: str) -> str:
    for item in trends.get("action_items", []):
        if item.get("rule_id") == rule_id:
            return _normalize_arr_basis_text(item.get("claim"))
    return ""


def _approval_gap_label(records: list[dict[str, str]]) -> str:
    labels: list[str] = []
    for row in records:
        text = f"{row.get('Account', '')} {row.get('Opportunity', '')}"
        if "Mandiri" in text:
            label = "Bank Mandiri"
        elif "Temasek" in text or str(row.get("Opportunity", "")).startswith("TEM"):
            label = "Temasek"
        else:
            label = row.get("Opportunity", "") or row.get("Account", "") or "unnamed deal"
        if label not in labels:
            labels.append(label)
    return " + ".join(labels) if labels else "No current Q2 approval gaps"


def _approval_gap_phrase(records: list[dict[str, str]]) -> str:
    count = len(records)
    if count == 0:
        return "no current-quarter approval gaps"
    if count == 1:
        return "one current-quarter approval gap"
    return f"{count} current-quarter approval gaps"


def _styled_table(
    slide: Any,
    *,
    left: Emu,
    top: Emu,
    width: Emu,
    height: Emu,
    headers: list[str],
    rows: list[list[str]],
    widths: tuple[float, ...],
    right_align_cols: tuple[int, ...] = (),
    body_font_size: Pt | None = None,
) -> None:
    n_rows = len(rows) + 1
    shape = slide.shapes.add_table(n_rows, len(headers), left, top, width, height)
    table = shape.table

    total = sum(widths)
    for idx, fraction in enumerate(widths):
        table.columns[idx].width = Emu(int(int(width) * fraction / total))

    row_height = Emu(int(int(height) / max(1, n_rows)))
    for row in table.rows:
        row.height = row_height

    body_font = Pt(10)
    if len(headers) >= 9:
        body_font = Pt(8.5)
    if len(rows) >= 10:
        body_font = Pt(8.5)
    if len(headers) >= 11:
        body_font = Pt(7.5)
    if len(rows) >= 14:
        body_font = Pt(7.5)
    if body_font_size is not None:
        body_font = body_font_size
    header_font = Pt(max(8, body_font.pt + 1))

    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER_FILL
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.margin_left = Emu(38100)
        tf.margin_right = Emu(38100)
        tf.margin_top = Emu(19050)
        tf.margin_bottom = Emu(19050)
        tf.clear()
        _set_paragraph_text(
            tf.paragraphs[0],
            header,
            font_size=header_font,
            bold=True,
            color=COLOR_HEADER_TEXT,
            align=PP_ALIGN.RIGHT if col_idx in right_align_cols else PP_ALIGN.LEFT,
        )

    for row_idx, row_values in enumerate(rows, start=1):
        for col_idx, value in enumerate(row_values):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_ROW_ALT if row_idx % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            tf = cell.text_frame
            tf.margin_left = Emu(38100)
            tf.margin_right = Emu(38100)
            tf.margin_top = Emu(19050)
            tf.margin_bottom = Emu(19050)
            tf.clear()
            _set_paragraph_text(
                tf.paragraphs[0],
                value,
                font_size=body_font,
                color=COLOR_BODY_TEXT,
                align=PP_ALIGN.RIGHT if col_idx in right_align_cols else PP_ALIGN.LEFT,
            )


def _trim_lines(lines: list[str]) -> list[str]:
    out = lines[:]
    while out and not out[0].strip():
        out.pop(0)
    while out and not out[-1].strip():
        out.pop()
    return out


def _strip_markdown_inline(text: str) -> str:
    return text.replace("**", "").replace("__", "").replace("`", "").replace("_", "").strip()


def _extract_headline_bullets(brief_sections: dict[str, list[str]]) -> list[str]:
    bullets: list[str] = []
    for line in brief_sections.get("Headline", []):
        stripped = line.strip()
        if stripped.startswith("- "):
            bullets.append(_strip_markdown_inline(stripped[2:].strip()))
    return bullets


def _extract_risk_claims(trends: dict[str, Any]) -> list[str]:
    claims: list[str] = []
    for risk in trends.get("risks", []):
        claim = str(risk.get("claim") or "").strip()
        if claim:
            claims.append(claim)
    return claims


def _extract_action_item_claims(trends: dict[str, Any], *, limit: int | None = None) -> list[str]:
    claims: list[str] = []
    for item in trends.get("action_items", []):
        claim = _normalize_arr_basis_text(item.get("claim")).strip()
        if claim:
            claims.append(claim)
    return claims if limit is None else claims[:limit]


def _exec_summary_lists(
    trends: dict[str, Any],
    brief_sections: dict[str, list[str]],
) -> tuple[list[str], list[str]]:
    left_items = _extract_headline_bullets(brief_sections)
    if not left_items:
        left_items = _extract_action_item_claims(trends, limit=3)

    right_items = _extract_risk_claims(trends)
    if not right_items:
        right_items = _extract_action_item_claims(trends, limit=3)

    return left_items or ["No highlights generated."], right_items or ["No risks generated."]


def _risks_outlook_lines(trends: dict[str, Any], brief_sections: dict[str, list[str]]) -> list[str]:
    risks = _extract_risk_claims(trends)
    if not risks:
        risks = _extract_action_item_claims(trends, limit=4)
    headlines = _extract_headline_bullets(brief_sections)
    if headlines:
        risks.append(f"Outlook: {headlines[0]}")
    return risks or ["No risks or outlook generated."]


def _fill_cover_slide(slide: Any, *, artifacts: Any) -> None:
    _set_shape_text(
        _find_shape(slide, "Text Placeholder 1"),
        f"{DECK_MONTH_LABEL} LAND territory review · {artifacts.scope_label}",
    )
    _set_shape_text(_find_shape(slide, "Text Placeholder 3"), artifacts.name)
    _set_shape_text(
        _find_shape(slide, "Text Placeholder 2"),
        f"{artifacts.period} operating pack · {SNAPSHOT_LABEL} · {DECK_DATE_LABEL}",
    )


def _fill_exec_summary(
    slide: Any,
    *,
    artifacts: Any,
    trends: dict[str, Any],
    model: Any,
    legacy: Any,
    is_apac: bool,
) -> None:
    _set_slide_header(
        slide,
        title="May operating summary",
        subtitle=f"{artifacts.scope_label} review · Salesforce snapshot 2026-04-30 · May 1 kickoff.",
    )
    _set_shape_text(_find_shape(slide, "Text Placeholder 2"), "")
    _clear_exec_placeholder(slide)
    _remove_shape(_find_shape(slide, "Content Placeholder 1"))
    _remove_shape(_find_shape(slide, "Content Placeholder 3"))

    total_arr = _kpi_value(trends, "total_pipeline_arr")
    renewal_acv = _kpi_value(trends, "total_renewal_acv")
    beyond_arr = _kpi_value(trends, "pipeline_arr_beyond_cfq")
    forecast = {row.get("Category", ""): row for row in _forecast_category_records(legacy)}
    commit = forecast.get("Commit", {})
    pipeline = forecast.get("Pipeline", {})
    approval_records = _pending_approval_records(legacy)
    approval_label = _approval_gap_label(approval_records)
    approval_phrase = _approval_gap_phrase(approval_records)
    wins_losses = _matrix_rows(legacy, "Wins_Losses_QTD", "A1:D3")
    won_arr = wins_losses[1][2] if len(wins_losses) > 1 else ""
    lost_arr = wins_losses[2][2] if len(wins_losses) > 2 else ""

    left = _inches(0.72)
    top = _inches(1.28)
    tile_w = _inches(2.95)
    gap = _inches(0.18)
    tile_h = _inches(0.98)
    _add_textbox(
        slide,
        _inches(2.18),
        _inches(6.92),
        _inches(9.3),
        _inches(0.24),
        "Metric basis: Land+Expand ARR is unweighted unless explicitly labeled weighted; Renewal values are ACV.",
        font_size=Pt(8.5),
        color=COLOR_MUTED_TEXT,
    )
    _add_kpi_tile(
        slide,
        left=left,
        top=top,
        width=tile_w,
        height=tile_h,
        label="Q2 closeable unweighted ARR",
        value=_format_meur(total_arr),
        note=("Land+Expand only; 26 opportunities." if is_apac else "Land+Expand current-quarter workbook."),
    )
    _add_kpi_tile(
        slide,
        left=left + tile_w + gap,
        top=top,
        width=tile_w,
        height=tile_h,
        label="Q2 renewal ACV",
        value=_format_meur(renewal_acv),
        note="Separate renewal motion.",
        accent=RGBColor(0x3B, 0x7F, 0x8C),
    )
    _add_kpi_tile(
        slide,
        left=left + (tile_w + gap) * 2,
        top=top,
        width=tile_w,
        height=tile_h,
        label="Beyond-Q2 unweighted ARR",
        value=_format_meur(beyond_arr),
        note="Context only; not Q2 coverage.",
        accent=RGBColor(0x7B, 0x6B, 0xA8),
    )
    _add_kpi_tile(
        slide,
        left=left + (tile_w + gap) * 3,
        top=top,
        width=tile_w,
        height=tile_h,
        label="QTD closed unweighted ARR",
        value=(
            f"Won {_format_meur(won_arr).replace('EUR ', '')} / "
            f"Lost {_format_meur(lost_arr).replace('EUR ', '')}"
        ),
        note="Land+Expand closed-won vs lost.",
        accent=COLOR_ACCENT,
    )

    _add_bullet_block(
        slide,
        left=left,
        top=_inches(2.65),
        width=_inches(5.95),
        height=_inches(3.45),
        title="What matters now",
        bullets=(
            [
                f"Forecast quality is narrow: Commit is {_format_meur(commit.get('ARR (EUR)'))} unweighted ARR across {commit.get('# Opps', '')} opps; Pipeline is {_format_meur(pipeline.get('ARR (EUR)'))} unweighted ARR across {pipeline.get('# Opps', '')} opps.",
                "The close call depends on named deal inspection: Danantara, LTH, Krungthai, Bank Mandiri, Temasek, and HKMA carry the story.",
                "Keep the original APAC target spine in the discussion: Q1 delivery, Q1 losses, Q1 slips, owner push coaching, commercial approvals, and Q2 readiness.",
            ]
            if is_apac
            else [
                f"Forecast quality is narrow: Commit is {_format_meur(commit.get('ARR (EUR)'))} unweighted ARR across {commit.get('# Opps', '')} opps; Pipeline is {_format_meur(pipeline.get('ARR (EUR)'))} unweighted ARR across {pipeline.get('# Opps', '')} opps.",
                "Use the current workbook rows as the source of truth for named deal inspection, owner coaching, approval gaps, and Q2 readiness.",
                "Keep Land+Expand ARR separate from Renewal ACV in every headline, chart, and operating action.",
            ]
        ),
        font_size=Pt(11.5),
    )
    _add_bullet_block(
        slide,
        left=_inches(7.05),
        top=_inches(2.65),
        width=_inches(5.55),
        height=_inches(3.45),
        title="May operating actions",
        bullets=(
            [
                "Refresh next-step evidence for Danantara, Krungthai, Temasek, and HKMA; flag overdue-close exposure and August timing risk where next steps miss Q2.",
                f"Submit {approval_phrase}: {approval_label}. Preserve original approved-2026 context: Coolabah and Danantara were already commercially approved.",
                "Validate the FY26 Renewal ACV basis before it is quoted outside this current-state review.",
            ]
            if is_apac
            else [
                "Refresh next-step evidence for the highest-value Q2 rows and require dated customer proof before accepting forecast confidence.",
                f"Submit {approval_phrase}: {approval_label}. Keep approval claims tied to supported current-quarter workbook rows.",
                "Validate the FY26 Renewal ACV basis before it is quoted outside this current-state review.",
            ]
        ),
        font_size=Pt(11.5),
    )


def _fill_review_delta_slide(
    slide: Any,
    *,
    artifacts: Any,
    trends: dict[str, Any],
    legacy: Any,
    is_apac: bool,
) -> None:
    _set_slide_header(
        slide,
        title=("Original APAC targets vs current state" if is_apac else "Original ETL targets vs current state"),
        subtitle=(
            "Preserve the Q1 accountability spine: Q1 opened EUR 17.1M, Q1 lost EUR 5.5M, Q1 slips EUR 9.4M."
            if is_apac
            else "Use this director's original ETL sidecar as the prior baseline and the current workbook for May values."
        ),
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    wins_losses = _matrix_rows(legacy, "Wins_Losses_QTD", "A1:D3")
    won_arr = wins_losses[1][2] if len(wins_losses) > 1 else ""
    lost_arr = wins_losses[2][2] if len(wins_losses) > 2 else ""

    if is_apac:
        original_rows = JESPER_ORIGINAL_Q1_TARGET_ROWS
        slip_rows = JESPER_ORIGINAL_Q1_SLIP_ROWS
    else:
        original_rows = [
            ["Original ETL sidecar", "Attached", "Use this director's own prior baseline; do not borrow another territory's pack."],
        ]
        slip_rows = []

    current_bullets = (
        [
            f"Current state: {_format_meur(_kpi_value(trends, 'total_pipeline_arr'))} Q2 closeable L+E unweighted ARR; {_format_meur(_kpi_value(trends, 'total_renewal_acv'))} Q2 Renewal ACV.",
            f"QTD Land+Expand unweighted ARR is already loss-heavy: {_format_meur(won_arr)} won versus {_format_meur(lost_arr)} lost.",
            "Original since-last-review delta: open Land deals moved 6 -> 12, open Land unweighted ARR rose by EUR 4.6M, and approved-2026 deals moved 1 -> 2.",
            "Original Q2 activity baseline: 6 deals / EUR 5.0M and zero recent activity; use this as activity-recovery coaching context.",
            "May review question: which named Q2 deals have customer proof, approval status, and dated next steps for the commit call?",
        ]
        if is_apac
        else [
            f"Current state: {_format_meur(_kpi_value(trends, 'total_pipeline_arr'))} Q2 closeable L+E unweighted ARR; {_format_meur(_kpi_value(trends, 'total_renewal_acv'))} Q2 Renewal ACV.",
            f"QTD Land+Expand unweighted ARR: {_format_meur(won_arr)} won versus {_format_meur(lost_arr)} lost.",
            "Use the original ETL sidecar as the accountability baseline, with all live May actions tied to current workbook rows.",
            "May review question: which named Q2 deals have customer proof, approval status, and dated next steps for the commit call?",
        ]
    )

    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=Emu(1750000),
        headers=["Original target", "Value", "How to use it"],
        rows=original_rows,
        widths=(0.20, 0.24, 0.56),
    )
    _add_bullet_block(
        slide,
        left=left + Emu(63500),
        top=top + Emu(2000000),
        width=Emu(int(int(width) * 0.43)),
        height=height - Emu(2070000),
        title="Current state",
        bullets=current_bullets,
        font_size=Pt(10.2),
    )


def _fill_forecast_quality_slide(slide: Any, *, trends: dict[str, Any], legacy: Any, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="May forecast quality",
        subtitle="May 2026 view · Salesforce snapshot as of 2026-04-30. Omitted and out-of-quarter pipe are context, not coverage.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    forecast = _forecast_category_records(legacy)
    by_category = {row.get("Category", ""): row for row in forecast}
    commit = by_category.get("Commit", {})
    pipeline = by_category.get("Pipeline", {})
    best_case = by_category.get("Best Case", {})

    _add_kpi_tile(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=Emu(2820000),
        height=Emu(930000),
        label="Closeable L+E unweighted ARR",
        value=_format_meur(_kpi_value(trends, "total_pipeline_arr")),
        note="This is the CFQ numerator.",
    )
    _add_kpi_tile(
        slide,
        left=left + Emu(3100000),
        top=top + Emu(63500),
        width=Emu(2820000),
        height=Emu(930000),
        label="Commit unweighted ARR",
        value=_format_meur(commit.get("ARR (EUR)")),
        note=f"{commit.get('# Opps', '')} opps; unweighted ARR.",
        accent=RGBColor(0x1F, 0x7A, 0x4D),
    )
    _add_kpi_tile(
        slide,
        left=left + Emu(6130000),
        top=top + Emu(63500),
        width=Emu(2820000),
        height=Emu(930000),
        label="Pipeline unweighted ARR",
        value=_format_meur(pipeline.get("ARR (EUR)")),
        note=f"{pipeline.get('# Opps', '')} opps; unweighted ARR.",
        accent=RGBColor(0x7B, 0x6B, 0xA8),
    )
    _add_kpi_tile(
        slide,
        left=left + Emu(9160000),
        top=top + Emu(63500),
        width=Emu(2820000),
        height=Emu(930000),
        label="Best Case unweighted ARR",
        value=_format_meur(best_case.get("ARR (EUR)")),
        note="No material Best Case cushion.",
        accent=COLOR_ACCENT,
    )

    table_rows = [
        [
            row.get("Category", ""),
            row.get("# Opps", ""),
            _format_meur(row.get("ARR (EUR)")),
        ]
        for row in forecast
        if row.get("Category") not in ("Closed",)
    ]
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(1250000),
        width=Emu(5200000),
        height=Emu(2450000),
        headers=["Category", "#", "Unweighted ARR"],
        rows=table_rows,
        widths=(0.45, 0.18, 0.37),
        right_align_cols=(1, 2),
    )
    _add_bullet_block(
        slide,
        left=left + Emu(5650000),
        top=top + Emu(1280000),
        width=width - Emu(5850000),
        height=Emu(2400000),
        title="Review focus",
        bullets=[
            "Commit is not enough by itself; the rest of the quarter depends on Stage 3 Pipeline conversion.",
            "Best Case is effectively absent, so upside needs explicit deal movement rather than forecast optimism.",
            (
                JESPER_ORIGINAL_FORECAST_MIX_NOTE
                if is_apac
                else "Compare category movement only against this director's workbook-supported current and prior extracts."
            ),
            "Current Q2 category values in this deck are unweighted ARR unless explicitly labeled weighted.",
            "Use the readiness page as the operating source of truth for the Q2 call.",
        ],
        font_size=Pt(11),
    )


def _fill_pipeline_aging_slide(slide: Any, legacy: Any) -> None:
    _set_slide_header(
        slide,
        title="Q2 activity hygiene",
        subtitle="Close inspection uses LastActivityDate, next-step quality, push count, and approval status.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _readiness_records(legacy, limit=10)
    signals = [
        ("Silent 90d+", lambda r: "Silent 90d+" in r.get("Readiness", "")),
        ("No next step", lambda r: "No next step" in r.get("Readiness", "")),
        ("4+ pushes", lambda r: "pushes" in r.get("Readiness", "")),
        ("Approval gap", lambda r: "Commercial approval gap" in r.get("Readiness", "")),
    ]
    signal_rows = []
    for label, predicate in signals:
        matched = [row for row in records if predicate(row)]
        signal_rows.append(
            [
                label,
                str(len(matched)),
                _format_meur(sum(_number_from_display(row.get("ARR")) or 0 for row in matched)),
                _clip_cell(", ".join(row.get("Opportunity", "") for row in matched[:4]), 82),
            ]
        )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=Emu(1850000),
        headers=["Signal", "# deals", "Unweighted ARR", "Named deals"],
        rows=signal_rows,
        widths=(0.18, 0.10, 0.14, 0.58),
        right_align_cols=(1, 2),
    )
    watch_rows = [
        [
            _clip_cell(row.get("Opportunity", ""), 42),
            _clip_cell(row.get("Owner", ""), 22),
            row.get("ARR", ""),
            row.get("Forecast", ""),
            row.get("Last Activity", "") or "None",
            _clip_cell(row.get("Readiness", ""), 50),
            _clip_cell(row.get("Next Step", "") or "No next step", 58),
        ]
        for row in records[:5]
    ]
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(2180000),
        width=width - Emu(127000),
        height=height - Emu(2250000),
        headers=["Deal", "Owner", "Unweighted ARR", "Fcst", "Last activity", "Signal", "Next step"],
        rows=watch_rows,
        widths=(0.22, 0.10, 0.08, 0.08, 0.12, 0.18, 0.22),
        right_align_cols=(2,),
    )


def _fill_forecast_category_detail_slide(slide: Any, legacy: Any) -> None:
    _set_slide_header(
        slide,
        title="Forecast category detail",
        subtitle="Native detail table: Q2 Land+Expand unweighted ARR and opportunity count by SF forecast category.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _forecast_category_records(legacy)
    rows_for_bars = []
    table_rows = []
    for row in records:
        category = row.get("Category", "")
        if category == "TOTAL":
            continue
        arr = _number_from_display(row.get("ARR (EUR)")) or 0
        rows_for_bars.append((category, arr, _format_meur(arr), f"{row.get('# Opps', '')} opps"))
        table_rows.append([category, row.get("# Opps", ""), _format_meur(arr)])
    _add_horizontal_bar_panel(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=Emu(int(int(width) * 0.55)),
        height=Emu(2600000),
        rows=rows_for_bars,
        color=RGBColor(0x7B, 0x6B, 0xA8),
    )
    _styled_table(
        slide,
        left=left + Emu(int(int(width) * 0.62)),
        top=top + Emu(63500),
        width=Emu(int(int(width) * 0.34)),
        height=Emu(2200000),
        headers=["Category", "#", "Unweighted ARR"],
        rows=table_rows,
        widths=(0.46, 0.18, 0.36),
        right_align_cols=(1, 2),
    )
    _add_bullet_block(
        slide,
        left=left + Emu(63500),
        top=top + Emu(2920000),
        width=width - Emu(127000),
        height=height - Emu(3000000),
        title="Review focus",
        bullets=[
            "Pipeline plus Commit carry essentially the entire Q2 unweighted ARR story; Best Case is not a meaningful cushion.",
            "Omitted remains visible for governance but is excluded from active coverage claims.",
        ],
        font_size=Pt(12),
    )


def _fill_by_owner_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="Owner pipeline coverage",
        subtitle="Owner-level Land+Expand unweighted ARR, filtered to nonzero value for readability.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _sheet_records(legacy, "By_Owner", "A1:C11")
    rows = []
    for row in records:
        arr = _number_from_display(row.get("Open ARR")) or 0
        if arr <= 0:
            continue
        rows.append(
            (
                row.get("Owner", ""),
                arr,
                _format_meur(arr),
                f"{row.get('# Open opps', '')} opps",
            )
        )
    _add_horizontal_bar_panel(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=Emu(int(int(width) * 0.58)),
        height=height - Emu(127000),
        rows=rows,
    )
    _add_bullet_block(
        slide,
        left=left + Emu(int(int(width) * 0.64)),
        top=top + Emu(95000),
        width=Emu(int(int(width) * 0.32)),
        height=height - Emu(190000),
        title="Coaching order",
        bullets=(
            [
                "Edwina carries the largest visible Q2 unweighted ARR and owns several stale/pushed deals.",
                "Jesper owns Danantara, the largest Q2 Pipeline deal and a no-next-step risk.",
                "Hahn/Ronald late-stage entries need activity and next-step cleanup before forecast review.",
            ]
            if is_apac
            else [
                "Rank owner coaching by visible Q2 unweighted ARR, stale activity, push count, and approval status.",
                "Use the highest-value owner rows as the first operating review, then confirm next steps and close proof.",
            ]
        ),
        font_size=Pt(11.5),
    )


def _fill_territory_pipeline_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="Territory pipeline mix",
        subtitle="May Land+Expand unweighted ARR by country/sub-region, filtered to nonzero value.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _sheet_records(legacy, "Territory_Performance", "A1:C10")
    rows = []
    for row in records:
        arr = _number_from_display(row.get("Open ARR")) or 0
        if arr <= 0:
            continue
        rows.append(
            (
                row.get("Country / sub-region", ""),
                arr,
                _format_meur(arr),
                f"{row.get('# Opps', '')} opps",
            )
        )
    _add_horizontal_bar_panel(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=Emu(int(int(width) * 0.58)),
        height=height - Emu(127000),
        rows=rows,
        color=RGBColor(0x3B, 0x7F, 0x8C),
    )
    _add_bullet_block(
        slide,
        left=left + Emu(int(int(width) * 0.64)),
        top=top + Emu(95000),
        width=Emu(int(int(width) * 0.32)),
        height=height - Emu(190000),
        title="Review focus",
        bullets=(
            [
                "Indonesia, Malaysia, and Thailand carry the meaningful Q2 value.",
                "Use geography only as context; the decision pages should still anchor on named deals and owners.",
            ]
            if is_apac
            else [
                "Use geography only as context; the decision pages should still anchor on named deals and owners.",
                "Escalate any territory concentration only when the current workbook shows material Q2 value.",
            ]
        ),
        font_size=Pt(11.5),
    )


def _last_nonempty_row(
    workbook: Any,
    sheet_name: str,
    *,
    start_row: int,
    columns: list[str],
    max_row: int | None = None,
) -> int:
    ws = workbook.workbook[sheet_name]
    limit = max_row or ws.max_row
    last = start_row
    for row_idx in range(start_row, limit + 1):
        if any(workbook.cell_value(sheet_name, f"{column}{row_idx}") not in (None, "") for column in columns):
            last = row_idx
    return last


def _matrix_rows(workbook: Any, sheet_name: str, cell_range: str) -> list[list[Any]]:
    return workbook.matrix(sheet_name, cell_range)


def _formatted_rows(matrix: list[list[Any]]) -> list[list[str]]:
    headers = [str(cell or "") for cell in matrix[0]]
    rows: list[list[str]] = []
    for row in matrix[1:]:
        rows.append([_format_table_cell(header, value) for header, value in zip(headers, row)])
    return rows


def _project_matrix(matrix: list[list[Any]], indices: tuple[int, ...]) -> list[list[Any]]:
    return [[row[idx] if idx < len(row) else "" for idx in indices] for row in matrix]


def _clip_cell(value: Any, limit: int = 58) -> str:
    text = str(value or "")
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "..."


def _named_deals_widths() -> tuple[float, ...]:
    return (0.05, 0.18, 0.30, 0.12, 0.10, 0.11, 0.06, 0.08)


def _fill_legacy_table_slide(
    slide: Any,
    legacy: Any,
    sheet_name: str,
    cell_range: str,
    widths: tuple[float, ...],
) -> None:
    left, top, width, height = _clear_standard_placeholder(slide)
    matrix = _matrix_rows(legacy, sheet_name, cell_range)
    headers = [str(cell or "") for cell in matrix[0]]
    rows = _formatted_rows(matrix)
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=height - Emu(127000),
        headers=headers,
        rows=[row for row in rows if any(cell for cell in row)],
        widths=widths,
        right_align_cols=(0, 6, 7),
    )


def _fill_q2_readiness_slide(slide: Any, legacy: Any) -> None:
    subtitle = "May 2026 operating view: snapshot as of 2026-04-30 for the May 1 kickoff."
    _set_slide_header(slide, title="May deal readiness", subtitle="")
    _set_shape_text(_find_shape(slide, "Text Placeholder 4"), subtitle)
    _remove_shape(_find_shape(slide, "Text Placeholder 2"))
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _readiness_records(legacy, limit=8)
    rows = [
        [
            row.get("#", ""),
            _clip_cell(f"{row.get('Account', '')} / {row.get('Opportunity', '')}", 64),
            _clip_cell(row.get("Owner", ""), 24),
            _clip_cell(row.get("Stage", ""), 20),
            row.get("Close Date", ""),
            row.get("ARR", ""),
            f"{row.get('Forecast', '')} / {row.get('Prob', '')}",
            _clip_cell(
                (
                    (row.get("Readiness", "") + "; overdue-close").strip("; ")
                    if str(row.get("Close Date", "")) and str(row.get("Close Date", "")) < "2026-05-01"
                    else row.get("Readiness", "")
                ),
                48,
            ),
            _clip_cell(row.get("Next Step", "") or "No next step", 58),
        ]
        for row in records
    ]
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=height - Emu(127000),
        headers=["#", "Account / opportunity", "Owner", "Stage", "Close", "Unweighted ARR", "Fcst / prob", "Flags", "Next step"],
        rows=rows,
        widths=(0.04, 0.22, 0.10, 0.09, 0.09, 0.08, 0.10, 0.14, 0.14),
        right_align_cols=(0, 5),
    )


def _fill_close_plan_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="May close plan by deal",
        subtitle=(
            "May operating view for top APAC deals: Q2 close evidence, timing resets, and owner actions."
            if is_apac
            else "May operating view for top Q2 deals: close evidence, timing resets, and owner actions."
        ),
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _readiness_records(legacy, limit=6)
    rows: list[list[str]] = []
    for row in records:
        flags = row.get("Readiness", "")
        next_step = row.get("Next Step", "")
        close_date = str(row.get("Close Date", ""))
        if close_date and close_date < "2026-05-01":
            action = "Overdue-close: move close date or provide signed close proof."
        elif "August" in next_step:
            action = "Q2 timing check: August decision timing needs reset or customer proof."
        elif "Commercial approval gap" in flags:
            action = "Submit Commercial Approval and confirm decision forum."
        elif "No next step" in flags:
            action = "Add customer-owned next step with date and sponsor."
        elif "pushes" in flags:
            action = "Run close-date reset; anchor Q2 timing in customer proof."
        else:
            action = "Confirm close plan evidence and next milestone."
        rows.append(
            [
                _clip_cell(row.get("Opportunity", ""), 42),
                _clip_cell(row.get("Owner", ""), 20),
                close_date,
                row.get("ARR", ""),
                _clip_cell(flags, 54),
                _clip_cell(next_step or "No next step recorded", 62),
                action,
            ]
        )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=height - Emu(127000),
        headers=["Deal", "Owner", "Close", "Unweighted ARR", "Signal", "Current next step", "Director action"],
        rows=rows,
        widths=(0.17, 0.09, 0.09, 0.07, 0.17, 0.22, 0.19),
        right_align_cols=(3,),
    )


def _fill_pending_approval_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    records = _pending_approval_records(legacy)
    gap_label = _approval_gap_label(records)
    gap_phrase = _approval_gap_phrase(records)
    total = sum(_number_from_display(row.get("ARR (EUR)")) or 0 for row in records)
    _set_slide_header(
        slide,
        title="Commercial approval gaps",
        subtitle=(
            (
                f"Current-quarter row-level support: {gap_phrase} missing the gate. "
                "Original APAC approved-2026 deals stay visible."
            )
            if is_apac
            else f"Current-quarter row-level support: {gap_phrase} missing the gate."
        ),
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    _add_kpi_tile(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=Emu(3600000),
        height=Emu(930000),
        label="Q2 approval gaps",
        value=f"{len(records)} deals / {_format_meur(total)}",
        note="Supported current-quarter unweighted ARR.",
        accent=COLOR_ACCENT,
    )
    _add_bullet_block(
        slide,
        left=left + Emu(3850000),
        top=top + Emu(63500),
        width=width - Emu(4050000),
        height=Emu(780000),
        title="Scope note",
        bullets=(
            [
                "Quote Q2 supported gaps separately from original approved deals and all-open action claims.",
                "Original APAC Approved 2026 context: Coolabah and Danantara / EUR 3.8M; all-open approval scope was 7 gaps / EUR 10.7M; Amova remained pending/candidate.",
            ]
            if is_apac
            else [
                "Quote Q2 supported gaps only from current-quarter workbook rows.",
                "Do not infer prior approved deals from another territory pack.",
            ]
        ),
        font_size=Pt(9.5),
    )
    scope_rows = (
        [
            ["Q2 supported gaps", gap_label, f"{len(records)} / {_format_meur(total)}", "Current-state claim; unweighted ARR."],
            ["Original approved 2026", "Coolabah + Danantara", "2 / EUR 3.8M", "Preserve as prior APAC governance spine."],
            ["All-open action claim", "7 gaps / EUR 10.7M", "Needs appendix", "Keep internal until row-level support is attached."],
        ]
        if is_apac
        else [
            ["Q2 supported gaps", gap_label, f"{len(records)} / {_format_meur(total)}", "Current-state claim; unweighted ARR."],
            ["Prior approved deals", "Not attached", "Do not infer", "Use only this director's workbook or source pack."],
            ["All-open action claim", "Needs appendix", "Needs support", "Keep internal until row-level support is attached."],
        ]
    )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(1180000),
        width=width - Emu(127000),
        height=Emu(950000),
        headers=["Scope", "Rows", "Amount", "Data guardrail"],
        rows=scope_rows,
        widths=(0.20, 0.34, 0.16, 0.30),
        right_align_cols=(2,),
    )
    gap_rows = [
        [
            row.get("#", ""),
            _clip_cell(row.get("Account", ""), 34),
            _clip_cell(row.get("Opportunity", ""), 42),
            _clip_cell(row.get("Owner", ""), 22),
            _clip_cell(row.get("Stage", ""), 20),
            row.get("Close Date", ""),
            row.get("Type", ""),
            row.get("ARR (EUR)", ""),
        ]
        for row in records
    ]
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(2300000),
        width=Emu(int(int(width) * 0.47)),
        height=height - Emu(2380000),
        headers=["#", "Current Q2 gap", "Owner", "Close", "Unweighted ARR"],
        rows=[
            [
                row[0],
                _clip_cell(f"{row[1]} / {row[2]}", 58),
                row[3],
                row[5],
                row[7],
            ]
            for row in gap_rows
        ],
        widths=(0.07, 0.47, 0.17, 0.15, 0.14),
        right_align_cols=(0, 4),
        body_font_size=Pt(8.2),
    )
    if is_apac:
        _styled_table(
            slide,
            left=left + Emu(int(int(width) * 0.52)),
            top=top + Emu(2300000),
            width=Emu(int(int(width) * 0.44)),
            height=height - Emu(2380000),
            headers=["Original status", "Deal", "Value", "Date", "Owner / use"],
            rows=JESPER_ORIGINAL_COMMERCIAL_APPROVAL_ROWS,
            widths=(0.20, 0.28, 0.16, 0.13, 0.23),
            right_align_cols=(2,),
            body_font_size=Pt(8.0),
        )


def _fill_fy26_renewals_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="FY26 Renewal watchlist",
        subtitle=(
            "Operational watchlist plus original APAC baseline. ACV only; Land/Expand unweighted ARR is intentionally excluded."
            if is_apac
            else "Operational watchlist. ACV only; Land/Expand unweighted ARR is intentionally excluded."
        ),
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _renewal_records(legacy)
    total = sum(_number_from_display(row.get("ACV")) or 0 for row in records)
    _add_kpi_tile(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=Emu(3000000),
        height=Emu(930000),
        label=("Original APAC baseline" if is_apac else "Prior baseline"),
        value=("EUR 33.5M" if is_apac else "Not attached"),
        note=("3 FY26 renewals from Apr 20 ETL." if is_apac else "Use workbook rows until a prior source pack is attached."),
        accent=RGBColor(0x3B, 0x7F, 0x8C),
    )
    _add_kpi_tile(
        slide,
        left=left + Emu(3220000),
        top=top + Emu(63500),
        width=Emu(3000000),
        height=Emu(930000),
        label="Current state workbook",
        value=_format_meur(total),
        note=f"{len(records)} open renewal rows; basis differs.",
        accent=RGBColor(0x3B, 0x7F, 0x8C),
    )
    _add_bullet_block(
        slide,
        left=left + Emu(6300000),
        top=top + Emu(63500),
        width=width - Emu(6500000),
        height=Emu(930000),
        title="Renewal working view",
        bullets=[
            "Use the rows below to manage renewal owners and dates.",
            (
                "Hold external FY26 renewal ACV quotes until the Apr 20 baseline and refreshed workbook basis reconcile."
                if is_apac
                else "Hold external FY26 renewal ACV quotes until the refreshed workbook basis is validated."
            ),
        ],
        font_size=Pt(9.5),
    )
    filtered_rows = [
        [
            row.get("#", ""),
            row.get("Close Date", ""),
            _clip_cell(row.get("Account", ""), 34),
            _clip_cell(row.get("Opportunity", ""), 42),
            _clip_cell(row.get("Owner", ""), 24),
            _clip_cell(row.get("Stage", ""), 20),
            row.get("ACV", ""),
            row.get("Prob", ""),
            row.get("Risk", ""),
        ]
        for row in records
    ]
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(1280000),
        width=width - Emu(127000),
        height=height - Emu(1380000),
        headers=["#", "Close", "Account", "Opportunity", "Owner", "Stage", "ACV", "Prob", "Risk"],
        rows=filtered_rows,
        widths=(0.04, 0.10, 0.17, 0.23, 0.12, 0.12, 0.09, 0.06, 0.07),
        right_align_cols=(0, 6, 7),
    )


def _fill_renewal_guardrail_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="Renewal ACV reconciliation",
        subtitle=(
            "Original APAC renewal intelligence is preserved; current refreshed values require basis reconciliation."
            if is_apac
            else "Current refreshed renewal values require basis reconciliation before external quote."
        ),
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _renewal_records(legacy)
    total = sum(_number_from_display(row.get("ACV")) or 0 for row in records)
    q2 = sum(
        _number_from_display(row.get("ACV")) or 0
        for row in records
        if str(row.get("Close Date", "")) < "2026-07-01"
    )
    wins_losses = _matrix_rows(legacy, "Wins_Losses_QTD", "A1:D3")
    closed_renewal_acv = wins_losses[1][3] if len(wins_losses) > 1 and len(wins_losses[1]) > 3 else ""
    _add_kpi_tile(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=Emu(3000000),
        height=Emu(980000),
        label=("Original FY26 baseline" if is_apac else "Prior FY26 baseline"),
        value=("EUR 33.5M" if is_apac else "Not attached"),
        note=("3 open renewals; Apr 20 APAC ETL." if is_apac else "Use this director's source pack only."),
        accent=RGBColor(0x3B, 0x7F, 0x8C),
    )
    _add_kpi_tile(
        slide,
        left=left + Emu(3220000),
        top=top + Emu(63500),
        width=Emu(3000000),
        height=Emu(980000),
        label="Current FY26 watchlist",
        value=_format_meur(total),
        note="Refreshed workbook display basis.",
        accent=RGBColor(0x3B, 0x7F, 0x8C),
    )
    _add_kpi_tile(
        slide,
        left=left + Emu(6300000),
        top=top + Emu(63500),
        width=Emu(3000000),
        height=Emu(980000),
        label=("Original Q2 subset" if is_apac else "Prior Q2 subset"),
        value=("EUR 0.5M" if is_apac else "Not attached"),
        note=("1 renewal from original APAC ETL." if is_apac else "No prior subset applied."),
        accent=RGBColor(0x7B, 0x6B, 0xA8),
    )
    _add_kpi_tile(
        slide,
        left=left + Emu(9380000),
        top=top + Emu(63500),
        width=Emu(2600000),
        height=Emu(980000),
        label="QTD closed renewal",
        value=_format_meur(closed_renewal_acv),
        note="Closed renewal motion only.",
        accent=COLOR_ACCENT,
    )
    rows = (
        [
            ["Original APAC ETL", "FY26 open renewals", "3", "EUR 33.5M ACV", "Preserve as prior-deck baseline."],
            ["Original APAC ETL", "Q2 subset", "1", "EUR 0.455M ACV", "Fullerton renewal baseline."],
            ["Refreshed workbook", "FY26 watchlist", str(len(records)), _format_meur(total), "Operational current-state rows."],
            ["Refreshed workbook", "Q2 subset", "1", _format_meur(q2), "Hold the prior baseline until ACV-field reconciliation is complete."],
        ]
        if is_apac
        else [
            ["Prior source pack", "FY26 open renewals", "Not attached", "Not attached", "Do not infer prior renewal baseline."],
            ["Refreshed workbook", "FY26 watchlist", str(len(records)), _format_meur(total), "Operational current-state rows."],
            ["Refreshed workbook", "Q2 subset", "1", _format_meur(q2), "Validate ACV-field basis before external quote."],
        ]
    )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(1260000),
        width=width - Emu(127000),
        height=Emu(1620000),
        headers=["Source", "Scope", "#", "Amount", "Use in review"],
        rows=rows,
        widths=(0.18, 0.19, 0.06, 0.17, 0.40),
        right_align_cols=(2, 3),
    )
    _add_bullet_block(
        slide,
        left=left + Emu(63500),
        top=top + Emu(3050000),
        width=width - Emu(127000),
        height=height - Emu(3120000),
        title="Data guardrail",
        bullets=[
            "Renewal pages quote ACV only. Keep Renewal ACV separate from Land+Expand unweighted ARR in headlines, charts, and bridges.",
            (
                "The original EUR 33.5M baseline and the refreshed EUR 5.0M workbook basis conflict; keep both visible until Sales Ops reconciles field/extract basis."
                if is_apac
                else "Use only this director's current workbook basis unless a matching prior source pack is attached."
            ),
            "For the May 1 director meeting, use the row-level watchlist for ownership and dates; treat external renewal ACV as a reconciliation ask.",
        ],
        font_size=Pt(11.5),
    )


def _fill_grr_proxy_slide(slide: Any, model: Any) -> None:
    left, top, width, height = _clear_standard_placeholder(slide)
    matrix = _matrix_rows(model, "Retention", "A1:B5")
    headers = [str(matrix[0][0] or ""), str(matrix[0][1] or "")]
    rows = [
        [str(matrix[1][0] or ""), _format_currency(matrix[1][1])],
        [str(matrix[2][0] or ""), _format_currency(matrix[2][1])],
        [str(matrix[3][0] or ""), _format_percentage(matrix[3][1])],
    ]
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=Emu(1524000),
        headers=headers,
        rows=rows,
        widths=(0.66, 0.34),
        right_align_cols=(1,),
    )
    footnote = str(matrix[4][0] or "").strip()
    if footnote:
        _add_textbox(
            slide,
            left + Emu(63500),
            top + Emu(1714500),
            width - Emu(127000),
            height - Emu(1778000),
            footnote,
            font_size=Pt(10),
            color=COLOR_MUTED_TEXT,
            italic=True,
        )


def _fill_owner_coaching_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="Owner coaching focus",
        subtitle=(
            "Current Q2 readiness plus original APAC push discipline: 3 owners carry 50 pushes across EUR 22.0M."
            if is_apac
            else "Current Q2 readiness grouped by account owner, using workbook-supported rows only."
        ),
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    tile_gap = Emu(220000)
    tile_width = Emu(int((int(width) - 127000 - (2 * int(tile_gap))) / 3))
    tile_left = left + Emu(63500)
    _add_kpi_tile(
        slide,
        left=tile_left,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(880000),
        label=("Original push spine" if is_apac else "Current push spine"),
        value=("3 owners carry 50 pushes" if is_apac else "Workbook rows"),
        note=("EUR 22.0M exposed in prior APAC pack." if is_apac else "Use current readiness flags; no prior pack applied."),
        accent=RGBColor(0x7B, 0x6B, 0xA8),
    )
    _add_kpi_tile(
        slide,
        left=tile_left + tile_width + tile_gap,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(880000),
        label="Pushed deals",
        value=("11 open deals pushed" if is_apac else "Current flagged rows"),
        note=("Edwina Chow owns 5; pattern review remains necessary." if is_apac else "Coach push root causes from workbook flags."),
        accent=COLOR_ACCENT,
    )
    _add_kpi_tile(
        slide,
        left=tile_left + (tile_width + tile_gap) * 2,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(880000),
        label="High-push tail",
        value=("2 at 5+ pushes" if is_apac else "Current high-push rows"),
        note=("EUR 1.0M highest risk; 5 at 3-4 pushes = EUR 3.0M." if is_apac else "Review push intensity from current readiness flags."),
        accent=RGBColor(0x3B, 0x7F, 0x8C),
    )
    summary: dict[str, dict[str, Any]] = {}
    for row in _readiness_records(legacy, limit=12):
        owner = row.get("Owner", "") or "(unknown)"
        item = summary.setdefault(owner, {"arr": 0.0, "deals": 0, "signals": []})
        item["arr"] += _number_from_display(row.get("ARR")) or 0
        item["deals"] += 1
        signal = row.get("Readiness", "")
        if signal:
            item["signals"].append(signal)
    rows: list[list[str]] = []
    for owner, item in sorted(summary.items(), key=lambda kv: kv[1]["arr"], reverse=True)[:7]:
        signals = "; ".join(dict.fromkeys(item["signals"]))
        if "Commercial approval gap" in signals:
            action = "Approval submission + deal desk timing."
        elif "No next step" in signals:
            action = "Next-step evidence and customer sponsor."
        elif "pushes" in signals:
            action = "Close-date realism and push root cause."
        else:
            action = "Confirm close proof."
        rows.append(
            [
                owner,
                str(item["deals"]),
                _format_meur(item["arr"]),
                _clip_cell(signals, 80),
                action,
            ]
        )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(1120000),
        width=width - Emu(127000),
        height=height - Emu(1190000),
        headers=["Owner", "# Q2 deals", "Q2 unweighted ARR", "Dominant signal", "Coaching ask"],
        rows=rows,
        widths=(0.16, 0.10, 0.12, 0.36, 0.26),
        right_align_cols=(1, 2),
    )


def _fill_qtd_wins_losses_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="QTD wins and losses",
        subtitle="Separate Land+Expand unweighted ARR from Renewal ACV; keep the motions distinct.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    matrix = _matrix_rows(legacy, "Wins_Losses_QTD", "A1:D3")
    headers = [str(cell or "") for cell in matrix[0]]
    values = [
        {header: _format_table_cell(header, value) for header, value in zip(headers, row)}
        for row in _filtered_body_rows(matrix)
    ]
    by_outcome = {row.get("Outcome", ""): row for row in values}
    won = by_outcome.get("Won", {})
    lost = by_outcome.get("Lost", {})
    won_arr = _number_from_display(won.get("ARR (Land+Expand)")) or 0
    lost_arr = _number_from_display(lost.get("ARR (Land+Expand)")) or 0
    tile_gap = Emu(220000)
    tile_width = Emu(int((int(width) - 127000 - (3 * int(tile_gap))) / 4))
    tile_left = left + Emu(63500)
    _add_kpi_tile(
        slide,
        left=tile_left,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(980000),
        label="Won L+E unweighted ARR",
        value=_format_meur(won.get("ARR (Land+Expand)")),
        note="Land+Expand closed-won unweighted ARR only.",
        accent=RGBColor(0x1F, 0x7A, 0x4D),
    )
    _add_kpi_tile(
        slide,
        left=tile_left + tile_width + tile_gap,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(980000),
        label="Lost L+E unweighted ARR",
        value=_format_meur(lost.get("ARR (Land+Expand)")),
        note=f"{lost.get('Count', '')} closed-lost/no-opportunity.",
        accent=COLOR_ACCENT,
    )
    _add_kpi_tile(
        slide,
        left=tile_left + (tile_width + tile_gap) * 2,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(980000),
        label="Net L+E unweighted ARR",
        value=_format_meur(won_arr - lost_arr),
        note="Loss pressure exceeds won unweighted ARR.",
        accent=COLOR_ACCENT,
    )
    _add_kpi_tile(
        slide,
        left=tile_left + (tile_width + tile_gap) * 3,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(980000),
        label="Won Renewal ACV",
        value=_format_meur(won.get("ACV (Renewal)")),
        note="Renewal ACV, separate motion.",
        accent=RGBColor(0x3B, 0x7F, 0x8C),
    )
    _add_bullet_block(
        slide,
        left=left + Emu(63500),
        top=top + Emu(1450000),
        width=Emu(int(int(width) * 0.43)),
        height=height - Emu(1520000),
        title="Q2 operating signal",
        bullets=[
            "The earlier mixed chart hid the most important fact: Land+Expand losses are roughly 3.7x won unweighted ARR by value.",
            "Renewal ACV is a separate closed-renewal motion; Land+Expand unweighted ARR loss pressure remains its own signal.",
            "Original Q1 Land losses: 14 / EUR 5.5M; Missing reason code and Stage-at-loss mix need hygiene review; forecast accuracy was 1W / 16L, EUR 8.6M lost, 6% win rate.",
        ],
        font_size=Pt(11.5),
    )
    _styled_table(
        slide,
        left=left + Emu(int(int(width) * 0.49)),
        top=top + Emu(1450000),
        width=Emu(int(int(width) * 0.47)),
        height=height - Emu(1520000),
        headers=[("Original Q1 loss spine" if is_apac else "Prior Q1 loss spine"), "#", "Unweighted ARR", "Action"],
        rows=(
            JESPER_ORIGINAL_Q1_LOSS_ROWS
            if is_apac
            else [["Prior territory loss pack", "Not attached", "Not attached", "Do not infer Q1 losses from another director pack."]]
        ),
        widths=(0.28, 0.09, 0.14, 0.49),
        right_align_cols=(1, 2),
    )


def _fill_deal_hygiene_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="Deal hygiene scorecard",
        subtitle="Use actual readiness flags from the Q2 deal list.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _readiness_records(legacy, limit=12)
    signals = [
        ("Silent 90d+", lambda r: "Silent 90d+" in r.get("Readiness", "")),
        ("No next step", lambda r: "No next step" in r.get("Readiness", "")),
        ("4+ pushes", lambda r: "pushes" in r.get("Readiness", "")),
        ("Approval gap", lambda r: "Commercial approval gap" in r.get("Readiness", "")),
    ]
    rows = []
    for label, predicate in signals:
        matched = [row for row in records if predicate(row)]
        rows.append(
            [
                label,
                str(len(matched)),
                _format_meur(sum(_number_from_display(row.get("ARR")) or 0 for row in matched)),
                _clip_cell(", ".join(row.get("Opportunity", "") for row in matched[:4]), 82),
            ]
        )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=Emu(2350000),
        headers=["Signal", "# deals", "Unweighted ARR", "Named deals"],
        rows=rows,
        widths=(0.20, 0.12, 0.16, 0.52),
        right_align_cols=(1, 2),
    )
    _add_bullet_block(
        slide,
        left=left + Emu(63500),
        top=top + Emu(2620000),
        width=width - Emu(127000),
        height=height - Emu(2700000),
        title="Weekly inspection gate",
        bullets=(
            [
                "Original APAC Q2 book: 6 deals / EUR 5.0M and zero recent activity; this remains the baseline for activity-recovery coaching.",
                "Every Q2 deal needs a dated customer step, a current activity signal, and an approval status before the director meeting.",
                "Deals failing two or more hygiene checks should be removed from Commit/Best Case unless the rep provides customer evidence.",
            ]
            if is_apac
            else [
                "Every Q2 deal needs a dated customer step, a current activity signal, and an approval status before the director meeting.",
                "Deals failing two or more hygiene checks should be removed from Commit/Best Case unless the rep provides customer evidence.",
                "Use this director's original ETL sidecar for prior context; do not borrow another territory's baseline.",
            ]
        ),
        font_size=Pt(11.5),
    )


def _fill_concentration_slide(slide: Any, model: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="Concentration risk",
        subtitle=(
            "Current all-open exposure plus original APAC concentration spine: top 7 open deals / EUR 9.6M, top 5 = 89%, Amova Asset Management."
            if is_apac
            else "Current all-open exposure plus this director's original ETL/gold concentration spine."
        ),
    )
    _remove_shape(_find_shape(slide, "Rectangle 5"))
    _remove_shape(_find_shape(slide, "Content Placeholder 1"))
    left = Emu(839787)
    top = Emu(1143000)
    metrics = [
        ("Largest account", _clip_cell(model.cell_value("Concentration", "B5"), 23)),
        ("Largest unweighted ARR", _format_currency(model.cell_value("Concentration", "B6"))),
        ("All-open share", _format_percentage(model.cell_value("Concentration", "B7"))),
        ("25% threshold", str(model.cell_value("Concentration", "B8") or "")),
    ]
    tile_width = Emu(2520000)
    for idx, (label, value) in enumerate(metrics):
        x = left + Emu(idx * 2520000)
        _add_textbox(slide, x, top, tile_width, Emu(180000), label, font_size=Pt(10), color=COLOR_MUTED_TEXT)
        _add_textbox(slide, x, top + Emu(185000), tile_width, Emu(320000), value, font_size=Pt(15), bold=True)
        if idx < len(metrics) - 1:
            _add_rule(slide, x + tile_width - Emu(80000), top + Emu(63500), Emu(12700), color=COLOR_RULE)
    _add_bullet_block(
        slide,
        left=Emu(839787),
        top=Emu(2050000),
        width=Emu(4700000),
        height=Emu(2300000),
        title="Concentration watch",
        bullets=[
            "Use concentration as a call-prep filter: the largest deal should have explicit sponsor, approval, decision date, and fallback plan.",
            "If the largest-deal share is above threshold, headline forecast confidence should cite that dependency directly.",
        ],
        font_size=Pt(11.5),
    )
    _styled_table(
        slide,
        left=Emu(5950000),
        top=Emu(2050000),
        width=Emu(5550000),
        height=Emu(2300000),
        headers=[("Original APAC baseline" if is_apac else "Original ETL baseline"), "Value", "How to use it"],
        rows=(
            JESPER_ORIGINAL_CONCENTRATION_ROWS
            if is_apac
            else [["Original concentration pack", "Attached", "Compare current largest-deal exposure to the director's source pack."]]
        ),
        widths=(0.30, 0.28, 0.42),
        right_align_cols=(1,),
    )


def _fill_named_stale_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="Q2-Q3 risk triage",
        subtitle=(
            "Original APAC risk spine preserved: top four Q2-Q3 exposed deals total EUR 4.9M."
            if is_apac
            else "Current workbook risk triage plus this director's original territory risk spine."
        ),
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    risk_rows = (
        JESPER_Q2_Q3_RISK_ROWS
        if is_apac
        else [
            [
                _clip_cell(row.get("Account", ""), 24),
                _clip_cell(row.get("Opportunity", ""), 42),
                _clip_cell(row.get("Stage", ""), 20),
                row.get("Close Date", ""),
                row.get("ARR", ""),
                _clip_cell(row.get("Readiness", "") or "Current workbook row", 58),
            ]
            for row in _readiness_records(legacy, limit=4)
        ]
    )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=Emu(2350000),
        headers=["Account", "Opportunity", "Stage", "Close", "Unweighted ARR", "Reason"],
        rows=risk_rows,
        widths=(0.15, 0.25, 0.13, 0.10, 0.10, 0.27),
        right_align_cols=(4,),
    )
    _add_bullet_block(
        slide,
        left=left + Emu(63500),
        top=top + Emu(2650000),
        width=width - Emu(127000),
        height=height - Emu(2720000),
        title="Follow-through",
        bullets=(
            [
                "This preserves the original APAC lens: current Q2 readiness plus Q3 push risk.",
                "Coolabah and BOCI are out-of-quarter context, but they explain why APAC needs push discipline beyond the May close call.",
                "Mandiri and Danantara remain current-quarter operating actions.",
            ]
            if is_apac
            else [
                "Use the current workbook's highest-risk rows as the operating watchlist.",
                "Out-of-quarter context needs a director-specific source pack before it is added to this page.",
                "Current-quarter operating actions should tie to named rows, dated next steps, and approval status.",
            ]
        ),
        font_size=Pt(12),
    )


def _fill_sales_velocity_slide(slide: Any, legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="Operating rhythm",
        subtitle="Weekly operating cadence for the director review.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    records = _readiness_records(legacy, limit=8)
    rows = [
        ["1", "Deal evidence", "Every Commit/Pipeline deal", "Customer-owned next step, dated decision, and sponsor proof."],
        ["2", "Forecast hygiene", "Commit + Pipeline", "Downgrade deals with stale activity or no next step unless proof is supplied."],
        ["3", "Governance", "Stage 3+ approval gaps", "Submit Commercial Approval before the next forecast call."],
        ["4", "Loss discipline", "Closed-lost QTD", "Attach reason codes and preserve director-specific lesson-learned actions."],
        (
            ["5", "Push discipline", "Original APAC pushed-deal spine", "Review 11 open deals pushed; Edwina Chow owns 5; 2 at 5+ pushes / EUR 1.0M."]
            if is_apac
            else ["5", "Push discipline", "Current workbook push flags", "Review pushed deals from this director's readiness rows."]
        ),
        ["6", "Renewal ACV", "FY26 renewal watchlist", "Validate current FY26 ACV basis before quoting outside this review."],
    ]
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=Emu(2650000),
        headers=["Cadence", "Gate", "Population", "Required evidence"],
        rows=rows,
        widths=(0.08, 0.18, 0.24, 0.50),
        right_align_cols=(0,),
    )
    spotlight = ", ".join(row.get("Opportunity", "") for row in records[:3])
    _add_bullet_block(
        slide,
        left=left + Emu(63500),
        top=top + Emu(2950000),
        width=width - Emu(127000),
        height=height - Emu(3020000),
        title="This week's first inspection",
        bullets=[f"Start with {spotlight}. These are the highest-value readiness rows in the current Q2 book."],
        font_size=Pt(12.5),
    )


def _fill_account_expansion_slide(slide: Any, trends: dict[str, Any], legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="May territory plan",
        subtitle="May 2026 territory plan for Friday, May 1: close evidence, pipeline creation, governance, and renewal ACV reconciliation.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    forecast = {row.get("Category", ""): row for row in _forecast_category_records(legacy)}
    commit = forecast.get("Commit", {})
    pipeline = forecast.get("Pipeline", {})
    wins_losses = _matrix_rows(legacy, "Wins_Losses_QTD", "A1:D3")
    won_arr = wins_losses[1][2] if len(wins_losses) > 1 else ""
    lost_arr = wins_losses[2][2] if len(wins_losses) > 2 else ""
    won_numeric = _number_from_display(won_arr) or 0
    lost_numeric = _number_from_display(lost_arr) or 0
    renewals = _renewal_records(legacy)
    renewal_total = sum(_number_from_display(row.get("ACV")) or 0 for row in renewals)
    approvals = _pending_approval_records(legacy)
    approval_label = _approval_gap_label(approvals)
    approval_total = sum(_number_from_display(row.get("ARR (EUR)")) or 0 for row in approvals)

    tile_gap = Emu(220000)
    tile_width = Emu(int((int(width) - 127000 - (3 * int(tile_gap))) / 4))
    tile_left = left + Emu(63500)
    _add_kpi_tile(
        slide,
        left=tile_left,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(940000),
        label="Q2 L+E unweighted ARR",
        value=_format_meur(_kpi_value(trends, "total_pipeline_arr")),
        note="Current-quarter unweighted ARR only.",
    )
    _add_kpi_tile(
        slide,
        left=tile_left + tile_width + tile_gap,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(940000),
        label="Commit / Pipeline unweighted ARR",
        value=f"{_format_meur(commit.get('ARR (EUR)')).replace('EUR ', '')} / {_format_meur(pipeline.get('ARR (EUR)')).replace('EUR ', '')}",
        note="Best Case is not material.",
        accent=RGBColor(0x7B, 0x6B, 0xA8),
    )
    _add_kpi_tile(
        slide,
        left=tile_left + (tile_width + tile_gap) * 2,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(940000),
        label="QTD net L+E unweighted ARR",
        value=_format_meur(won_numeric - lost_numeric),
        note="Loss pressure exceeds won unweighted ARR.",
        accent=COLOR_ACCENT,
    )
    _add_kpi_tile(
        slide,
        left=tile_left + (tile_width + tile_gap) * 3,
        top=top + Emu(63500),
        width=tile_width,
        height=Emu(940000),
        label="Creation pool",
        value=("194 Tier-1" if is_apac else "Validate"),
        note=("No open L+E opp in 90d." if is_apac else "Use this director's account coverage source."),
        accent=RGBColor(0x3B, 0x7F, 0x8C),
    )
    rows = (
        [
            [
                "Confirm or reset Q2 close",
                "Danantara, LTH, Krungthai, Mandiri, Temasek, HKMA",
                "Dated customer next step, sponsor proof, approval status; Krungthai is overdue-close and LTH conflicts with August decision timing.",
            ],
            ["Create future pipe", "194 Tier-1 accounts with no open Land/Expand opp in 90d", "Rep-owned coverage list and first-touch plan by May 8."],
            [
                "Clean forecast",
                f"26 Q2 opps with no activity in 30d; {approval_label} approval gap totals {_format_meur(approval_total)} unweighted ARR",
                "Refresh activity/NextStep, submit approvals, and downgrade deals without customer proof.",
            ],
            [
                "Protect Renewal view",
                f"Original EUR 33.5M vs refreshed {_format_meur(renewal_total)} FY26 Renewal ACV",
                "Reconcile ACV basis and keep Renewal ACV separate from Land+Expand unweighted ARR.",
            ],
        ]
        if is_apac
        else [
            [
                "Confirm or reset Q2 close",
                "Top current-quarter readiness rows",
                "Dated customer next step, sponsor proof, approval status, and realistic close timing.",
            ],
            ["Create future pipe", "Current account coverage gap", "Rep-owned coverage list and first-touch plan by May 8."],
            [
                "Clean forecast",
                f"{approval_label} approval gap totals {_format_meur(approval_total)} unweighted ARR",
                "Refresh activity/NextStep, submit approvals, and downgrade deals without customer proof.",
            ],
            [
                "Protect Renewal view",
                f"Refreshed {_format_meur(renewal_total)} FY26 Renewal ACV",
                "Validate ACV basis and keep Renewal ACV separate from Land+Expand unweighted ARR.",
            ],
        ]
    )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(1200000),
        width=width - Emu(127000),
        height=height - Emu(1270000),
        headers=["May lane", "Fact base", "May output"],
        rows=rows,
        widths=(0.19, 0.35, 0.46),
    )


def _fill_next_14_days_slide(slide: Any, trends: dict[str, Any], legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="May 1-15 operating cadence",
        subtitle="First two-week director cadence for Q2: launch, inspect, clean, and reset the forecast.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    readiness = _readiness_records(legacy, limit=6)
    approvals = _pending_approval_records(legacy)
    approval_label = _approval_gap_label(approvals)
    rows = [
        ["May 1", "Director kickoff", "Confirm Q2 Commit/Pipeline roster and owner asks.", "Top six deals have named owners and required proof."],
        (
            ["May 4-6", "Deal inspection", "Danantara, LTH, Krungthai, Mandiri, Temasek, HKMA.", "Customer next step, sponsor, close proof, approval status; flag Krungthai overdue-close and LTH August timing."]
            if is_apac
            else ["May 4-6", "Deal inspection", "Highest-value current readiness rows.", "Customer next step, sponsor, close proof, approval status, and close-date realism."]
        ),
        ["May 7-8", "Governance gate", f"{approval_label}; {len(approvals)} supported current-state rows.", "Submissions logged; stale/no-next-step rows refreshed."],
        ["May 11-13", "Forecast reset", "Deals missing customer proof or realistic timing.", "Downgrade, re-date, or confirm with evidence."],
        ["May 14-15", "May checkpoint", "Q2 close view, Tier-1 creation plan, Renewal ACV basis.", "One evidence-backed May forecast and next pipeline-creation list."],
    ]
    if readiness:
        rows.append(
            [
                "Every week",
                "Owner hygiene",
                "All Q2 readiness rows.",
                "LastActivityDate and NextStep updated before forecast call.",
            ]
        )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=height - Emu(127000),
        headers=["Date", "Forum", "Focus", "Output"],
        rows=rows,
        widths=(0.12, 0.18, 0.32, 0.38),
    )


def _action_items_table(trends: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    for idx, item in enumerate(trends.get("action_items", []), start=1):
        rows.append(
            [
                str(idx),
                str(item.get("priority") or "").upper(),
                str(item.get("rule_id") or ""),
                str(item.get("claim") or ""),
                str(item.get("suggested_action") or ""),
                str(item.get("due_date") or ""),
                str(item.get("owner") or ""),
            ]
        )
    return rows or [["", "", "", "(no action items generated)", "", "", ""]]


def _fill_action_items_slide(slide: Any, trends: dict[str, Any], legacy: Any, *, is_apac: bool) -> None:
    _set_slide_header(
        slide,
        title="May action register",
        subtitle="Named May actions tied to source evidence; broader rules remain horizon-labeled.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    approvals = _pending_approval_records(legacy)
    approval_label = _approval_gap_label(approvals)
    approval_phrase = _approval_gap_phrase(approvals)
    activity_claim = _action_item_claim(trends, "activity_drought")
    rows = (
        [
            ["1", "HIGH", "Danantara", "EUR 1.8M unweighted ARR Pipeline, 3 pushes, silent 90d, no next step", "Jesper to confirm sponsor, dated customer meeting, and approval path."],
            ["2", "HIGH", "LTH / Krungthai", "EUR 2.0M unweighted ARR Commit but stale activity; LTH next step points after Q2", "Edwina to confirm Q2 close timing with customer proof or reset forecast."],
            ["3", "HIGH", approval_label, f"{approval_phrase}, {_format_meur(sum(_number_from_display(row.get('ARR (EUR)')) or 0 for row in approvals))} unweighted ARR; all-open 7 / EUR 10.7M needs appendix", "Submit current Q2 approval gaps; keep all-open claim internal until row-level support is attached."],
            ["4", "HIGH", "Full Q2 book", activity_claim or "Q2 activity drought requires refresh", "Every owner updates activity and next step before director review."],
            ["5", "MED", "Tier-1 creation", "194 Tier-1 accounts have no open Land/Expand opp in 90d", "Assign rep-owned coverage list and first-touch plan by May 8."],
            ["6", "MED", "Push discipline", "Original APAC: 11 open deals pushed; Edwina Chow owns 5; EUR 22.0M unweighted ARR owner-push exposure", "Coach push root causes before accepting Q2/Q3 close-date confidence."],
            ["7", "MED", "FY26 renewals", "Original EUR 33.5M vs refreshed EUR 5.0M FY26 Renewal ACV", "Sales Ops reconciles ACV basis before external quote."],
        ]
        if is_apac
        else [
            ["1", "HIGH", "Top Q2 deal", "Highest-value readiness row from current workbook", "Confirm sponsor, dated customer meeting, and approval path."],
            ["2", "HIGH", "Close timing", "Current Commit/Pipeline rows with stale activity, no next step, or close-date risk", "Confirm Q2 close timing with customer proof or reset forecast."],
            ["3", "HIGH", approval_label, f"{approval_phrase}, {_format_meur(sum(_number_from_display(row.get('ARR (EUR)')) or 0 for row in approvals))} unweighted ARR", "Submit current Q2 approval gaps; keep unsupported all-open claims internal."],
            ["4", "HIGH", "Full Q2 book", activity_claim or "Q2 activity drought requires refresh", "Every owner updates activity and next step before director review."],
            ["5", "MED", "Tier-1 creation", "Current account coverage gap", "Assign rep-owned coverage list and first-touch plan by May 8."],
            ["6", "MED", "Push discipline", "Current readiness rows with push flags", "Coach push root causes before accepting Q2/Q3 close-date confidence."],
            ["7", "MED", "FY26 renewals", "Refreshed FY26 Renewal ACV requires basis validation", "Sales Ops reconciles ACV basis before external quote."],
        ]
    )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=height - Emu(127000),
        headers=["#", "Priority", "Scope", "Evidence", "Action"],
        rows=rows,
        widths=(0.04, 0.08, 0.17, 0.36, 0.35),
        right_align_cols=(0,),
    )


def _fill_risks_outlook_slide(
    slide: Any,
    *,
    trends: dict[str, Any],
    legacy: Any,
    is_apac: bool,
) -> None:
    _set_slide_header(
        slide,
        title="May decision checklist",
        subtitle="Open decisions for the May territory review.",
    )
    left, top, width, height = _clear_standard_placeholder(slide)
    approvals = _pending_approval_records(legacy)
    approval_label = _approval_gap_label(approvals)
    approval_question = (
        "Is the supported Q2 approval gap submitted?"
        if len(approvals) == 1
        else f"Are the {len(approvals)} supported Q2 approval gaps submitted?"
    )
    rows = (
        [
            ["1", "Forecast call", "Is there named customer evidence behind EUR 2.8M Commit and EUR 2.4M Pipeline?", "Deal owners refresh next steps and close proof."],
            ["2", "Governance", approval_question, f"Track {approval_label}; keep Coolabah and Danantara as approved-2026 context."],
            ["3", "Q1 learning", "Are Q1 losses, missing reason codes, and slipped deals included in the operating readout?", "Use the original APAC targets as the accountability baseline."],
            ["4", "Push discipline", "Are the original APAC push patterns owned as coaching actions?", "Review 3 owners / 50 pushes and Edwina's 5 pushed deals."],
            ["5", "Renewals", "Is the current FY26 Renewal ACV basis validated for external quote?", "Keep renewal ACV separate from Land+Expand unweighted ARR."],
        ]
        if is_apac
        else [
            ["1", "Forecast call", "Is there named customer evidence behind current Commit and Pipeline?", "Deal owners refresh next steps and close proof."],
            ["2", "Governance", approval_question, f"Track {approval_label}; do not import approval context from another territory pack."],
            ["3", "Q1 learning", "Are director-specific losses, reason codes, and slipped deals included in the operating readout?", "Use only this director's prior pack when available."],
            ["4", "Push discipline", "Are current workbook push patterns owned as coaching actions?", "Review current readiness rows with push flags."],
            ["5", "Renewals", "Is the current FY26 Renewal ACV basis validated for external quote?", "Keep renewal ACV separate from Land+Expand unweighted ARR."],
        ]
    )
    _styled_table(
        slide,
        left=left + Emu(63500),
        top=top + Emu(63500),
        width=width - Emu(127000),
        height=height - Emu(127000),
        headers=["#", "Decision", "Question", "Required closure"],
        rows=rows,
        widths=(0.05, 0.16, 0.44, 0.35),
        right_align_cols=(0,),
    )


def _inject_donor_charts(pptx_path: Path) -> None:
    with ZipFile(pptx_path) as zf:
        target_entries = {name: zf.read(name) for name in zf.namelist()}

    ct_root = ET.fromstring(target_entries["[Content_Types].xml"])

    for injection in CHART_INJECTIONS:
        _inject_single_chart(
            target_entries=target_entries,
            ct_root=ct_root,
            injection=injection,
        )

    target_entries["[Content_Types].xml"] = ET.tostring(ct_root, encoding="utf-8", xml_declaration=True)

    tmp_path = pptx_path.with_suffix(".patched.pptx")
    with ZipFile(tmp_path, "w", compression=ZIP_DEFLATED) as zf:
        for name, data in target_entries.items():
            zf.writestr(name, data)
    tmp_path.replace(pptx_path)


def _inject_single_chart(
    *,
    target_entries: dict[str, bytes],
    ct_root: ET.Element,
    injection: ChartInjection,
) -> None:
    slide_name = f"ppt/slides/slide{injection.slide_number}.xml"
    rels_name = _rels_part_name(slide_name)

    target_slide_root = ET.fromstring(target_entries[slide_name])
    target_rels_root = ET.fromstring(target_entries[rels_name])
    target_sp_tree = target_slide_root.find(f".//{{{P_NS}}}spTree")
    if target_sp_tree is None:
        raise RuntimeError(f"{slide_name} is missing spTree")

    _remove_target_placeholder_shapes(target_sp_tree)
    next_shape_id = _max_shape_id(target_sp_tree) + 1

    with ZipFile(injection.donor.template_path) as donor_zip:
        donor_ct_root = ET.fromstring(donor_zip.read("[Content_Types].xml"))
        donor_slide_name = f"ppt/slides/slide{injection.donor.slide_number}.xml"
        donor_rels_name = _rels_part_name(donor_slide_name)
        donor_slide_root = ET.fromstring(donor_zip.read(donor_slide_name))
        donor_rels_root = ET.fromstring(donor_zip.read(donor_rels_name))
        donor_rel_map = _relationship_map(donor_rels_root)
        copier = _PartCopier(
            target_entries=target_entries,
            ct_root=ct_root,
            donor_zip=donor_zip,
            donor_ct_root=donor_ct_root,
            chart_name=injection.chart_name,
        )

        donor_sp_tree = donor_slide_root.find(f".//{{{P_NS}}}spTree")
        if donor_sp_tree is None:
            raise RuntimeError(f"{donor_slide_name} is missing spTree")

        for donor_child in list(donor_sp_tree)[2:]:
            shape_name = _shape_name(donor_child)
            shape_text = _shape_text(donor_child)

            cloned = copy.deepcopy(donor_child)
            if _exclude_donor_shape(injection.donor, shape_name, shape_text):
                continue
            _replace_shape_text(cloned, dict(injection.text_replacements))
            _rewrite_relationships(
                cloned,
                donor_slide_name=donor_slide_name,
                target_slide_name=slide_name,
                target_rels_root=target_rels_root,
                donor_rel_map=donor_rel_map,
                copier=copier,
            )
            _set_shape_id(cloned, next_shape_id)
            next_shape_id += 1
            target_sp_tree.append(cloned)

    target_entries[slide_name] = ET.tostring(target_slide_root, encoding="utf-8", xml_declaration=True)
    target_entries[rels_name] = ET.tostring(target_rels_root, encoding="utf-8", xml_declaration=True)


def _remove_target_placeholder_shapes(sp_tree: ET.Element) -> None:
    for child in list(sp_tree):
        name = _shape_name(child)
        if name in {"Content Placeholder 1", "Rectangle 5"}:
            sp_tree.remove(child)


def _shape_name(element: ET.Element) -> str:
    c_nv_pr = element.find(f".//{{{P_NS}}}cNvPr")
    if c_nv_pr is None:
        c_nv_pr = element.find(f".//{{{A_NS}}}cNvPr")
    return "" if c_nv_pr is None else c_nv_pr.get("name", "")


def _shape_text(element: ET.Element) -> str:
    return "".join(node.text or "" for node in element.findall(f".//{{{A_NS}}}t"))


def _exclude_donor_shape(donor: DonorSlide, shape_name: str, shape_text: str) -> bool:
    if shape_name in donor.exclude_names:
        return True
    return any(shape_text.startswith(prefix) for prefix in donor.exclude_text_prefixes)


def _replace_shape_text(element: ET.Element, replacements: dict[str, str]) -> None:
    for text_node in element.findall(f".//{{{A_NS}}}t"):
        text = text_node.text or ""
        if text in replacements:
            text_node.text = replacements[text]


def _max_shape_id(sp_tree: ET.Element) -> int:
    max_id = 0
    for c_nv_pr in sp_tree.findall(f".//{{{P_NS}}}cNvPr"):
        value = c_nv_pr.get("id")
        if value and value.isdigit():
            max_id = max(max_id, int(value))
    return max_id


def _set_shape_id(element: ET.Element, new_id: int) -> None:
    c_nv_pr = element.find(f".//{{{P_NS}}}cNvPr")
    if c_nv_pr is None:
        c_nv_pr = element.find(f".//{{{A_NS}}}cNvPr")
    if c_nv_pr is not None:
        c_nv_pr.set("id", str(new_id))


def _hide_shape(element: ET.Element) -> None:
    c_nv_pr = element.find(f".//{{{P_NS}}}cNvPr")
    if c_nv_pr is None:
        c_nv_pr = element.find(f".//{{{A_NS}}}cNvPr")
    if c_nv_pr is not None:
        c_nv_pr.set("hidden", "1")


def _relationship_map(rels_root: ET.Element) -> dict[str, ET.Element]:
    rels: dict[str, ET.Element] = {}
    for rel in rels_root.findall(f".//{{{PKG_REL_NS}}}Relationship"):
        rel_id = rel.get("Id")
        if rel_id:
            rels[rel_id] = rel
    return rels


class _PartCopier:
    def __init__(
        self,
        *,
        target_entries: dict[str, bytes],
        ct_root: ET.Element,
        donor_zip: ZipFile,
        donor_ct_root: ET.Element,
        chart_name: str,
    ) -> None:
        self.target_entries = target_entries
        self.ct_root = ct_root
        self.donor_zip = donor_zip
        self.donor_ct_root = donor_ct_root
        self.chart_name = chart_name
        self.memo: dict[str, str] = {}

    def copy_part(self, donor_part_name: str) -> str:
        donor_part_name = donor_part_name.lstrip("/")
        if donor_part_name in self.memo:
            return self.memo[donor_part_name]

        target_part_name = _next_available_name(self.target_entries, donor_part_name)
        data = self.donor_zip.read(donor_part_name)
        if target_part_name.endswith(".bin"):
            data = _patch_ole_name(data, self.chart_name)
        self.target_entries[target_part_name] = data
        self.memo[donor_part_name] = target_part_name
        _copy_content_type(self.ct_root, self.donor_ct_root, donor_part_name, target_part_name)

        donor_rels_name = _rels_part_name(donor_part_name)
        if donor_rels_name in self.donor_zip.namelist():
            donor_rels_root = ET.fromstring(self.donor_zip.read(donor_rels_name))
            for rel in donor_rels_root.findall(f".//{{{PKG_REL_NS}}}Relationship"):
                if rel.get("TargetMode") == "External":
                    continue
                donor_child = _resolve_target(donor_part_name, rel.get("Target", ""))
                target_child = self.copy_part(donor_child)
                rel.set("Target", _relative_target(target_part_name, target_child))
            self.target_entries[_rels_part_name(target_part_name)] = ET.tostring(
                donor_rels_root,
                encoding="utf-8",
                xml_declaration=True,
            )
        return target_part_name


def _rewrite_relationships(
    element: ET.Element,
    *,
    donor_slide_name: str,
    target_slide_name: str,
    target_rels_root: ET.Element,
    donor_rel_map: dict[str, ET.Element],
    copier: _PartCopier,
) -> None:
    rewritten_rids: dict[str, str] = {}
    for node in element.iter():
        for attr_name in (f"{{{R_NS}}}id", f"{{{R_NS}}}embed", f"{{{R_NS}}}link"):
            old_rid = node.get(attr_name)
            if not old_rid or old_rid not in donor_rel_map:
                continue
            if old_rid in rewritten_rids:
                node.set(attr_name, rewritten_rids[old_rid])
                continue
            donor_rel = donor_rel_map[old_rid]
            target_mode = donor_rel.get("TargetMode")
            rel_type = donor_rel.get("Type")
            if not rel_type:
                continue
            new_rid = _next_relationship_id(target_rels_root)
            new_rel = ET.SubElement(target_rels_root, f"{{{PKG_REL_NS}}}Relationship")
            new_rel.set("Id", new_rid)
            new_rel.set("Type", rel_type)
            if target_mode == "External":
                new_rel.set("TargetMode", "External")
                new_rel.set("Target", donor_rel.get("Target", ""))
            else:
                donor_target = _resolve_target(donor_slide_name, donor_rel.get("Target", ""))
                copied_target = copier.copy_part(donor_target)
                new_rel.set("Target", _relative_target(target_slide_name, copied_target))
            rewritten_rids[old_rid] = new_rid
            node.set(attr_name, new_rid)


def _patch_ole_name(data: bytes, chart_name: str) -> bytes:
    if b"<m_strName>" not in data:
        return data
    try:
        result = replace_cfb_stream_data(
            data,
            (
                CfbStreamEdit(
                    name="think-cellXML",
                    replacements=(
                        (
                            b"<m_strName></m_strName>",
                            f"<m_strName>{chart_name}</m_strName>".encode("utf-8"),
                        ),
                    ),
                ),
            ),
        )
        if result.replacements_made:
            return result.data
    except ValueError:
        pass
    return _patch_cfb_stream(
        data,
        "think-cellXML",
        {
            rb"<m_strName>.*?</m_strName>": f"<m_strName>{chart_name}</m_strName>".encode(
                "utf-8"
            )
        },
        regex=True,
    )


@dataclass(frozen=True)
class _CfbStream:
    name: str
    data: bytes


def _cfb_sector_slice(sector_id: int, sector_size: int) -> slice:
    start = (sector_id + 1) * sector_size
    return slice(start, start + sector_size)


def _cfb_chain(fat: list[int], start_sector: int) -> list[int]:
    chain: list[int] = []
    sector_id = start_sector
    seen: set[int] = set()
    while (
        sector_id not in (CFB_FREE, CFB_END_OF_CHAIN)
        and sector_id < len(fat)
        and sector_id not in seen
    ):
        seen.add(sector_id)
        chain.append(sector_id)
        sector_id = fat[sector_id]
    return chain


def _cfb_fat(data: bytes, sector_size: int, fat_sector_count: int) -> list[int]:
    difat = list(struct.unpack_from("<109I", data, 76))
    fat_sectors = [
        sector_id
        for sector_id in difat
        if sector_id not in (CFB_FREE, CFB_END_OF_CHAIN, CFB_FAT_SECTOR, CFB_DIFAT_SECTOR)
    ][:fat_sector_count]
    fat: list[int] = []
    for sector_id in fat_sectors:
        sector = data[_cfb_sector_slice(sector_id, sector_size)]
        fat.extend(struct.unpack(f"<{sector_size // 4}I", sector))
    return fat


def _cfb_directory(data: bytes) -> tuple[int, list[int], bytearray, list[slice]]:
    sector_size = 1 << struct.unpack_from("<H", data, 30)[0]
    fat_sector_count = struct.unpack_from("<I", data, 44)[0]
    first_directory_sector = struct.unpack_from("<I", data, 48)[0]
    fat = _cfb_fat(data, sector_size, fat_sector_count)
    directory_chain = _cfb_chain(fat, first_directory_sector)
    directory_bytes = bytearray()
    directory_slices: list[slice] = []
    for sector_id in directory_chain:
        sector_slice = _cfb_sector_slice(sector_id, sector_size)
        directory_slices.append(sector_slice)
        directory_bytes.extend(data[sector_slice])
    return sector_size, fat, directory_bytes, directory_slices


def _cfb_streams(data: bytes) -> list[_CfbStream]:
    if data[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return []

    sector_size, fat, directory_bytes, _ = _cfb_directory(data)
    streams: list[_CfbStream] = []
    for entry_offset in range(0, len(directory_bytes), 128):
        entry = directory_bytes[entry_offset : entry_offset + 128]
        name_length = struct.unpack_from("<H", entry, 64)[0]
        if name_length < 2 or entry[66] != 2:
            continue
        name = entry[: name_length - 2].decode("utf-16le", "ignore")
        start_sector = struct.unpack_from("<I", entry, 116)[0]
        stream_size = struct.unpack_from("<Q", entry, 120)[0]
        stream_chain = _cfb_chain(fat, start_sector)
        stream = bytearray()
        for sector_id in stream_chain:
            stream.extend(data[_cfb_sector_slice(sector_id, sector_size)])
        streams.append(_CfbStream(name=name, data=bytes(stream[:stream_size])))
    return streams


def _patch_cfb_stream(
    data: bytes,
    stream_name: str,
    replacements: dict[bytes, bytes],
    *,
    regex: bool = False,
) -> bytes:
    if data[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return data

    out = bytearray(data)
    sector_size, fat, directory_bytes, directory_slices = _cfb_directory(data)

    for entry_offset in range(0, len(directory_bytes), 128):
        entry = directory_bytes[entry_offset : entry_offset + 128]
        name_length = struct.unpack_from("<H", entry, 64)[0]
        if name_length < 2 or entry[66] != 2:
            continue
        name = entry[: name_length - 2].decode("utf-16le", "ignore")
        if name != stream_name:
            continue

        start_sector = struct.unpack_from("<I", entry, 116)[0]
        stream_size = struct.unpack_from("<Q", entry, 120)[0]
        stream_chain = _cfb_chain(fat, start_sector)
        capacity = len(stream_chain) * sector_size
        stream = bytearray()
        for sector_id in stream_chain:
            stream.extend(out[_cfb_sector_slice(sector_id, sector_size)])
        patched = bytes(stream[:stream_size])

        for old, new in replacements.items():
            if regex:
                patched = re.sub(old, new, patched, flags=re.DOTALL)
            else:
                patched = patched.replace(old, new)

        if len(patched) < stream_size:
            patched = patched + b" " * (stream_size - len(patched))
        elif len(patched) > stream_size:
            if len(patched) > capacity:
                raise RuntimeError(
                    f"{stream_name} grew from {stream_size} to {len(patched)}, "
                    f"but the allocated CFB sector chain only holds {capacity} bytes"
                )
            struct.pack_into("<Q", directory_bytes, entry_offset + 120, len(patched))

        padded = patched.ljust(capacity, b"\x00")
        cursor = 0
        for sector_id in stream_chain:
            sector_slice = _cfb_sector_slice(sector_id, sector_size)
            out[sector_slice] = padded[cursor : cursor + sector_size]
            cursor += sector_size

        cursor = 0
        for directory_slice in directory_slices:
            out[directory_slice] = directory_bytes[cursor : cursor + sector_size]
            cursor += sector_size
        return bytes(out)

    return data


def _rels_part_name(part_name: str) -> str:
    directory, filename = posixpath.split(part_name)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def _resolve_target(base_part_name: str, target: str) -> str:
    base_dir = posixpath.dirname("/" + base_part_name)
    resolved = posixpath.normpath(posixpath.join(base_dir, target))
    return resolved.lstrip("/")


def _relative_target(from_part_name: str, to_part_name: str) -> str:
    from_dir = posixpath.dirname("/" + from_part_name)
    return posixpath.relpath("/" + to_part_name, from_dir)


def _next_available_name(entries: dict[str, bytes], donor_part_name: str) -> str:
    donor_part_name = donor_part_name.lstrip("/")
    if donor_part_name not in entries:
        return donor_part_name

    directory, filename = posixpath.split(donor_part_name)
    stem, suffix = _split_filename(filename)
    match = re.match(r"^(.*?)(\d+)$", stem)
    prefix = stem
    counter = 1
    if match:
        prefix = match.group(1)
        counter = int(match.group(2))
    while True:
        counter += 1
        candidate = posixpath.join(directory, f"{prefix}{counter}{suffix}")
        if candidate not in entries:
            return candidate


def _split_filename(filename: str) -> tuple[str, str]:
    if "." not in filename:
        return filename, ""
    stem, suffix = filename.rsplit(".", 1)
    return stem, f".{suffix}"


def _copy_content_type(
    target_ct_root: ET.Element,
    donor_ct_root: ET.Element,
    donor_part_name: str,
    target_part_name: str,
) -> None:
    donor_override = donor_ct_root.find(
        f".//{{{CT_NS}}}Override[@PartName='/{donor_part_name}']"
    )
    if donor_override is not None:
        existing = target_ct_root.find(f".//{{{CT_NS}}}Override[@PartName='/{target_part_name}']")
        if existing is None:
            override = ET.SubElement(target_ct_root, f"{{{CT_NS}}}Override")
            override.set("PartName", f"/{target_part_name}")
            override.set("ContentType", donor_override.get("ContentType", ""))
        return

    extension = target_part_name.rsplit(".", 1)[-1]
    existing_default = target_ct_root.find(f".//{{{CT_NS}}}Default[@Extension='{extension}']")
    if existing_default is not None:
        return
    donor_default = donor_ct_root.find(f".//{{{CT_NS}}}Default[@Extension='{extension}']")
    if donor_default is None:
        return
    default = ET.SubElement(target_ct_root, f"{{{CT_NS}}}Default")
    default.set("Extension", extension)
    default.set("ContentType", donor_default.get("ContentType", ""))


def _next_relationship_id(rels_root: ET.Element) -> str:
    max_id = 0
    for rel in rels_root.findall(f".//{{{PKG_REL_NS}}}Relationship"):
        rel_id = rel.get("Id", "")
        match = re.fullmatch(r"rId(\d+)", rel_id)
        if match:
            max_id = max(max_id, int(match.group(1)))
    return f"rId{max_id + 1}"
