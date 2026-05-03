#!/usr/bin/env python3
"""Build source-backed regional intelligence specs for May Sales Director decks.

The deck gate uses these specs to distinguish a technically linked deck from a
director-ready deck. The spec is intentionally small and auditable: original ETL
sidecar facts, current May/Q2 facts, ARR/ACV basis, and the first two-week action
contract.
"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from _directors import canonical_directors
from build_connected_factory_workbook import original_slug


ROOT = Path(__file__).resolve().parent.parent
CRM_OUTPUT = Path("/Users/test/crm-analytics/output")
DEFAULT_PERIOD = "2026-Q2"
ORIGINAL_SIDECAR_DIR = CRM_OUTPUT / "simcorp_director_decks/2026-04-20/land-only"
GOLD_ANALYTICS_DIR = CRM_OUTPUT / "director_gold_analytics/2026-04-23"
ETL_AUDIT_DIR = CRM_OUTPUT / "etl_intelligence_audit/2026-04-23"


@dataclass(frozen=True)
class DirectorContext:
    name: str
    territory: str
    slug: str
    original_slug: str
    director_dir: Path


def slugify(value: str) -> str:
    return value.replace(" ", "-")


def _fmt_meur(value: Any) -> str:
    try:
        return f"EUR {float(value or 0) / 1_000_000:.1f}M"
    except (TypeError, ValueError):
        return "EUR 0.0M"


def _json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _sheet_rows(path: Path, sheet: str) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=False)
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    output: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not any(value not in (None, "") for value in row):
            continue
        output.append({headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))})
    return output


def _raw_original_intel(path: Path) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for row in _sheet_rows(path, "Raw_Original_Intel"):
        key = str(row.get("Key") or "").strip()
        if key:
            values[key] = row.get("Value")
    return values


def _top_actions(path: Path) -> list[dict[str, Any]]:
    rows = _sheet_rows(path, "Raw_Current_Action_Items")
    actions: list[dict[str, Any]] = []
    for row in rows[:6]:
        claim = str(row.get("Claim") or "").strip()
        action = str(row.get("Suggested action") or "").strip()
        if not claim and not action:
            continue
        actions.append(
            {
                "priority": str(row.get("Priority") or "").strip(),
                "claim": claim,
                "suggested_action": action,
                "owner": str(row.get("Owner") or "").strip() or "Director",
                "due_date": "2026-05-15",
                "source": "Raw_Current_Action_Items",
            }
        )
    return actions


def _top_q2_deals(path: Path) -> list[dict[str, Any]]:
    rows = _sheet_rows(path, "Raw_Current_Q2_Readiness")
    deals: list[dict[str, Any]] = []
    for row in rows[:5]:
        account = str(row.get("Account") or "").strip()
        opportunity = str(row.get("Opportunity") or "").strip()
        if not account and not opportunity:
            continue
        deals.append(
            {
                "account": account,
                "opportunity": opportunity,
                "owner": str(row.get("Owner") or "").strip(),
                "stage": str(row.get("Stage") or "").strip(),
                "forecast": str(row.get("Forecast") or "").strip(),
                "close_date": str(row.get("Close Date") or "").strip(),
                "arr_eur": row.get("ARR"),
                "readiness": str(row.get("Readiness") or "").strip(),
            }
        )
    return deals


def _current_kpis(trends: dict[str, Any]) -> dict[str, Any]:
    kpis = {str(item.get("name")): item.get("value") for item in trends.get("kpis", []) if isinstance(item, dict)}
    return {
        "q2_closeable_land_expand_unweighted_arr_eur": kpis.get("total_pipeline_arr", 0),
        "q2_renewal_acv_eur": kpis.get("total_renewal_acv", 0),
        "beyond_q2_land_expand_unweighted_arr_eur": kpis.get("pipeline_arr_beyond_cfq", 0),
    }


def _gold_context(path: Path) -> dict[str, Any]:
    data = _json(path)
    analytics = data.get("analytics", {})
    concentration = analytics.get("pipeline_concentration", {}) if isinstance(analytics, dict) else {}
    bands = concentration.get("bands", []) if isinstance(concentration, dict) else []
    largest = concentration.get("largest_deals", []) if isinstance(concentration, dict) else []
    return {
        "source": str(path),
        "total_open_arr_eur": concentration.get("total_open_arr"),
        "top_5_share_pct": next((band.get("pct_of_open_pipeline") for band in bands if band.get("top_n") == 5), None),
        "top_10_share_pct": next((band.get("pct_of_open_pipeline") for band in bands if band.get("top_n") == 10), None),
        "largest_deals": largest[:5],
    }


def _build_one(period: str, director: dict[str, Any]) -> tuple[Path, Path]:
    name = str(director["name"])
    slug = slugify(name)
    ctx = DirectorContext(
        name=name,
        territory=str(director["scope_label"]),
        slug=slug,
        original_slug=original_slug(name),
        director_dir=ROOT / "state" / period / slug,
    )
    connected = ctx.director_dir / "factory/connected/connected_factory.xlsx"
    trends_path = ctx.director_dir / "trends.json"
    sidecar_path = ORIGINAL_SIDECAR_DIR / f"{ctx.original_slug}-LAND.json"
    gold_path = GOLD_ANALYTICS_DIR / ctx.original_slug / "gold_analytics.json"
    audit_path = ETL_AUDIT_DIR / ctx.original_slug / "etl_intelligence_audit.json"

    original = _raw_original_intel(connected)
    sidecar = _json(sidecar_path)
    trends = _json(trends_path)
    current = _current_kpis(trends)
    actions = _top_actions(connected)
    q2_deals = _top_q2_deals(connected)
    gold = _gold_context(gold_path)
    audit = _json(audit_path)
    audit_summary = audit.get("summary", {}) if isinstance(audit.get("summary"), dict) else {}
    readiness_checks = {
        "connected_factory_exists": connected.exists(),
        "trends_exists": trends_path.exists(),
        "has_original_sidecar": sidecar_path.exists(),
        "has_original_intelligence_rows": bool(original),
        "has_gold_or_audit_context": gold_path.exists() or audit_path.exists(),
        "has_first_two_week_actions": len(actions) >= 3,
        "has_named_q2_deals": len(q2_deals) >= 1,
        "arr_acv_separated": True,
        "no_generic_apac_label_for_non_apac": ctx.territory == "APAC" or "APAC" not in json.dumps(
            {
                "original_intelligence": original,
                "actions": actions,
                "q2_deals": q2_deals,
            },
            default=str,
        ),
    }
    source_gaps = sorted(key for key, value in readiness_checks.items() if value is not True)

    spec = {
        "schema": "regional-intelligence-spec/v1",
        "period": period,
        "orientation_date": "2026-05-01",
        "director": ctx.name,
        "territory": ctx.territory,
        "status": "pass" if not source_gaps else "fail",
        "source_gaps": source_gaps,
        "source_coverage": {
            "connected_factory": str(connected),
            "original_etl_sidecar": str(sidecar_path),
            "gold_analytics": str(gold_path) if gold_path.exists() else "",
            "etl_intelligence_audit": str(audit_path) if audit_path.exists() else "",
            "connected_factory_exists": connected.exists(),
            "original_sidecar_exists": sidecar_path.exists(),
            "gold_analytics_exists": gold_path.exists(),
            "etl_audit_exists": audit_path.exists(),
            "etl_coverage_gap_count": audit_summary.get("coverage_gap_count"),
        },
        "metric_contract": {
            "arr_basis": "Land+Expand unweighted ARR in EUR unless explicitly labeled weighted.",
            "renewal_basis": "Renewal ACV in EUR, never blended with ARR.",
            "headline_currency": "EUR",
            "arr_acv_separated": True,
        },
        "original_intelligence": {
            "scope_original": original.get("scope_original"),
            "open_land_deals": original.get("open_land_deals_sidecar", sidecar.get("open_land_deals")),
            "open_land_arr_eur": original.get("open_land_arr_eur_sidecar", sidecar.get("open_land_arr")),
            "open_land_weighted_arr_eur": original.get("open_land_arr_wtd_eur_sidecar", sidecar.get("open_land_arr_wtd")),
            "q1_land_wins": original.get("q1_land_wins", sidecar.get("q1_land_wins")),
            "q1_land_won_arr_eur": original.get("q1_land_wins_arr_eur", sidecar.get("q1_land_wins_arr")),
            "q1_land_losses": original.get("q1_land_lost", sidecar.get("q1_land_lost")),
            "q1_land_lost_arr_eur": original.get("q1_land_lost_arr_eur", sidecar.get("q1_land_lost_arr")),
            "q2_q3_renewal_acv_eur": float(original.get("q2_renewals_acv_eur", 0) or 0)
            + float(original.get("q3_renewals_acv_eur", 0) or 0),
            "approved_2026": original.get("approved_2026", sidecar.get("approved_2026")),
            "missing_stage3": original.get("missing_stage3", sidecar.get("missing_stage3")),
        },
        "current_may_orientation": current,
        "gold_context": gold,
        "top_q2_deals": q2_deals,
        "first_two_week_actions": actions,
        "deck_injection_targets": [
            "S02 May operating summary",
            "S07 Q2 readiness table",
            "S09 Commercial approval table",
            "S24 May plan table",
            "S26 Action items table",
            "S27 Risks and outlook",
        ],
        "readiness_checks": readiness_checks,
    }

    out_json = ctx.director_dir / "factory/regional_intelligence_spec.json"
    out_md = ctx.director_dir / "factory/regional_intelligence_spec.md"
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(spec, indent=2, default=str) + "\n", encoding="utf-8")
    lines = [
        f"# {ctx.name} - May 2026 Regional Intelligence Spec",
        "",
        f"- Status: {spec['status']}",
        f"- Territory: {ctx.territory}",
        "- Basis: Land+Expand unweighted ARR; Renewal ACV separate.",
        f"- Original open Land: {spec['original_intelligence']['open_land_deals']} / {_fmt_meur(spec['original_intelligence']['open_land_arr_eur'])}",
        f"- Q1 Land wins/losses: {spec['original_intelligence']['q1_land_wins']} / {spec['original_intelligence']['q1_land_losses']}",
        f"- Current Q2 closeable L+E ARR: {_fmt_meur(current['q2_closeable_land_expand_unweighted_arr_eur'])}",
        f"- Current Q2 Renewal ACV: {_fmt_meur(current['q2_renewal_acv_eur'])}",
        "",
        "## First Two-Week Actions",
        "",
        "| Priority | Claim | May 1-15 action |",
        "|---|---|---|",
    ]
    for action in actions:
        lines.append(
            f"| {action['priority']} | {action['claim']} | {action['suggested_action']} |"
        )
    if source_gaps:
        lines.extend(["", "## Source Gaps", ""])
        lines.extend(f"- {gap}" for gap in source_gaps)
    out_md.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return out_json, out_md


def _selected(args: argparse.Namespace) -> list[dict[str, Any]]:
    directors = canonical_directors()
    if not args.director_slug:
        return directors
    return [director for director in directors if slugify(str(director["name"])) == args.director_slug]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--period", default=DEFAULT_PERIOD)
    parser.add_argument("--director-slug")
    args = parser.parse_args()

    outputs: list[tuple[Path, Path]] = []
    for director in _selected(args):
        outputs.append(_build_one(args.period, director))
    for json_path, md_path in outputs:
        print(f"{json_path}")
        print(f"{md_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
