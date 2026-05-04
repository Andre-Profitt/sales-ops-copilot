"""Tests for binding-level render verification.

The verifier loads:
  - rendered .pptx
  - registry
  - render_evidence_manifest.json (from build_ppttc)
And produces a per-binding pass/fail report. Every required binding
must have evidence on the corresponding slide.
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

REPO = Path(__file__).resolve().parents[1]
PY = sys.executable
SCRIPT = REPO / "scripts" / "verify_render_bindings.py"


def _make_rendered(path: Path, slide_texts: dict[int, list[str]]) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    max_slide = max(slide_texts) if slide_texts else 0
    for i in range(1, max_slide + 1):
        slide = prs.slides.add_slide(blank)
        for j, t in enumerate(slide_texts.get(i, [])):
            tb = slide.shapes.add_textbox(Inches(1), Inches(1 + j * 0.5), Inches(5), Inches(0.4))
            tb.text_frame.text = t
    prs.save(str(path))


def test_required_text_present(tmp_path: Path) -> None:
    rendered = tmp_path / "rendered.pptx"
    _make_rendered(rendered, {1: ["Patrick Gaughan"], 5: ["Late-stage coverage is thin"]})

    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "bindings": [
                    {
                        "name": "S01_DirectorName",
                        "kind": "text",
                        "lane": "ppttc_text",
                        "required": True,
                        "status": "bound",
                        "expected_text": "Patrick Gaughan",
                        "slide": 1,
                    },
                    {
                        "name": "S05_Title",
                        "kind": "text",
                        "lane": "ppttc_text",
                        "required": True,
                        "status": "bound",
                        "expected_text": "Late-stage coverage is thin",
                        "slide": 5,
                    },
                ]
            }
        )
    )
    out = tmp_path / "report.json"

    proc = subprocess.run(
        [PY, str(SCRIPT), "--pptx", str(rendered), "--manifest", str(manifest), "--out", str(out)],
        capture_output=True,
        text=True,
        cwd=REPO,
    )
    assert proc.returncode == 0, proc.stdout + proc.stderr
    report = json.loads(out.read_text())
    assert report["status"] == "pass"
    assert all(b["status"] == "pass" for b in report["bindings"] if b["required"])


def test_missing_required_text_fails(tmp_path: Path) -> None:
    rendered = tmp_path / "rendered.pptx"
    _make_rendered(rendered, {1: ["Patrick Gaughan"]})  # S05 title missing

    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "bindings": [
                    {
                        "name": "S01_DirectorName",
                        "kind": "text",
                        "lane": "ppttc_text",
                        "required": True,
                        "status": "bound",
                        "expected_text": "Patrick Gaughan",
                        "slide": 1,
                    },
                    {
                        "name": "S05_Title",
                        "kind": "text",
                        "lane": "ppttc_text",
                        "required": True,
                        "status": "bound",
                        "expected_text": "Late-stage coverage is thin",
                        "slide": 5,
                    },
                ]
            }
        )
    )
    out = tmp_path / "report.json"

    proc = subprocess.run(
        [PY, str(SCRIPT), "--pptx", str(rendered), "--manifest", str(manifest), "--out", str(out)],
        capture_output=True,
        text=True,
        cwd=REPO,
    )
    assert proc.returncode != 0
    report = json.loads(out.read_text())
    assert report["status"] == "fail"
    failed = [b for b in report["bindings"] if b["status"] == "fail"]
    assert any(b["name"] == "S05_Title" for b in failed)
