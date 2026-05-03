from __future__ import annotations

import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from scrub_stale_thinkcell_metadata import scrub_deck  # noqa: E402


def test_scrub_removes_stale_thinkcell_named_shape(tmp_path: Path) -> None:
    deck = tmp_path / "stale.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    stale = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(0.4))
    stale.name = "think-cell data - do not delete"
    keeper = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(2), Inches(0.4))
    keeper.text_frame.text = "Keep me"
    prs.save(deck)

    result = scrub_deck(deck)

    assert result.status == "pass"
    assert result.removed_data_shapes == 1
    assert result.residual_stale_tokens == {}
    with ZipFile(deck) as zf:
        slide_xml = zf.read("ppt/slides/slide1.xml").decode("utf-8")
    assert "think-cell data - do not delete" not in slide_xml
    assert "Keep me" in slide_xml


def test_scrub_preserves_valid_presentation_level_relationships(tmp_path: Path) -> None:
    deck = tmp_path / "presentation-rels.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    stale = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(0.4))
    stale.name = "think-cell data - do not delete"
    prs.save(deck)

    with ZipFile(deck) as zin:
        entries = {item.filename: zin.read(item.filename) for item in zin.infolist()}
    rels = entries["ppt/_rels/presentation.xml.rels"].decode("utf-8")
    rels = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdPresProps" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/presProps" '
        'Target="presProps.xml"/></Relationships>',
    )
    entries["ppt/_rels/presentation.xml.rels"] = rels.encode("utf-8")
    entries["ppt/presProps.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:presentationPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>'
    ).encode("utf-8")
    with ZipFile(deck, "w", ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            zout.writestr(name, data)

    result = scrub_deck(deck)

    assert result.status == "pass"
    with ZipFile(deck) as zf:
        assert "ppt/presProps.xml" in zf.namelist()
        rels_after = zf.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
    assert "presProps.xml" in rels_after
