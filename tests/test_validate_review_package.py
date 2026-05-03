"""Tests for the May 2026 meeting-spine review-package validator hardening.

Covers A-3 (case-insensitive forbidden-text scan) and the C-2 / A-4 expansion of
the FORBIDDEN_TEXT set, mirrored into the review-package validator so the same
pollution and Excel-error tokens fail closed regardless of casing.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

# scripts/ is not a package; load like the other top-level tests do.
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from validate_may_review_package import (  # noqa: E402  - sys.path mutation above
    FORBIDDEN_TEXT,
    _deck_check,
)


def _build_deck(path: Path, body_text: str) -> None:
    """Build a tiny valid PPTX so _deck_check can scan it."""

    prs = Presentation()
    layout = prs.slide_layouts[5]  # title-only layout (always present in the default master)
    slide = prs.slides.add_slide(layout)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.text = body_text
    prs.save(path)


# ---------- A-3 case-insensitive matching ------------------------------------


@pytest.mark.parametrize(
    "body_text, expected_canonical",
    [
        ("Account: sc test pty ltd", "SC Test"),
        ("Project owner: TEST ACCOUNT (US)", "Test Account"),
        ("masb pilot — TeST mAsB scope", "Test MASB"),
    ],
)
def test_deck_check_forbidden_text_is_case_insensitive(
    tmp_path: Path, body_text: str, expected_canonical: str
) -> None:
    """A-3 (P1): the meeting-spine validator must catch forbidden text regardless of case."""

    deck = tmp_path / "case-insensitive.pptx"
    _build_deck(deck, body_text)
    result = _deck_check(deck, required_text=set())
    assert expected_canonical in result.forbidden_text, (
        f"case-insensitive match should report canonical needle for {body_text!r}; "
        f"got {result.forbidden_text!r}"
    )


# ---------- C-2 / A-4 forbidden text expansion in the validator ---------------


def test_validator_forbidden_text_includes_expanded_error_tokens() -> None:
    needles_lower = {needle.lower() for needle in FORBIDDEN_TEXT}
    for token in ("#name", "#null", "#ref", "#value", "#div/0", "#n/a"):
        assert token in needles_lower, f"expected {token!r} in validator FORBIDDEN_TEXT"


def test_validator_forbidden_text_includes_pollution_patterns() -> None:
    needles_lower = {needle.lower() for needle in FORBIDDEN_TEXT}
    for needle in (
        "simcorp test",
        "demo account",
        "sample account",
        "sample co",
        "lorem ipsum",
    ):
        assert needle in needles_lower, f"expected {needle!r} in validator FORBIDDEN_TEXT"


def test_validator_forbidden_text_does_not_block_simcorp_brand_phrases() -> None:
    """Regression: every needle must be specific enough to not collide with footer text."""

    legit_footer = "© 2026 SimCorp. SimCorp Dimension. SimCorp Axioma."
    for needle in FORBIDDEN_TEXT:
        assert needle.lower() not in legit_footer.lower(), needle


def test_deck_check_expanded_error_tokens_caught_in_deck_text(tmp_path: Path) -> None:
    """C-2 mirror: deck text containing #DIV/0! must be reported as forbidden."""

    deck = tmp_path / "div0.pptx"
    _build_deck(deck, "Forecast variance: #DIV/0!")
    result = _deck_check(deck, required_text=set())
    joined = " ".join(result.forbidden_text).lower()
    assert "#div/0" in joined, f"expected '#DIV/0' to be forbidden, got {result.forbidden_text!r}"


def test_deck_check_pollution_phrase_caught_in_deck_text(tmp_path: Path) -> None:
    """A-1 conservative: 'SimCorp Test' deck text must surface the pollution needle."""

    deck = tmp_path / "pollution.pptx"
    _build_deck(deck, "Account name placeholder: SimCorp Test sandbox row")
    result = _deck_check(deck, required_text=set())
    joined = " ".join(result.forbidden_text).lower()
    assert "simcorp test" in joined, (
        f"expected 'SimCorp Test' to be flagged, got {result.forbidden_text!r}"
    )
