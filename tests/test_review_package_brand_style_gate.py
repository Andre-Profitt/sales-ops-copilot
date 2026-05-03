from __future__ import annotations

import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from run_review_package_brand_style_gate import inspect_deck  # noqa: E402


def _build_deck(
    path: Path,
    *,
    font_size: float = 6.0,
    font_name: str = "Aptos",
    rounded: bool = False,
    stale_thinkcell_name: bool = False,
) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if rounded:
        slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(0.5), Inches(1.0), Inches(0.5))
    if stale_thinkcell_name:
        stale = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(1.0), Inches(0.3))
        stale.name = "think-cell data - do not delete"
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.5))
    paragraph = box.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Operating review"
    run.font.name = font_name
    run.font.size = Pt(font_size)
    prs.save(path)


def test_brand_style_gate_passes_square_approved_type_scale(tmp_path: Path) -> None:
    deck = tmp_path / "approved.pptx"
    _build_deck(deck)

    result = inspect_deck(deck)
    assert result.status == "pass"
    assert result.rounded_geometry_count == 0
    assert result.findings == []


def test_brand_style_gate_blocks_rounded_geometry(tmp_path: Path) -> None:
    deck = tmp_path / "rounded.pptx"
    _build_deck(deck, rounded=True)

    result = inspect_deck(deck)
    assert result.status == "fail"
    assert result.rounded_geometry_count > 0
    assert any("rounded/curved" in finding for finding in result.findings)


def test_brand_style_gate_blocks_tiny_text(tmp_path: Path) -> None:
    deck = tmp_path / "tiny.pptx"
    _build_deck(deck, font_size=5.0)

    result = inspect_deck(deck)
    assert result.status == "fail"
    assert any("below 5.5pt" in finding for finding in result.findings)


def test_brand_style_gate_blocks_off_scale_text_sizes(tmp_path: Path) -> None:
    deck = tmp_path / "off-scale.pptx"
    _build_deck(deck, font_size=6.3)

    result = inspect_deck(deck)
    assert result.status == "fail"
    assert any("off-scale" in finding for finding in result.findings)


def test_brand_style_gate_blocks_non_brand_fonts(tmp_path: Path) -> None:
    deck = tmp_path / "font.pptx"
    _build_deck(deck, font_name="Comic Sans MS")

    result = inspect_deck(deck)
    assert result.status == "fail"
    assert any("non-brand" in finding for finding in result.findings)


def test_brand_style_gate_blocks_stale_thinkcell_ownership(tmp_path: Path) -> None:
    deck = tmp_path / "stale-tc.pptx"
    _build_deck(deck, stale_thinkcell_name=True)

    result = inspect_deck(deck)
    assert result.status == "fail"
    assert result.stale_thinkcell_ownership_count > 0
    assert any("stale think-cell" in finding for finding in result.findings)


def test_brand_style_gate_blocks_missing_powerpoint_support_parts(tmp_path: Path) -> None:
    deck = tmp_path / "missing-support.pptx"
    _build_deck(deck)

    with ZipFile(deck) as zin:
        entries = {item.filename: zin.read(item.filename) for item in zin.infolist()}
    entries.pop("ppt/presProps.xml", None)
    rels = entries["ppt/_rels/presentation.xml.rels"].decode("utf-8")
    rels = rels.replace(
        '<Relationship Id="rId4" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/presProps" '
        'Target="presProps.xml"/>',
        "",
    )
    entries["ppt/_rels/presentation.xml.rels"] = rels.encode("utf-8")
    with ZipFile(deck, "w", ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            zout.writestr(name, data)

    result = inspect_deck(deck)
    assert result.status == "fail"
    assert "ppt/presProps.xml" in result.missing_package_support
    assert any("missing PowerPoint support" in finding for finding in result.findings)
