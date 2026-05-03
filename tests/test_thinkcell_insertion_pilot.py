"""Tests for scripts/run_thinkcell_insertion_pilot.py.

The CLI runs in two modes, both forbidden from invoking Office/PowerPoint, the
Windows VM bridge, or SharePoint:

- ``--plan-only``: reads the L5 build scaffold + per-contract proof JSON,
  encodes per-director eligibility, runs deterministic preflight (ppttc JSON
  parse, bound deck zip readability, optional base-deck), emits a manifest +
  ``vm_insertion_commands.sh`` for an operator to run on the VM bridge.
- ``--promote``: reads the existing run manifest, looks for actual VM-produced
  candidate files at the intended output paths, validates them (zip readable,
  slides present, no native PPT tables), and updates statuses to
  ``candidate_created`` / ``awaiting_vm_insertion`` / ``candidate_invalid``.
  Always preserves ``publishable: false``.

These tests pin down preflight + planning + promote logic against tmp_path
fixtures so no real state under state/ is touched.
"""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches  # noqa: F401  - re-exported for fixture helpers below

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "scripts"))


# --------------------------------------------------------------------------
# PPTX fixture helpers (real zips — preflight checks zip readability)
# --------------------------------------------------------------------------


def _real_minimal_pptx_bytes() -> bytes:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    prs.slides.add_slide(blank)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _real_pptx_with_native_table_bytes() -> bytes:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_table(
        rows=2, cols=2, left=Inches(1), top=Inches(1), width=Inches(4), height=Inches(1)
    )
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _corrupt_pptx_bytes() -> bytes:
    return b"PK\x03\x04not-a-real-zip-payload"


# --------------------------------------------------------------------------
# Fixture builders
# --------------------------------------------------------------------------


def _qtr_contract(
    *,
    name: str,
    family: str,
    readiness: str = "l5_proven",
    proof_status: str = "pass",
    ready_directors: list[str] | None = None,
    fallback_directors: list[str] | None = None,
    fallback: str = "table fallback",
    supported_lane: str = "native think-cell chart via named seed + ppttc",
) -> dict:
    return {
        "name": name,
        "family": family,
        "readiness": readiness,
        "proof_status": proof_status,
        "salesforce_fit": {
            "gate": family,
            "ready_directors": ready_directors or [],
            "fallback_directors": fallback_directors or [],
            "rows": [],
        },
        "fallback": fallback,
        "supported_lane": supported_lane,
    }


def _build_scaffold_dict() -> dict:
    return {
        "schema": "thinkcell-build-scaffold/v1",
        "period": "2026-Q2",
        "contracts": [
            _qtr_contract(
                name="QTR04_DealRisk_Scatter",
                family="scatter_bubble",
                ready_directors=[
                    "Megan Miceli",
                    "Jesper Tyrer",
                    "Sarah Pittroff",
                    "Francois Thaury",
                    "Dan Peppett",
                    "Christian Ebbesen",
                    "Mourad",
                    "Adam Steinhouse",
                ],
                fallback_directors=["Patrick Gaughan"],
                fallback="ranked risk table when probability or ARR has no spread",
                supported_lane="stock-donor scatter via patched name + ppttc",
            ),
            _qtr_contract(
                name="QTR05_FY26RenewalTimeline_Gantt",
                family="timeline_gantt_fy26_renewals",
                ready_directors=[
                    "Megan Miceli",
                    "Patrick Gaughan",
                    "Jesper Tyrer",
                    "Sarah Pittroff",
                    "Francois Thaury",
                    "Dan Peppett",
                    "Christian Ebbesen",
                    "Mourad",
                ],
                fallback_directors=["Adam Steinhouse"],
                fallback="renewal watchlist table",
                supported_lane="stock-donor Gantt via patched name + ppttc",
            ),
            _qtr_contract(
                name="QTR06_Q2RenewalTimeline_Gantt",
                family="timeline_gantt_fy26_renewals",
                ready_directors=[
                    "Sarah Pittroff",
                    "Dan Peppett",
                    "Christian Ebbesen",
                ],
                fallback_directors=[
                    "Adam Steinhouse",
                    "Francois Thaury",
                    "Jesper Tyrer",
                    "Megan Miceli",
                    "Mourad",
                    "Patrick Gaughan",
                ],
                fallback="renewal watchlist table",
                supported_lane="stock-donor Gantt via patched name + ppttc",
            ),
            _qtr_contract(
                name="QTR99_NotProven_Bar",
                family="bar_column",
                readiness="l4_payload",
                proof_status="not_run",
                ready_directors=["Jesper Tyrer"],
            ),
        ],
    }


def _write_proof_artifacts(repo: Path, contract_id: str, director_slug: str) -> dict:
    """Create a fake proof.json + bound pptx + ppttc + seed pptx pair on disk.

    The bound + seed PPTX bytes are produced via python-pptx so they pass deep
    preflight (zipfile readability, slide count > 0). The ppttc payload is
    valid JSON so JSON parse passes.
    """
    work = repo / "state" / "thinkcell_bridge" / "build_scaffold" / "2026-Q2" / "work" / contract_id
    work.mkdir(parents=True, exist_ok=True)
    bound = work / f"{contract_id}-stock-donor-2026-Q2-bound.pptx"
    bound.write_bytes(_real_minimal_pptx_bytes())
    seed = work / f"{contract_id}-stock-donor-seed.pptx"
    seed.write_bytes(_real_minimal_pptx_bytes())
    ppttc = work / f"{contract_id}-stock-donor-2026-Q2.ppttc"
    # Real ppttc payload — array of one bind entry that the VM bridge consumes.
    ppttc.write_text(
        json.dumps(
            [
                {
                    "template": str(seed),
                    "data": [{"name": "chart01", "table": [["x", "y"], [1, 2]]}],
                }
            ]
        )
    )
    render_dir = work / "rendered_stock_donor"
    render_dir.mkdir(exist_ok=True)
    proof = {
        "schema": "thinkcell-stock-donor-proof/v1",
        "status": "pass",
        "contract": contract_id,
        "period": "2026-Q2",
        "director_slug": director_slug,
        "deck": str(bound),
        "seed_pptx": str(seed),
        "ppttc": str(ppttc),
        "render_dir": str(render_dir),
        "donor_template": "/Library/Application Support/Microsoft/think-cell/templates/fake.potx",
        "workbook": str(work / "fake_workbook.xlsx"),
        "targets": [
            {
                "name": contract_id,
                "bound_required_terms": ["Bound Term", contract_id],
            }
        ],
    }
    proof_path = work / f"{contract_id}-proof.json"
    proof_path.write_text(json.dumps(proof))
    return proof


@pytest.fixture()
def fixture_repo(tmp_path: Path) -> Path:
    """A pretend repo containing scaffold + proof JSON for QTR04 + QTR05 (Jesper)."""
    repo = tmp_path / "repo"
    state = repo / "state" / "thinkcell_bridge" / "build_scaffold" / "2026-Q2"
    state.mkdir(parents=True)
    (state / "thinkcell_build_scaffold.json").write_text(json.dumps(_build_scaffold_dict()))
    _write_proof_artifacts(repo, "QTR04_DealRisk_Scatter", "Jesper-Tyrer")
    _write_proof_artifacts(repo, "QTR05_FY26RenewalTimeline_Gantt", "Jesper-Tyrer")
    return repo


# --------------------------------------------------------------------------
# Pure helper tests
# --------------------------------------------------------------------------


def test_slug_round_trip_for_two_word_name():
    import run_thinkcell_insertion_pilot as mod

    assert mod.slug_for_director("Jesper Tyrer") == "Jesper-Tyrer"
    assert mod.name_for_slug("Jesper-Tyrer") == "Jesper Tyrer"


def test_slug_round_trip_for_single_word_name():
    import run_thinkcell_insertion_pilot as mod

    assert mod.slug_for_director("Mourad") == "Mourad"
    assert mod.name_for_slug("Mourad") == "Mourad"


def test_is_l5_proven_pass_true_for_l5_pass():
    import run_thinkcell_insertion_pilot as mod

    contract = _qtr_contract(name="QTR01_StageMix_Bar", family="bar_column")
    assert mod.is_l5_proven_pass(contract) is True


def test_is_l5_proven_pass_false_when_proof_status_fails():
    import run_thinkcell_insertion_pilot as mod

    contract = _qtr_contract(name="QTRX", family="bar_column", proof_status="fail")
    assert mod.is_l5_proven_pass(contract) is False


def test_is_l5_proven_pass_false_when_readiness_below_l5():
    import run_thinkcell_insertion_pilot as mod

    contract = _qtr_contract(name="QTRX", family="bar_column", readiness="l4_payload")
    assert mod.is_l5_proven_pass(contract) is False


# --------------------------------------------------------------------------
# Eligibility tests (Patrick/QTR04, Adam/QTR05, QTR06 only Sarah/Dan/Christian)
# --------------------------------------------------------------------------


def test_eligibility_jesper_qtr04_ready():
    import run_thinkcell_insertion_pilot as mod

    scaffold = _build_scaffold_dict()
    qtr04 = mod.find_contract(scaffold, "QTR04_DealRisk_Scatter")
    decision = mod.eligibility_for(qtr04, "Jesper Tyrer")
    assert decision["decision"] == "ready"


def test_eligibility_patrick_qtr04_fallback():
    import run_thinkcell_insertion_pilot as mod

    scaffold = _build_scaffold_dict()
    qtr04 = mod.find_contract(scaffold, "QTR04_DealRisk_Scatter")
    decision = mod.eligibility_for(qtr04, "Patrick Gaughan")
    assert decision["decision"] == "fallback"
    assert "ranked risk table" in decision["fallback_lane"]


def test_eligibility_adam_qtr05_fallback():
    import run_thinkcell_insertion_pilot as mod

    scaffold = _build_scaffold_dict()
    qtr05 = mod.find_contract(scaffold, "QTR05_FY26RenewalTimeline_Gantt")
    decision = mod.eligibility_for(qtr05, "Adam Steinhouse")
    assert decision["decision"] == "fallback"
    assert "renewal watchlist" in decision["fallback_lane"].lower()


def test_eligibility_jesper_qtr05_ready():
    import run_thinkcell_insertion_pilot as mod

    scaffold = _build_scaffold_dict()
    qtr05 = mod.find_contract(scaffold, "QTR05_FY26RenewalTimeline_Gantt")
    decision = mod.eligibility_for(qtr05, "Jesper Tyrer")
    assert decision["decision"] == "ready"


@pytest.mark.parametrize("director", ["Sarah Pittroff", "Dan Peppett", "Christian Ebbesen"])
def test_eligibility_qtr06_ready_for_renewal_heavy_directors(director):
    import run_thinkcell_insertion_pilot as mod

    scaffold = _build_scaffold_dict()
    qtr06 = mod.find_contract(scaffold, "QTR06_Q2RenewalTimeline_Gantt")
    decision = mod.eligibility_for(qtr06, director)
    assert decision["decision"] == "ready"


@pytest.mark.parametrize(
    "director",
    [
        "Adam Steinhouse",
        "Francois Thaury",
        "Jesper Tyrer",
        "Megan Miceli",
        "Mourad",
        "Patrick Gaughan",
    ],
)
def test_eligibility_qtr06_fallback_for_non_renewal_directors(director):
    import run_thinkcell_insertion_pilot as mod

    scaffold = _build_scaffold_dict()
    qtr06 = mod.find_contract(scaffold, "QTR06_Q2RenewalTimeline_Gantt")
    decision = mod.eligibility_for(qtr06, director)
    assert decision["decision"] == "fallback"


def test_eligibility_unknown_director_returns_unknown():
    import run_thinkcell_insertion_pilot as mod

    scaffold = _build_scaffold_dict()
    qtr04 = mod.find_contract(scaffold, "QTR04_DealRisk_Scatter")
    decision = mod.eligibility_for(qtr04, "Nobody Real")
    assert decision["decision"] == "unknown"


def test_eligibility_baked_rules_match_scaffold_for_pilot_contracts():
    """Defensive: the script encodes the three exception rules from HANDOFF.

    Patrick → QTR04 fallback. Adam → QTR05 fallback. QTR06 only Sarah/Dan/Christian.
    The script must surface a discrepancy if the scaffold drifts.
    """
    import run_thinkcell_insertion_pilot as mod

    scaffold = _build_scaffold_dict()
    discrepancies = mod.check_baked_eligibility_rules(scaffold)
    assert discrepancies == []


# --------------------------------------------------------------------------
# Proof artifact tests
# --------------------------------------------------------------------------


def test_proof_artifacts_loaded_when_present(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    proof = mod.load_proof_for_contract(fixture_repo, "2026-Q2", "QTR04_DealRisk_Scatter")
    assert proof["status"] == "pass"
    assert Path(proof["deck"]).exists()


def test_proof_artifacts_missing_returns_none(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    proof = mod.load_proof_for_contract(fixture_repo, "2026-Q2", "QTR99_NotProven_Bar")
    assert proof is None


def test_intended_candidate_output_path_under_pilot_state_dir(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    candidate = mod.intended_candidate_output_path(
        fixture_repo,
        "2026-Q2",
        "Jesper-Tyrer",
        "QTR04_DealRisk_Scatter",
        run_id="20260501-2125Z",
    )
    rel = candidate.relative_to(fixture_repo)
    assert rel.parts[:5] == (
        "state",
        "2026-Q2",
        "__regional__",
        "thinkcell_insertion_pilot",
        "20260501-2125Z",
    )
    assert candidate.suffix == ".pptx"


# --------------------------------------------------------------------------
# Plan assembly tests
# --------------------------------------------------------------------------


def test_plan_pilot_jesper_qtr04_qtr05_marks_ready_for_manual(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=[
            "QTR04_DealRisk_Scatter",
            "QTR05_FY26RenewalTimeline_Gantt",
        ],
        run_id="20260501-2125Z",
    )
    assert manifest["overall_status"] == "ready_for_vm_or_manual_insertion"
    assert manifest["publishable"] is False
    assert manifest["plan_only"] is True
    assert manifest["director_slug"] == "Jesper-Tyrer"
    assert manifest["director_name"] == "Jesper Tyrer"

    statuses = {c["contract_id"]: c["status"] for c in manifest["contracts"]}
    assert statuses == {
        "QTR04_DealRisk_Scatter": "ready_for_vm_or_manual_insertion",
        "QTR05_FY26RenewalTimeline_Gantt": "ready_for_vm_or_manual_insertion",
    }


def test_plan_pilot_lists_exact_proof_artifact_paths(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z",
    )
    qtr04 = manifest["contracts"][0]
    artifacts = qtr04["source_artifacts"]
    for key in ("proof_json", "bound_deck", "ppttc", "seed_pptx", "render_dir"):
        assert key in artifacts, key
        assert Path(artifacts[key]).exists(), f"{key} should exist on disk"
    assert "intended_candidate_output_path" in qtr04
    assert qtr04["intended_candidate_output_path"].endswith(".pptx")


def test_plan_pilot_patrick_qtr04_blocks_with_fallback(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    # Patrick's QTR04 proof JSON does not exist, but eligibility blocks first.
    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Patrick-Gaughan",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z",
    )
    assert manifest["overall_status"] == "blocked"
    qtr04 = manifest["contracts"][0]
    assert qtr04["status"] == "not_eligible_director_fallback"
    assert qtr04["eligibility"]["decision"] == "fallback"
    assert "ranked risk table" in qtr04["eligibility"]["fallback_lane"]
    assert qtr04["status"] != "publishable"


def test_plan_pilot_adam_qtr05_blocks_with_fallback(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Adam-Steinhouse",
        contract_ids=["QTR05_FY26RenewalTimeline_Gantt"],
        run_id="20260501-2125Z",
    )
    assert manifest["overall_status"] == "blocked"
    assert manifest["contracts"][0]["status"] == "not_eligible_director_fallback"


def test_plan_pilot_jesper_qtr06_blocks_because_jesper_not_renewal_heavy(
    fixture_repo: Path,
):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR06_Q2RenewalTimeline_Gantt"],
        run_id="20260501-2125Z",
    )
    assert manifest["overall_status"] == "blocked"
    assert manifest["contracts"][0]["status"] == "not_eligible_director_fallback"


def test_plan_pilot_unknown_contract_blocks(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR99_DoesNotExist"],
        run_id="20260501-2125Z",
    )
    assert manifest["overall_status"] == "blocked"
    assert manifest["contracts"][0]["status"] == "unknown_contract"


def test_plan_pilot_unproven_contract_blocks(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR99_NotProven_Bar"],
        run_id="20260501-2125Z",
    )
    assert manifest["overall_status"] == "blocked"
    assert manifest["contracts"][0]["status"] == "not_l5_proven"


def test_plan_pilot_unknown_director_blocks(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Nobody-Real",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z",
    )
    assert manifest["overall_status"] == "blocked"
    # All contracts blocked for unknown director.
    assert manifest["contracts"][0]["status"] in {
        "unknown_director",
        "not_eligible_director_fallback",
    }


def test_plan_pilot_missing_proof_json_blocks(tmp_path: Path):
    """Scaffold says l5_proven/pass but proof JSON is absent on disk → fail closed."""
    import run_thinkcell_insertion_pilot as mod

    repo = tmp_path / "repo"
    state = repo / "state" / "thinkcell_bridge" / "build_scaffold" / "2026-Q2"
    state.mkdir(parents=True)
    (state / "thinkcell_build_scaffold.json").write_text(json.dumps(_build_scaffold_dict()))
    # Intentionally do NOT write any proof artifacts.

    manifest = mod.plan_pilot(
        repo_root=repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z",
    )
    assert manifest["overall_status"] == "blocked"
    assert manifest["contracts"][0]["status"] == "missing_proof_artifact"


def test_plan_pilot_never_emits_publishable_status(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=[
            "QTR04_DealRisk_Scatter",
            "QTR05_FY26RenewalTimeline_Gantt",
        ],
        run_id="20260501-2125Z",
    )
    assert manifest["publishable"] is False
    for contract in manifest["contracts"]:
        assert contract["status"] != "publishable"


def test_plan_pilot_lists_fallback_directors_in_eligibility_block(fixture_repo: Path):
    """Manifest must surface that Patrick is the fallback director for QTR04 even
    when planning for Jesper, so reviewers see the universality rule."""
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z",
    )
    qtr04 = manifest["contracts"][0]
    assert "Patrick Gaughan" in qtr04["eligibility"]["all_fallback_directors"]


# --------------------------------------------------------------------------
# Run-dir output tests
# --------------------------------------------------------------------------


def test_write_run_dir_emits_manifest_and_five_markdown_files(fixture_repo: Path, tmp_path: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=[
            "QTR04_DealRisk_Scatter",
            "QTR05_FY26RenewalTimeline_Gantt",
        ],
        run_id="20260501-2125Z",
    )
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    mod.write_run_dir(run_dir, manifest)

    expected = {
        "manifest.json",
        "RUN_LOG.md",
        "CHANGES.md",
        "VERIFICATION.md",
        "NEXT_FOR_CODEX_AUDIT.md",
        "PRODUCTION_LINE_STATUS.md",
    }
    assert expected.issubset({p.name for p in run_dir.iterdir()})

    parsed = json.loads((run_dir / "manifest.json").read_text())
    assert parsed["overall_status"] == "ready_for_vm_or_manual_insertion"


def test_production_line_status_includes_percent_complete_and_blockers(
    fixture_repo: Path, tmp_path: Path
):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Patrick-Gaughan",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z",
    )
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    mod.write_run_dir(run_dir, manifest)

    text = (run_dir / "PRODUCTION_LINE_STATUS.md").read_text()
    assert "%" in text  # percent-complete reported
    assert "Blocker" in text or "blocker" in text


# --------------------------------------------------------------------------
# CLI tests (via main())
# --------------------------------------------------------------------------


def test_cli_unsupported_period_fails_closed_argparse(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    with pytest.raises(SystemExit) as exc_info:
        mod.main(
            [
                "--period",
                "2025-Q1",
                "--director-slug",
                "Jesper-Tyrer",
                "--contracts",
                "QTR04_DealRisk_Scatter",
                "--plan-only",
                "--repo-root",
                str(fixture_repo),
                "--run-id",
                "20260501-2125Z",
            ]
        )
    assert exc_info.value.code != 0


def test_cli_writes_run_dir_and_returns_zero_for_jesper(
    fixture_repo: Path, capsys: pytest.CaptureFixture
):
    import run_thinkcell_insertion_pilot as mod

    rc = mod.main(
        [
            "--period",
            "2026-Q2",
            "--director-slug",
            "Jesper-Tyrer",
            "--contracts",
            "QTR04_DealRisk_Scatter",
            "QTR05_FY26RenewalTimeline_Gantt",
            "--plan-only",
            "--repo-root",
            str(fixture_repo),
            "--run-id",
            "20260501-2125Z",
        ]
    )
    assert rc == 0

    run_dir = (
        fixture_repo
        / "state"
        / "2026-Q2"
        / "__regional__"
        / "thinkcell_insertion_pilot"
        / "20260501-2125Z"
    )
    assert (run_dir / "manifest.json").exists()
    assert (run_dir / "RUN_LOG.md").exists()
    assert (run_dir / "PRODUCTION_LINE_STATUS.md").exists()
    parsed = json.loads((run_dir / "manifest.json").read_text())
    assert parsed["overall_status"] == "ready_for_vm_or_manual_insertion"


def test_cli_blocked_director_returns_nonzero(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    rc = mod.main(
        [
            "--period",
            "2026-Q2",
            "--director-slug",
            "Patrick-Gaughan",
            "--contracts",
            "QTR04_DealRisk_Scatter",
            "--plan-only",
            "--repo-root",
            str(fixture_repo),
            "--run-id",
            "20260501-2125Z-Patrick",
        ]
    )
    assert rc != 0


def test_cli_blocked_run_still_writes_manifest_for_forensics(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    mod.main(
        [
            "--period",
            "2026-Q2",
            "--director-slug",
            "Patrick-Gaughan",
            "--contracts",
            "QTR04_DealRisk_Scatter",
            "--plan-only",
            "--repo-root",
            str(fixture_repo),
            "--run-id",
            "20260501-2125Z-Patrick2",
        ]
    )
    run_dir = (
        fixture_repo
        / "state"
        / "2026-Q2"
        / "__regional__"
        / "thinkcell_insertion_pilot"
        / "20260501-2125Z-Patrick2"
    )
    parsed = json.loads((run_dir / "manifest.json").read_text())
    assert parsed["overall_status"] == "blocked"


def test_cli_requires_plan_only_or_promote_flag(fixture_repo: Path):
    """Either --plan-only or --promote is required; this CLI never auto-runs Office."""
    import run_thinkcell_insertion_pilot as mod

    with pytest.raises(SystemExit) as exc_info:
        mod.main(
            [
                "--period",
                "2026-Q2",
                "--director-slug",
                "Jesper-Tyrer",
                "--contracts",
                "QTR04_DealRisk_Scatter",
                "--repo-root",
                str(fixture_repo),
                "--run-id",
                "20260501-2125Z-noflag",
            ]
        )
    assert exc_info.value.code != 0


def test_cli_plan_only_and_promote_are_mutually_exclusive(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    with pytest.raises(SystemExit) as exc_info:
        mod.main(
            [
                "--period",
                "2026-Q2",
                "--director-slug",
                "Jesper-Tyrer",
                "--contracts",
                "QTR04_DealRisk_Scatter",
                "--plan-only",
                "--promote",
                "--repo-root",
                str(fixture_repo),
                "--run-id",
                "20260501-2125Z-both",
            ]
        )
    assert exc_info.value.code != 0


# --------------------------------------------------------------------------
# Deep preflight tests: ppttc JSON parse + bound deck zip readability
# --------------------------------------------------------------------------


def _corrupt_proof_artifact(repo: Path, contract_id: str, kind: str) -> None:
    """Replace one proof-companion artifact with a corrupt payload after fixture setup."""
    work = repo / "state" / "thinkcell_bridge" / "build_scaffold" / "2026-Q2" / "work" / contract_id
    if kind == "bound_deck":
        (work / f"{contract_id}-stock-donor-2026-Q2-bound.pptx").write_bytes(_corrupt_pptx_bytes())
    elif kind == "seed_pptx":
        (work / f"{contract_id}-stock-donor-seed.pptx").write_bytes(_corrupt_pptx_bytes())
    elif kind == "ppttc":
        (work / f"{contract_id}-stock-donor-2026-Q2.ppttc").write_text("{not: 'valid', json")
    else:  # pragma: no cover - defensive
        raise ValueError(f"unknown kind {kind}")


@pytest.mark.parametrize("kind", ["bound_deck", "seed_pptx", "ppttc"])
def test_plan_pilot_deep_preflight_fails_when_artifact_is_corrupt(fixture_repo: Path, kind: str):
    """Deep preflight surfaces zip/JSON breakage in source artifacts and blocks."""
    import run_thinkcell_insertion_pilot as mod

    _corrupt_proof_artifact(fixture_repo, "QTR04_DealRisk_Scatter", kind)
    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z-deepfail",
    )
    assert manifest["overall_status"] == "blocked"
    qtr04 = manifest["contracts"][0]
    assert qtr04["status"] == "preflight_failed"
    assert qtr04["preflight"]["ok"] is False
    assert any(failure["kind"] == kind for failure in qtr04["preflight"]["failures"])


def test_plan_pilot_deep_preflight_passes_records_each_check(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z-preflight",
    )
    qtr04 = manifest["contracts"][0]
    assert qtr04["status"] == "ready_for_vm_or_manual_insertion"
    assert qtr04["preflight"]["ok"] is True
    kinds = {check["kind"] for check in qtr04["preflight"]["checks"]}
    assert {"bound_deck", "seed_pptx", "ppttc"}.issubset(kinds)


# --------------------------------------------------------------------------
# Per-contract guardrail tests (ARR/ACV split, Type filter, native-table block)
# --------------------------------------------------------------------------


def test_qtr04_arr_acv_axis_calls_out_arr_and_land_expand(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z-guard",
    )
    qtr04 = manifest["contracts"][0]
    axis = qtr04["arr_acv_guardrails"]
    assert axis["axis"] == "ARR"
    assert axis["salesforce_field"] == "APTS_Opportunity_ARR__c"
    assert "Land" in axis["type_filter"] and "Expand" in axis["type_filter"]
    assert axis["must_not_blend_with"] == "Renewal ACV"


def test_qtr05_arr_acv_axis_calls_out_renewal_acv_and_renewal_type(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR05_FY26RenewalTimeline_Gantt"],
        run_id="20260501-2125Z-guard",
    )
    qtr05 = manifest["contracts"][0]
    axis = qtr05["arr_acv_guardrails"]
    assert axis["axis"] == "Renewal ACV"
    assert axis["salesforce_field"] == "APTS_Renewal_ACV__c"
    assert "Renewal" in axis["type_filter"]
    assert axis["must_not_blend_with"] == "ARR"


def test_per_contract_guardrails_block_native_thinkcell_tables(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z-guard",
    )
    guardrails = manifest["contracts"][0]["guardrails"]
    joined = " | ".join(guardrails).lower()
    assert "native think-cell editable tables" in joined or "native think-cell tables" in joined


# --------------------------------------------------------------------------
# vm_insertion_commands.sh + MANUAL_POWERPOINT_STEPS.md emission
# --------------------------------------------------------------------------


def test_write_run_dir_emits_executable_vm_insertion_commands_sh(
    fixture_repo: Path, tmp_path: Path
):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=[
            "QTR04_DealRisk_Scatter",
            "QTR05_FY26RenewalTimeline_Gantt",
        ],
        run_id="20260501-2125Z-vmsh",
    )
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    mod.write_run_dir(run_dir, manifest)

    sh_path = run_dir / "vm_insertion_commands.sh"
    assert sh_path.exists()
    text = sh_path.read_text()
    assert text.startswith("#!/")
    # Each per-contract VM bridge invocation is present and uses the proven ppttc.
    for contract_id in ("QTR04_DealRisk_Scatter", "QTR05_FY26RenewalTimeline_Gantt"):
        assert contract_id in text
        assert f"{contract_id}-stock-donor-2026-Q2.ppttc" in text
        assert f"{contract_id}-stock-donor-seed.pptx" in text
    assert "--template" in text
    assert "--expect-text" in text
    assert "Bound Term" in text
    # Operator must be told this is non-publishable.
    assert "publishable" in text.lower()
    # Script should be executable for the owner so the operator can run it directly.
    assert sh_path.stat().st_mode & 0o100


def test_write_run_dir_emits_manual_powerpoint_steps_md(fixture_repo: Path, tmp_path: Path):
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=[
            "QTR04_DealRisk_Scatter",
            "QTR05_FY26RenewalTimeline_Gantt",
        ],
        run_id="20260501-2125Z-manual",
    )
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    mod.write_run_dir(run_dir, manifest)

    md_path = run_dir / "MANUAL_POWERPOINT_STEPS.md"
    assert md_path.exists()
    text = md_path.read_text()
    assert "QTR04_DealRisk_Scatter" in text
    assert "QTR05_FY26RenewalTimeline_Gantt" in text
    # Manual fallback must remind the operator about the publishable + native-table rules.
    assert "publishable: false" in text.lower() or "not publishable" in text.lower()
    assert "native" in text.lower() and "table" in text.lower()


def test_blocked_run_does_not_emit_vm_commands_sh(fixture_repo: Path, tmp_path: Path):
    """If even one contract is blocked, do not hand the operator a script that would
    silently skip the bad contract — fail closed instead.
    """
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        director_slug="Patrick-Gaughan",
        contract_ids=["QTR04_DealRisk_Scatter"],
        run_id="20260501-2125Z-blockedsh",
    )
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    mod.write_run_dir(run_dir, manifest)
    assert manifest["overall_status"] == "blocked"
    assert not (run_dir / "vm_insertion_commands.sh").exists()


# --------------------------------------------------------------------------
# Candidate validator tests
# --------------------------------------------------------------------------


def test_validate_candidate_pptx_accepts_real_minimal_pptx(tmp_path: Path):
    import run_thinkcell_insertion_pilot as mod

    candidate = tmp_path / "candidate.pptx"
    candidate.write_bytes(_real_minimal_pptx_bytes())
    result = mod.validate_candidate_pptx(candidate)
    assert result["ok"] is True
    assert result["slide_count"] >= 1


def test_validate_candidate_pptx_rejects_corrupt_zip(tmp_path: Path):
    import run_thinkcell_insertion_pilot as mod

    candidate = tmp_path / "candidate.pptx"
    candidate.write_bytes(_corrupt_pptx_bytes())
    result = mod.validate_candidate_pptx(candidate)
    assert result["ok"] is False
    assert any("zip" in r.lower() or "read" in r.lower() for r in result["reasons"])


def test_validate_candidate_pptx_rejects_native_table(tmp_path: Path):
    import run_thinkcell_insertion_pilot as mod

    candidate = tmp_path / "candidate.pptx"
    candidate.write_bytes(_real_pptx_with_native_table_bytes())
    result = mod.validate_candidate_pptx(candidate)
    assert result["ok"] is False
    assert any("native" in r.lower() and "table" in r.lower() for r in result["reasons"])


def test_validate_candidate_pptx_rejects_missing_file(tmp_path: Path):
    import run_thinkcell_insertion_pilot as mod

    result = mod.validate_candidate_pptx(tmp_path / "nope.pptx")
    assert result["ok"] is False


# --------------------------------------------------------------------------
# --promote mode tests
# --------------------------------------------------------------------------


def _plan_only_manifest_for_jesper(repo: Path, run_id: str) -> dict:
    """Helper: run plan_pilot for Jesper QTR04+QTR05, write manifest, return parsed manifest."""
    import run_thinkcell_insertion_pilot as mod

    manifest = mod.plan_pilot(
        repo_root=repo,
        period="2026-Q2",
        director_slug="Jesper-Tyrer",
        contract_ids=[
            "QTR04_DealRisk_Scatter",
            "QTR05_FY26RenewalTimeline_Gantt",
        ],
        run_id=run_id,
    )
    run_dir = repo / "state" / "2026-Q2" / "__regional__" / "thinkcell_insertion_pilot" / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    mod.write_run_dir(run_dir, manifest)
    return manifest


def test_promote_with_no_candidates_marks_awaiting_vm_insertion(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    run_id = "20260501-2125Z-promote-empty"
    _plan_only_manifest_for_jesper(fixture_repo, run_id)

    promoted = mod.promote_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        run_id=run_id,
    )
    assert promoted["publishable"] is False
    assert promoted["overall_status"] == "ready_for_vm_or_manual_insertion"
    statuses = {c["contract_id"]: c["status"] for c in promoted["contracts"]}
    assert statuses == {
        "QTR04_DealRisk_Scatter": "awaiting_vm_insertion",
        "QTR05_FY26RenewalTimeline_Gantt": "awaiting_vm_insertion",
    }


def test_promote_with_valid_candidates_marks_candidate_created(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    run_id = "20260501-2125Z-promote-good"
    plan = _plan_only_manifest_for_jesper(fixture_repo, run_id)
    for contract in plan["contracts"]:
        candidate_path = Path(contract["intended_candidate_output_path"])
        candidate_path.parent.mkdir(parents=True, exist_ok=True)
        candidate_path.write_bytes(_real_minimal_pptx_bytes())

    promoted = mod.promote_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        run_id=run_id,
    )
    assert promoted["publishable"] is False
    assert promoted["overall_status"] == "candidate_created"
    statuses = {c["contract_id"]: c["status"] for c in promoted["contracts"]}
    assert statuses == {
        "QTR04_DealRisk_Scatter": "candidate_created",
        "QTR05_FY26RenewalTimeline_Gantt": "candidate_created",
    }
    for c in promoted["contracts"]:
        assert c["candidate_validation"]["ok"] is True
        assert Path(c["candidate_path"]).exists()


def test_promote_with_corrupt_candidate_marks_candidate_invalid(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    run_id = "20260501-2125Z-promote-corrupt"
    plan = _plan_only_manifest_for_jesper(fixture_repo, run_id)
    qtr04 = next(c for c in plan["contracts"] if c["contract_id"] == "QTR04_DealRisk_Scatter")
    qtr04_path = Path(qtr04["intended_candidate_output_path"])
    qtr04_path.parent.mkdir(parents=True, exist_ok=True)
    qtr04_path.write_bytes(_corrupt_pptx_bytes())
    qtr05 = next(
        c for c in plan["contracts"] if c["contract_id"] == "QTR05_FY26RenewalTimeline_Gantt"
    )
    qtr05_path = Path(qtr05["intended_candidate_output_path"])
    qtr05_path.parent.mkdir(parents=True, exist_ok=True)
    qtr05_path.write_bytes(_real_minimal_pptx_bytes())

    promoted = mod.promote_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        run_id=run_id,
    )
    assert promoted["publishable"] is False
    assert promoted["overall_status"] == "blocked"
    statuses = {c["contract_id"]: c["status"] for c in promoted["contracts"]}
    assert statuses["QTR04_DealRisk_Scatter"] == "candidate_invalid"
    assert statuses["QTR05_FY26RenewalTimeline_Gantt"] == "candidate_created"


def test_promote_rejects_native_thinkcell_table_in_candidate(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    run_id = "20260501-2125Z-promote-table"
    plan = _plan_only_manifest_for_jesper(fixture_repo, run_id)
    bad = next(c for c in plan["contracts"] if c["contract_id"] == "QTR04_DealRisk_Scatter")
    bad_path = Path(bad["intended_candidate_output_path"])
    bad_path.parent.mkdir(parents=True, exist_ok=True)
    bad_path.write_bytes(_real_pptx_with_native_table_bytes())
    good = next(
        c for c in plan["contracts"] if c["contract_id"] == "QTR05_FY26RenewalTimeline_Gantt"
    )
    good_path = Path(good["intended_candidate_output_path"])
    good_path.parent.mkdir(parents=True, exist_ok=True)
    good_path.write_bytes(_real_minimal_pptx_bytes())

    promoted = mod.promote_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        run_id=run_id,
    )
    qtr04 = next(c for c in promoted["contracts"] if c["contract_id"] == "QTR04_DealRisk_Scatter")
    assert qtr04["status"] == "candidate_invalid"
    assert any(
        "native" in r.lower() and "table" in r.lower()
        for r in qtr04["candidate_validation"]["reasons"]
    )
    assert promoted["overall_status"] == "blocked"
    assert promoted["publishable"] is False


def test_promote_without_existing_manifest_fails_closed(tmp_path: Path):
    import run_thinkcell_insertion_pilot as mod

    repo = tmp_path / "repo"
    state = repo / "state" / "thinkcell_bridge" / "build_scaffold" / "2026-Q2"
    state.mkdir(parents=True)
    (state / "thinkcell_build_scaffold.json").write_text(json.dumps(_build_scaffold_dict()))

    with pytest.raises(FileNotFoundError):
        mod.promote_pilot(
            repo_root=repo,
            period="2026-Q2",
            run_id="20260501-no-such-run",
        )


def test_promote_writes_back_manifest_and_run_log(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    run_id = "20260501-2125Z-promote-writeback"
    plan = _plan_only_manifest_for_jesper(fixture_repo, run_id)
    for contract in plan["contracts"]:
        candidate_path = Path(contract["intended_candidate_output_path"])
        candidate_path.parent.mkdir(parents=True, exist_ok=True)
        candidate_path.write_bytes(_real_minimal_pptx_bytes())

    mod.promote_pilot(
        repo_root=fixture_repo,
        period="2026-Q2",
        run_id=run_id,
    )
    run_dir = (
        fixture_repo / "state" / "2026-Q2" / "__regional__" / "thinkcell_insertion_pilot" / run_id
    )
    on_disk = json.loads((run_dir / "manifest.json").read_text())
    assert on_disk["overall_status"] == "candidate_created"
    # promote_pilot should re-emit RUN_LOG.md so the latest status is reflected.
    assert "candidate_created" in (run_dir / "RUN_LOG.md").read_text()


def test_cli_promote_smoke(fixture_repo: Path):
    import run_thinkcell_insertion_pilot as mod

    run_id = "20260501-2125Z-cli-promote"
    plan = _plan_only_manifest_for_jesper(fixture_repo, run_id)
    for contract in plan["contracts"]:
        candidate_path = Path(contract["intended_candidate_output_path"])
        candidate_path.parent.mkdir(parents=True, exist_ok=True)
        candidate_path.write_bytes(_real_minimal_pptx_bytes())

    rc = mod.main(
        [
            "--period",
            "2026-Q2",
            "--director-slug",
            "Jesper-Tyrer",
            "--contracts",
            "QTR04_DealRisk_Scatter",
            "QTR05_FY26RenewalTimeline_Gantt",
            "--promote",
            "--repo-root",
            str(fixture_repo),
            "--run-id",
            run_id,
        ]
    )
    assert rc == 0
    parsed = json.loads(
        (
            fixture_repo
            / "state"
            / "2026-Q2"
            / "__regional__"
            / "thinkcell_insertion_pilot"
            / run_id
            / "manifest.json"
        ).read_text()
    )
    assert parsed["overall_status"] == "candidate_created"
    assert parsed["publishable"] is False
