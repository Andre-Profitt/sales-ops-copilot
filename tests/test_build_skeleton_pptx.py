"""Tests for scripts/build_skeleton_pptx.py.

The skeleton builder copies the canonical SimCorp LAND template and
makes it 28 slides matching the registry slide IDs. No Think-Cell
charts are inserted here — that happens manually on the VM.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from pptx import Presentation

REPO = Path(__file__).resolve().parents[1]
PY = sys.executable
SCRIPT = REPO / "scripts" / "build_skeleton_pptx.py"


def test_skeleton_has_28_slides(tmp_path: Path) -> None:
    out = tmp_path / "skel.pptx"
    proc = subprocess.run(
        [PY, str(SCRIPT), "--out", str(out)],
        capture_output=True,
        text=True,
        cwd=REPO,
    )
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert out.exists()
    prs = Presentation(str(out))
    assert len(prs.slides) == 28, f"expected 28 slides, got {len(prs.slides)}"


def test_skeleton_has_no_forbidden_strings(tmp_path: Path) -> None:
    out = tmp_path / "skel.pptx"
    subprocess.run([PY, str(SCRIPT), "--out", str(out)], check=True, cwd=REPO)
    prs = Presentation(str(out))
    bad = ["[think-cell", "Lorem ipsum", "paste from", "User count [K]"]
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text  # type: ignore[attr-defined]
                for b in bad:
                    assert b not in txt, f"forbidden text in skeleton: {b!r}"
