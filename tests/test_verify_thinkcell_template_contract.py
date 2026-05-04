"""Tests for verify_thinkcell_template_contract.py.

The contract verifier compares a (possibly-Think-Cell-wired) PPTX
against a binding registry. It checks named-shape coverage, debris
strings, and surfaces a JSON report.
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

REPO = Path(__file__).resolve().parents[1]
PY = sys.executable
SCRIPT = REPO / "scripts" / "verify_thinkcell_template_contract.py"
FIXTURES = REPO / "tests" / "fixtures"


def _make_pptx_with_named_shapes(
    path: Path, names: list[str], debris: list[str] | None = None
) -> None:
    """Create a synthetic pptx with one shape per name, and optional debris textboxes."""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    for i, name in enumerate(names):
        slide = prs.slides.add_slide(blank_layout)
        tb = slide.shapes.add_textbox(Inches(1), Inches(1 + i * 0.1), Inches(2), Inches(0.5))
        tb.name = name
        tb.text_frame.text = f"placeholder for {name}"
    if debris:
        slide = prs.slides.add_slide(blank_layout)
        for d in debris:
            tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.5))
            tb.text_frame.text = d
    prs.save(str(path))


def _registry_path() -> Path:
    return REPO / "config" / "thinkcell" / "land_review_full_28.binding_registry.yml"


def _run(template: Path, registry: Path, out: Path) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            PY,
            str(SCRIPT),
            "--template",
            str(template),
            "--registry",
            str(registry),
            "--out",
            str(out),
        ],
        capture_output=True,
        text=True,
        cwd=REPO,
    )


def test_pptx_with_all_named_shapes_passes(tmp_path: Path) -> None:
    """A synthetic pptx that has every required shape name passes the contract."""
    import yaml

    registry = yaml.safe_load(_registry_path().read_text())
    required = [
        el["name"] for s in registry["slides"] for el in s.get("elements", []) if el.get("required")
    ]
    template = tmp_path / "all_present.pptx"
    _make_pptx_with_named_shapes(template, required)
    out = tmp_path / "report.json"

    proc = _run(template, _registry_path(), out)
    assert proc.returncode == 0, proc.stdout + proc.stderr

    report = json.loads(out.read_text())
    assert report["status"] == "pass"
    assert report["required_elements_missing"] == []


def test_pptx_missing_required_shape_fails(tmp_path: Path) -> None:
    template = tmp_path / "missing.pptx"
    _make_pptx_with_named_shapes(template, ["S01_DirectorName"])
    out = tmp_path / "report.json"

    proc = _run(template, _registry_path(), out)
    assert proc.returncode != 0
    report = json.loads(out.read_text())
    assert report["status"] == "fail"
    assert len(report["required_elements_missing"]) > 0
    assert "S05_PipelineByStage" in report["required_elements_missing"]


def test_pptx_with_debris_text_fails(tmp_path: Path) -> None:
    import yaml

    registry = yaml.safe_load(_registry_path().read_text())
    required = [
        el["name"] for s in registry["slides"] for el in s.get("elements", []) if el.get("required")
    ]
    template = tmp_path / "debris.pptx"
    _make_pptx_with_named_shapes(
        template,
        required,
        debris=["[think-cell TABLE WITH FORMATTING — datalinked]", "Lorem ipsum dolor sit amet"],
    )
    out = tmp_path / "report.json"

    proc = _run(template, _registry_path(), out)
    assert proc.returncode != 0
    report = json.loads(out.read_text())
    assert report["status"] == "fail"
    assert len(report["forbidden_text"]) >= 1
