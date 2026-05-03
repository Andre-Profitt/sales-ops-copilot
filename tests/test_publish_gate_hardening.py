"""Tests for the May 2026 / 2026-Q2 publish-gate P0/P1 hardening lane.

Covers:
- Expanded error-token list in FORBIDDEN_TEXT_NEEDLES (#REF, #VALUE, #DIV/0, #N/A).
- Conservative internal/test pollution patterns (SimCorp Test, Demo Account, ...).
- Stale/non-May 2026 period-label rejection driven by period_context.
- Workbook cached-error scanner (#NAME / #NULL / #REF / #VALUE / #DIV/0 / #N/A).
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Mapping

import pytest
from openpyxl import Workbook

# scripts/ is not a package for the top-level orchestrator; load it like other tests.
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from run_regional_deck_publish_gate import (  # noqa: E402  - sys.path mutation above
    FORBIDDEN_TEXT_NEEDLES,
    forbidden_period_label_findings,
    workbook_error_findings,
)


# ---------- A-3 / A-4 / C-2 forbidden-text expansion --------------------------


def test_forbidden_text_needles_includes_expanded_error_tokens() -> None:
    """C-2 (P0): every Excel error token must be deck-text forbidden."""

    needles_lower = {needle.lower() for needle in FORBIDDEN_TEXT_NEEDLES}
    for token in ("#name", "#null", "#ref", "#value", "#div/0", "#n/a"):
        assert token in needles_lower, f"expected {token!r} in FORBIDDEN_TEXT_NEEDLES"


def test_forbidden_text_needles_includes_conservative_pollution_patterns() -> None:
    """A-1 / A-4 (P0): block obvious internal/demo deck pollution.

    Conservative scope: multi-word patterns that cannot legitimately appear in
    SimCorp logo/footer text or in a real customer name (per Andre's guidance).
    """

    needles_lower = {needle.lower() for needle in FORBIDDEN_TEXT_NEEDLES}
    for needle in (
        "simcorp test",
        "demo account",
        "sample account",
        "sample co",
        "lorem ipsum",
    ):
        assert needle in needles_lower, f"expected {needle!r} in FORBIDDEN_TEXT_NEEDLES"


def test_forbidden_text_needles_does_not_block_simcorp_logo_phrases() -> None:
    """Regression guard: pure 'SimCorp' brand mentions must remain allowed.

    The publish gate's forbidden-text matching is case-insensitive substring,
    so a needle like 'SimCorp' alone would false-positive every footer.
    """

    needles_lower = {needle.lower() for needle in FORBIDDEN_TEXT_NEEDLES}
    assert "simcorp" not in needles_lower
    # Common legit footer fragments must not collide with any needle.
    legit_footer = "© 2026 SimCorp. SimCorp Dimension. SimCorp logo."
    for needle in FORBIDDEN_TEXT_NEEDLES:
        assert needle.lower() not in legit_footer.lower(), needle


# ---------- D-1 forbidden period labels --------------------------------------


def test_period_label_findings_returns_empty_for_clean_text() -> None:
    text = "May 2026 Land+Expand Renewal ACV unweighted forecast"
    assert forbidden_period_label_findings(text, period="2026-Q2") == []


def test_period_label_findings_blocks_other_2026_months() -> None:
    findings = forbidden_period_label_findings(
        "April 2026 close-out, January 2026 wins", period="2026-Q2"
    )
    joined = " | ".join(findings)
    assert "April 2026" in joined
    assert "January 2026" in joined


def test_period_label_findings_blocks_off_quarter_labels() -> None:
    findings = forbidden_period_label_findings(
        "Q1 2026 deals, Q3 2026 forecast, Q4 2026 carry", period="2026-Q2"
    )
    joined = " | ".join(findings)
    assert "Q1 2026" in joined
    assert "Q3 2026" in joined
    assert "Q4 2026" in joined


def test_period_label_findings_allows_may_2026_label() -> None:
    text = "Friday, May 1, 2026 kickoff. May 2026 review. Q2 2026 Land+Expand."
    assert forbidden_period_label_findings(text, period="2026-Q2") == []


def test_period_label_findings_is_case_insensitive() -> None:
    findings = forbidden_period_label_findings("january 2026 deals", period="2026-Q2")
    assert findings, "expected case-insensitive match for 'january 2026'"
    assert any("january 2026" in f.lower() for f in findings)


def test_period_label_findings_unsupported_period_raises() -> None:
    """Period coverage must stay loud: an unsupported period rejects, not silently passes."""

    with pytest.raises(ValueError):
        forbidden_period_label_findings("anything", period="2026-Q3")


# ---------- C-1 workbook cached-error scanner --------------------------------


def _save_wb(tmp_path: Path, name: str, sheet_title: str, cells: Mapping[str, object]) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # always present on a fresh Workbook; narrows for the type checker
    ws.title = sheet_title
    for coord, value in cells.items():
        ws[coord] = value  # type: ignore[assignment]
    path = tmp_path / name
    wb.save(str(path))
    return path


def test_workbook_error_findings_clean_workbook_returns_empty(tmp_path: Path) -> None:
    path = _save_wb(
        tmp_path,
        "clean.xlsx",
        "Out_S04",
        {"A1": "Header", "A2": "Real Asset Manager", "B2": 1234.56},
    )
    assert workbook_error_findings(path) == []


@pytest.mark.parametrize(
    "token",
    ["#REF!", "#NAME?", "#NULL!", "#VALUE!", "#DIV/0!", "#N/A"],
)
def test_workbook_error_findings_finds_cached_error_value(tmp_path: Path, token: str) -> None:
    """C-1 (P0): cached error values in any sheet must be reported as blockers."""

    safe = token.replace("/", "_").replace("!", "").replace("?", "").replace("#", "h")
    path = _save_wb(tmp_path, f"err_{safe}.xlsx", "Out_S07", {"A1": "Header", "A2": token})
    findings = workbook_error_findings(path)
    assert findings, f"expected at least one finding for {token}"
    assert any(token in finding for finding in findings)


def test_workbook_error_findings_does_not_match_token_inside_prose(tmp_path: Path) -> None:
    """Prose mentioning '#N/A' (not the cell value itself) must not trip the gate."""

    path = _save_wb(
        tmp_path,
        "prose.xlsx",
        "Notes",
        {
            "A1": "Header",
            "A2": "We treat missing values as not-applicable (use #N/A in commentary only).",
        },
    )
    assert workbook_error_findings(path) == []


def test_workbook_error_findings_caps_findings(tmp_path: Path) -> None:
    """Bound output so a wholly broken workbook does not generate megabytes of blockers."""

    cells = {f"A{i}": "#REF!" for i in range(1, 100)}
    path = _save_wb(tmp_path, "all_broken.xlsx", "Out_S07", cells)
    findings = workbook_error_findings(path)
    assert findings, "expected non-empty findings"
    assert len(findings) <= 12, "findings should be capped to keep the gate report readable"


def test_workbook_error_findings_scans_all_sheets(tmp_path: Path) -> None:
    """A multi-sheet workbook with the error on a non-active sheet must still trigger."""

    wb = Workbook()
    ws_first: Worksheet = wb.active  # type: ignore[assignment]
    ws_first.title = "Out_S04"
    ws_first["A1"] = "Clean header"
    ws_second = wb.create_sheet("Out_S07")
    ws_second["A1"] = "Header"
    ws_second["B5"] = "#REF!"
    path = tmp_path / "multi.xlsx"
    wb.save(str(path))
    findings = workbook_error_findings(path)
    assert findings
    assert any("Out_S07" in finding for finding in findings)
