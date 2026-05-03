from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from run_template_contract_gate import TemplateContract, inspect_template  # noqa: E402


def _deck(path: Path, *, stale: bool = False) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5))
    box.text_frame.text = "Template"
    if stale:
        stale_shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(0.5))
        stale_shape.name = "think-cell data - do not delete"
    prs.save(path)


def test_clean_shell_contract_passes(tmp_path: Path) -> None:
    asset = tmp_path / "assets" / "LAND_template.pptx"
    asset.parent.mkdir()
    _deck(asset)

    result = inspect_template(
        TemplateContract(
            path="assets/LAND_template.pptx",
            role="clean_shell",
            expected_slides=1,
            expected_names=0,
            max_embedded_ole=0,
            max_stale_tokens=0,
        ),
        root=tmp_path,
    )

    assert result.status == "pass"
    assert result.stale_thinkcell_token_count == 0


def test_clean_shell_contract_blocks_stale_thinkcell_tokens(tmp_path: Path) -> None:
    asset = tmp_path / "assets" / "LAND_template.pptx"
    asset.parent.mkdir()
    _deck(asset, stale=True)

    result = inspect_template(
        TemplateContract(
            path="assets/LAND_template.pptx",
            role="clean_shell",
            expected_slides=1,
            expected_names=0,
            max_embedded_ole=0,
            max_stale_tokens=0,
        ),
        root=tmp_path,
    )

    assert result.status == "fail"
    assert result.stale_thinkcell_token_count > 0
    assert any("stale_thinkcell" in finding for finding in result.findings)
