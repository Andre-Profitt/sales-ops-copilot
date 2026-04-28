"""Excel companion — 15-sheet workbook from trends envelope."""

from __future__ import annotations
import json
from pathlib import Path
from openpyxl import load_workbook


FIXTURES = Path(__file__).parent / "fixtures"


def test_excel_has_15_sheets(tmp_path):
    from scripts.excel_companion import build_director_excel
    from scripts.land_brief import build_trends_envelope, derive_highlights_risks

    sf = json.loads((FIXTURES / "sample_director.json").read_text())
    director = {
        "name": "Adam Steinhouse",
        "book_codes": ["P&I"],
        "where_clause": "Region__c IN ('Americas')",
        "scope": "us_only",
    }
    env = derive_highlights_risks(build_trends_envelope(sf, director, "2026-Q2"))

    out_path = tmp_path / "land.xlsx"
    build_director_excel(env, out_path)

    wb = load_workbook(out_path)
    assert len(wb.sheetnames) == 15, (
        f"expected 15 sheets, got {len(wb.sheetnames)}: {wb.sheetnames}"
    )
    assert "Cover" in wb.sheetnames
    assert "Methodology" in wb.sheetnames
    assert any("Pipeline" in s for s in wb.sheetnames)
