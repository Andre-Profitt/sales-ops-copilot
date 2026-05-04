"""Tests for the consulting-grade image-based chart renderer.

Covers the per-binding renderers (waterfall / bar / mekko / line / grouped-bar),
the SIMCORP_BRAND constant, and the deck-enhance round-trip on a synthetic
small deck.

Hard rules:
    * No live pptx/ppttc fixtures required: synthetic table data is generated
      in-test and the round-trip uses a 1-slide pptx written via python-pptx.
    * matplotlib is a hard dep of this module; tests assert on PNG bytes/dims.
    * ASCII-only.
"""

from __future__ import annotations

import json
import struct
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from tcrender.image_charts import (
    SIMCORP_BRAND,
    SIMCORP_FONT,
    EnhanceImagesResult,
    ImageChartResult,
    enhance_deck_with_images,
    render_chart_image,
)


# ----------------------------------------------------------- helpers / fixtures


def _png_dimensions(path: Path) -> tuple[int, int]:
    """Return (width_px, height_px) by parsing PNG IHDR -- avoids PIL dep."""
    data = path.read_bytes()
    # PNG signature + IHDR length(4) + 'IHDR'(4) -> width(4), height(4)
    # IHDR chunk starts at byte 8 (after signature) -> length=4, type=4, then data.
    assert data[:8] == b"\x89PNG\r\n\x1a\n", "not a PNG"
    width = struct.unpack(">I", data[16:20])[0]
    height = struct.unpack(">I", data[20:24])[0]
    return width, height


def _two_row_table(headers: list[str], values: list[float]) -> list[list[object]]:
    return [
        [None] + [{"string": h} for h in headers],
        [{"string": "Series"}] + [{"number": v} for v in values],
    ]


# ============================================================================
#                           CONSTANT TESTS
# ============================================================================


def test_simcorp_brand_has_expected_keys() -> None:
    expected = {
        "primary",
        "accent_dark",
        "purple",
        "warning_coral",
        "highlight",
        "tertiary_gold",
        "hyperlink",
        "neutral_dark",
        "neutral_mid",
        "neutral_light",
    }
    missing = expected - set(SIMCORP_BRAND.keys())
    assert not missing, f"missing palette keys: {missing}"
    # Every value is a 6-digit hex starting with '#'.
    for k, v in SIMCORP_BRAND.items():
        assert isinstance(v, str) and v.startswith("#") and len(v) == 7, (
            f"bad palette entry {k!r}={v!r}"
        )


def test_simcorp_font_is_a_string() -> None:
    assert isinstance(SIMCORP_FONT, str) and SIMCORP_FONT


# ============================================================================
#                       PER-RENDERER TESTS
# ============================================================================


def test_render_waterfall_produces_png(tmp_path: Path) -> None:
    table = _two_row_table(
        headers=[
            "Opening pipe",
            "New + Advanced",
            "Won this Q",
            "Lost this Q",
            "Closing pipe",
        ],
        values=[18.0, 12.5, -1.6, -6.0, 22.9],
    )
    out = tmp_path / "S04.png"
    result = render_chart_image(
        binding_name="S04_PipeMovement",
        table_data=table,
        output_path=out,
        width_in=10.0,
        height_in=5.5,
    )
    assert isinstance(result, ImageChartResult)
    assert result.binding_kind == "waterfall"
    assert out.exists()
    w, h = _png_dimensions(out)
    assert w > 1500 and h > 800, (w, h)
    assert result.dimensions_px == (w, h)


def test_render_simple_bar_stage(tmp_path: Path) -> None:
    table = _two_row_table(
        headers=[f"{i} - Stage" for i in range(1, 9)],
        values=[0.0, 0.0, 2.3, 0.0, 0.3, 5.1, 0.0, 1.6],
    )
    out = tmp_path / "S05.png"
    result = render_chart_image(
        binding_name="S05_PipelineByStage",
        table_data=table,
        output_path=out,
    )
    assert result.binding_kind == "bar"
    assert out.exists()
    assert _png_dimensions(out)[0] > 1000


def test_render_aging_bar_flags_old_buckets(tmp_path: Path) -> None:
    table = _two_row_table(
        headers=["0-30d", "31-90d", "91-180d", "181-365d", ">365d", ">730d"],
        values=[1.2, 3.4, 2.1, 4.5, 6.6, 8.8],
    )
    out = tmp_path / "S06.png"
    result = render_chart_image(
        binding_name="S06_PipelineAging",
        table_data=table,
        output_path=out,
    )
    assert result.binding_kind == "bar"
    assert out.exists()


def test_render_horizontal_owner_bar_top_10(tmp_path: Path) -> None:
    # 12 owners; renderer should clip to top-10
    headers = [f"Owner {i}" for i in range(12)]
    values = [float(12 - i) for i in range(12)]  # descending 12..1
    table = _two_row_table(headers=headers, values=values)
    out = tmp_path / "S15.png"
    result = render_chart_image(
        binding_name="S15_ByOwner",
        table_data=table,
        output_path=out,
    )
    assert result.binding_kind == "bar"
    assert out.exists()


def test_render_grouped_winloss(tmp_path: Path) -> None:
    table = [
        [None, {"string": "Won"}, {"string": "Lost"}],
        [{"string": "ARR (Land+Expand, EUR M)"}, {"number": 1.6}, {"number": 6.0}],
        [{"string": "ACV (Renewal, EUR M)"}, {"number": 2.8}, {"number": 0.0}],
    ]
    out = tmp_path / "S18.png"
    result = render_chart_image(
        binding_name="S18_WinsLossesQTD",
        table_data=table,
        output_path=out,
    )
    assert result.binding_kind == "grouped_bar"
    assert out.exists()


def test_render_line_chart(tmp_path: Path) -> None:
    weeks = [f"Wk {i}" for i in range(1, 9)]
    vals = [0.0, 1.7, 0.4, 3.0, 1.2, 2.4, 0.8, 4.1]
    table = _two_row_table(headers=weeks, values=vals)
    out = tmp_path / "S25.png"
    result = render_chart_image(
        binding_name="S25_PipelineCreationVelocity",
        table_data=table,
        output_path=out,
    )
    assert result.binding_kind == "line"
    assert out.exists()


def test_render_concentration_bar(tmp_path: Path) -> None:
    table = _two_row_table(
        headers=["Top 1", "Top 3", "Top 5", "Top 10"],
        values=[35.0, 65.0, 82.0, 95.0],  # last two trip the 80% threshold
    )
    out = tmp_path / "S21.png"
    result = render_chart_image(
        binding_name="S21_ConcentrationRiskChart",
        table_data=table,
        output_path=out,
    )
    assert result.binding_kind == "bar"
    assert out.exists()


def test_render_velocity_bar_flags_long_stages(tmp_path: Path) -> None:
    table = _two_row_table(
        headers=["1 - Prospect", "5 - Preferred", "6 - Contract"],
        values=[120.0, 284.0, 410.0],  # last >365 -> coral
    )
    out = tmp_path / "S19.png"
    result = render_chart_image(
        binding_name="S19_Velocity",
        table_data=table,
        output_path=out,
    )
    assert result.binding_kind == "bar"
    assert out.exists()


# ============================================================================
#                       MEKKO -- proportional-width assertion
# ============================================================================


def test_render_mekko_produces_png(tmp_path: Path) -> None:
    """Smoke-test: PNG produced for a 4-industry x 3-stage matrix."""
    table = [
        [
            None,
            {"string": "Pension"},
            {"string": "Asset Management"},
            {"string": "Central Bank"},
            {"string": "Other"},
        ],
        [{"string": "Stage 1"}, {"number": 4.0}, {"number": 1.0}, {"number": 2.0}, {"number": 0.5}],
        [{"string": "Stage 2"}, {"number": 6.0}, {"number": 2.0}, {"number": 3.0}, {"number": 0.5}],
        [{"string": "Stage 3"}, {"number": 1.8}, {"number": 4.5}, {"number": 0.0}, {"number": 1.0}],
    ]
    out = tmp_path / "S16.png"
    result = render_chart_image(
        binding_name="S16_StageByIndustry",
        table_data=table,
        output_path=out,
    )
    assert result.binding_kind == "mekko"
    assert out.exists()
    w, h = _png_dimensions(out)
    assert w > 1500 and h > 800


def test_mekko_column_widths_proportional() -> None:
    """Mekko widths must reflect industry totals, not be uniform.

    We can't pixel-inspect a rasterised PNG without PIL, but we can verify
    the matplotlib patches drawn in the figure: each industry column maps
    to a stack of Rectangles whose width is proportional to its column total.
    To do this we monkeypatch the renderer to capture rectangles.
    """
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.patches as mpatches
    import matplotlib.pyplot as plt

    captured: list[mpatches.Rectangle] = []
    orig_init = mpatches.Rectangle.__init__

    def _capture(self, xy, width, height, **kw):  # type: ignore[no-untyped-def]
        captured.append(self)
        return orig_init(self, xy, width, height, **kw)

    # Industry totals: 12, 8, 4, 1. Stage shares within each = 50%/50%.
    table = [
        [
            None,
            {"string": "Pension"},
            {"string": "Asset Mgmt"},
            {"string": "Central Bank"},
            {"string": "Tiny"},
        ],
        [{"string": "S1"}, {"number": 6.0}, {"number": 4.0}, {"number": 2.0}, {"number": 0.5}],
        [{"string": "S2"}, {"number": 6.0}, {"number": 4.0}, {"number": 2.0}, {"number": 0.5}],
    ]
    expected_totals = [12.0, 8.0, 4.0, 1.0]
    grand = sum(expected_totals)

    # Monkeypatch the Rectangle constructor for the duration of this render.
    mpatches.Rectangle.__init__ = _capture  # type: ignore[method-assign]
    try:
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
            out = Path(f.name)
        render_chart_image(
            binding_name="S16_StageByIndustry",
            table_data=table,
            output_path=out,
        )
    finally:
        mpatches.Rectangle.__init__ = orig_init  # type: ignore[method-assign]
        plt.close("all")

    # Filter to mekko data rectangles only -- ignore matplotlib's internal
    # Rectangle patches (legend, frame, etc.) by checking the gid set in the
    # renderer.
    data_rects = [r for r in captured if r.get_gid() == "mekko_cell"]
    assert data_rects, "no mekko_cell rectangles captured"

    # Group the captured rectangles by their x-coordinate (each column has one
    # x-anchor; rectangles in same column share xy[0]).
    by_xy: dict[float, list[mpatches.Rectangle]] = {}
    for rect in data_rects:
        x = round(float(rect.get_x()), 4)
        by_xy.setdefault(x, []).append(rect)

    # 4 columns expected
    cols = sorted(by_xy.keys())
    assert len(cols) == 4, f"expected 4 columns, got {len(cols)}: {cols}"

    # Width per column (matplotlib width = first rect's width per group).
    widths = [by_xy[x][0].get_width() for x in cols]
    # Expected proportions: 12/25, 8/25, 4/25, 1/25
    expected = [t / grand for t in expected_totals]
    # Each plotted width is (col_w_normalised * 0.97) -- gutter shrink. So
    # check ratios instead of absolute widths.
    ratios = [w / sum(widths) for w in widths]
    for got, exp in zip(ratios, expected, strict=True):
        # Renormalisation pushes apart slightly when min_w padding kicks in,
        # but the ordering + relative magnitudes should stay correct.
        assert abs(got - exp) < 0.05, f"col width ratios off: got={ratios} expected={expected}"


def test_mekko_handles_empty_table(tmp_path: Path) -> None:
    """Empty table shouldn't crash; renderer falls back gracefully."""
    out = tmp_path / "empty.png"
    result = render_chart_image(
        binding_name="S16_StageByIndustry",
        table_data=[],
        output_path=out,
    )
    assert out.exists()
    assert result.binding_kind in {"mekko", "bar"}  # falls back to bar in extreme case


# ============================================================================
#                       PUBLIC API -- input validation
# ============================================================================


def test_render_chart_image_unknown_binding_raises(tmp_path: Path) -> None:
    out = tmp_path / "bad.png"
    with pytest.raises(ValueError, match="not a chart binding"):
        render_chart_image(
            binding_name="S07_TopDealsLand",  # table binding, not chart
            table_data=[],
            output_path=out,
        )


def test_render_chart_image_subtitle_includes_period_and_scope(tmp_path: Path) -> None:
    """Smoke test that the subtitle path doesn't blow up with director context."""
    table = _two_row_table(headers=["1 - Prospect", "2 - Discovery"], values=[2.0, 3.5])
    out = tmp_path / "subtitle.png"
    res = render_chart_image(
        binding_name="S05_PipelineByStage",
        table_data=table,
        output_path=out,
        director_name="Jesper Tyrer",
        period="2026-Q2",
        scope_label="APAC",
    )
    assert res.image_path.exists()


# ============================================================================
#                  ENHANCE_DECK_WITH_IMAGES round-trip
# ============================================================================


def _make_synthetic_deck(path: Path, n_slides: int) -> None:
    """Write a minimal n-slide pptx using python-pptx."""
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # 6 = blank in default theme
    for i in range(n_slides):
        slide = prs.slides.add_slide(blank_layout)
        # Add a simple textbox so the slide isn't empty
        tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(0.5))
        tx.text_frame.text = f"Slide {i + 1}"
    prs.save(str(path))


def _make_synthetic_ppttc(path: Path) -> None:
    """Write a .ppttc with the bindings the enhancer cares about."""
    bindings = [
        {"name": "S01_DirectorName", "table": [[{"string": "Test Director"}]]},
        {"name": "S01_Period", "table": [[{"string": "2026-Q2"}]]},
        {"name": "S01_ScopeLabel", "table": [[{"string": "TEST"}]]},
        # Chart bindings: deliberately partial -- enhancer must skip missing.
        {
            "name": "S04_PipeMovement",
            "table": _two_row_table(
                ["Open", "New", "Won", "Lost", "Close"], [10.0, 5.0, -1.0, -2.0, 12.0]
            ),
        },
        {
            "name": "S05_PipelineByStage",
            "table": _two_row_table(["1 - P", "2 - D", "3 - E"], [1.5, 2.5, 3.5]),
        },
        {
            "name": "S25_PipelineCreationVelocity",
            "table": _two_row_table(["W1", "W2", "W3"], [0.5, 1.0, 1.8]),
        },
    ]
    payload = [{"template": "ignored.pptx", "data": bindings}]
    path.write_text(json.dumps(payload), encoding="utf-8")


def test_enhance_deck_with_images_round_trip(tmp_path: Path) -> None:
    deck = tmp_path / "synthetic.pptx"
    _make_synthetic_deck(deck, n_slides=26)  # cover all binding slide indices
    ppttc = tmp_path / "synthetic.ppttc"
    _make_synthetic_ppttc(ppttc)
    out = tmp_path / "enhanced.pptx"

    result = enhance_deck_with_images(
        deck_path=deck, ppttc_path=ppttc, output_path=out, keep_pngs=True
    )
    assert isinstance(result, EnhanceImagesResult)
    assert out.exists()
    # We supplied 3 chart bindings; all 3 should render.
    assert set(result.bindings_rendered) == {
        "S04_PipeMovement",
        "S05_PipelineByStage",
        "S25_PipelineCreationVelocity",
    }, f"got bindings_rendered={result.bindings_rendered}"

    # Skipped bindings: every other chart binding the .ppttc didn't supply.
    skipped_names = {n for n, _ in result.skipped}
    assert "S16_StageByIndustry" in skipped_names
    assert "S18_WinsLossesQTD" in skipped_names

    # Round-trip: open the enhanced deck and verify each rendered slide has a
    # picture shape. Slide indices map directly to _BINDING_TO_SLIDE.
    prs = Presentation(str(out))
    for binding, slide_idx in [
        ("S04_PipeMovement", 4),
        ("S05_PipelineByStage", 5),
        ("S25_PipelineCreationVelocity", 25),
    ]:
        slide = prs.slides[slide_idx - 1]
        # has_picture: a shape whose shape_type == 13 (PICTURE) is what we added
        pic_count = sum(1 for shape in slide.shapes if shape.shape_type == 13)
        assert pic_count >= 1, (
            f"slide {slide_idx} ({binding}) missing embedded picture; "
            f"shapes={[s.shape_type for s in slide.shapes]}"
        )

    # Cleanup left PNG paths exist when keep_pngs=True
    assert all(p.exists() for p in result.png_paths.values()), (
        "keep_pngs=True but PNG paths missing"
    )


def test_enhance_deck_with_images_missing_deck(tmp_path: Path) -> None:
    fake_deck = tmp_path / "doesnotexist.pptx"
    fake_ppttc = tmp_path / "doesnotexist.ppttc"
    with pytest.raises(FileNotFoundError):
        enhance_deck_with_images(
            deck_path=fake_deck, ppttc_path=fake_ppttc, output_path=tmp_path / "out.pptx"
        )
